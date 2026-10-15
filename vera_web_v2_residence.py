"""Persistent residence roster, projected from employee profiles without bank data."""
from datetime import datetime
from io import BytesIO

from fastapi import Depends, HTTPException
from fastapi.responses import StreamingResponse, Response
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from pydantic import BaseModel, Field
from sqlalchemy import text

from vera_web_v2_hr import registry, department_code
from vera_auto_check import _norm

COLUMNS = [
    ('username', 'Tên nhân viên'), ('full_name', 'Họ và tên'), ('department', 'Bộ phận'),
    ('birth_date', 'Ngày sinh'), ('gender', 'Giới tính'), ('ethnicity', 'Dân tộc'),
    ('phone', 'Số điện thoại'), ('email', 'Email'), ('cccd_number', 'Số CCCD'),
    ('cccd_issue_date', 'Ngày cấp CCCD'), ('cccd_issue_place', 'Nơi cấp CCCD'),
    ('address', 'Địa chỉ'), ('province', 'Tỉnh/Thành phố'), ('district', 'Quận/Huyện'),
    ('ward', 'Xã/Phường'), ('address_detail', 'Địa chỉ chi tiết'),
    ('employment_status', 'Trạng thái làm việc'), ('employment_start_date', 'Ngày bắt đầu làm'),
    ('employment_end_date', 'Ngày nghỉ việc'), ('work_shift', 'Ca làm việc'),
    ('shift_start_date', 'Ngày bắt đầu ca'), ('rotation_cycle', 'Chu kỳ ca'),
    ('monthly_generated', 'Phép phát sinh tháng'), ('monthly_leave', 'Phép tháng'),
    ('annual_leave', 'Phép năm'),
]
PAYLOAD_FIELDS = {
    'gender': 'Giới tính', 'ethnicity': 'Dân tộc', 'cccd_number': 'Số CCCD',
    'cccd_issue_date': 'Ngày cấp CCCD', 'cccd_issue_place': 'Nơi cấp CCCD',
    'province': 'Tỉnh/Thành phố', 'district': 'Quận/Huyện', 'ward': 'Xã/Phường',
    'address_detail': 'Địa chỉ chi tiết', 'employment_end_date': 'Ngày nghỉ việc',
}
DATE_FIELDS = {'birth_date', 'cccd_issue_date', 'employment_start_date', 'employment_end_date', 'shift_start_date'}


class Membership(BaseModel):
    username: str = Field(min_length=1, max_length=300)
    included: bool = True


class LaborReportRequest(BaseModel):
    usernames: list[str] = Field(min_length=1, max_length=1000)


def display_date(value):
    raw = str(value or '')
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y'):
        try:
            return datetime.strptime(raw[:10], fmt).strftime('%d-%m-%Y')
        except ValueError:
            pass
    return raw


def rows(conn, ident):
    state = registry(conn)
    # Explicit SQL allowlist: neither bank fields nor credentials are loaded.
    fields = ', '.join('payload->>\'' + label + '\' AS ' + key for key, label in PAYLOAD_FIELDS.items())
    records = conn.execute(text('''SELECT username, full_name, role, birth_date, phone, email, address,
        employment_start_date, work_shift, shift_start_date, rotation_cycle,
        monthly_generated, monthly_leave, annual_leave,
        COALESCE(payload->>'Trạng thái làm việc', payload->>'employment_status', 'Đang làm việc') AS employment_status,
        COALESCE(payload->>'vera_residence', 'false') = 'true' AS included,
        COALESCE(payload->>'Ẩn nhân viên', '') AS profile_hidden, ''' + fields + '''
        FROM employees WHERE COALESCE(payload->>'__deleted','false') <> 'true'
        ORDER BY COALESCE(stt,2147483647), username''')).mappings()
    output = []
    for record in records:
        row = dict(record)
        if str(row.get('role') or '').lower() == 'admin':
            continue
        if ident.role != 'admin' and (row.get('role') == 'quanly' or _norm(row.get('profile_hidden')) in {'true','1','yes','co','x'}):
            continue
        code = department_code(row, state)
        row['department'] = state['departments'].get(code, {}).get('name', code)
        projected = {key: row.get(key) if row.get(key) is not None else '' for key, _ in COLUMNS}
        for field in DATE_FIELDS:
            projected[field] = display_date(projected[field])
        output.append({**projected, 'department_code': code, 'included': bool(row['included'])})
    return output


def filtered(records, department='', search=''):
    needle = _norm(search)
    return [row for row in records if row['included']
            and (not department or row['department_code'] == department)
            and (not needle or needle in _norm(row['username'] + ' ' + row['full_name']))]


def workbook(records):
    from openpyxl.styles import Border, Side
    from openpyxl.worksheet.page import PageMargins
    from openpyxl.utils import get_column_letter
    from vera_web_v2_residence_export import EXPORT_COLUMNS, export_rows, TITLE, BUSINESS, ADDRESS, WIDTHS
    wb = Workbook()
    ws = wb.active
    ws.title = 'Tạm trú'
    from vera_web_v2_residence_export import NATIONAL, MOTTO, INDUSTRY
    for cell_ref,value in [('A1',BUSINESS),('A2',ADDRESS),('A3',INDUSTRY),('E1',NATIONAL),('E2',MOTTO),('A4',TITLE)]:
        cell=ws[cell_ref]; cell.value=value
        cell.font=Font(name='Times New Roman',size=14 if cell_ref=='A4' else 11,bold=cell_ref in {'A1','E1','E2','A4'})
        cell.alignment=Alignment(horizontal='center' if cell_ref in {'E1','E2','A4'} else 'left',vertical='center')
    for merged in ('E1:I1','E2:I2','A3:E3','A4:I4'): ws.merge_cells(merged)
    for number in range(1,5): ws.row_dimensions[number].height=30 if number==4 else 20.1
    ws.append(['STT']+[label for _,label in EXPORT_COLUMNS])
    side=Side(style='thin',color='000000')
    border=Border(left=side,right=side,top=side,bottom=side)
    for index, record in enumerate(export_rows(records),1):
        ws.append([index]+[record[k] for k,_ in EXPORT_COLUMNS])
    last=ws.max_row
    for cells in ws.iter_rows(min_row=5,max_row=last):
        for cell in cells:
            cell.font=Font(name='Times New Roman',size=11,bold=cell.row==5)
            cell.border=border
            cell.alignment=Alignment(vertical='center',wrap_text=True,
                horizontal='center' if cell.row==5 or cell.column in {1,4,5,6,7} else 'left')
            if isinstance(cell.value,str): cell.data_type='s'; cell.number_format='@'
        # Excel does not reliably autofit wrapped rows at print time.
        import math
        lines=max(sum(max(1,math.ceil(len(part)/(WIDTHS[c.column-1]))) for part in str(c.value or '').split('\n')) for c in cells)
        ws.row_dimensions[cells[0].row].height=max(32,lines*15+8)
    for index,w in enumerate(WIDTHS,1): ws.column_dimensions[get_column_letter(index)].width=w
    for offset,value in ((2,'Ngày..........tháng..........năm..........'),(4,'Xác nhận của đơn vị')):
        n=last+offset
        ws.merge_cells(start_row=n,start_column=8,end_row=n,end_column=9)
        c=ws.cell(n,8,value); c.font=Font(name='Times New Roman',size=11,bold=offset==4)
        c.alignment=Alignment(horizontal='center')
    ws.freeze_panes='C6'
    ws.auto_filter.ref=f'A5:I{last}'
    ws.print_title_rows='1:5'
    ws.print_area=f'A1:I{last+6}'
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape'
    ws.page_setup.paperSize=ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth=1
    ws.page_setup.fitToHeight=0
    ws.page_margins=PageMargins(left=.25,right=.25,top=.75,bottom=.75,header=.3,footer=.3)
    stream=BytesIO(); wb.save(stream); stream.seek(0)
    return stream


def install_residence_routes(app, *, engine_instance, current_identity, require_feature, identity_type):
    @app.get('/v2/residence')
    def overview(ident: identity_type = Depends(current_identity)):
        with engine_instance().connect() as conn:
            require_feature(conn, ident, 'residence_view')
            records = rows(conn, ident)
        departments = {r['department_code']: r['department'] for r in records}
        return {'columns': [{'key': k, 'label': label} for k, label in COLUMNS],
                'records': [r for r in records if r['included']],
                'employees': [{k: r[k] for k in ('username','full_name','department_code','department','included')} for r in records],
                'departments': [{'code': k, 'name': v} for k, v in sorted(departments.items())]}

    @app.put('/v2/residence')
    def membership(body: Membership, ident: identity_type = Depends(current_identity)):
        with engine_instance().begin() as conn:
            require_feature(conn, ident, 'residence_view')
            require_feature(conn, ident, 'residence_manage')
            # Serialize with staff profile row writes; change only the roster flag.
            conn.execute(text("SELECT username FROM employees WHERE username=:username FOR UPDATE"), {'username': body.username})
            if body.username not in {r['username'] for r in rows(conn, ident)}:
                raise HTTPException(404, 'Không tìm thấy nhân viên phù hợp.')
            changed = conn.execute(text('''UPDATE employees SET
                payload=jsonb_set(COALESCE(payload,'{}'::jsonb), '{vera_residence}', CAST(:value AS jsonb)),
                updated_at=NOW() WHERE username=:username AND COALESCE(payload->>'__deleted','false') <> 'true'
                RETURNING username'''), {'username': body.username, 'value': 'true' if body.included else 'false'}).first()
            if not changed:
                raise HTTPException(409, 'Hồ sơ đã thay đổi. Vui lòng tải lại.')
        return {'ok': True, 'message': 'Đã thêm nhân viên vào bảng tạm trú.' if body.included else 'Đã bỏ nhân viên khỏi bảng tạm trú.'}

    @app.get('/v2/residence/export.xlsx')
    def export(department: str = '', search: str = '', ident: identity_type = Depends(current_identity)):
        with engine_instance().connect() as conn:
            require_feature(conn, ident, 'residence_view')
            require_feature(conn, ident, 'residence_export')
            records = filtered(rows(conn, ident), department, search)
        return StreamingResponse(workbook(records),
            media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={'Content-Disposition': 'attachment; filename="VERA_TamTru.xlsx"'})

    @app.get('/v2/residence/export')
    def export_document(format: str = 'pdf', department: str = '', search: str = '',
                        ident: identity_type = Depends(current_identity)):
        from vera_web_v2_residence_export import render_residence
        from vera_web_v2_training_report_export import ReportExportTooLarge, ReportExportUnavailable
        with engine_instance().connect() as conn:
            require_feature(conn, ident, 'residence_view')
            require_feature(conn, ident, 'residence_export')
            if format not in {'pdf', 'png'}:
                raise HTTPException(422, 'Chọn định dạng PDF hoặc PNG.')
            records = filtered(rows(conn, ident), department, search)
        try:
            content = render_residence(records, COLUMNS, format)
        except ReportExportTooLarge:
            raise HTTPException(413, 'Ảnh quá lớn. Hãy lọc theo bộ phận hoặc tải PDF.')
        except ReportExportUnavailable as exc:
            raise HTTPException(503, str(exc))
        return Response(content, media_type='application/pdf' if format == 'pdf' else 'image/png',
                        headers={'Content-Disposition': f'inline; filename="VERA_TamTru.{format}"',
                                 'Cache-Control': 'no-store'})

    @app.post('/v2/staff/labor-report/export')
    def labor_report(body: LaborReportRequest, format: str = 'pdf', ident: identity_type = Depends(current_identity)):
        from vera_web_v2_labor_report import labor_rows, render_labor
        from vera_web_v2_training_report_export import ReportExportTooLarge, ReportExportUnavailable
        if ident.role != 'admin':
            raise HTTPException(403, 'Chỉ Admin được xuất danh sách khai trình lao động.')
        if format not in {'pdf','png'}:
            raise HTTPException(422, 'Chọn định dạng PDF hoặc PNG.')
        names = list(dict.fromkeys(name.strip() for name in body.usernames if name.strip()))
        if not names:
            raise HTTPException(422, 'Chưa chọn nhân viên cần xuất.')
        with engine_instance().connect() as conn:
            require_feature(conn, ident, 'staff_list')
            require_feature(conn, ident, 'staff_export')
            records = labor_rows(conn, ident, names)
        if len(records) != len(names):
            raise HTTPException(404, 'Danh sách nhân viên đã thay đổi. Vui lòng tải lại trước khi xuất.')
        try:
            content = render_labor(records, format)
        except ReportExportTooLarge:
            raise HTTPException(413, 'Ảnh quá lớn. Hãy chọn ít nhân viên hơn hoặc tải PDF.')
        except ReportExportUnavailable as exc:
            raise HTTPException(503, str(exc))
        return Response(content, media_type='application/pdf' if format=='pdf' else 'image/png',
                        headers={'Content-Disposition': f'inline; filename="VERA_KhaiTrinhLaoDong.{format}"',
                                 'Cache-Control':'private, no-store'})
