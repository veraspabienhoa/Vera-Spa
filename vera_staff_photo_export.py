"""Optional employee portraits in Excel, formatted without changing stored photos."""
from io import BytesIO

from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils import get_column_letter
from PIL import Image, ImageOps
from sqlalchemy import text


def read_portraits(conn, usernames):
    names = list(dict.fromkeys(usernames))
    if len(names) > 200:
        raise HTTPException(400, 'Xuất kèm ảnh tối đa 200 nhân viên/lần. Hãy thu hẹp bộ lọc.')
    if not names or not conn.execute(text("SELECT to_regclass('vera_employee_identity_document')")).scalar():
        return {}
    params = {'names': names}
    size = conn.execute(text("""SELECT COALESCE(SUM(octet_length(content)),0)
        FROM vera_employee_identity_document WHERE side='portrait' AND employee_username=ANY(:names)"""), params).scalar()
    if size > 40 * 1024 * 1024:
        raise HTTPException(400, 'Tổng ảnh vượt 40 MB. Hãy thu hẹp bộ lọc nhân viên để xuất.')
    return {row['employee_username']: bytes(row['content']) for row in conn.execute(text("""
        SELECT employee_username,content FROM vera_employee_identity_document
        WHERE side='portrait' AND employee_username=ANY(:names)"""), params).mappings()}


def portrait_frame(content):
    """Preserve the full image inside a 3:4 white frame, with no face cropping."""
    with Image.open(BytesIO(content)) as image:
        if image.width > 6000 or image.height > 6000 or image.width * image.height > 24_000_000:
            raise ValueError('image too large')
        image = ImageOps.exif_transpose(image).convert('RGBA')
        background = Image.new('RGBA', image.size, 'white')
        background.alpha_composite(image)
        framed = ImageOps.pad(background.convert('RGB'), (300, 400), method=Image.Resampling.LANCZOS, color='white')
        output = BytesIO()
        framed.save(output, format='PNG')
        return output.getvalue()


def with_portraits(workbook_bytes, rows, portraits):
    workbook = load_workbook(BytesIO(workbook_bytes))
    sheet = workbook['DanhSachNhanSu']
    column = sheet.max_column + 1
    letter = get_column_letter(column)
    sheet.cell(1, column, 'Ảnh nhân viên (3 × 4 cm)')
    sheet.column_dimensions[letter].width = 19
    for row_number, employee in enumerate(rows, start=2):
        content = portraits.get(employee['username'])
        if not content:
            continue
        try:
            image = ExcelImage(BytesIO(portrait_frame(content)))
        except (ValueError, OSError, Image.DecompressionBombError):
            sheet.cell(row_number, column, 'Ảnh không đọc được')
            continue
        # 360,000 EMUs per centimetre; anchor dimensions survive style middleware.
        image.anchor = OneCellAnchor(_from=AnchorMarker(col=column-1, row=row_number-1, colOff=19050, rowOff=19050),
                                    ext=XDRPositiveSize2D(cx=1080000, cy=1440000))
        sheet.add_image(image)
        sheet.row_dimensions[row_number].height = 4 / 2.54 * 72 + 6
    sheet.auto_filter.ref = f'A1:{letter}{max(1,sheet.max_row)}'
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
