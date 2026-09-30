from copy import deepcopy
from datetime import date
from io import BytesIO
import asyncio
import json
import os
from types import SimpleNamespace
from uuid import uuid4

import pytest
from fastapi import FastAPI, HTTPException
from openpyxl import load_workbook
from sqlalchemy import create_engine, text

import vera_employee_names as names
import vera_web_v2_system_name as rename
import vera_web_v2_work_schedule as schedule
import vera_web_v2_department_payroll as payroll
from vera_web_v2_live_tour_roster import reconcile


def employee(username='Gia Anh', **extra):
    return {'username': username, 'full_name': 'Nguyễn Gia Anh', 'payload': {}, **extra}


def test_confirmed_history_only_and_ambiguous_aliases():
    rows = [employee()]
    before = deepcopy(rows)
    index = names.identity_index(rows)
    assert names.canonical_username(index, 'Anh Nguyễn') == 'Gia Anh'
    assert names.canonical_username(index, 'Anh Nguyen') == 'Gia Anh'
    assert names.canonical_username(index, 'Nguyễn Anh') == 'Nguyễn Anh'
    assert rows == before
    for competing in [employee('Anh Nguyễn'), employee('Anh Nguyễn', payload={'__deleted': True})]:
        assert names.canonical_username(names.identity_index(rows + [competing]), 'Anh Nguyễn') == 'Anh Nguyễn'
    other = employee('Other', payload={names.HISTORY_KEY: ['Anh Nguyễn']})
    assert names.canonical_username(names.identity_index(rows + [other]), 'Anh Nguyễn') == 'Anh Nguyễn'
    assert names.canonical_username(names.identity_index([employee(full_name='Different')]), 'Anh Nguyễn') == 'Anh Nguyễn'


def test_live_tour_keeps_worker_id_and_open_work_across_rename():
    directory = [employee('New', role='nhanvien', payload={names.HISTORY_KEY: ['Old']})]
    state = {'employees': [{'id': 'stable-id', 'username': 'Old', 'name': 'Old', 'service': 'Existing', 'booking_id': 'b1'}],
             'pending': [{'entries': [{'employee_id': 'stable-id', 'employee_name': 'Old', 'total': 123}]}]}
    assert reconcile(state, directory, lambda *args: pytest.fail('must not create a new worker'))
    assert len(state['employees']) == 1
    assert state['employees'][0]['username'] == 'New'
    assert state['employees'][0]['id'] == 'stable-id'
    assert state['employees'][0]['booking_id'] == 'b1'
    assert state['pending'][0]['entries'][0] == {'employee_id': 'stable-id', 'employee_name': 'New', 'total': 123}


@pytest.fixture
def database():
    url = os.getenv('VERA_TEST_POSTGRES_URL')
    if not url:
        pytest.skip('VERA_TEST_POSTGRES_URL required')
    schema = 'employee_names_' + uuid4().hex
    admin = create_engine(url)
    with admin.begin() as conn:
        conn.execute(text(f'CREATE SCHEMA {schema}'))
    engine = create_engine(url, pool_size=1, max_overflow=0, pool_timeout=1,
                           connect_args={'options': f'-csearch_path={schema}'})
    try:
        with engine.begin() as conn:
            conn.execute(text("CREATE TABLE employees(username text PRIMARY KEY, full_name text, role text, payload jsonb DEFAULT '{}'::jsonb)"))
            conn.execute(text("INSERT INTO employees VALUES('Gia Anh','Nguyễn Gia Anh','letan','{}')"))
            schedule._ensure_schema(conn)
            for i, username in enumerate(['Anh Nguyễn', 'Anh Nguyen', 'Gia Anh']):
                conn.execute(text("""INSERT INTO vera_work_schedule_combo_sale
                    (id,sale_date,employee_username,employee_name,department,customer_name,combo_ticket)
                    VALUES(:id,'2026-09-30',:username,:username,'letan','Synthetic customer','Synthetic combo')"""),
                    {'id': str(i), 'username': username})
        yield engine
    finally:
        engine.dispose()
        with admin.begin() as conn:
            conn.execute(text(f'DROP SCHEMA {schema} CASCADE'))
        admin.dispose()


def endpoints(engine, monkeypatch):
    # Installer wraps the global catalog; restore it at teardown.
    monkeypatch.setattr(schedule, '_employee_catalog', rename._ORIGINAL_EMPLOYEE_CATALOG)
    app = FastAPI()
    who = SimpleNamespace(role='admin', employee_username='operator')
    schedule.install_work_schedule_routes(app, engine_instance=lambda: engine,
        current_identity=lambda: who, feature_allowed=lambda *args: True)
    rename.install_system_name_routes(app, engine_instance=lambda: engine,
        current_identity=lambda: who, identity_type=SimpleNamespace)
    return {(route.path, method): route.endpoint for route in app.routes for method in (route.methods or [])}, who


def test_combo_read_export_and_count_share_canonical_owner_without_writes(database, monkeypatch):
    routes, who = endpoints(database, monkeypatch)
    args = (date(2026,9,1), date(2026,9,30), 'letan', who)
    result = routes['/v2/work-schedule/combo-sales','GET'](*args)
    assert len(result['rows']) == 3
    assert {row['employee_username'] for row in result['rows']} == {'Gia Anh'}
    with database.begin() as conn:
        assert payroll._combo_sale_counts(conn, args[0], args[1]) == {'gia anh': 3}
        assert conn.execute(text('SELECT COUNT(DISTINCT employee_username) FROM vera_work_schedule_combo_sale')).scalar_one() == 3
    response = routes['/v2/work-schedule/combo-sales/export.xlsx','GET'](*args)
    async def collect():
        return b''.join([chunk async for chunk in response.body_iterator])
    workbook = load_workbook(BytesIO(asyncio.run(collect())))
    assert len(workbook.worksheets) == 1
    sheet = workbook.active
    assert [sheet.cell(row, 3).value for row in range(2,5)] == ['Gia Anh'] * 3
    assert sheet.cell(2,10).value == 'Gia Anh'
    workbook.close()


def test_combo_stale_form_saves_only_current_username(database, monkeypatch):
    routes, who = endpoints(database, monkeypatch)
    monkeypatch.setattr(schedule, '_resolve_combo_customer', lambda conn, body, *args: body)
    body = schedule.ComboSaleSave(sale_date=date(2026,9,30), employee_username='Anh Nguyễn',
        department='letan', customer_name='Synthetic', combo_ticket='Synthetic')
    result = routes['/v2/work-schedule/combo-sales','POST'](body, who)
    with database.begin() as conn:
        assert conn.execute(text('SELECT employee_username FROM vera_work_schedule_combo_sale WHERE id=:id'), result).scalar_one() == 'Gia Anh'


def test_future_rename_is_transactional_preserves_photo_and_alias_chain(database, monkeypatch):
    routes, who = endpoints(database, monkeypatch)
    with database.begin() as conn:
        conn.execute(text('CREATE TABLE vera_employee_face_id(employee_username text PRIMARY KEY, content bytea)'))
        conn.execute(text("INSERT INTO vera_employee_face_id VALUES('Gia Anh',decode('010203','hex'))"))
        conn.execute(text('CREATE TABLE vera_v2_user_profile(employee_username text REFERENCES employees(username) ON UPDATE CASCADE, auth_user_id text)'))
        conn.execute(text("INSERT INTO vera_v2_user_profile VALUES('Gia Anh','same-auth-id')"))
    change = routes['/v2/staff/{username}/system-name','PATCH']
    change('Gia Anh', rename.SystemNameUpdate(system_name='New'), who)
    change('New', rename.SystemNameUpdate(system_name='Newest'), who)
    with database.begin() as conn:
        row = conn.execute(text('SELECT * FROM employees')).mappings().one()
        assert row['username'] == 'Newest'
        assert set(row['payload'][names.HISTORY_KEY]) == {'Gia Anh','Anh Nguyễn','Anh Nguyen','New'}
        assert bytes(conn.execute(text('SELECT content FROM vera_employee_face_id')).scalar_one()) == b'\x01\x02\x03'
        assert conn.execute(text('SELECT employee_username FROM vera_employee_face_id')).scalar_one() == 'Newest'
        assert tuple(conn.execute(text('SELECT employee_username,auth_user_id FROM vera_v2_user_profile')).one()) == ('Newest','same-auth-id')
        assert payroll._combo_sale_counts(conn, date(2026,9,1), date(2026,9,30)) == {'newest': 3}


def test_rename_conflict_rolls_back_every_reference(database, monkeypatch):
    routes, who = endpoints(database, monkeypatch)
    with database.begin() as conn:
        conn.execute(text('CREATE TABLE vera_employee_face_id(employee_username text PRIMARY KEY, content bytea)'))
        conn.execute(text("INSERT INTO vera_employee_face_id VALUES('Gia Anh',NULL),('New',NULL)"))
    with pytest.raises(HTTPException) as exc:
        routes['/v2/staff/{username}/system-name','PATCH']('Gia Anh', rename.SystemNameUpdate(system_name='New'), who)
    assert exc.value.status_code == 409
    with database.begin() as conn:
        row = conn.execute(text('SELECT username,payload FROM employees')).mappings().one()
        assert row['username'] == 'Gia Anh' and names.HISTORY_KEY not in row['payload']
        assert conn.execute(text("SELECT employee_username FROM vera_work_schedule_combo_sale WHERE id='2'")).scalar_one() == 'Gia Anh'


def test_rename_permission_and_reserved_history(database, monkeypatch):
    routes, who = endpoints(database, monkeypatch)
    change = routes['/v2/staff/{username}/system-name','PATCH']
    for role in ['letan','quanly','leader','nhanvien']:
        with pytest.raises(HTTPException) as exc:
            change('Gia Anh', rename.SystemNameUpdate(system_name='New'), SimpleNamespace(role=role))
        assert exc.value.status_code == 403
    with database.begin() as conn:
        conn.execute(text("INSERT INTO employees VALUES('Other','Other','letan','{}')"))
    with pytest.raises(HTTPException) as exc:
        change('Other', rename.SystemNameUpdate(system_name='Anh Nguyễn'), who)
    assert exc.value.status_code == 409
