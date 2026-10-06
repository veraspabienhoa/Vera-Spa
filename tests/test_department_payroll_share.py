import ast
from io import BytesIO
from pathlib import Path
import unittest
from typing import Any
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image
from vera_department_payroll_share import export_values, payroll_pdf, payroll_png

source = ast.parse((Path(__file__).resolve().parents[1] / 'vera_web_v2_department_payroll.py').read_text())
columns_node = next(n for n in source.body if isinstance(n, ast.Assign) and any(isinstance(t,ast.Name) and t.id=='EXCEL_COLUMNS' for t in n.targets))
COLUMNS = ast.literal_eval(columns_node.value)
function = next(n for n in source.body if isinstance(n,ast.FunctionDef) and n.name=='_combined_workbook')
namespace = dict(Any=Any, BytesIO=BytesIO, Workbook=Workbook, Alignment=Alignment, Font=Font, PatternFill=PatternFill, EXCEL_COLUMNS=COLUMNS, ROW_MONEY_FIELDS=tuple(key for key,_ in COLUMNS if key not in {'tt','employee_name','department_label','work_days','hours_ca1','hours_ca2_before_22','hours_ca2_after_22'}))
exec(compile(ast.Module(body=[function], type_ignores=[]), '<workbook>', 'exec'), namespace)

class PayrollShareTests(unittest.TestCase):
    def setUp(self):
        self.rows = [dict(tt=i+1,employee_name='Ngô Sĩ Đạt',department_label='Lễ tân',work_days=2.5,salary=100000,net_salary=80000,other_income_1=999999) for i in range(2)]

    def test_columns_and_excel_totals(self):
        self.assertEqual(len(COLUMNS),19)
        self.assertFalse(any(key.startswith('other_income') for key,_ in COLUMNS))
        ws=load_workbook(BytesIO(namespace['_combined_workbook'](self.rows,'10-2026'))).active
        self.assertEqual(ws.max_column,19)
        for column in range(5,20):
            letter=ws.cell(5,column).column_letter
            self.assertEqual(ws.cell(5,column).value,f'=SUM({letter}3:{letter}4)')
        self.assertEqual(ws.page_setup.orientation,'landscape')
        self.assertEqual(ws.page_setup.fitToWidth,1)
        data=export_values(self.rows,COLUMNS)
        self.assertEqual(data[-1][4],5)
        self.assertEqual(data[-1][-1],160000)

    def test_png_and_multipage_pdf_render(self):
        png=payroll_png(self.rows,'10-2026',COLUMNS)
        with Image.open(BytesIO(png)) as image:
            self.assertEqual(image.format,'PNG')
            self.assertGreater(image.width,image.height)
        pdf=payroll_pdf(self.rows*30,'10-2026',COLUMNS)
        self.assertTrue(pdf.startswith(b'%PDF'))

if __name__ == '__main__':
    unittest.main()
