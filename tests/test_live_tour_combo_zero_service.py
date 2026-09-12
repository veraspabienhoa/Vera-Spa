from io import BytesIO

import pytest
from openpyxl import load_workbook

import vera_web_v2_live_tour as live
from test_live_tour_backend import NOW
from test_live_tour_combo_booking import setup, booking
from test_service_catalog import action


@pytest.mark.parametrize('legacy', [False, True])
@pytest.mark.parametrize('tip', [0, 50000])
def test_prepaid_redemption_has_zero_service_in_invoice_export_and_edits(legacy, tip):
    state, skin, _, customer, owned = setup()
    sale_total = state['invoices'][0]['total']
    if legacy:
        owned.pop('component_balances')
        next(row for row in state['services'] if row['id'] == skin['id'])['ticket_units'] = 1
    action(state, 'booking', booking(customer, owned, skin))
    action(state, 'start', {'employee_id': 'e1'})
    pending = action(state, 'finish_to_pending', {'employee_id': 'e1'})['pending']
    invoice = action(state, 'checkout', {'pending_id': pending['id'], 'combo_purchase_id': owned['id'],
        'payment_method': 'COMBO', 'tip': tip})['invoice']
    assert invoice['subtotal'] == invoice['discount'] == 0
    assert invoice['total'] == invoice['tip'] == tip
    assert invoice['combo_units'] == 1
    assert state['invoices'][0]['total'] == sale_total > 0
    assert state['invoices'][0]['subtotal'] == sale_total
    assert state['customers'][0]['combo_purchases'][0]['remaining'] == 2
    # Old saved invoice subtotals must also export as zero without rewriting history.
    invoice['subtotal'] = 200000
    content, _ = live._excel_bytes(state, 'revenue', NOW)
    sheet = load_workbook(BytesIO(content))['Doanh_thu']
    headers = [cell.value for cell in sheet[1]]
    assert sheet.cell(3, headers.index('Tiền dịch vụ') + 1).value == 0
    assert sheet.cell(2, headers.index('Tiền dịch vụ') + 1).value == sale_total
    action(state, 'paid_invoice_update', {'invoice_id': invoice['id'], 'reason': 'Cập nhật Tip', 'tip': 70000})
    saved = state['invoices'][-1]
    assert saved['subtotal'] == 0 and saved['total'] == 70000
    report = next(row for row in state['reports'] if row['invoice_id'] == saved['id'])
    assert report['total'] == report['tip'] == 70000
