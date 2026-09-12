"""Regression gates for Live Tour integration with main (September 2026)."""
from copy import deepcopy
from datetime import timedelta
from io import BytesIO

import pytest
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import vera_web_v2_live_tour as live
from test_live_tour_backend import NOW, RouteEngine, RouteIdentity, employee, state_with


def source_record(name="An"):
    return {"Tên nhân viên": name, "Dịch vụ": "Body 90", "Phòng": "1.1",
            "Trạng thái": "CHO THANH TOÁN", "Đi làm": "Đi làm", "Vào ca": "Ca 1"}


def payable_state():
    state = state_with(employee("e1", "An"))
    worker = state["employees"][0]
    worker.update(service="Body 90", service_price=100, service_price_source="catalog",
                  room="1.1", status="CHO THANH TOÁN", payment_status="CHO THANH TOÁN")
    next(item for item in state["services"] if item["name"] == "Body 90")["price"] = 100
    return state


def api_client(monkeypatch, state=None):
    shared = {"state": deepcopy(state or state_with(employee("e1", "An"))), "revision": 1}
    monkeypatch.setattr(live, "_read_state", lambda *_args, **_kwargs: (deepcopy(shared["state"]), shared["revision"]))

    def write(_conn, state, revision, _actor):
        assert revision == shared["revision"]
        shared.update(state=deepcopy(state), revision=revision + 1)
        return shared["revision"]

    monkeypatch.setattr(live, "_write_state", write)
    app = FastAPI()
    live.install_live_tour_routes(
        app, engine_instance=RouteEngine, current_identity=lambda: RouteIdentity(),
        require_feature=lambda *_args: None, feature_allowed=lambda *_args: True,
        identity_type=RouteIdentity,
    )
    return TestClient(app), shared


def test_restore_rejects_paid_service_replay_and_preserves_ledger():
    state = payable_state()
    backup = live._apply_action(state, "backup", {}, "admin", NOW)["backup"]
    live._apply_action(state, "checkout", {"employee_id": "e1", "payment_method": "TIỀN MẶT"}, "admin", NOW)
    before = deepcopy(state)
    with pytest.raises(HTTPException, match="Chỉ khôi phục"):
        live._apply_action(state, "restore", {"backup_id": backup["id"]}, "admin", NOW)
    assert state == before


def test_restore_rejects_active_break_even_with_idle_snapshot():
    state = state_with(employee("e1", "An"))
    backup = live._apply_action(state, "backup", {}, "admin", NOW)["backup"]
    live._apply_action(state, "start_break", {"employee_id": "e1"}, "admin", NOW)
    with pytest.raises(HTTPException):
        live._apply_action(state, "restore", {"backup_id": backup["id"]}, "admin", NOW)
    assert len(state["break_events"]) == 1


def test_employee_without_work_status_and_shift_cannot_be_booked():
    worker = employee("e1", "An")
    worker.update(work_status="Nghỉ", shift="")
    state = state_with(worker)
    with pytest.raises(HTTPException):
        live._booking(state, {"employee_id": worker["id"], "service": "Body 90", "room": "1.1"}, NOW)



def test_counter_rollover_at_ten_is_separate_from_financial_day():
    prior = NOW.replace(hour=9, minute=59)
    state = live._empty_state(prior)
    worker = employee("e1", "An")
    worker.update(tour_count=4, request_count=2, break_count=1, service="Body 90", room="1.1", status="Đang thực hiện")
    state["employees"] = [worker]
    before = live._state_response(state, 1, prior)
    after = live._state_response(state, 1, prior + timedelta(minutes=1))
    assert before["records"][0]["Tổng SL"] == 6
    assert after["records"][0]["Tổng SL"] == 0
    assert after["records"][0]["Dịch vụ"] == "Body 90"
    assert after["metrics_business_date"] == "2026-09-05"
    assert live._business_date(prior + timedelta(minutes=1)).isoformat() == "2026-09-04"
    assert worker["tour_count"] == 4  # response must not mutate persisted state


def test_expected_revenue_is_separate_and_follows_pending_without_double_count():
    state = payable_state()
    before = live._state_response(state, 1, NOW, can_payment=True)
    assert before["reports"]["total_revenue"] == 0
    assert before["reports"]["expected_unbilled_revenue"] == 100
    pending = live._apply_action(state, "move_pending", {"employee_id": "e1"}, "admin", NOW)["pending"]
    assert live._state_response(state, 2, NOW, can_payment=True)["reports"]["expected_unbilled_revenue"] == 100
    content, _ = live._excel_bytes(state, "revenue", NOW)
    workbook = load_workbook(BytesIO(content))
    assert workbook["Doanh_thu"].max_row == 2
    assert workbook["Doanh_thu"].cell(2, 1).value == "Tổng cộng"
    assert workbook["Doanh_thu"].cell(2, 9).value == 0
    assert workbook["Du_kien_chua_xuat_bill"].cell(2, 5).value == 100
    live._apply_action(state, "checkout", {"pending_id": pending["id"], "payment_method": "TIỀN MẶT"}, "admin", NOW)
    reports = live._state_response(state, 3, NOW, can_payment=True)["reports"]
    assert reports["total_revenue"] == 100 and reports["expected_unbilled_revenue"] == 0


def test_catalog_private_flag_locks_room_and_cannot_be_removed_while_unpaid():
    state = state_with(employee("e1", "An"), employee("e2", "Bình"))
    service = next(item for item in state["services"] if item["name"] == "Body 90")
    service["private"] = True
    live._booking(state, {"employee_id": "e1", "service": "Body 90", "room": "19.1"}, NOW)
    assert not live._room_available(state, "19.2")
    assert live._state_response(state, 1, NOW)["records"][0]["_private_service"] is True
    with pytest.raises(HTTPException):
        live._apply_action(state, "service_upsert", {**service, "private": False}, "admin", NOW)
    with pytest.raises(HTTPException):
        live._apply_action(state, "service_delete", {"id": service["id"]}, "admin", NOW)


def test_catalog_yc_eligibility_and_duration_are_server_enforced():
    state = state_with(employee("e1", "An"))
    service = next(item for item in state["services"] if item["name"] == "Body 90")
    service.update(non_request_eligible=False, request_duration=110)
    with pytest.raises(HTTPException):
        live._booking(state, {"employee_id": "e1", "service": "Body 90", "room": "1.1"}, NOW)
    worker = live._booking(state, {"employee_id": "e1", "service": "Body 90", "room": "1.1", "request": "YC"}, NOW)
    assert worker["duration"] == 110


def test_actual_http_requires_revision_and_validates_nested_idempotency(monkeypatch):
    client, shared = api_client(monkeypatch)
    assert client.get("/v2/live-tour").status_code == 200
    body = {"action": "set_vip", "idempotency_key": "vip-request", "payload": {"employee_id": "e1", "vip": True}}
    assert client.post("/v2/live-tour/action", json=body).status_code == 428
    body["expected_revision"] = 1
    del body["idempotency_key"]
    body["payload"]["idempotency_key"] = "short"
    assert client.post("/v2/live-tour/action", json=body).status_code == 400
    body["payload"]["idempotency_key"] = "valid-request"
    assert client.post("/v2/live-tour/action", json=body).status_code == 200
    replay = client.post("/v2/live-tour/action", json=body)
    assert replay.status_code == 200 and replay.json()["duplicate"] is True
    assert shared["revision"] == 2


def test_failed_sync_idempotency_markers_are_prunable_without_evicting_new_key(monkeypatch):
    state = state_with()
    monkeypatch.setattr(live, "MAX_IDEMPOTENCY", 2)
    for key in ("old-one", "old-two"):
        state["idempotency"][key] = {"action": "sync_leaves", "status": "failed", "at": live._iso(NOW)}
    live._remember_idempotency(state, "new-request", action="start", actor="admin", payload_hash="hash", result={}, now=NOW)
    assert len(state["idempotency"]) == 2 and "new-request" in state["idempotency"]


def test_financial_receipts_never_expire_due_to_request_count(monkeypatch):
    state = state_with()
    monkeypatch.setattr(live, "MAX_IDEMPOTENCY", 1)
    for key in ("old-payment", "new-payment"):
        live._remember_idempotency(state, key, action="checkout", actor="admin", payload_hash=key, result={}, now=NOW)
    assert set(state["idempotency"]) == {"old-payment", "new-payment"}


def test_physical_room_counts_are_separate_from_bookable_beds():
    state = state_with(employee("e1", "An"))
    state["physical_rooms"] = ["3", "19"]
    response = live._state_response(state, 1, NOW)
    assert "3" in response["rooms"]["all"]
    assert "19.1" not in response["rooms"]["all"]
    assert "19.1" in response["available_beds"]
    assert response["rooms"]["total_count"] < len(response["catalogs"]["rooms"])


def test_catalog_rename_collision_and_unpaid_ticket_change_are_blocked():
    state = payable_state()
    service = next(item for item in state["services"] if item["name"] == "Body 90")
    with pytest.raises(HTTPException):
        live._apply_action(state, "service_upsert", {**service, "name": "Body 70"}, "admin", NOW)
    with pytest.raises(HTTPException):
        live._apply_action(state, "service_upsert", {**service, "ticket_units": 2}, "admin", NOW)
    live._apply_action(state, "move_pending", {"employee_id": "e1"}, "admin", NOW)
    with pytest.raises(HTTPException):
        live._apply_action(state, "service_delete", {"id": service["id"]}, "admin", NOW)


def test_catalog_can_initialize_price_without_changing_active_service_semantics():
    state = payable_state()
    service = next(item for item in state["services"] if item["name"] == "Body 90")
    result = live._apply_action(state, "service_upsert", {**service, "duration": "90", "price": 200,
                               "request_eligible": True, "non_request_eligible": True, "request_duration": None}, "admin", NOW)
    assert result["service"]["price"] == 200
