import json
from pathlib import Path

from vera_employee_search import employee_name_matches


def test_employee_matching_agrees_with_frontend_cases():
    cases = json.loads((Path(__file__).parent / "fixtures/employee_search_cases.json").read_text())
    for case in cases:
        assert employee_name_matches(case["name"], case["query"]) is case["matches"], case


def test_attendance_display_and_export_filter_excludes_nearby_names():
    from vera_web_v2_operations_v41 import _snapshot_filter, _snapshot_workbook
    from io import BytesIO
    from openpyxl import load_workbook

    rows = [
        {"employee_name": "Cẩm Vân", "break_department": "Nhân viên", "shift": "Ca 1"},
        {"employee_name": "Cẩm Vân Anh", "break_department": "Nhân viên", "shift": "Ca 1"},
        {"employee_name": "Đỗ Ánh", "break_department": "Locker", "shift": "Ca 2"},
    ]
    filtered = _snapshot_filter(rows, "cam van", "nhan", "1")
    assert [item["employee_name"] for item in filtered] == ["Cẩm Vân"]
    workbook = load_workbook(BytesIO(_snapshot_workbook(filtered)))
    assert workbook.active.max_row == 2
    assert workbook.active.cell(2, 3).value == "Cẩm Vân"
    assert _snapshot_filter(rows, "Vân", "", "") == []


def test_payroll_history_matches_short_names_without_crossing_batches():
    from vera_employee_search import normalize_employee_search
    from vera_web_v2_payroll import _filter_rows

    rows = [
        {"Tên Hệ thống": "Linh Đan - KTV", "Mã bản lưu": "A"},
        {"Tên Hệ thống": "Linh Đan Anh", "Mã bản lưu": "A"},
        {"Tên Hệ thống": "Linh Đan - KTV", "Mã bản lưu": "B"},
    ]
    assert _filter_rows(rows, "A", "linh dan", normalize_employee_search) == rows[:1]
