"""Pure server trust, rate allocation and no-double-deduction regressions."""
from copy import deepcopy
from datetime import date
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

import vera_web_v2_department_payroll as dep

DAY = date(2026, 10, 1)
MONTH = '2026-10'
norm = lambda value: str(value or '').strip().casefold()
DEFINITIONS = {'letan': {'Ca 1': {'start': '09:00', 'end': '17:00'}, 'Ca 2': {'start': '17:00', 'end': '01:00'}}}


def config(source='schedule', mode='hourly'):
    return {**dep.DEFAULT_CONFIG['letan'], 'calculation_mode': mode, 'salary_source': source,
            'default_base_salary': 7_800_000, 'configured_rate_keys': list(dep.RATE_KEYS)}


def schedule(name='a', shift='Ca 1', **extra):
    return {'employee_username': name, 'work_date': DAY, 'shift_code': shift, 'schedule_department': 'letan', **extra}


def record(name='a', arrival='09:15', departure='17:00', shift='Ca 1', **extra):
    return {'employee_name': name, 'date': '01/10/2026', 'check_in': arrival,
            'check_out': departure, 'shift': shift, 'attendance_expected': True, **extra}


def event(cfg, approved=False, start='09:00:00', end='09:15:00', shift='Ca 1'):
    from vera_hc_late_approval import price_basis
    prefix = '2026-10-01T'
    basis = {'kind': 'late', 'config': deepcopy(cfg), 'start': prefix + start + '+07:00',
             'end': prefix + end + '+07:00', 'rule': {'mode': 'multiplier', 'value': 2},
             'wage_segments': [{'start': prefix + start + '+07:00', 'end': prefix + end + '+07:00', 'shift': shift}]}
    return {'event_id': '101', 'day': DAY.isoformat(), 'basis': basis,
            'amount': price_basis(basis, DAY)[approved], 'approved': approved}


def payroll_row(totals, cfg, source, late):
    return {**{f'hours_{key}': float(totals[f'minutes_{key}']/60) for key in ('ca1', 'ca2_before_22', 'ca2_after_22')},
            'department': 'letan', 'employee_username': 'a', 'base_salary': cfg['default_base_salary'],
            'work_days': totals.get('work_days', sum(totals[f'minutes_{key}'] for key in ('ca1','ca2_before_22','ca2_after_22')) / 60 / cfg['standard_day_hours']),
            'payroll_time_basis': dep._payroll_time_basis(totals, cfg, source), 'late_penalty': late,
            'calculation_source': source, 'probation_rate': 1}


@pytest.mark.parametrize('value', [None, '', 'both', ['schedule'], {'attendance': True}])
def test_source_is_exclusive_and_invalid_values_do_not_default(value):
    with pytest.raises(HTTPException):
        dep._salary_source({'salary_source': value})
    assert dep._salary_source({}) == 'schedule'


def test_main_ca2_and_earlier_named_overtime_use_their_own_rates():
    cfg = config('attendance')
    planned = schedule(shift='Ca 2', overtime_shift='TC Ca 1')
    totals = dep._attendance_totals([record(arrival='09:00', departure='01:00', shift='Ca 2')], 'a', norm, cfg,
        schedule_records=[planned], definitions=DEFINITIONS, department='letan')
    assert [totals[key] for key in ('minutes_ca1','minutes_ca2_before_22','minutes_ca2_after_22')] == [480, 300, 180]
    row = dep._recalculate(payroll_row(totals, cfg, 'attendance', 0), cfg)
    assert row['salary'] == 465_000


def test_custom_after_midnight_overtime_continues_ca2_after22():
    cfg = config()
    planned = schedule(shift='Ca 2', overtime_shift='Từ giờ tới giờ', overtime_start_time='01:00', overtime_end_time='03:00')
    totals = dep._schedule_totals([planned], 'a', 'letan', DEFINITIONS, cfg, norm)
    assert [totals[key] for key in ('minutes_ca1','minutes_ca2_before_22','minutes_ca2_after_22')] == [0, 300, 300]
    assert dep._recalculate(payroll_row(totals,cfg,'schedule',0),cfg)['salary'] == 315_000


def test_named_overtime_overlap_is_not_paid_twice_and_missing_ot_is_not_invented():
    cfg = config()
    totals = dep._schedule_totals([schedule(overtime_shift='TC Ca 1')], 'a', 'letan', DEFINITIONS,cfg,norm)
    assert totals['minutes_ca1'] == 480
    totals = dep._schedule_totals([schedule(overtime_shift='TC Ca 2')], 'a','letan',{'letan':{'Ca 1':DEFINITIONS['letan']['Ca 1']}},cfg,norm)
    assert totals['configuration_issues'] == ['2026-10-01:missing_overtime_times']


def test_day_off_placeholders_are_not_missing_evidence_but_scheduled_absence_is():
    cfg = config('attendance')
    off = record(arrival='', departure='', shift='Nghỉ', attendance_expected=False)
    assert dep._attendance_totals([off], 'a',norm,cfg,schedule_records=[],definitions={})['pending_dates'] == []
    totals = dep._attendance_totals([], 'a',norm,cfg,schedule_records=[schedule()],definitions=DEFINITIONS)
    assert totals['pending_dates'] == ['2026-10-01']
    missing = dep._attendance_totals([record(departure='')], 'a',norm,cfg,schedule_records=[schedule()],definitions=DEFINITIONS)
    assert missing['pending_dates'] == ['2026-10-01']


@pytest.mark.parametrize('source', ['schedule', 'attendance'])
@pytest.mark.parametrize('approved', [False, True])
@pytest.mark.parametrize('mode', ['hourly', 'monthly'])
def test_late_is_unpaid_and_total_loss_is_one_or_two_times_for_both_sources(source, approved, mode):
    cfg = config(source, mode)
    late = event(cfg, approved)
    planned, actual = [schedule()], [record()]
    totals = dep._schedule_totals(planned,'a','letan',DEFINITIONS,cfg,norm) if source=='schedule' else dep._attendance_totals(actual,'a',norm,cfg,schedule_records=planned,definitions=DEFINITIONS,department='letan')
    row = dep._late_payroll_adjustment(payroll_row(totals,cfg,source,late['amount']),cfg,source,[late],actual,planned,DEFINITIONS,norm)
    one = 9_375 if mode=='monthly' else 6_750
    full_salary = 300_000 if mode=='monthly' else 216_000
    assert row['salary'] == full_salary-one
    assert row['unpaid_late_wage'] == one
    assert row['effective_late_penalty'] == (0 if approved else one)
    assert row['net_salary'] == full_salary-one*(1 if approved else 2)
    assert row['hours_ca1'] == 7.75


@pytest.mark.parametrize('source', ['schedule', 'attendance'])
def test_fractional_minute_does_not_round_display_hours_before_money(source):
    cfg = config(source)
    late = event(cfg, True, end='09:00:30')
    actual, planned = [record(arrival='09:00:30')], [schedule()]
    totals = dep._schedule_totals(planned,'a','letan',DEFINITIONS,cfg,norm) if source=='schedule' else dep._attendance_totals(actual,'a',norm,cfg,schedule_records=planned,definitions=DEFINITIONS,department='letan')
    row = dep._late_payroll_adjustment(payroll_row(totals,cfg,source,late['amount']),cfg,source,[late],actual,planned,DEFINITIONS,norm)
    assert row['hours_ca1'] == 7.99
    assert row['salary'] == 215_775
    assert row['unpaid_late_wage'] == 225 and row['effective_late_penalty'] == 0


def test_late_across_22_boundary_uses_both_rates():
    cfg = config()
    # A late main shift begins at21:45 with a check-in22:15.
    defs = {'letan': {'Ca 2': {'start': '21:45', 'end': '01:00'}}}
    planned, actual = [schedule(shift='Ca 2')], [record(arrival='22:15',departure='01:00',shift='Ca 2')]
    late = event(cfg,True,start='21:45:00',end='22:15:00',shift='Ca 2')
    totals = dep._schedule_totals(planned,'a','letan',defs,cfg,norm)
    row = dep._late_payroll_adjustment(payroll_row(totals,cfg,'schedule',late['amount']),cfg,'schedule',[late],actual,planned,defs,norm)
    assert row['unpaid_late_wage'] == 15_750
    assert row['hours_ca2_before_22'] == 0 and row['hours_ca2_after_22'] == 2.75
    assert row['effective_late_penalty'] == 0


@pytest.mark.parametrize('source', ['schedule','attendance'])
@pytest.mark.parametrize('approved', [False,True])
@pytest.mark.parametrize('change', ['rate','probation','monthly_base'])
def test_actual_payroll_entitlement_keeps_one_and_two_times_after_wage_changes(source,approved,change):
    cfg = config(source, 'monthly' if change=='monthly_base' else 'hourly')
    late = event(cfg,approved)
    if change=='rate': cfg['rate_ca1'] *= 2
    planned,actual = [schedule()],[record()]
    totals = dep._schedule_totals(planned,'a','letan',DEFINITIONS,cfg,norm) if source=='schedule' else dep._attendance_totals(actual,'a',norm,cfg,schedule_records=planned,definitions=DEFINITIONS,department='letan')
    row = payroll_row(totals,cfg,source,late['amount'])
    if change=='probation': row['probation_rate'] = .75
    if change=='monthly_base': row['base_salary'] *= 2
    result = dep._late_payroll_adjustment(row,cfg,source,[late],actual,planned,DEFINITIONS,norm)
    assert not result['configuration_pending']
    expected = {'rate':13_500,'probation':5_063,'monthly_base':18_750}[change]
    assert result['unpaid_late_wage'] == expected
    assert result['effective_late_penalty'] == (0 if approved else expected)
    assert result['late_penalty'] == late['amount'], 'stored original amount stays audit-only'


@pytest.mark.parametrize('source', ['schedule','attendance'])
@pytest.mark.parametrize('approved', [False,True])
def test_seven_second_half_up_rounding_never_becomes_an_extra_excused_fine(source,approved):
    cfg=config(source);late=event(cfg,approved,end='09:00:07')
    planned,actual=[schedule()],[record(arrival='09:00:07')]
    totals=dep._schedule_totals(planned,'a','letan',DEFINITIONS,cfg,norm) if source=='schedule' else dep._attendance_totals(actual,'a',norm,cfg,schedule_records=planned,definitions=DEFINITIONS,department='letan')
    result=dep._late_payroll_adjustment(payroll_row(totals,cfg,source,late['amount']),cfg,source,[late],actual,planned,DEFINITIONS,norm)
    assert result['salary']==215_947 and result['unpaid_late_wage']==53
    assert result['effective_late_penalty']==(0 if approved else 53)
    assert result['late_wage_rounding_adjustment']==1


@pytest.mark.parametrize('change',['missing_schedule','earlier_arrival','later_arrival','changed_shift','event_issue'])
def test_stale_same_event_evidence_blocks_instead_of_partial_credit(change):
    cfg=config();late=event(cfg)
    planned,actual=[schedule()],[record()]
    totals=dep._schedule_totals(planned,'a','letan',DEFINITIONS,cfg,norm)
    if change=='missing_schedule': planned=[]
    if change=='earlier_arrival': actual=[record(arrival='09:10')]
    if change=='later_arrival': actual=[record(arrival='09:20')]
    if change=='changed_shift': planned=[schedule(start_time='10:00',end_time='18:00')]
    if change=='event_issue': late['issue']='unverified_late_event:101'
    result=dep._late_payroll_adjustment(payroll_row(totals,cfg,'schedule',late['amount']),cfg,'schedule',[late],actual,planned,DEFINITIONS,norm)
    assert result['configuration_pending'] and result['salary'] is None
    with pytest.raises(HTTPException): dep._require_complete_attendance(object(),MONTH,[result],norm)


def test_missing_default_rate_blocks_but_explicit_zero_is_valid():
    cfg = {**dep.DEFAULT_CONFIG['support'], 'configured_rate_keys': dep._configured_rates('support',{})}
    row = dep._recalculate({'hours_ca1':8},cfg)
    assert row['salary'] is None and row['configuration_pending']
    cfg['configured_rate_keys'] = dep._configured_rates('support',{'rate_ca1':0})
    assert dep._recalculate({'hours_ca1':8},cfg)['salary'] == 0


@pytest.fixture
def mixed(monkeypatch):
    state = {'configs': {'a':config(), 'b':config('attendance')}, 'saved': {}, 'revisions': {}, 'records': [record('b',arrival='09:00')], 'attendance_reads': 0, 'schedule_reads': 0, 'late_events': {}}
    class Connection:
        def execute(self, statement, params=None):
            sql = str(statement)
            if 'pg_advisory_xact_lock_shared' in sql:
                return None
            if 'FROM vera_work_schedule ' in sql:
                state['schedule_reads'] += 1
                rows = [schedule('a'), schedule('b')]
            elif 'FROM vera_work_shift_definition' in sql:
                rows = [{'department':'letan','shift_code':key,'start_time':v['start'],'end_time':v['end']} for key,v in DEFINITIONS['letan'].items()]
            else:
                raise AssertionError(sql)
            return SimpleNamespace(mappings=lambda:SimpleNamespace(all=lambda:rows))
    def records(*args):
        state['attendance_reads'] += 1
        return state['records']
    monkeypatch.setattr(dep.work_schedule,'_ensure_schema',lambda conn:None)
    monkeypatch.setattr(dep,'_probation_policies',lambda *a,**k:{})
    monkeypatch.setattr(dep,'_combo_sale_counts',lambda *a:{})
    monkeypatch.setattr(dep,'_penalty_maps',lambda *a:({},{}))
    monkeypatch.setattr(dep,'_hc_late_events',lambda *a:state['late_events'])
    monkeypatch.setattr(dep,'_employee_config_map',lambda c:deepcopy(state['configs']))
    monkeypatch.setattr(dep,'_employees',lambda c,d:[{'username':'a'},{'username':'b'}])
    monkeypatch.setattr(dep,'_settings',lambda c,d:{'config':config(),'department_label':'Lễ tân'})
    monkeypatch.setattr(dep.hr,'admin_departments',lambda c:{'letan':{'name':'Lễ tân','salary_mode':'hourly'}})
    monkeypatch.setattr(dep.attendance,'_records',records)
    monkeypatch.setattr(dep.payroll,'_setting',lambda c,key,default:deepcopy(state['saved'].get(key,default)))
    monkeypatch.setattr(dep,'_setting_snapshot',lambda c,key,default:(deepcopy(state['saved'].get(key,default)),state['revisions'].get(key,0)))
    monkeypatch.setattr(dep,'_combined_employee_catalog',lambda c:{name:{'username':name,'role':'letan'} for name in ['a','b']})
    monkeypatch.setattr(dep,'_visible_payroll_rows',lambda c,rows,**kwargs:rows)
    state['conn'] = Connection()
    return state


@pytest.mark.parametrize('global_source', ['employee','schedule','attendance'])
def test_mixed_calculation_reads_each_source_once_and_legacy_query_cannot_override(mixed,global_source):
    result = dep._combined_calculation(mixed['conn'], MONTH,norm,global_source)
    assert [row['calculation_source'] for row in result['rows']] == ['schedule','attendance']
    assert mixed['attendance_reads'] == mixed['schedule_reads'] == 1
    assert all(row['calculation_snapshot']['month']==MONTH for row in result['rows'])


def test_mixed_missing_attendance_cannot_be_unblocked_by_schedule_employee(mixed):
    mixed['records'] = []
    rows = dep._combined_calculation(mixed['conn'],MONTH,norm)['rows']
    assert rows[0]['salary'] is not None and rows[1]['salary'] is None
    with pytest.raises(HTTPException): dep._require_complete_attendance(mixed['conn'],MONTH,rows,norm)
    rows[1]['attendance_pending'] = False
    clean = dep._clean_combined_rows(mixed['conn'],rows,norm,month=MONTH)
    assert clean[1]['attendance_pending'] is True


@pytest.mark.parametrize('field,value', [('calculation_source','schedule'), ('calculation_config',config()), ('hours_ca1',1000)])
def test_forged_attendance_row_provenance_is_rejected(mixed,field,value):
    rows = dep._combined_calculation(mixed['conn'],MONTH,norm)['rows']
    rows[1][field] = value
    with pytest.raises(HTTPException) as error: dep._clean_combined_rows(mixed['conn'],rows,norm,month=MONTH)
    assert error.value.status_code==409


def test_injected_credit_and_precise_basis_are_replaced_from_server(mixed):
    rows = dep._combined_calculation(mixed['conn'],MONTH,norm)['rows']
    rows[1].update(late_penalty_credit=999999,unpaid_late_wage=999999,payroll_time_basis={'hours_ca1':'100000','work_days':'10000'})
    clean = dep._clean_combined_rows(mixed['conn'],rows,norm,month=MONTH)
    assert clean[1]['late_penalty_credit']==0 and clean[1]['salary']==216_000


def test_saved_draft_cannot_keep_schedule_after_employee_switches_to_attendance(mixed):
    rows = dep._combined_calculation(mixed['conn'],MONTH,norm)['rows']
    mixed['saved'][dep._draft_key(MONTH)] = deepcopy(rows)
    mixed['configs']['a']['salary_source']='attendance'
    with pytest.raises(HTTPException) as error: dep._clean_combined_rows(mixed['conn'],rows,norm,month=MONTH)
    assert error.value.status_code==409


def test_cross_month_snapshot_cannot_be_relabelled(mixed):
    rows = dep._combined_calculation(mixed['conn'],MONTH,norm)['rows']
    with pytest.raises(HTTPException): dep._clean_combined_rows(mixed['conn'],rows,norm,month='2026-09')


def test_frozen_history_export_preserves_money_source_and_config(mixed):
    rows = dep._combined_calculation(mixed['conn'],MONTH,norm)['rows']
    mixed['saved']['department_payroll_combined_history']=[{'id':'h1','month':MONTH,'rows':deepcopy(rows)}]
    mixed['configs']['a']['salary_source']='attendance'
    mixed['configs']['a']['rate_ca1']=99999
    body = dep.CombinedPayrollDraft(month=MONTH,rows=deepcopy(rows),history_id='h1')
    body.rows[0]['_opened_history_id']='h1'
    frozen = dep._frozen_export_rows(mixed['conn'],body)
    assert frozen == rows
    body.rows[0]['advance']=100
    assert dep._frozen_export_rows(mixed['conn'],body) is None


def test_excel_contains_source_snapshot_and_effective_fine():
    row = {'employee_username':'a','employee_name':'A','calculation_source':'attendance',
           'late_penalty':20_000,'effective_late_penalty':10_000,'unpaid_late_wage':10_000,
           'calculation_snapshot':{'month':MONTH,'source':'attendance'}}
    workbook = load_workbook(BytesIO(dep._combined_workbook([row],'10/2026')))
    assert workbook.active['R3'].value==10_000
    sheet=workbook['Nguồn và khấu trừ']
    assert sheet['B2'].value=='Chấm công' and sheet['C2'].value==10_000
    assert 'attendance' in sheet['F2'].value


def test_source_only_config_uses_department_rates_without_creating_numeric_override(monkeypatch):
    values = {'existing': {'rate_ca1': 99_000}}
    sources = {'a': 'attendance'}
    monkeypatch.setattr(dep,'_employee_settings_snapshot',lambda c:(values,sources,7))
    monkeypatch.setattr(dep,'_employee_config_rows',lambda *args:[])
    monkeypatch.setattr(dep,'_salary_employee_catalog',lambda *args:[{'employee_username':'a','department':'custom','salary_source':'attendance'}])
    inherited = {**config(), 'rate_ca1':12345}
    monkeypatch.setattr(dep,'_settings',lambda *a:{'config':inherited})
    result=dep._employee_config_map(object())
    assert result['a']['rate_ca1']==12345 and result['a']['salary_source']=='attendance'
    assert values=={'existing': {'rate_ca1':99_000}}, 'lookup must not materialize inherited rates'


def test_setting_snapshot_fetches_value_and_revision_in_one_statement():
    calls=[]
    class Connection:
        def execute(self,sql,params):
            calls.append(str(sql))
            return SimpleNamespace(mappings=lambda:SimpleNamespace(first=lambda:{'value_json':{'a':'attendance'},'revision':11}))
    assert dep._setting_snapshot(Connection(),'key',{})==({'a':'attendance'},11)
    assert len(calls)==1 and 'value_json' in calls[0] and 'revision' in calls[0]


@pytest.mark.parametrize('expected',[None,1,3])
def test_stale_draft_revision_fails_closed(monkeypatch,expected):
    monkeypatch.setattr(dep,'_setting_snapshot',lambda *args:([],2))
    connection=SimpleNamespace(execute=lambda *args:None)
    with pytest.raises(HTTPException) as exc: dep._require_draft_revision(connection,MONTH,expected)
    assert exc.value.status_code==409
    assert dep._require_draft_revision(connection,MONTH,2)==([],2)


def test_unchanged_numeric_rows_do_not_materialize_inherited_rates_on_source_save():
    inherited=config()
    previous={'default_combo_sales':42}
    row={**inherited, **previous, 'salary_source':'attendance'}
    assert dep._employee_numeric_update('letan',row,previous,inherited)==previous
    row['rate_ca1']=31000
    assert dep._employee_numeric_update('letan',row,previous,inherited)=={**previous,'rate_ca1':31000}


@pytest.mark.parametrize('approved',[True,False])
@pytest.mark.parametrize('reverse',[False,True])
def test_multiple_fractional_late_events_telescope_and_are_source_order_independent(approved,reverse):
    result_by_source={}
    for source in ('schedule','attendance'):
        cfg=config(source)
        events=[];planned=[];actual=[]
        for day in (6,7):
            item=event(cfg,approved,end='09:00:07')
            item['event_id']=str(day)
            item['day']=f'2026-10-{day:02d}'
            for key in ('start','end'): item['basis'][key]=item['basis'][key].replace('2026-10-01',item['day'])
            for segment in item['basis']['wage_segments']:
                for key in ('start','end'): segment[key]=segment[key].replace('2026-10-01',item['day'])
            events.append(item)
            planned.append({**schedule(),'work_date':date(2026,10,day)})
            actual.append({**record(arrival='09:00:07'),'date':f'{day:02d}/10/2026'})
        if reverse:events.reverse()
        totals=dep._schedule_totals(planned,'a','letan',DEFINITIONS,cfg,norm) if source=='schedule' else dep._attendance_totals(actual,'a',norm,cfg,schedule_records=planned,definitions=DEFINITIONS,department='letan')
        result=dep._late_payroll_adjustment(payroll_row(totals,cfg,source,sum(e['amount'] for e in events)),cfg,source,events,actual,planned,DEFINITIONS,norm)
        result_by_source[source]=result
        assert result['salary']==431_894 and result['unpaid_late_wage']==106
        assert result['effective_late_penalty']==(0 if approved else 106)
    assert result_by_source['schedule']['net_salary']==result_by_source['attendance']['net_salary']


def test_monthly_scheduled_lateness_cannot_claim_more_than_its_paid_day():
    cfg=config('schedule','monthly')
    planned=[schedule(overtime_shift='TC Ca 2')]
    actual=[record(arrival='18:00',departure='01:00')]
    late=event(cfg,True,end='18:00:00')
    late['basis']['wage_segments']=[
        {'start':'2026-10-01T09:00:00+07:00','end':'2026-10-01T17:00:00+07:00','shift':'Ca 1'},
        {'start':'2026-10-01T17:00:00+07:00','end':'2026-10-01T18:00:00+07:00','shift':'Ca 2'}]
    totals=dep._schedule_totals(planned,'a','letan',DEFINITIONS,cfg,norm)
    result=dep._late_payroll_adjustment(payroll_row(totals,cfg,'schedule',late['amount']),cfg,'schedule',[late],actual,planned,DEFINITIONS,norm)
    assert result['configuration_pending'] and result['salary'] is None
    assert result['late_wage_rounding_adjustment']==0


@pytest.mark.parametrize('seconds',[1,3,7,11,15,19,23,27,31,59])
def test_fractional_wage_multiplies_exact_seconds_before_division(seconds):
    cfg=config();late=event(cfg,True,end=f'09:00:{seconds:02d}')
    actual,planned=[record(arrival=f'09:00:{seconds:02d}')],[schedule()]
    totals=dep._schedule_totals(planned,'a','letan',DEFINITIONS,cfg,norm)
    row=dep._late_payroll_adjustment(payroll_row(totals,cfg,'schedule',late['amount']),cfg,'schedule',[late],actual,planned,DEFINITIONS,norm)
    expected=(seconds*15+1)//2
    assert row['unpaid_late_wage']==expected
    assert row['salary']==216_000-expected and row['effective_late_penalty']==0


def test_historical_payroll_email_uses_current_directory_recipient_without_rewriting_snapshot(monkeypatch):
    monkeypatch.setattr(dep,'_combined_employee_catalog',lambda conn:{'a':{'username':'a','email':'current@example.test'}})
    historical={'employee_username':'a','email':'obsolete@example.test','salary':12345,'calculation_source':'schedule'}
    assert dep._resolve_payroll_email_rows(object(),[historical])==[{**historical,'email':'current@example.test'}]
    assert historical['email']=='obsolete@example.test'
    with pytest.raises(HTTPException):dep._resolve_payroll_email_rows(object(),[dict(historical,employee_username='retired')])


@pytest.mark.parametrize('seconds',[125,129,131])
def test_paid_salary_without_lateness_uses_exact_seconds(seconds):
    cfg=config('attendance')
    actual=[record(arrival='09:00:00',departure=f'09:{seconds//60:02d}:{seconds%60:02d}')]
    totals=dep._attendance_totals(actual,'a',norm,cfg,schedule_records=[schedule()],definitions=DEFINITIONS,department='letan')
    row=dep._recalculate(payroll_row(totals,cfg,'attendance',0),cfg)
    assert row['salary']==(seconds*15+1)//2
    assert Decimal(row['payroll_time_basis']['seconds_ca1'])==seconds
