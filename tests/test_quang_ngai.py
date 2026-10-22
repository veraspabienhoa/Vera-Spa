"""HTTP role enforcement and real PostgreSQL ledger isolation."""
from datetime import date
from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4
import os
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, text
import vera_web_v2_quang_ngai as routes
import vera_revenue_store as original
import vera_purchase_store as original_purchases


def client(engine, role):
    app=FastAPI()
    def identity():
        return SimpleNamespace(role=role, employee_username=role, full_name=role, email='test@example.invalid')
    routes.install_quang_ngai_routes(app,engine_instance=lambda:engine,current_identity=identity,norm=lambda v:str(v or '').lower())
    return TestClient(app),app


@pytest.mark.parametrize('role',['nhanvien','quanly','letan','locker','', 'unknown'])
def test_every_endpoint_denies_other_roles_before_database(role):
    c,app=client(None,role)
    paths=app.openapi()['paths']
    assert paths
    for path,methods in paths.items():
        for method in methods:
            r=c.request(method,path.replace('{entry_id}','1'),json={})
            assert r.status_code==403,(method,path,r.text)


def test_registers_no_revenue_or_tip_routes():
    _,app=client(None,'admin')
    paths=list(app.openapi()['paths'])
    assert paths and all(p.startswith('/v2/quang-ngai') for p in paths)
    assert not any('tip' in p or 'source' in p for p in paths)


@pytest.fixture
def db():
    url=os.getenv('VERA_TEST_POSTGRES_URL')
    if not url: pytest.skip('Real PostgreSQL runs in CI')
    admin=create_engine(url)
    schema='quang_ngai_'+uuid4().hex
    with admin.begin() as c: c.execute(text(f'CREATE SCHEMA {schema}'))
    engine=create_engine(url,connect_args={'options':f'-csearch_path={schema}'})
    try:
        with engine.begin() as c:
            c.execute(text('CREATE TABLE vera_schema_version(component text PRIMARY KEY,version int,updated_at timestamptz)'))
            original.ensure_schema(c);original_purchases.ensure_schema(c)
            routes.prepare_schema(c)
            c.execute(text("INSERT INTO vera_revenue_entry(transaction_type,amount,transaction_date,note) VALUES('Thu',999999,'2026-10-22','original only')"))
            c.execute(text("INSERT INTO vera_purchase_entry(purchase_date,item,quantity,unit_price,amount) VALUES('2026-10-22','original only','1',888,888)"))
        yield engine
    finally:
        engine.dispose()
        with admin.begin() as c:c.execute(text(f'DROP SCHEMA {schema} CASCADE'))
        admin.dispose()


@pytest.mark.parametrize('role',['admin','giamdoc'])
def test_independent_crud_purchase_audit_import_export(db,role):
    c,_=client(db,role)
    assert c.get('/v2/quang-ngai/summary').json()['total_income']==0
    token=c.get('/v2/quang-ngai/revision').json()['revision']
    body=dict(transaction_date='2026-10-22',income_amount=1000000,expense_amount=200000,income_note='Thu QN',expense_note='Mua đồ')
    assert c.post('/v2/quang-ngai/entry',json=body).status_code==200
    assert c.post('/v2/quang-ngai/entry',json=body).status_code==409
    summary=c.get('/v2/quang-ngai/summary').json()
    assert summary['balance']==800000
    assert not any('tip' in key for key in summary)
    assert c.get('/v2/quang-ngai/revision').json()['revision']!=token
    purchase=dict(request_id=str(uuid4()),rows=[dict(purchase_date='2026-10-22',item='QN towels',quantity=2,unit_price=100000,note='QN')])
    for _ in range(2):assert c.post('/v2/quang-ngai/purchases',json=purchase).status_code==200
    assert len(c.get('/v2/quang-ngai/purchases?preset=all').json()['rows'])==1
    result=c.get('/v2/quang-ngai/purchase-reconcile?preset=all&live_ledger=true').json()
    assert result['purchase_total']==200000
    assert result['ledger_income']==1000000
    assert result['overall_status']=='KHỚP'
    exported=c.get('/v2/quang-ngai/ledger/export.xlsx?preset=all&live_ledger=true')
    assert exported.status_code==200
    wb=load_workbook(BytesIO(exported.content));assert wb.active.max_row==3
    assert c.post('/v2/quang-ngai/import.xlsx?mode=append',content=exported.content).json()['inserted']==0
    entry=next(r for r in summary['entries'] if r['type']=='Thu')
    edit=dict(transaction_type='Thu',amount=1500000,transaction_date='2026-10-21',note='Edited QN')
    assert c.patch(f"/v2/quang-ngai/entries/{entry['id']}",json=edit).status_code==200
    assert c.get('/v2/quang-ngai/summary').json()['balance']==1300000
    assert c.delete(f"/v2/quang-ngai/entries/{entry['id']}").status_code==200
    assert len(c.get('/v2/quang-ngai/audit?time_range=all').json()['rows'])==2
    assert c.post('/v2/quang-ngai/import.xlsx?mode=replace',content=b'invalid').status_code==400
    assert c.get('/v2/quang-ngai/summary').json()['balance']==-200000
    assert c.post('/v2/quang-ngai/import.xlsx?mode=replace',content=exported.content).status_code==200
    assert c.get('/v2/quang-ngai/summary').json()['balance']==800000
    # Original ledgers remain untouched even by replace and matching row IDs.
    with db.connect() as conn:
        assert conn.execute(text('SELECT amount FROM vera_revenue_entry')).scalar_one()==999999
        assert conn.execute(text('SELECT amount FROM vera_purchase_entry')).scalar_one()==888
