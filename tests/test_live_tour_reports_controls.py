from copy import deepcopy
from datetime import timedelta
from io import BytesIO
import json
from pathlib import Path
import subprocess

import pytest
from fastapi import HTTPException
from PIL import Image
from openpyxl import load_workbook
import vera_web_v2_live_tour as live
from test_live_tour_backend import NOW, employee, state_with
from test_live_tour_invoice_permissions import ALL, pending_state, paid_state, scoped_client, post
from vera_web_v2_live_tour_permissions import CAPABILITY_FEATURES


def test_pending_carries_booking_time_across_employee_reuse_and_checkout():
    worker = employee('e1', 'An')
    worker.update(service='Body 90', service_price=100, service_price_source='catalog', room='1.1', status='CHO THANH TOÁN', booked_at=(NOW-timedelta(days=2)).isoformat())
    state = state_with(worker)
    pending = live._apply_action(state, 'move_pending', {'employee_id':'e1'}, 'admin', NOW)['pending']
    assert pending['effective_at'] == (NOW-timedelta(days=2)).isoformat()
    assert pending['entries'][0]['booked_at'] == pending['effective_at']
    # Employee assignment cleared/reused; the pending source timestamp survives.
    worker['booked_at'] = NOW.isoformat()
    invoice = live._apply_action(state, 'checkout', {'pending_id':pending['id'], 'payment_method':'TIỀN MẶT'}, 'admin', NOW)['invoice']
    assert invoice['effective_at'] == pending['effective_at']
    assert invoice['recorded_at'] == NOW.isoformat()


@pytest.mark.parametrize('paid', [False, True])
def test_invoice_date_edit_requires_separate_grant_and_updates_ledger(monkeypatch, paid):
    state, item = paid_state() if paid else pending_state()
    action = 'paid_invoice_update' if paid else 'pending_update'
    key = 'invoice_id' if paid else 'pending_id'
    payload = {key:item['id'], 'reason':'Đối soát giờ booking', 'invoice_at':'2026-09-02T09:14:00+07:00'}
    client, shared = scoped_client(monkeypatch, state, ALL - {'live_tour_invoice_date_edit'})
    before = deepcopy(shared)
    assert post(client, shared, action, payload).status_code == 403
    assert shared == before
    client, shared = scoped_client(monkeypatch, state, ALL)
    response = post(client, shared, action, payload)
    assert response.status_code == 200, response.text
    changed = shared['state']['invoices' if paid else 'pending'][0]
    assert changed['effective_at'] == payload['invoice_at'] and changed['business_date'] == '2026-09-02'
    assert changed['created_at'] == item['created_at']
    assert shared['state']['invoice_changes' if paid else 'pending_changes'][-1]['before'] == item
    if paid:
        assert shared['state']['reports'][0]['effective_at'] == payload['invoice_at']
    else:
        invoice = live._checkout(shared['state'], {'pending_id': item['id'], 'payment_method':'TIỀN MẶT'}, 'admin', NOW, False)
        assert invoice['effective_at'] == payload['invoice_at']


@pytest.mark.parametrize('value', ['', 'bad', '2200-01-01T00:00', None])
def test_invalid_date_is_atomic(value):
    state, paid = paid_state(); before = deepcopy(state)
    with pytest.raises(HTTPException):
        live._apply_action(state, 'paid_invoice_update', {'invoice_id':paid['id'], 'reason':'Sửa giờ', 'invoice_at':value}, 'admin', NOW)
    assert state == before


def customer_state():
    state = state_with(employee('e1', 'An'))
    state['customers'] = [{'id':'c1', 'name':'Khách', 'phone':'0901234567', 'combo_purchases':[
        {'id':'p1', 'combo_name':'Body', 'total':10, 'used':3, 'remaining':7}]}]
    return state


@pytest.mark.parametrize('action,feature,payload', [
    ('customer_combo_update','live_tour_customer_combo_edit',{'purchase_id':'p1','remaining':5}),
    ('customer_combo_delete','live_tour_customer_combo_delete',{'purchase_id':'p1'}),
    ('customer_delete','live_tour_customers_delete',{}),
    ('customer_upsert','live_tour_customers_edit',{'customer_name':'Tên sửa','customer_phone':'0901234567'}),
])
def test_customer_write_grants_do_not_inherit_payment(monkeypatch, action, feature, payload):
    state = customer_state()
    client, shared = scoped_client(monkeypatch,state,{'live_tour_payment','live_tour_customers_view','live_tour_admin'})
    before=deepcopy(shared)
    assert post(client,shared,action,{'customer_id':'c1','reason':'Điều chỉnh',**payload}).status_code==403
    assert shared==before
    assert feature in CAPABILITY_FEATURES.values()


def test_combo_adjustment_keeps_usage_and_sales_and_is_audited():
    state=customer_state(); before=deepcopy(state)
    live._apply_action(state,'customer_combo_update',{'customer_id':'c1','purchase_id':'p1','reason':'Đối soát vé','remaining':5},'admin',NOW)
    purchase=state['customers'][0]['combo_purchases'][0]
    assert (purchase['total'],purchase['used'],purchase['remaining'])==(8,3,5)
    assert state['invoices']==before['invoices'] and state['combo_usage']==before['combo_usage']
    assert state['customer_changes'][0]['before']==before['customers'][0]['combo_purchases'][0]
    live._apply_action(state,'customer_combo_delete',{'customer_id':'c1','purchase_id':'p1','reason':'Hủy vé còn lại'},'admin',NOW)
    assert live._available_combo(state,'c1',state['customers'][0]['combo_purchases'][0])['remaining']==0
    assert live._state_response(state,1,NOW,can_customers_view=True)['customers'][0]['combos']==[]
    live._apply_action(state,'customer_delete',{'customer_id':'c1','reason':'Gỡ hồ sơ'},'admin',NOW)
    assert not live._state_response(state,1,NOW,can_customers_view=True)['customers']
    with pytest.raises(HTTPException):live._customer(state,{'customer_id':'c1'})


def test_reserved_combo_cannot_be_adjusted_or_deleted():
    for action in ['customer_combo_update','customer_combo_delete']:
        state=customer_state();state['employees'][0].update(service='Body 90',customer_id='c1',combo_purchase_id='p1')
        before=deepcopy(state)
        payload={'customer_id':'c1','purchase_id':'p1','reason':'Kiểm tra'}
        if action.endswith('update'):payload['remaining']=2
        with pytest.raises(HTTPException):live._apply_action(state,action,payload,'admin',NOW)
        assert state==before


def test_report_route_has_independent_access_and_corrections_use_invoice_ledger(monkeypatch):
    state,invoice=paid_state()
    grants={'live_tour_reports_view','live_tour_reports_edit','live_tour_reports_delete','live_tour_paid_invoice_view'}
    client,shared=scoped_client(monkeypatch,state,grants)
    response=client.get('/v2/live-tour/reports')
    assert response.status_code==200 and len(response.json()['invoices'])==1
    assert 'employees' not in response.json()
    assert client.get('/v2/live-tour').status_code==403
    payload={'invoice_id':invoice['id'],'reason':'Đối soát TIP','tip':50}
    updated=post(client,shared,'report_invoice_update',payload)
    assert updated.status_code==200,updated.text
    assert shared['state']['invoices'][0]['tip']==50 and sum(r['tip'] for r in shared['state']['reports'])==50
    result=post(client,shared,'report_invoice_delete',{'invoice_id':invoice['id'],'reason':'Hủy giao dịch trùng'})
    assert result.status_code==200 and not shared['state']['invoices'] and not shared['state']['reports']
    assert len(shared['state']['invoice_changes'])==2


def test_report_export_selector_matches_same_invoice_line(monkeypatch):
    state,invoice=paid_state()
    invoice['entries']=[{'employee_name':'An','service':'Body 90','price':10},{'employee_name':'Bình','service':'Foot','price':20}]
    invoice.update(effective_at='2026-09-05T01:00:00+07:00',business_date='2026-09-04')
    client,_=scoped_client(monkeypatch,state,ALL)
    base={'kind':'revenue','date_from':'2026-09-05','date_to':'2026-09-05','employee':'An'}
    data=client.get('/v2/live-tour/export.xlsx',params={**base,'service':'Foot'})
    assert data.status_code==200 and load_workbook(BytesIO(data.content)).active.max_row==1
    data=client.get('/v2/live-tour/export.xlsx',params={**base,'service':'Body'})
    assert load_workbook(BytesIO(data.content)).active.max_row==2


def test_full_board_image_starts_at_header_and_grows_for_every_row():
    state=state_with(employee('e1','Thiên Kim'))
    state['employees'][0]['note']='Ghi chú dài '*30
    first=Image.open(BytesIO(live._png_bytes(state,NOW)))
    assert first.width > 4000
    assert first.getpixel((10,2)) == (23,78,59)
    state['employees'] += [employee(f'e{i}',f'Nhân viên {i}') for i in range(2,45)]
    second=Image.open(BytesIO(live._png_bytes(state,NOW)))
    assert second.width==first.width and second.height>first.height+1000


def test_date_presets_and_filters_in_vietnam_timezone():
    module=(Path(__file__).resolve().parents[1]/'web-v2/src/lib/liveTourFilters.js').as_uri()
    code=f"""import {{tourDateRange,filterTourRows}} from {json.dumps(module)};
const now=new Date('2026-09-06T18:00:00Z');
const row={{effective_at:'2026-09-06T18:00:00Z',customer_name:'Khách Á',entries:[{{employee_name:'An',service:'Body'}},{{employee_name:'Bình',service:'Foot'}}]}};
console.log(JSON.stringify([tourDateRange('week',now),tourDateRange('last-month',new Date('2026-01-02T00:00Z')),filterTourRows([row],{{date_from:'2026-09-07',date_to:'2026-09-07',employee:'an',service:'Foot'}}).length,filterTourRows([row],{{date_from:'2026-09-07',customer:'khach a',employee:'an',service:'body'}}).length]));"""
    result=subprocess.run(['node','--input-type=module','-e',code],check=True,text=True,capture_output=True)
    assert json.loads(result.stdout)==[{'date_from':'2026-09-07','date_to':'2026-09-13'},{'date_from':'2025-12-01','date_to':'2025-12-31'},0,1]


def test_component_combo_adjustment_preserves_used_and_reconciles_totals():
    state=customer_state()
    purchase=state['customers'][0]['combo_purchases'][0]
    purchase['component_balances']=[{'service_id':'s1','used':1,'remaining':3,'total':4},{'service_id':'s2','used':2,'remaining':4,'total':6}]
    payload={'customer_id':'c1','purchase_id':'p1','reason':'Đối soát từng dịch vụ','components':[{'service_id':'s1','remaining':2},{'service_id':'s2','remaining':1}]}
    live._apply_action(state,'customer_combo_update',payload,'admin',NOW)
    target=state['customers'][0]['combo_purchases'][0]
    assert target['remaining']==3 and target['total']==6 and target['used']==3
    assert [p['used'] for p in target['component_balances']]==[1,2]
    before=deepcopy(state)
    payload['components'][1]['service_id']='s1'
    with pytest.raises(HTTPException):live._apply_action(state,'customer_combo_update',payload,'admin',NOW)
    assert state==before


def test_new_controls_and_dialogs_are_mounted_in_both_workspaces():
    root=Path(__file__).resolve().parents[1]/'web-v2/src'
    board=(root/'pages/LiveTourPage.jsx').read_text()
    for obsolete in ["openModal('add_employee')", "runSelected('hide_employee')", "runSelected('delete_employee')", '>Đặt lịch nhanh</button>', '>Đặt lịch hàng loạt</button>', '>+30 phút</button>']:
        assert obsolete not in board
    assert 'Chụp hình bảng tua' in board and 'Nghỉ giữa ca' in board
    assert 'LiveTourBookingDialog' in board and 'setBookingContext' in board
    for page in ['LiveTourPage.jsx','SpaManagementPage.jsx']:
        source=(root/'pages'/page).read_text()
        assert '<LiveTourCustomerDialog' in source and 'customer_combo_edit' in source and 'customer_combo_delete' in source
    reports=(root/'pages/LiveTourReportsPage.jsx').read_text()
    assert 'live_tour_reports_view' in reports and 'reports_edit' in reports and 'reports_delete' in reports
    for filename in ['LiveTourPaidInvoiceDialog.jsx','LiveTourPendingDialog.jsx']:
        source=(root/'components'/filename).read_text()
        assert 'canEditDate' in source and 'invoice_at:' in source and 'expectedRevision: revision' in source


def test_report_excel_preserves_money_filters_and_customer_read_grants(monkeypatch):
    state, invoice = paid_state()
    report = state['reports'][0]
    report.update(customer_name='Khách Á', customer_phone='0901234567', employee_name='An',
                  service='Body', effective_at='2026-09-07T01:00:00+07:00', business_date='2026-09-07')
    invoice['purchased_combo_id'] = 'combo-sale'
    other = dict(report, id='other', invoice_id='other', employee_name='Bình', service='Foot', total=999)
    state['reports'].append(other)
    client, _ = scoped_client(monkeypatch, state, ALL)
    query = {'kind':'reports', 'date_from':'2026-09-07', 'date_to':'2026-09-07',
             'employee':'an', 'customer':'khach a', 'service':'body', 'report_kind':'combos'}
    response = client.get('/v2/live-tour/export.xlsx', params=query)
    assert response.status_code == 200
    rows = list(load_workbook(BytesIO(response.content)).active.values)
    assert len(rows) == 2
    actual = dict(zip(rows[0], rows[1]))
    assert actual['Tổng tiền'] == report['total'] and actual['Tip'] == report['tip']
    assert actual['Khách hàng'] == 'Khách Á' and actual['Ngày giờ hóa đơn'] == report['effective_at']
    grants = {'live_tour_reports_view', 'live_tour_export'}
    client, _ = scoped_client(monkeypatch, state, grants)
    response = client.get('/v2/live-tour/export.xlsx', params={'kind':'reports'})
    rows = list(load_workbook(BytesIO(response.content)).active.values)
    for values in rows[1:]:
        actual = dict(zip(rows[0], values))
        assert actual['Khách hàng'] is None and actual['Điện thoại'] is None
    client, _ = scoped_client(monkeypatch, state, {'live_tour_export'})
    assert client.get('/v2/live-tour/export.xlsx', params={'kind':'reports'}).status_code == 403


def test_pending_excel_uses_corrected_booking_time(monkeypatch):
    state, pending = pending_state()
    live._apply_action(state, 'pending_update', {'pending_id': pending['id'], 'reason':'Đối soát',
                      'invoice_at':'2026-09-02T09:14:00+07:00'}, 'admin', NOW)
    client, _ = scoped_client(monkeypatch, state, ALL)
    response = client.get('/v2/live-tour/export.xlsx', params={'kind':'pending', 'date_from':'2026-09-02', 'date_to':'2026-09-02'})
    rows = list(load_workbook(BytesIO(response.content)).active.values)
    assert len(rows) == 2 and rows[1][0] == '2026-09-02T09:14:00+07:00'


def test_customer_export_omits_deleted_profiles_and_purchases():
    state = customer_state()
    live._apply_action(state, 'customer_combo_delete', {'customer_id':'c1', 'purchase_id':'p1', 'reason':'Gỡ vé'}, 'admin', NOW)
    _, _, rows = live._export_rows(state, 'customers', NOW)
    assert len(rows) == 1 and rows[0][0] == 'Khách' and rows[0][2] is None
    live._apply_action(state, 'customer_delete', {'customer_id':'c1', 'reason':'Gỡ hồ sơ'}, 'admin', NOW)
    assert live._export_rows(state, 'customers', NOW)[2] == []
    assert len(state['customer_changes']) == 2
