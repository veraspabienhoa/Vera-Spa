from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from vera_web_v2_excel_export_style import style_workbook_bytes
from vera_web_v2_payroll import _new_payroll_workbook


FIELDS = [
    "TT",
    "Tên Hệ thống",
    "Họ và tên",
    "Tiền Lương",
    "Tiền Hỗ Trợ Hoàn Lại",
    "Hoàn trả tiền tích lũy",
    "Tích lũy",
    "Chi Phí Sinh Hoạt",
    "Tiền phạt trong tháng",
    "Vi phạm kỳ trước",
    "Tiền ứng lương",
    "Tiền hỗ trợ Locker",
    "Số tiền thực nhận",
    "Số tài khoản ngân hàng",
    "Tên ngân hàng",
]


def test_new_payroll_excel_matches_attached_template_after_global_styling():
    record = {
        "TT": 1,
        "Tên Hệ thống": "An An",
        "Họ và tên": "Đinh Thúy An",
        "Tiền Lương": 28_800_000,
        "Tiền Hỗ Trợ Hoàn Lại": 0,
        "Hoàn trả tiền tích lũy": 0,
        "Tích lũy": 0,
        "Chi Phí Sinh Hoạt": 150_000,
        "Tiền phạt trong tháng": 500_000,
        "Vi phạm kỳ trước": 0,
        "Tiền ứng lương": 0,
        "Tiền hỗ trợ Locker": 80_000,
        "Số tiền thực nhận": 28_070_000,
        "Số tài khoản ngân hàng": "0019074348179017",
        "Tên ngân hàng": "Ngân hàng TMCP Kỹ thương Việt Nam",
    }

    raw = _new_payroll_workbook(
        [record],
        FIELDS,
        date(2026, 9, 1),
        date(2026, 9, 15),
    )
    workbook = load_workbook(BytesIO(style_workbook_bytes(raw)))
    worksheet = workbook.active

    assert worksheet.title == "Bảng lương bản mới K1 Tháng 9"
    assert not worksheet.merged_cells.ranges
    assert worksheet.max_column == 15
    assert worksheet["A1"].value == "BẢNG LƯƠNG NHÂN VIÊN"
    assert worksheet["A2"].value == "KỲ LƯƠNG"
    assert worksheet["B2"].value == "Từ ngày 01/09/2026 đến 15/09/2026"
    assert [worksheet.cell(3, column).value for column in range(1, 16)] == FIELDS
    assert worksheet["A4"].value == 1
    assert worksheet["B4"].value == "An An"
    assert worksheet.freeze_panes == "A4"
    assert worksheet.auto_filter.ref == "A3:O4"

    for coordinate in ("A1", "O1", "A2", "B2", "O2", "A3", "O3"):
        cell = worksheet[coordinate]
        assert cell.fill.fgColor.rgb[-6:] == "1F513F"
        assert cell.font.bold is True
        assert cell.font.color.rgb[-6:] == "FFFFFF"

    assert worksheet["D4"].number_format == "#,##0"
    assert worksheet["M4"].number_format == "#,##0"
    assert worksheet["N4"].value == "0019074348179017"
    assert worksheet["N4"].number_format == "@"
    assert worksheet.column_dimensions["A"].width == 13.44140625
    assert worksheet.column_dimensions["O"].width == 46
    assert worksheet.row_dimensions[4].height == 21
    workbook.close()
