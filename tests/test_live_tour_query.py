"""Cross-language report contract and pre-migration HTTP fallback tests."""
from copy import deepcopy
import json
from pathlib import Path
import shutil
import subprocess

import pytest

import vera_live_tour_query as query
import vera_web_v2_live_tour as live
from test_live_tour_backend import NOW, state_with, employee
from test_live_tour_invoice_permissions import ALL, scoped_client

ROOT = Path(__file__).resolve().parents[1]
FILTERS = dict(date_from='', date_to='', date='', employee='', customer='', service='', bill_no='', total_amount=None, tip_amount=None, performance_timing='all')
GRANTS = {f'can_{name}': True for name in live.CAPABILITY_FEATURES}


def edge_state():
    state = state_with(employee('e1', 'An'))
    state['invoices'] = [dict(id='i1', total=120, tip=20, discount=10, entries=[], purchased_combo_id='')]
    state['reports'] = [
        dict(id='r1', invoice_id='i1', bill_no='B1', effective_at='1/10/2026', customer_name='Đặng Ánh', customer_phone='84901234567', employee_name='An_An', service='Body_90', total=120, tip=20, request='YC', combo_sale=True),
        dict(id='r2', invoice_id='missing', bill_no='B2', effective_at='2026-10-01T00:30:00', customer_name='Đặng Ánh', customer_phone='0901234567', employee_name='Bình', service='Body 90', total=10.125, tip=1.005, invoice_total=0, invoice_discount=8),
        dict(id='r3', invoice_id='missing', bill_no='B2', effective_at='2026-09-30T17:00:00Z', customer_name='An', employee_name='Bình', service='Combo 3', total=20.555, tip=2.555, invoice_discount=11),
        dict(id='r4', bill_no='', business_date='2026-10-01', employee_name='a\u1ab0b', service='Body 90', total=0, tip=0),
        dict(id='r5', bill_no='B5', effective_at='not-a-date', business_date='2026-10-01', employee_name='An', service='Body', total=1, tip=0),
    ]
    return state


def javascript_report(public, tab, filters):
    if not shutil.which('node'):
        pytest.skip('node required for cross-language report parity')
    script = """
      import {selectReportRows} from './web-v2/src/lib/liveTourReportSelection.js';
      import {reportInvoiceMetrics} from './web-v2/src/lib/liveTourReportMetrics.js';
      import {summarizeEmployeeRevenue} from './web-v2/src/lib/liveTourEmployeeRevenue.js';
      import {summarizeTourRevenue} from './web-v2/src/lib/liveTourRevenue.js';
      let input='';for await(const chunk of process.stdin)input+=chunk;
      const {data,tab,filters}=JSON.parse(input);
      const rows=selectReportRows(data,tab,filters);
      console.log(JSON.stringify({ids:rows.map(row=>row.id),
        summary:{...summarizeTourRevenue(rows),...reportInvoiceMetrics(rows,new Map(data.invoices.map(i=>[i.id,i]))),
        invoiceCount:new Set(rows.map(r=>String(r.invoice_id||r.bill_no||'').trim()).filter(Boolean)).size,
        tipEmployeeCount:new Set(rows.map(r=>r.employee_id||r.employee_name)).size},
        employee_totals:summarizeEmployeeRevenue(rows).sort((a,b)=>a.employee<b.employee?-1:a.employee>b.employee?1:0)}));
    """
    data = dict(invoices=public['state']['invoices'], reports=public['report_rows'], performance=[], pending=[])
    return json.loads(subprocess.check_output(['node', '--input-type=module', '-e', script], cwd=ROOT,
                     input=json.dumps(dict(data=data, tab=tab, filters=filters)).encode()))


@pytest.mark.parametrize('tab', ['revenue', 'employee', 'tip', 'combos'])
@pytest.mark.parametrize('filters', [{}, {'date': '2026-10-01'}, {'employee': 'an an'}, {'service': 'body 9'},
    {'employee': 'a b'}, {'customer': 'Dang +84 (901) 234-567'}, {'customer': '0901234567,'}, {'bill_no': ' '}])
def test_fallback_matches_original_javascript_filter_and_totals(tab, filters):
    state = edge_state()
    public = live._state_response(state, 1, NOW, **GRANTS)
    filters = {**FILTERS, **filters}
    expected = javascript_report(public, tab, filters)
    result = query.report_page(public, state, tab=tab, page=1, page_size=100, filters=filters, performance=[])
    assert [row['id'] for row in result['rows']] == expected['ids']
    for key, value in expected['summary'].items():
        assert result['summary'][key] == pytest.approx(value)
    assert sorted(result['employee_totals'], key=lambda row: row['employee']) == expected['employee_totals']


def test_bounded_http_fallback_keeps_complete_totals_and_exports(monkeypatch):
    state = edge_state()
    client, _ = scoped_client(monkeypatch, state, ALL)
    response = client.get('/v2/live-tour/reports', params=dict(tab='revenue', page=2, page_size=1, date='2026-10-01'))
    assert response.status_code == 200, response.text
    data = response.json()
    assert data['total'] == 4 and data['pages'] == 4 and len(data['rows']) == 1
    assert data['summary']['totalRevenue'] == pytest.approx(150.68)
    assert data['summary']['discount'] == 21
    assert client.get('/v2/live-tour/reports').json()['reports'] == live._state_response(state, 1, NOW, **GRANTS)['report_rows']
    assert client.get('/v2/live-tour/reports', params={'tab':'no-such-tab'}).status_code == 400
    assert client.get('/v2/live-tour/reports', params={'tab':'revenue','date':'2026-02-30'}).status_code == 400
    assert client.get('/v2/live-tour/reports', params={'tab':'revenue','page_size':101}).status_code == 422


def test_report_export_predicate_shares_exact_screen_filters():
    state = edge_state()
    for row in live._report_rows_with_combo_kind(state):
        for filters in [dict(date='2026-10-01'), dict(employee='an an'), dict(customer='Dang +84 (901) 234-567')]:
            bounds = {**live._parse_export_bounds(), **filters, 'calendar_date':True, 'invoice_dates':True, 'report_filter':True}
            assert live._event_in_export_bounds(row, bounds) == query.matches_report(row, filters)


def test_report_customer_search_keeps_invoice_view_visibility(monkeypatch):
    state = edge_state()
    client, _ = scoped_client(monkeypatch, state, {'live_tour_reports_view', 'live_tour_invoice_view'})
    data = client.get('/v2/live-tour/reports', params={'tab':'revenue','customer':'dang'}).json()
    assert data['total'] == 2 and data['rows'][0]['customer_name'] == 'Đặng Ánh'
    client, _ = scoped_client(monkeypatch, state, {'live_tour_reports_view'})
    data = client.get('/v2/live-tour/reports', params={'tab':'revenue','customer':'dang'}).json()
    assert data['total'] == 0


def test_customer_pdf_daily_buckets_match_report_calendar_for_legacy_rows():
    from vera_customer_count_pdf import customer_counts
    rows = live._report_rows_with_combo_kind(edge_state())
    rows = [row for row in rows if query.matches_report(row, {'date':'2026-10-01'})]
    summary = customer_counts(rows, date_from='2026-10-01', date_to='2026-10-01')
    assert summary['daily'] == [('2026-10-01', 2)]
    assert summary['total'] == 2 and summary['undated'] == 0


@pytest.mark.parametrize('kind', ['reports', 'tip', 'employee'])
def test_report_based_exports_filter_after_same_pii_redaction(monkeypatch, kind):
    from io import BytesIO
    from openpyxl import load_workbook
    client, _ = scoped_client(monkeypatch, edge_state(), {'live_tour_reports_view', 'live_tour_export'})
    response = client.get('/v2/live-tour/export.xlsx', params={'kind':kind, 'customer':'dang'})
    assert response.status_code == 200, response.text
    workbook = load_workbook(BytesIO(response.content), data_only=True)
    assert (workbook['Tip'] if kind == 'tip' else workbook.active).max_row == 1


def test_performance_export_rejects_non_admin_before_business_reads(monkeypatch):
    client, _ = scoped_client(monkeypatch, edge_state(), {'live_tour_reports_view', 'live_tour_export'})
    response = client.get('/v2/live-tour/export.xlsx', params={'kind':'performance'})
    assert response.status_code == 403


def test_combo_export_uses_same_report_row_predicate():
    for row in [{'purchased_combo_id':'old', 'service':'Body'}, {'entries':[{'service':'Combo'}], 'service':'Body'}, {'service':'Combo'}]:
        bounds = {'report_filter':True, 'report_kind':'combos', 'invoice_dates':True}
        assert live._event_in_export_bounds(row, bounds) == query.is_combo_report(row)
