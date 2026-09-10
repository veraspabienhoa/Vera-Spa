from contextlib import contextmanager
from datetime import datetime
from io import BytesIO

import pytest
from fastapi import FastAPI, HTTPException
from openpyxl import load_workbook
from pydantic import BaseModel

import vera_web_v2_live_tour as live


NOW = datetime(2026, 9, 5, 15, 0, tzinfo=live.VN_TZ)


def history_state():
    state = live._empty_state(NOW)
    state["customers"] = [
        {
            "id": "c-exact", "name": "=Khách A", "phone": "+84901",
            "combo_purchases": [{
                "id": "cp1", "combo_name": "Combo 13", "total": 13,
                "used": 2, "remaining": 11, "price": 2_500_000,
                "purchased_at": "2026-09-04T23:59:00+07:00",
            }],
        },
        {"id": "c-other", "name": "=Khách A", "phone": "+84901", "combo_purchases": []},
    ]
    state["invoices"] = [
        {
            "id": "i1", "customer_id": "c-exact", "bill_no": "LIVE-1",
            "business_date": "2026-09-04", "created_at": "2026-09-05T15:00:00+07:00",
            "effective_at": "2026-09-04T23:59:00+07:00", "recorded_at": "2026-09-05T15:00:00+07:00",
            "payment_method": "TIỀN MẶT", "subtotal": 100, "discount": 10,
            "tip": 5, "total": 95, "actor": "admin", "note": "@ghi chú",
            "entries": [{"employee_id": "e1", "employee_name": "An", "service": "Body 90", "room": "1.1", "price": 100}],
        },
        {"id": "i2", "customer_id": "c-other", "bill_no": "OTHER", "total": 999, "created_at": "2026-09-05T15:00:00+07:00"},
    ]
    state["reports"] = [
        {"id": "r1", "customer_id": "c-exact", "invoice_id": "i1", "service": "Body 90", "total": 95, "tip": 5, "created_at": "2026-09-04T23:59:00+07:00"},
        {"id": "r2", "customer_id": "c-other", "invoice_id": "i2", "total": 999, "created_at": "2026-09-04T23:59:00+07:00"},
    ]
    state["combo_usage"] = [
        {"id": "u1", "customer_id": "c-exact", "combo_purchase_id": "cp1", "units": 2, "created_at": "2026-09-04T23:59:00+07:00"},
        {"id": "u2", "customer_id": "c-other", "units": 9, "created_at": "2026-09-04T23:59:00+07:00"},
    ]
    state["pending"] = [
        {"id": "p1", "customer_id": "c-exact", "created_at": "2026-09-05T15:00:00+07:00", "entries": [{"employee_name": "An", "service": "Body 70", "room": "2.1"}]},
        {"id": "p2", "customer_id": "c-other", "created_at": "2026-09-05T15:00:00+07:00", "entries": []},
    ]
    return state


def test_customer_history_joins_only_exact_stable_id_and_summarizes_ledgers():
    result = live._customer_history(history_state(), "c-exact")

    assert result["customer"]["id"] == "c-exact"
    assert [item["id"] for item in result["invoices"]] == ["i1"]
    assert [item["id"] for item in result["reports"]] == ["r1"]
    assert [item["id"] for item in result["combo_usage"]] == ["u1"]
    assert [item["id"] for item in result["pending"]] == ["p1"]
    assert result["services"][0]["service"] == "Body 90"
    assert result["combo_purchases"][0]["id"] == "cp1"
    assert result["summary"] == {
        "invoice_count": 1, "service_count": 1, "combo_purchase_count": 1,
        "combo_usage_count": 1, "pending_count": 1, "total_revenue": 95,
        "total_tip": 5, "combo_purchased_units": 13, "combo_used_units": 2,
        "combo_remaining_units": 11,
    }


def test_customer_history_rejects_missing_or_unknown_id_without_fuzzy_fallback():
    state = history_state()
    with pytest.raises(HTTPException) as missing:
        live._customer_history(state, "")
    with pytest.raises(HTTPException) as unknown:
        live._customer_history(state, "same-name-is-not-an-id")
    assert missing.value.status_code == 400
    assert unknown.value.status_code == 404


def test_customer_detail_excel_is_multi_sheet_filtered_and_formula_safe():
    bounds = live._parse_export_bounds(date_from="2026-09-04", date_to="2026-09-04")
    content, filename = live._customer_detail_excel_bytes(
        history_state(), "c-exact", NOW, bounds=bounds,
    )
    workbook = load_workbook(BytesIO(content), data_only=False)

    assert workbook.sheetnames == [
        "Tong_quan", "Hoa_don", "Dich_vu", "Combo_da_mua", "Combo_su_dung", "Cho_thanh_toan",
    ]
    assert workbook["Tong_quan"]["B4"].value == "'+84901"
    assert workbook["Tong_quan"]["B3"].value == "'=Khách A"
    assert workbook["Hoa_don"].max_row == 2
    assert workbook["Cho_thanh_toan"].max_row == 1
    assert "c-exact" in filename


class Identity(BaseModel):
    employee_username: str = "tester"
    full_name: str = "Tester"


class Connection:
    def execute(self, *_args, **_kwargs):
        return object()


class Engine:
    @contextmanager
    def begin(self):
        yield Connection()


def test_customer_history_route_requires_customers_permission_before_read(monkeypatch):
    app = FastAPI()
    reads = []

    def require(_conn, _ident, feature):
        assert feature == "live_tour_customers_view"
        raise HTTPException(403, "Không có quyền")

    monkeypatch.setattr(live, "_read_state", lambda *_args, **_kwargs: reads.append(True))
    live.install_live_tour_routes(
        app, engine_instance=Engine, current_identity=lambda: Identity(),
        require_feature=require, feature_allowed=lambda *_args: False,
        identity_type=Identity,
    )
    endpoint = next(
        route.endpoint for route in app.routes
        if getattr(route, "path", "") == "/v2/live-tour/customers/{customer_id}/history"
    )
    with pytest.raises(HTTPException) as error:
        endpoint("c-exact", Identity())
    assert error.value.status_code == 403
    assert reads == []


def test_customer_detail_export_requires_export_and_customer_access(monkeypatch):
    app = FastAPI()
    required = []
    state = history_state()

    def require(_conn, _ident, feature):
        required.append(feature)
        if feature == "live_tour_customers_view":
            raise HTTPException(403, "Không có quyền")

    monkeypatch.setattr(live, "_read_state", lambda *_args, **_kwargs: (state, 1))
    live.install_live_tour_routes(
        app, engine_instance=Engine, current_identity=lambda: Identity(),
        require_feature=require,
        feature_allowed=lambda _conn, _ident, feature: feature == "live_tour_export",
        identity_type=Identity,
    )
    endpoint = next(
        route.endpoint for route in app.routes
        if getattr(route, "path", "") == "/v2/live-tour/export.xlsx"
    )
    with pytest.raises(HTTPException) as error:
        endpoint(
            kind="customer_detail", include_hidden=False, customer_id="c-exact",
            date_from="", date_to="", time_from="", time_to="", ident=Identity(),
        )
    assert error.value.status_code == 403
    assert required == ["live_tour_export", "live_tour_customers_view"]
