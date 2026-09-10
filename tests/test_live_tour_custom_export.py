from copy import deepcopy
from io import BytesIO

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

import vera_web_v2_live_tour as live
from test_live_tour_backend import NOW, employee, state_with
from test_live_tour_safety import api_client


def test_custom_export_keeps_requested_column_order_and_selected_employee():
    state = state_with(employee("e1", "An"), employee("e2", "Bình"))
    state["employees"][1]["note"] = '=HYPERLINK("https://example.com")'
    before = deepcopy(state)
    content, _ = live._excel_bytes(state, "custom", NOW, selected_columns=["Ghi chú", "Tên nhân viên"], employee_ids=["e2"])
    sheet = load_workbook(BytesIO(content)).active
    assert list(sheet.values)[0] == ("Ghi chú", "Tên nhân viên")
    assert sheet.max_row == 2 and sheet.max_column == 2
    assert sheet["B2"].value == "Bình"
    assert sheet["A2"].data_type != "f"
    assert state == before


@pytest.mark.parametrize("columns", [[], ["customer_phone"], ["Tên nhân viên", "Tên nhân viên"]])
def test_custom_export_rejects_unknown_empty_or_duplicate_columns(columns):
    with pytest.raises(HTTPException) as error:
        live._excel_bytes(state_with(employee("e1", "An")), "custom", NOW, selected_columns=columns)
    assert error.value.status_code == 400


@pytest.mark.parametrize("ids", [[], ["deleted"]])
def test_custom_export_rejects_empty_missing_or_hidden_selection(ids):
    state = state_with(employee("e1", "An"), employee("e2", "Bình"))
    state["employees"][1]["hidden"] = True
    with pytest.raises(HTTPException) as error:
        live._excel_bytes(state, "custom", NOW, employee_ids=ids)
    assert error.value.status_code in (400, 409)


def test_http_custom_export_parses_repeated_columns_and_ids(monkeypatch):
    client, _ = api_client(monkeypatch, state_with(employee("e1", "An"), employee("e2", "Bình")))
    result = client.get("/v2/live-tour/export.xlsx", params=[("kind", "custom"), ("columns", "Tên nhân viên"), ("columns", "Phòng"), ("employee_ids", "e2")])
    assert result.status_code == 200
    assert list(load_workbook(BytesIO(result.content)).active.values) == [("Tên nhân viên", "Phòng"), ("Bình", None)]


def test_other_reports_reject_custom_selection_instead_of_ignoring_it():
    with pytest.raises(HTTPException) as error:
        live._excel_bytes(state_with(employee("e1", "An")), "revenue", NOW, employee_ids=["e1"])
    assert error.value.status_code == 400
