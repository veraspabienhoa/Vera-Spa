from contextlib import contextmanager
from copy import deepcopy
from datetime import date
from io import BytesIO
import json
from types import SimpleNamespace
from unittest.mock import patch

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

import vera_web_v2_devices as devices


class Store:
    def __init__(self):
        self.row = None
        self.open = False
        self.calls = []

    @contextmanager
    def begin(self):
        before = deepcopy(self.row)
        self.open = True
        try:
            yield self
        except Exception:
            self.row = before
            raise
        finally:
            self.open = False

    connect = begin

    def execute(self, statement, params=None):
        sql = str(statement); self.calls.append(sql)
        if sql.startswith('INSERT') and self.row is None:
            self.row = {'value_json': json.loads(params['initial']), 'revision': 0}
        if sql.startswith('UPDATE'):
            self.row = {'value_json': json.loads(params['value']), 'revision': self.row['revision'] + 1}
        return SimpleNamespace(mappings=lambda: SimpleNamespace(first=lambda: deepcopy(self.row)))


def fixture(role='admin', rows=None):
    store = Store()
    app = FastAPI()
    identity = SimpleNamespace(role=role, employee_username='admin-test')
    devices.install_device_routes(app, engine_instance=lambda: store, current_identity=lambda: identity,
        require_feature=lambda *a: None, identity_type=SimpleNamespace, read_timesoft=lambda *a: rows or [])
    return TestClient(app), store


BASE = 'start=2026-09-01&end=2026-09-30'
RECORDS = [
    {'event_id': 1, 'occurred_at': '2026-09-23T18:16:43+07:00', 'device_name': 'Đặng Ánh', 'status_code': 0, 'type_code': 'A'},
    {'event_id': 2, 'occurred_at': '2026-09-24T01:16:43Z', 'device_name': '=HYPERLINK("evil")', 'status_code': '1', 'type_code': 'B'},
]


def test_all_routes_deny_non_admin_before_database_or_device_io():
    client, store = fixture('nhanvien')
    with patch('vera_facegate_control_log.fetch_control_log', side_effect=AssertionError('must not call')):
        for path in ['/registry', '/checkin-history?' + BASE, '/checkin-history/export.xlsx?' + BASE]:
            assert client.get('/v2/devices' + path).status_code == 403
        assert client.put('/v2/devices/registry', json={'expected_revision': 0, 'devices': devices.default_devices()}).status_code == 403
    assert not store.calls


def test_inventory_roundtrip_conflict_and_no_secret_leak():
    client, store = fixture()
    data = client.get('/v2/devices/registry').json()
    assert data['revision'] == 0
    data['devices'].append(devices.Device(id='printer-1', name='Lễ tân', kind='printer', connection='usb').model_dump())
    body = {'expected_revision': 0, 'devices': data['devices']}
    assert client.put('/v2/devices/registry', json=body).status_code == 200
    assert client.put('/v2/devices/registry', json=body).status_code == 409
    assert any('FOR UPDATE' in sql for sql in store.calls)
    with patch.dict('os.environ', {'VERA_FACEGATE_BASE_URL': 'http://private-host', 'VERA_FACEGATE_USERNAME': 'secret-user', 'VERA_FACEGATE_PASSWORD': 'secret-value'}):
        response = client.get('/v2/devices/registry')
    assert response.json()['revision'] == 1
    assert response.json()['devices'][1]['connection'] == 'usb'
    assert 'secret-value' not in response.text and 'private-host' not in response.text
    assert response.json()['adapters']['pending']['capabilities'] == []


@pytest.mark.parametrize('patch_value', [{'adapter': 'facegate_server'}, {'address': 'https://user:pass@example.test'}, {'port': 0}, {'password': 'bad'}])
def test_inventory_rejects_invalid_connections_and_secrets(patch_value):
    client, _ = fixture()
    entry = devices.Device(id='new', name='New', kind='printer').model_dump() | patch_value
    assert client.put('/v2/devices/registry', json={'expected_revision': 0, 'devices': devices.default_devices() + [entry]}).status_code == 422


def test_duplicate_device_and_serial_rejected():
    client, _ = fixture()
    for entry in [devices.default_devices()[0], devices.Device(id='new', name='New', kind='faceid', serial='2023044').model_dump()]:
        assert client.put('/v2/devices/registry', json={'expected_revision': 0, 'devices': devices.default_devices() + [entry]}).status_code == 422


def test_filters_and_excel_have_same_rows_and_keep_zero_status():
    client, store = fixture()
    def load(*args):
        assert not store.open
        return {'records': RECORDS, 'truncated': False}
    query = BASE + '&employee=dang+anh&event_id=1&event_date=2026-09-23&status=0&event_type=A'
    with patch('vera_facegate_control_log.fetch_control_log', side_effect=load):
        response = client.get('/v2/devices/checkin-history?' + query)
        export = client.get('/v2/devices/checkin-history/export.xlsx?' + query)
    assert response.status_code == 200, response.text
    assert response.json()['records'] == RECORDS[:1]
    assert response.json()['options']['statuses'] == ['0', '1']
    wb = load_workbook(BytesIO(export.content))
    assert wb.active.max_row == 2
    assert wb.active['B2'].value == '23-09-2026 18:16:43'
    assert wb.active['D2'].value == '0'


def test_excel_formula_strings_and_leading_zeros_are_safe():
    client, _ = fixture(rows=[{'date': '23/09/2026', 'employee_code': '001', 'employee_name': '=1+1', 'punch_times': ['09:00', '18:00']}])
    response = client.get('/v2/devices/checkin-history/export.xlsx?' + BASE + '&source=timesoft&employee=%3D1')
    assert response.status_code == 200
    wb = load_workbook(BytesIO(response.content))
    assert wb.active['A2'].value == '23-09-2026'
    assert wb.active['B2'].value == '001'
    assert wb.active['C2'].data_type == 's'
    assert wb.active['C2'].value == '=1+1'
    assert wb['Thông tin']['B7'].data_type == 's'


def test_capture_filters_export_without_image_pointer():
    client, _ = fixture()
    records = [{'event_id': 5, 'occurred_at': '2026-09-23T18:16:43+07:00', 'event_text': 'Snap', 'event_status': 'Not Processed', 'image_available': True, 'image_ref': {'secret_pointer': 900}}]
    with patch('vera_facegate_control_log.fetch_capture_log', return_value={'records': records, 'truncated': False}):
        r = client.get('/v2/devices/checkin-history/export.xlsx?' + BASE + '&source=capture&status=Not+Processed')
    assert r.status_code == 200
    wb = load_workbook(BytesIO(r.content))
    assert wb.active.max_row == 2
    assert 'secret_pointer' not in str(list(wb.active.values))
    assert wb.active['E2'].value == 'Có'


def test_truncation_blocks_export_and_invalid_dates_do_not_call_device():
    client, _ = fixture()
    with patch('vera_facegate_control_log.fetch_control_log', return_value={'records': RECORDS, 'truncated': True}) as fetch:
        assert client.get('/v2/devices/checkin-history/export.xlsx?' + BASE).status_code == 409
        assert client.get('/v2/devices/checkin-history?start=2026-01-01&end=2026-09-30').status_code == 400
        assert client.get('/v2/devices/checkin-history?' + BASE + '&event_date=2026-10-01').status_code == 400
        assert fetch.call_count == 1


def test_vietnam_day_boundary():
    assert devices.record_day({'occurred_at': '2026-09-23T18:00:00Z'}) == '2026-09-24'
    assert devices.filtered_records(RECORDS, 'facegate', event_date=date(2026, 9, 24)) == RECORDS[1:]
