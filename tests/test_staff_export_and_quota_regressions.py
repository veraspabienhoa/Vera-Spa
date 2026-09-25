from contextlib import contextmanager
from datetime import date
from inspect import getclosurevars, signature
from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile
from zoneinfo import ZoneInfo

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

import vera_web_v2_staff as staff
import vera_leave_quota_alerts as quota
from vera_web_v2_excel_export_style import style_workbook_bytes


@pytest.fixture
def staff_app():
    app = FastAPI()
    dependencies = {name: (lambda *args, **kwargs: None)
                    for name in signature(staff.install_staff_routes).parameters if name != 'app'}
    dependencies.update(identity_type=object, current_identity=lambda:SimpleNamespace(role='admin'), vn_tz=ZoneInfo('Asia/Ho_Chi_Minh'), leave_sheet_id='')
    staff.install_staff_routes(app, **dependencies)
    return app


@pytest.fixture
def workbook_builder(staff_app):
    app = staff_app
    route = next(r for r in app.routes if r.path == '/v2/staff/export.xlsx')
    return getclosurevars(route.endpoint).nonlocals['build_staff_workbook']


def test_export_reads_directory_once_and_never_fetches_identity_images(staff_app):
    endpoint = next(r.endpoint for r in staff_app.routes if r.path == '/v2/staff/export.xlsx')
    cells = dict(zip(endpoint.__code__.co_freevars, endpoint.__closure__))
    calls = []
    class Connection:
        def execute(self, *args, **kwargs):
            raise AssertionError('Export must not query identity documents')
    class Engine:
        @contextmanager
        def connect(self):
            yield Connection()
    row = staff._public_employee({'username':'Test','role':'nhanvien'}, 'Đang làm việc')
    def directory(conn, ident):
        calls.append('directory')
        return {'employees':[row], 'shifts_by_department':{}}
    cells['engine_instance'].cell_contents = Engine
    cells['staff_result'].cell_contents = directory
    response = TestClient(staff_app).get('/v2/staff/export.xlsx')
    assert response.status_code == 200, response.text
    assert calls == ['directory']
    with ZipFile(BytesIO(response.content)) as archive:
        assert not any(path.startswith('xl/media/') for path in archive.namelist())


def test_export_contains_only_employee_data_after_style_middleware(workbook_builder):
    employee = staff._public_employee({'username': 'Test', 'role': 'nhanvien', 'phone': '0123456789'}, 'Đang làm việc')
    payload = workbook_builder([employee], {})
    styled = style_workbook_bytes(payload)
    workbook = load_workbook(BytesIO(styled))
    ws = workbook['DanhSachNhanSu']
    assert not ws._images
    assert 'Ảnh nhân viên' not in [cell.value for cell in ws[1]]
    assert ws['A2'].value == 'Test'
    assert ws.cell(2, staff.STAFF_EXPORT_COLUMNS.index('Điện thoại') + 1).value == '0123456789'
    with ZipFile(BytesIO(styled)) as archive:
        media = [p for p in archive.namelist() if p.startswith('xl/media/')]
        assert not media
        assert not any('/drawings/' in path for path in archive.namelist())


def test_staff_export_retains_text_and_dates_without_portrait_column(workbook_builder):
    employee = staff._public_employee({'username': 'Test', 'role': 'nhanvien'}, 'Đang làm việc')
    payload = workbook_builder([employee], {})
    workbook = load_workbook(BytesIO(payload))
    sheet = workbook['DanhSachNhanSu']
    assert sheet['A2'].value == 'Test'
    assert sheet.row_dimensions[2].height == 22
    assert sheet.cell(2, staff.STAFF_EXPORT_COLUMNS.index('Ngày sinh') + 1).number_format == 'dd-mm-yyyy'


def leave(day, reason, days=0):
    return dict(employee_name='Test', leave_date=date.fromisoformat(day), leave_reason=reason,
                leave_type='Phát sinh' if 'PHÁT SINH' in reason else 'Có phép', calculated_days=days)


def test_quota_includes_older_zero_day_and_excluded_months():
    rows = [leave('2026-07-01', 'Quay video'), leave('2026-08-01', 'Nghỉ PHÉP NĂM', 1),
            leave('2026-09-01', 'Nghỉ CÓ phép', 6)]
    items = quota.summarize(rows)
    assert [(i['month'], i['days'], i['exceeded']) for i in items] == [('2026-09', 6, ['days'])]


def test_quota_reports_older_generated_month_even_without_ordinary_leave():
    rows = [leave('2026-08-01', 'Đi trễ PHÁT SINH') for _ in range(3)]
    rows += [leave('2026-09-01', 'Quay video')]
    items = quota.summarize(rows)
    assert len(items) == 1
    assert items[0]['month'] == '2026-08'
    assert items[0]['generated'] == 3
    assert items[0]['exceeded'] == ['generated']


def test_quota_endpoint_selected_month_handles_older_excluded_history():
    rows = [leave('2026-08-01', 'Quay video'), leave('2026-09-01', 'Nghỉ CÓ phép', 6)]
    class Connection:
        def execute(self, sql, params):
            return SimpleNamespace(mappings=lambda: SimpleNamespace(all=lambda: rows))
    class Engine:
        @contextmanager
        def connect(self):
            yield Connection()
    app = FastAPI()
    quota.install(app, engine_instance=Engine, current_identity=lambda: SimpleNamespace(role='admin'), identity_type=object)
    response = TestClient(app).get('/v2/leave/quota-check?start=2026-09-25&end=2026-09-25')
    assert response.status_code == 200
    assert response.json()['items'][0]['month'] == '2026-09'
    assert response.json()['items'][0]['days'] == 6
