"""Attendance card exports select rows without changing source or access rules."""
from contextlib import contextmanager
from copy import deepcopy
from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from types import SimpleNamespace

import pytest
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import vera_web_v2_attendance_summary as summary
import vera_web_v2_operations_v41 as operations


NOW = datetime(2026, 10, 11, 10, tzinfo=timezone.utc)  # 17:00 Vietnam
DAY = date(2026, 10, 11)
BASE_QUERY = {"start": DAY.isoformat(), "end": DAY.isoformat()}
EXPORT = "/v2/snapshot/export.xlsx"


def row(name, **extra):
    return {
        "date": "11/10/2026", "employee_code": name, "employee_name": name,
        "break_department": "Lễ tân", "shift": "Ca 1", "break_enabled": True,
        "break_out": "", "break_in": "", "break_restricted_reason": "",
        "break_over_minutes": 0, "break_return_late_minutes": 0,
        "break_return_deadline": "17:00:00", "break_return_deadline_iso": "",
        "punch_times": ["09:00:00", "15:00:00"],
        "break_auto_penalty_amount": 125000,
        **extra,
    }


def examples():
    return [
        row("complete", break_out="15:00", break_in="16:00"),
        row("complete-over", break_out="15:00", break_in="16:00", break_over_minutes="2"),
        row("returned-late", break_out="15:00", break_in="16:00", break_return_late_minutes="0.25"),
        row("active-over", break_out="15:00", break_return_deadline_iso="2026-10-11T16:59:59+07:00"),
        row("active-boundary", break_out="15:00"),
        row("not-started"),
        row("disabled", break_enabled=False),
        row("restricted-out", break_restricted_reason="Đi trễ", break_out="15:00", break_in="16:00"),
        row("restricted-no-out", break_restricted_reason="Đi trễ", break_over_minutes=99),
        row("active-unknown", break_out="15:00", break_return_deadline="invalid"),
        row("missing-out", break_in="16:00"),
    ]


EXPECTED = {
    "all": [item["employee_name"] for item in examples()],
    "breaks": ["complete", "complete-over", "returned-late"],
    "over": ["complete-over", "returned-late", "active-over", "restricted-out"],
    "incomplete": ["active-over", "active-boundary", "not-started", "active-unknown", "missing-out"],
}


@pytest.mark.parametrize("kind", EXPECTED)
def test_card_membership_overlaps_and_never_mutates_rows(kind):
    records = examples()
    before = deepcopy(records)
    selected = summary.filter_attendance_summary(records, kind, now=NOW)
    assert [item["employee_name"] for item in selected] == EXPECTED[kind]
    assert all(any(item is original for original in records) for item in selected)
    assert records == before


@pytest.mark.parametrize("minutes", [None, "", "invalid", "NaN", "-1", 0])
def test_nonpositive_or_invalid_minutes_do_not_make_completed_pair_overdue(minutes):
    record = row("complete", break_out="15:00", break_in="16:00",
                 break_over_minutes=minutes, break_return_late_minutes=minutes)
    assert summary.filter_attendance_summary([record], "over", now=NOW) == []


@pytest.mark.parametrize("stamp", [
    "2026-10-11T16:59:59", "2026-10-11 16:59:59", "2026-10-11T16:59:59+07:00",
    "2026-10-11T09:59:59Z", "2026-10-11T02:59:59-07:00",
    "2026-10-11t09:59:59z", "2026-10-11T16:59:59+0700",
])
def test_deadline_iso_respects_offsets_and_naive_vietnam_time(stamp):
    record = row("active", break_out="15:00", break_return_deadline_iso=stamp)
    assert summary.filter_attendance_summary([record], "over", now=NOW) == [record]


@pytest.mark.parametrize("work_day", ["11/10/2026", "11-10-2026", "2026-10-11", "1/10/2026", "1-10-2026"])
@pytest.mark.parametrize("stamp", ["", "not-a-date"])
def test_deadline_fallback_supports_business_date_formats(work_day, stamp):
    record = row("active", date=work_day, break_out="15:00",
                 break_return_deadline="16:59:59", break_return_deadline_iso=stamp)
    assert summary.filter_attendance_summary([record], "over", now=NOW) == [record]


@pytest.mark.parametrize("offset,expected", [(-1, True), (0, False), (1, False)])
def test_deadline_is_strict_and_preserves_subsecond_boundary(offset, expected):
    stamp = (NOW + timedelta(milliseconds=offset)).isoformat()
    record = row("active", break_out="15:00", break_return_deadline_iso=stamp)
    assert bool(summary.filter_attendance_summary([record], "over", now=NOW)) is expected


@pytest.mark.parametrize("work_day,clock", [
    ("31/02/2026", "16:00:00"), ("29-02-2026", "16:00"),
    ("2026-13-11", "16:00"), ("11/10/2026", "24:00:00"),
    ("11/10/2026", "16:60:00"), ("", "16:00"), ("11/10/2026", ""),
])
def test_invalid_fallback_dates_and_clocks_cannot_be_overdue(work_day, clock):
    record = row("active", date=work_day, break_out="15:00", break_return_deadline=clock)
    assert summary.filter_attendance_summary([record], "over", now=NOW) == []


def test_valid_iso_deadline_takes_precedence_and_returned_or_unstarted_rows_do_not_expire():
    record = row("active", break_out="15:00", break_return_deadline="10:00",
                 break_return_deadline_iso="2026-10-11T17:00:00+07:00")
    assert summary.filter_attendance_summary([record], "over", now=NOW) == []
    expired = "2026-10-11T01:00:00Z"
    for changes in ({"break_in": "16:00"}, {"break_out": ""}):
        item = {**record, "break_return_deadline_iso": expired, **changes}
        assert summary.filter_attendance_summary([item], "over", now=NOW) == []


@pytest.mark.parametrize("stamp", ["2026-10-11", "20261011T01:00:00", "2026-10-11T01:00:00+07:60", "2026-10-11T24:00:00"])
def test_invalid_iso_shapes_fall_back_instead_of_normalizing(stamp):
    record = row("active", break_out="15:00", break_return_deadline_iso=stamp)
    assert summary.filter_attendance_summary([record], "over", now=NOW) == []


@pytest.mark.parametrize("kind", ["", "unknown", "BREAKS", None])
def test_unknown_summary_filter_is_rejected_by_pure_helper(kind):
    with pytest.raises(ValueError, match="Unknown attendance summary filter"):
        summary.filter_attendance_summary(examples(), kind, now=NOW)


@pytest.fixture
def api(monkeypatch):
    events = []
    state = {"rows": examples(), "grants": {"snapshot_today", "snapshot_export"},
             "authenticated": True, "transaction_open": False}
    identity = SimpleNamespace(role="letan", username="test-operator")

    class Engine:
        @contextmanager
        def begin(self):
            events.append("begin")
            assert not state["transaction_open"]
            state["transaction_open"] = True
            try:
                yield self
            finally:
                state["transaction_open"] = False
                events.append("close")

        def execute(self, *_args, **_kwargs):
            raise AssertionError("Summary filtering must not issue SQL or writes")

    engine = Engine()

    def current_identity():
        if not state["authenticated"]:
            raise HTTPException(401, "Authentication required")
        return identity

    def require_feature(conn, ident, feature):
        assert conn is engine and ident is identity
        events.append(feature)
        if feature not in state["grants"]:
            raise HTTPException(403, "Permission denied")

    def records(conn, start, end):
        assert conn is engine and state["transaction_open"]
        assert events[-1] in {"snapshot_today", "snapshot_export"}
        events.append(("records", start, end))
        return [item for item in state["rows"]
                if start <= datetime.strptime(item["date"], "%d/%m/%Y").date() <= end]

    class FrozenDatetime(datetime):
        @classmethod
        def now(cls, tz=None):
            return NOW.astimezone(tz) if tz else NOW.replace(tzinfo=None)

    monkeypatch.setattr(summary, "datetime", FrozenDatetime)
    monkeypatch.setattr(operations.snapshot_module, "_records", records)
    original_filter = operations.filter_attendance_summary

    def select(scoped, kind):
        assert not state["transaction_open"]
        events.append(("summary", kind, [item["employee_name"] for item in scoped]))
        return original_filter(scoped, kind)

    monkeypatch.setattr(operations, "filter_attendance_summary", select)
    app = FastAPI()
    operations.install_operations_v41(app, engine_instance=lambda: engine,
                                     current_identity=current_identity,
                                     require_feature=require_feature, identity_type=SimpleNamespace)
    with TestClient(app) as client:
        yield client, state, events


def workbook_values(response):
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
    try:
        assert workbook.sheetnames == ["Chấm công"]
        return list(workbook.active.values)
    finally:
        workbook.close()


@pytest.mark.parametrize("kind", EXPECTED)
def test_export_contains_selected_group_and_preserves_workbook_columns(api, kind):
    client, state, events = api
    before = deepcopy(state["rows"])
    values = workbook_values(client.get(EXPORT, params={**BASE_QUERY, "summary_filter": kind}))
    assert len(values[0]) == 25
    assert values[0][:5] == ("Ngày", "Mã nhân viên", "Nhân viên", "Bộ phận", "Ca")
    assert values[0][-5:] == ("Trạng thái ra", "Phút trễ", "Phút về sớm", "Tổng phút", "Số lần chấm")
    assert [item[2] for item in values[1:]] == EXPECTED[kind]
    assert all(item[18] == "09:00:00 · 15:00:00" for item in values[1:])
    assert state["rows"] == before
    assert events[:4] == ["begin", "snapshot_export", ("records", DAY, DAY), "close"]
    assert events[4] == ("summary", kind, EXPECTED["all"])


def test_default_export_retains_all_rows_and_summary_does_not_filter_snapshot_json(api):
    client, _, events = api
    default = workbook_values(client.get(EXPORT, params=BASE_QUERY))
    explicit = workbook_values(client.get(EXPORT, params={**BASE_QUERY, "summary_filter": "all"}))
    assert default == explicit
    events.clear()
    response = client.get("/v2/snapshot", params={**BASE_QUERY, "summary_filter": "breaks"})
    assert response.status_code == 200
    assert response.json()["count"] == len(EXPECTED["all"])
    assert [item["employee_name"] for item in response.json()["records"]] == EXPECTED["all"]
    assert response.json()["data_scope"] == "attendance_only"
    assert events == ["begin", "snapshot_today", ("records", DAY, DAY), "close"]


def test_summary_composes_after_date_employee_department_and_shift_filters(api):
    client, state, events = api
    state["rows"] = [
        row("An An", break_out="15:00", break_in="16:00"),
        row("An An", break_out="15:00", break_in="16:00", date="10/10/2026"),
        row("Vân Anh", break_out="15:00", break_in="16:00"),
        row("An An", break_out="15:00", break_in="16:00", shift="Ca 2"),
        row("An An", break_out="15:00", break_in="16:00", break_department="Locker"),
        row("An An", break_out="15:00", employee_code="incomplete"),
    ]
    query = {**BASE_QUERY, "employee": "an an", "department": "le tan", "shift": "ca 1", "summary_filter": "breaks"}
    values = workbook_values(client.get(EXPORT, params=query))
    assert len(values) == 2 and values[1][1:5] == ("An An", "An An", "Lễ tân", "Ca 1")
    assert events[-1] == ("summary", "breaks", ["An An", "An An"])
    response = client.get("/v2/snapshot", params=query)
    assert response.json()["count"] == 2


@pytest.mark.parametrize("kind", ["", "unknown", "BREAKS", "all,breaks", "../all"])
def test_unknown_export_group_fails_before_record_retrieval(api, kind):
    client, _, events = api
    response = client.get(EXPORT, params={**BASE_QUERY, "summary_filter": kind})
    assert response.status_code == 422
    assert events == []


@pytest.mark.parametrize("kind", EXPECTED)
def test_summary_filter_never_bypasses_existing_export_permission(api, kind):
    client, state, events = api
    state["grants"] = {"snapshot_today"}
    response = client.get(EXPORT, params={**BASE_QUERY, "summary_filter": kind})
    assert response.status_code == 403
    assert events == ["begin", "snapshot_export", "close"]


def test_unauthenticated_export_still_fails_before_database_access(api):
    client, state, events = api
    state["authenticated"] = False
    assert client.get(EXPORT, params={**BASE_QUERY, "summary_filter": "over"}).status_code == 401
    assert events == []


@pytest.mark.parametrize("start,end", [(DAY, DAY - timedelta(days=1)), (DAY, DAY + timedelta(days=63))])
def test_existing_date_range_limits_still_precede_database_access(api, start, end):
    client, _, events = api
    response = client.get(EXPORT, params={"start": start.isoformat(), "end": end.isoformat(), "summary_filter": "over"})
    assert response.status_code == 400
    assert events == []


def test_empty_group_retains_a_valid_header_only_workbook(api):
    client, state, _ = api
    state["rows"] = [row("not-started")]
    values = workbook_values(client.get(EXPORT, params={**BASE_QUERY, "summary_filter": "breaks"}))
    assert len(values) == 1 and len(values[0]) == 25
