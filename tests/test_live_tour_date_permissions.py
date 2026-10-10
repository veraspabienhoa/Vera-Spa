"""HTTP regression gates for independent, server-enforced Live Tour date grants.

These fixtures deliberately model explicitly configured effective permissions.
A missing date feature is denied, including for identities whose role name is
otherwise privileged. Legacy/default inheritance is tested at the resolver.
"""
from copy import deepcopy
from datetime import date, datetime
from io import BytesIO
import json

import pytest
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook

import vera_web_v2_live_tour as live
from vera_web_v2_live_tour_permissions import CAPABILITY_FEATURES
from test_live_tour_backend import NOW, RouteEngine, RouteIdentity, employee, state_with
from test_live_tour_invoice_permissions import combo_pending, post
from test_live_tour_safety import api_client, payable_state


SECTIONS = ("pending", "invoices", "reports", "history")
PRESETS = ("all", "today", "yesterday", "week", "last-week", "month", "last-month", "custom")
BASE_GRANTS = set(CAPABILITY_FEATURES.values()) | {
    "live_tour_view", "live_tour_operate", "live_tour_payment", "live_tour_admin", "live_tour_export",
}
DAYS = (
    "2026-08-01", "2026-08-24", "2026-08-30", "2026-08-31", "2026-09-01",
    "2026-09-04", "2026-09-05", "2026-09-06", "2026-09-30", "2026-10-01",
)
TODAY = "2026-09-05"
YESTERDAY = "2026-09-04"


def date_feature(section, preset):
    return f"live_tour_{section}_date_{preset.replace('-', '_')}"


def grants_for(**sections):
    grants = set(BASE_GRANTS)
    for section, presets in sections.items():
        for preset in (presets,) if isinstance(presets, str) else presets:
            grants.add(date_feature(section, preset))
    return grants


def dated_state():
    state = state_with(employee("e1", "An"))
    state["customers"] = [{"id": "c1", "name": "Date test customer", "phone": "0901234567", "combo_purchases": []}]
    for day in DAYS:
        stamp = f"{day}T15:00:00+07:00"
        common = {"business_date": day, "effective_at": stamp, "created_at": stamp,
                  "customer_id": "c1", "customer_name": "Date test customer", "total": 110, "tip": 10}
        entry = {"employee_id": "e1", "employee_name": "An", "service": "Body 90", "room": "1.1", "price": 100}
        state["pending"].append({**common, "id": f"p-{day}", "bill_no": f"P-{day}", "entries": [deepcopy(entry)]})
        state["invoices"].append({**common, "id": f"i-{day}", "bill_no": f"I-{day}", "entries": [deepcopy(entry)],
                                  "subtotal": 100, "discount": 0, "payment_method": "TIỀN MẶT"})
        state["reports"].append({**common, **entry, "id": f"r-{day}", "invoice_id": f"i-{day}", "bill_no": f"I-{day}"})
        state["audit"].append({"id": f"a-{day}", "at": stamp, "business_date": day,
                               "action": "set_vip", "actor": "tester", "detail": {"day": day}})
        state["break_events"].append({"id": f"b-{day}", "created_at": stamp, "business_date": day,
                                      "employee_id": "e1", "employee_name": "An", "event_type": "start_break"})
        state["backups"].append({"id": f"backup-{day}", "created_at": stamp, "snapshot": {"invoices": [], "pending": []}})
    return state


def date_client(monkeypatch, grants, state=None, *, role="letan", now=NOW, bulk=False):
    class FixedDateTime(datetime):
        @classmethod
        def now(cls, tz=None):
            return now.astimezone(tz) if tz else now.replace(tzinfo=None)

    class DateIdentity(RouteIdentity):
        role: str = "letan"

    monkeypatch.setattr(live, "datetime", FixedDateTime)
    monkeypatch.setattr(live.resource_store, "enabled", lambda: False)
    monkeypatch.setattr(live.query_store, "read_collection", lambda *_a, **_k: None)
    monkeypatch.setattr(live.query_store, "read_reports", lambda *_a, **_k: None)
    monkeypatch.setattr(live.query_store, "read_customer_history", lambda *_a, **_k: None)
    _, shared = api_client(monkeypatch, state if state is not None else dated_state())

    def require(_conn, _ident, feature):
        if feature not in grants:
            raise HTTPException(403, feature)

    app = FastAPI()
    live.install_live_tour_routes(
        app, engine_instance=RouteEngine, current_identity=lambda: DateIdentity(role=role),
        require_feature=require, feature_allowed=lambda _c, _i, feature: feature in grants,
        feature_map=(lambda _c, _i, keys: {key: key in grants for key in keys}) if bulk else None,
        identity_type=DateIdentity,
    )
    return TestClient(app), shared


def row_ids(rows):
    return {row["id"] for row in rows}


def section_rows(data, section):
    if section == "invoices":
        return data["state"]["invoices"]
    return data[{"pending": "pending", "reports": "report_rows", "history": "audit"}[section]]


def prefix(section):
    return {"pending": "p", "invoices": "i", "reports": "r", "history": "a"}[section]


# Each direct read/export must authorize its own date scope before reading data.
ROUTES = [
    *((section, f"/v2/live-tour/collections/{section}", {}) for section in SECTIONS),
    ("reports", "/v2/live-tour/reports", {}),
    ("reports", "/v2/live-tour/reports", {"tab": "revenue"}),
    ("reports", "/v2/live-tour/reports", {"tab": "employee"}),
    ("reports", "/v2/live-tour/reports", {"tab": "tip"}),
    ("reports", "/v2/live-tour/board-history", {}),
    ("reports", "/v2/live-tour/board-history/export.xlsx", {}),
    ("pending", "/v2/live-tour/export.xlsx", {"kind": "pending"}),
    ("invoices", "/v2/live-tour/export.xlsx", {"kind": "paid"}),
    *(("reports", "/v2/live-tour/export.xlsx", {"kind": kind}) for kind in ("reports", "revenue", "employee", "tip")),
    *(("history", "/v2/live-tour/export.xlsx", {"kind": kind}) for kind in ("history", "breaks")),
    ("reports", "/v2/live-tour/customer-count.pdf", {}),
    ("reports", "/v2/live-tour/customer-count.png", {}),
]


@pytest.mark.parametrize("section,path,params", ROUTES)
@pytest.mark.parametrize("query", [{}, {"preset": "all"}, {"preset": "today"},
                                    {"date_from": TODAY, "date_to": TODAY}])
def test_direct_routes_fail_closed_when_no_date_option_is_granted(monkeypatch, section, path, params, query):
    client, _ = date_client(monkeypatch, BASE_GRANTS)
    response = client.get(path, params={**params, **query})
    assert response.status_code == 403, (section, path, response.text)


@pytest.mark.parametrize("section", SECTIONS)
@pytest.mark.parametrize("preset,start,end", [
    ("all", DAYS[0], DAYS[-1]),
    ("today", TODAY, TODAY),
    ("yesterday", YESTERDAY, YESTERDAY),
    ("week", "2026-08-31", "2026-09-06"),
    ("last-week", "2026-08-24", "2026-08-30"),
    ("month", "2026-09-01", "2026-09-30"),
    ("last-month", "2026-08-01", "2026-08-31"),
    ("custom", "2026-08-24", "2026-09-04"),
])
def test_each_explicit_collection_preset_uses_server_vietnam_dates(monkeypatch, section, preset, start, end):
    client, _ = date_client(monkeypatch, grants_for(**{section: preset}))
    params = {"preset": preset, "page_size": 100}
    if preset == "custom":
        params.update(date_from=start, date_to=end)
    response = client.get(f"/v2/live-tour/collections/{section}", params=params)
    assert response.status_code == 200, response.text
    data = response.json()["data"]
    expected = {f"{prefix(section)}-{day}" for day in DAYS if start <= day <= end}
    assert row_ids(section_rows(data, section)) == expected


@pytest.mark.parametrize("section", SECTIONS)
def test_manual_range_is_covered_by_union_but_named_presets_need_the_exact_grant(monkeypatch, section):
    client, _ = date_client(monkeypatch, grants_for(**{section: ("today", "yesterday")}))
    path = f"/v2/live-tour/collections/{section}"
    allowed = client.get(path, params={"date_from": YESTERDAY, "date_to": TODAY})
    assert allowed.status_code == 200, allowed.text
    assert row_ids(section_rows(allowed.json()["data"], section)) == {
        f"{prefix(section)}-{YESTERDAY}", f"{prefix(section)}-{TODAY}",
    }
    for query in ({}, {"preset": "all"}, {"preset": "month"},
                  {"date_from": "2026-09-03", "date_to": TODAY},
                  {"date_from": YESTERDAY}, {"date_to": TODAY}):
        assert client.get(path, params=query).status_code == 403


@pytest.mark.parametrize("section", SECTIONS)
def test_month_grant_does_not_enable_today_button_but_covers_manual_today(monkeypatch, section):
    client, _ = date_client(monkeypatch, grants_for(**{section: "month"}))
    path = f"/v2/live-tour/collections/{section}"
    assert client.get(path, params={"preset": "today"}).status_code == 403
    response = client.get(path, params={"date_from": TODAY, "date_to": TODAY})
    assert response.status_code == 200, response.text
    assert row_ids(section_rows(response.json()["data"], section)) == {f"{prefix(section)}-{TODAY}"}


@pytest.mark.parametrize("section", SECTIONS)
def test_manual_range_cannot_bridge_a_gap_between_authorized_periods(monkeypatch, section):
    client, _ = date_client(monkeypatch, grants_for(**{section: ("last-month", "today")}))
    response = client.get(f"/v2/live-tour/collections/{section}",
                          params={"date_from": "2026-08-31", "date_to": TODAY})
    assert response.status_code == 403


@pytest.mark.parametrize("section,path,params", ROUTES)
def test_forged_named_preset_cannot_expand_its_authorized_range(monkeypatch, section, path, params):
    client, _ = date_client(monkeypatch, grants_for(**{section: "today"}))
    response = client.get(path, params={**params, "preset": "today", "date_from": "2020-01-01", "date_to": "2030-01-01"})
    assert response.status_code in (400, 403), response.text


@pytest.mark.parametrize("section", SECTIONS)
@pytest.mark.parametrize("query", [
    {"preset": "custom"}, {"preset": "custom", "date_from": TODAY},
    {"preset": "custom", "date_to": TODAY},
    {"preset": "custom", "date_from": "2026-02-30", "date_to": TODAY},
    {"preset": "custom", "date_from": TODAY, "date_to": YESTERDAY},
    {"preset": "unknown"},
])
def test_invalid_custom_or_unknown_ranges_are_rejected(monkeypatch, section, query):
    client, _ = date_client(monkeypatch, grants_for(**{section: PRESETS}))
    response = client.get(f"/v2/live-tour/collections/{section}", params=query)
    assert response.status_code in (400, 422), response.text


@pytest.mark.parametrize("section", SECTIONS)
def test_all_grant_does_not_silently_grant_an_explicit_custom_button(monkeypatch, section):
    client, _ = date_client(monkeypatch, grants_for(**{section: "all"}))
    path = f"/v2/live-tour/collections/{section}"
    assert client.get(path).status_code == 200
    assert client.get(path, params={"date_from": TODAY, "date_to": TODAY}).status_code == 200
    assert client.get(path, params={"preset": "custom", "date_from": TODAY, "date_to": TODAY}).status_code == 403


@pytest.mark.parametrize("section", SECTIONS)
@pytest.mark.parametrize("missing", ["live_tour_view", "section_parent"])
def test_date_grants_never_replace_view_or_section_read_permissions(monkeypatch, section, missing):
    parents = {"pending": "live_tour_pending_view", "invoices": "live_tour_paid_invoice_view",
               "reports": "live_tour_reports_view", "history": "live_tour_history_view"}
    grants = grants_for(**{section: "all"})
    grants.discard(parents[section] if missing == "section_parent" else missing)
    if section == "history" and missing == "section_parent":
        grants.discard("live_tour_backup")
    client, _ = date_client(monkeypatch, grants)
    assert client.get(f"/v2/live-tour/collections/{section}", params={"preset": "all"}).status_code == 403


def test_pending_date_grants_also_require_invoice_parent(monkeypatch):
    client, _ = date_client(monkeypatch, grants_for(pending="today") - {"live_tour_invoice_view"})
    assert client.get("/v2/live-tour/collections/pending", params={"preset": "today"}).status_code == 403


@pytest.mark.parametrize("role", ["admin", "quanly", "letan", "leader", "nhanvien"])
@pytest.mark.parametrize("bulk", [False, True])
def test_routes_obey_effective_date_flags_instead_of_role_names(monkeypatch, role, bulk):
    client, _ = date_client(monkeypatch, BASE_GRANTS, role=role, bulk=bulk)
    assert client.get("/v2/live-tour/collections/invoices", params={"preset": "today"}).status_code == 403
    data = client.get("/v2/live-tour").json()
    assert data["state"]["invoices"] == []


@pytest.mark.parametrize("preset", [None, "custom"])
@pytest.mark.parametrize("view", ["full", "board"])
def test_unscoped_snapshots_reveal_no_rows_without_finite_or_all_grants(monkeypatch, preset, view):
    grants = BASE_GRANTS if preset is None else grants_for(**{section: preset for section in SECTIONS})
    client, shared = date_client(monkeypatch, grants)
    before = deepcopy(shared)
    response = client.get("/v2/live-tour", params={"view": view})
    assert response.status_code == 200, response.text
    data = response.json()
    for key in ("pending", "pending_payments", "report_rows", "audit", "history", "backups", "break_events"):
        assert data[key] == [], key
    for key in ("pending", "invoices", "reports", "audit", "backups", "break_events"):
        assert data["state"][key] == [], key
    if view == "board":
        assert data["pending_count"] == 0
    assert data["records"]  # Operational board visibility is unchanged.
    assert shared == before  # Date filtering must never rewrite the ledger.


@pytest.mark.parametrize("bulk", [False, True])
def test_full_snapshot_redacts_each_section_using_its_own_union(monkeypatch, bulk):
    client, _ = date_client(monkeypatch, grants_for(pending="yesterday", invoices="today", reports="month", history="last-month"), bulk=bulk)
    data = client.get("/v2/live-tour").json()
    assert row_ids(data["pending"]) == {f"p-{YESTERDAY}"}
    assert row_ids(data["state"]["invoices"]) == {f"i-{TODAY}"}
    assert row_ids(data["report_rows"]) == {f"r-{day}" for day in DAYS if day.startswith("2026-09-")}
    assert row_ids(data["audit"]) == {f"a-{day}" for day in DAYS if day.startswith("2026-08-")}
    assert row_ids(data["backups"]) == {f"backup-{day}" for day in DAYS if day.startswith("2026-08-")}


def test_board_pending_count_matches_date_visible_pending(monkeypatch):
    client, _ = date_client(monkeypatch, grants_for(pending=("today", "yesterday")))
    data = client.get("/v2/live-tour", params={"view": "board"}).json()
    assert data["pending_count"] == len(data["pending"]) == 2
    assert row_ids(data["pending"]) == {f"p-{TODAY}", f"p-{YESTERDAY}"}


def test_customer_history_redacts_every_protected_ledger_and_recomputes_totals(monkeypatch):
    client, _ = date_client(monkeypatch, grants_for(pending="yesterday", invoices="today", reports="last-month"))
    response = client.get("/v2/live-tour/customers/c1/history")
    assert response.status_code == 200, response.text
    data = response.json()
    assert row_ids(data["pending"]) == {f"p-{YESTERDAY}"}
    assert row_ids(data["invoices"]) == {f"i-{TODAY}"}
    assert row_ids(data["reports"]) == {f"r-{day}" for day in DAYS if day.startswith("2026-08-")}
    assert {row["invoice_id"] for row in data["services"]} == {f"i-{TODAY}"}
    assert data["summary"]["invoice_count"] == data["summary"]["pending_count"] == 1
    assert data["summary"]["total_revenue"] == 110


def test_custom_only_does_not_unbound_customer_history(monkeypatch):
    client, _ = date_client(monkeypatch, grants_for(pending="custom", invoices="custom", reports="custom"))
    response = client.get("/v2/live-tour/customers/c1/history")
    assert response.status_code == 200, response.text
    assert response.json()["invoices"] == response.json()["reports"] == response.json()["pending"] == []


def test_hidden_pending_still_reserves_combo_inventory_in_customer_lists(monkeypatch):
    state, _, _, customer, owned, pending = combo_pending()
    for key in ("effective_at", "created_at", "booked_at"):
        pending[key] = "2026-08-01T15:00:00+07:00"
    pending["business_date"] = "2026-08-01"
    client, shared = date_client(monkeypatch, grants_for(pending="today"), state)
    before = deepcopy(shared)
    for path in ("/v2/live-tour", "/v2/live-tour/collections/customers"):
        response = client.get(path)
        assert response.status_code == 200, response.text
        data = response.json().get("data", response.json())
        current = next(row for row in data["customers"] if row["id"] == customer["id"])
        purchase = next(row for row in current["combo_purchases"] if row["id"] == owned["id"])
        assert purchase["booking_reserved"] == 1
        assert purchase["booking_remaining"] == 2
    assert shared == before


def test_history_backup_only_reader_still_gets_only_authorized_backup_dates(monkeypatch):
    grants = {"live_tour_view", "live_tour_backup", date_feature("history", "today")}
    client, _ = date_client(monkeypatch, grants)
    response = client.get("/v2/live-tour/collections/history", params={"preset": "today"})
    assert response.status_code == 200, response.text
    data = response.json()["data"]
    assert row_ids(data["backups"]) == {f"backup-{TODAY}"}
    assert data["audit"] == data["invoice_changes"] == data["pending_changes"] == []


def test_nested_history_and_backup_bill_numbers_do_not_bypass_section_dates(monkeypatch):
    state = dated_state()
    old_invoice = deepcopy(next(row for row in state["invoices"] if row["id"] == f"i-{YESTERDAY}"))
    new_invoice = deepcopy(next(row for row in state["invoices"] if row["id"] == f"i-{TODAY}"))
    old_pending = deepcopy(next(row for row in state["pending"] if row["id"] == f"p-{YESTERDAY}"))
    new_pending = deepcopy(next(row for row in state["pending"] if row["id"] == f"p-{TODAY}"))
    state["invoice_changes"] = [{"id": "invoice-change", "created_at": NOW.isoformat(), "before": old_invoice, "after": new_invoice,
                                   "reports_before": [state["reports"][5]], "reports_after": [state["reports"][6]]}]
    state["pending_changes"] = [{"id": "pending-change", "created_at": NOW.isoformat(), "before": old_pending, "after": new_pending}]
    state["backups"] = [{"id": "backup-today", "created_at": NOW.isoformat(),
                           "snapshot": {"invoices": [old_invoice, new_invoice], "pending": [old_pending, new_pending]}}]
    client, shared = date_client(monkeypatch, grants_for(history="today", invoices="today", pending="today", reports="today"), state)
    before = deepcopy(shared)
    for path, query in (("/v2/live-tour", {}), ("/v2/live-tour/collections/history", {"preset": "today"})):
        response = client.get(path, params=query)
        assert response.status_code == 200, response.text
        data = response.json().get("data", response.json())
        serialized = json.dumps({key: data[key] for key in ("invoice_changes", "pending_changes", "backups")})
        assert f"I-{YESTERDAY}" not in serialized and f"P-{YESTERDAY}" not in serialized
        assert data["backups"][0]["bill_numbers"] == sorted([f"I-{TODAY}", f"P-{TODAY}"])
        assert "snapshot" not in data["backups"][0]
    assert shared == before


@pytest.mark.parametrize("response_view", ["receipt", "full", "board"])
def test_payment_still_commits_without_date_reads_and_replay_uses_current_grants(monkeypatch, response_view):
    grants = set(BASE_GRANTS)
    client, shared = date_client(monkeypatch, grants, payable_state())
    payload = {"employee_id": "e1", "payment_method": "TIỀN MẶT", "tip": 20}
    response = post(client, shared, "checkout", payload, idempotency_key="date-payment-once", response_view=response_view)
    assert response.status_code == 200, response.text
    assert len(shared["state"]["invoices"]) == 1
    invoice = deepcopy(shared["state"]["invoices"][0])
    assert invoice["total"] == 120
    assert not response.json()["result"].get("invoice")
    before = deepcopy(shared)
    grants.add(date_feature("invoices", "today"))
    readable = post(client, shared, "checkout", payload, idempotency_key="date-payment-once", response_view=response_view)
    assert readable.status_code == 200 and readable.json()["duplicate"] is True
    assert readable.json()["result"]["invoice"]["id"] == invoice["id"]
    grants.remove(date_feature("invoices", "today"))
    hidden = post(client, shared, "checkout", payload, idempotency_key="date-payment-once", response_view=response_view)
    assert hidden.status_code == 200 and hidden.json()["duplicate"] is True
    assert not hidden.json()["result"].get("invoice")
    assert shared == before


def test_date_grants_do_not_authorize_checkout_without_payment_permission(monkeypatch):
    grants = grants_for(**{section: PRESETS for section in SECTIONS}) - {"live_tour_payment"}
    client, shared = date_client(monkeypatch, grants, payable_state())
    before = deepcopy(shared)
    response = post(client, shared, "checkout", {"employee_id": "e1", "payment_method": "TIỀN MẶT"})
    assert response.status_code == 403
    assert shared == before


@pytest.mark.parametrize("section,path,params", ROUTES)
def test_direct_route_success_with_today_scope_passes_only_authorized_rows(monkeypatch, section, path, params):
    client, _ = date_client(monkeypatch, grants_for(**{item: "today" for item in SECTIONS}))
    calls = []

    def board_rows(_conn, *, date_from=None, date_to=None, **_kwargs):
        calls.append((date_from, date_to))
        assert str(date_from) == str(date_to) == TODAY
        return [{"id": 1, "changed_at": NOW.isoformat(), "employee_name": "An",
                 "before": {}, "after": {}, "changed_columns": []}]

    monkeypatch.setattr(live, "_board_history_rows", board_rows)
    response = client.get(path, params={**params, "preset": "today"})
    assert response.status_code == 200, response.text
    if "board-history" in path:
        assert calls == [(date.fromisoformat(TODAY), date.fromisoformat(TODAY))]
        if not path.endswith("xlsx"):
            assert response.json()["count"] == 1
    elif "/collections/" in path:
        assert row_ids(section_rows(response.json()["data"], section)) == {f"{prefix(section)}-{TODAY}"}
    elif path == "/v2/live-tour/reports":
        data = response.json()
        rows = data["rows"] if params.get("tab") else data["reports"]
        assert row_ids(rows) == {f"r-{TODAY}"}
    elif path.endswith("xlsx"):
        workbook = load_workbook(BytesIO(response.content))
        values = json.dumps([list(sheet.values) for sheet in workbook], default=str, ensure_ascii=False)
        assert "I-2026-09-04" not in values
        assert "P-2026-09-04" not in values


@pytest.mark.parametrize("suffix,module,function", [
    ("pdf", "vera_customer_count_pdf", "customer_count_pdf"),
    ("png", "vera_customer_count_png", "customer_count_png"),
])
def test_customer_count_exports_use_authorized_rows_before_counting(monkeypatch, suffix, module, function):
    from importlib import import_module
    captured = []
    monkeypatch.setattr(import_module(module), function,
                        lambda summary, scope, **_kw: captured.append((deepcopy(summary), deepcopy(scope))) or b"rendered")
    client, _ = date_client(monkeypatch, grants_for(reports="today", invoices="all"))
    response = client.get(f"/v2/live-tour/customer-count.{suffix}", params={"preset": "today"})
    assert response.status_code == 200, response.text
    summary, scope = captured[0]
    assert summary["total"] == 1 and summary["daily"] == [(TODAY, 1)]
    assert scope["date_from"] == scope["date_to"] == TODAY


@pytest.mark.parametrize("path,params", [
    ("/v2/live-tour/reports", {"tab": "revenue"}),
    ("/v2/live-tour/customer-count.pdf", {}),
    ("/v2/live-tour/customer-count.png", {}),
    ("/v2/live-tour/export.xlsx", {"kind": "reports"}),
])
def test_selected_date_deeplinks_cannot_bypass_report_date_authorization(monkeypatch, path, params):
    client, _ = date_client(monkeypatch, grants_for(reports="today"))
    assert client.get(path, params={**params, "date": YESTERDAY}).status_code == 403
    assert client.get(path, params={**params, "date": TODAY}).status_code == 200


@pytest.mark.parametrize("path,params", [
    ("/v2/live-tour/reports", {"tab": "revenue"}),
    ("/v2/live-tour/customer-count.pdf", {}),
    ("/v2/live-tour/customer-count.png", {}),
    ("/v2/live-tour/export.xlsx", {"kind": "reports"}),
])
def test_selected_date_cannot_expand_a_named_preset(monkeypatch, path, params):
    client, _ = date_client(monkeypatch, grants_for(reports="today"))
    response = client.get(path, params={**params, "preset": "today", "date": YESTERDAY})
    assert response.status_code in (400, 403), response.text


@pytest.mark.parametrize("section", SECTIONS)
def test_filtered_collection_counts_and_pages_do_not_include_unauthorized_rows(monkeypatch, section):
    grants = grants_for(**{section: ("today", "yesterday")})
    client, _ = date_client(monkeypatch, grants)
    query = {"date_from": YESTERDAY, "date_to": TODAY, "page_size": 1}
    first = client.get(f"/v2/live-tour/collections/{section}", params=query)
    second = client.get(f"/v2/live-tour/collections/{section}", params={**query, "page": 2})
    assert first.status_code == second.status_code == 200
    rows = section_rows(first.json()["data"], section) + section_rows(second.json()["data"], section)
    assert row_ids(rows) == {f"{prefix(section)}-{TODAY}", f"{prefix(section)}-{YESTERDAY}"}
    assert first.json()["pages"] == second.json()["pages"] == 2
    # History's aggregate total covers audit, break and backup collections.
    expected_total = 6 if section == "history" else 2
    assert first.json()["total"] == second.json()["total"] == expected_total
    if section == "reports":
        assert first.json()["data"]["report_totals"]["invoiceCount"] == 2
        assert first.json()["data"]["report_totals"]["totalRevenue"] == 220


def test_report_screen_totals_and_invoice_tab_use_independent_section_scopes(monkeypatch):
    client, _ = date_client(monkeypatch, grants_for(reports="month", invoices="yesterday"))
    response = client.get("/v2/live-tour/reports", params={"tab": "invoices", "preset": "month"})
    assert response.status_code == 200, response.text
    assert response.json()["total"] == 1
    assert row_ids(response.json()["rows"]) == {f"i-{YESTERDAY}"}
    reports = client.get("/v2/live-tour/reports", params={"tab": "revenue", "preset": "month", "page_size": 1})
    assert reports.status_code == 200, reports.text
    assert reports.json()["total"] == 5
    assert reports.json()["summary"]["totalRevenue"] == 550
    assert all(row["id"] == f"i-{YESTERDAY}" for row in reports.json()["invoices"])


def test_customer_detail_workbook_requires_a_common_authorized_range(monkeypatch):
    client, _ = date_client(monkeypatch, grants_for(pending=("today", "yesterday"), invoices="month", reports="week"))
    params = {"kind": "customer_detail", "customer_id": "c1"}
    assert client.get("/v2/live-tour/export.xlsx", params=params).status_code == 403
    assert client.get("/v2/live-tour/export.xlsx", params={**params, "date_from": "2026-09-01", "date_to": TODAY}).status_code == 403
    response = client.get("/v2/live-tour/export.xlsx", params={**params, "date_from": TODAY, "date_to": TODAY})
    assert response.status_code == 200, response.text
    book = load_workbook(BytesIO(response.content))
    invoices = json.dumps(list(book["Hoa_don"].values), default=str)
    pending = json.dumps(list(book["Cho_thanh_toan"].values), default=str)
    assert f"I-{TODAY}" in invoices and f"I-{YESTERDAY}" not in invoices
    assert f"p-{TODAY}" in pending and f"p-{YESTERDAY}" not in pending
    assert book["Hoa_don"].max_row == book["Cho_thanh_toan"].max_row == 2


@pytest.mark.parametrize("section", ["pending", "invoices", "reports"])
def test_preset_uses_vietnam_midnight_and_effective_date_not_counter_day(monkeypatch, section):
    state = dated_state()
    key = section
    source = deepcopy(state[key][0])
    source.update(id="at-vietnam-midnight", business_date="2026-09-04", created_at="2026-09-04T15:00:00+07:00",
                  effective_at="2026-09-04T17:00:00Z")
    previous = deepcopy(source)
    previous.update(id="before-vietnam-midnight", effective_at="2026-09-04T16:59:59Z")
    state[key] = [previous, source]
    utc_now = datetime.fromisoformat("2026-09-04T17:05:00+00:00")
    client, _ = date_client(monkeypatch, grants_for(**{section: "today"}), state, now=utc_now)
    response = client.get(f"/v2/live-tour/collections/{section}", params={"preset": "today"})
    assert response.status_code == 200, response.text
    assert row_ids(section_rows(response.json()["data"], section)) == {"at-vietnam-midnight"}


def test_standalone_reports_do_not_gain_an_unrelated_live_board_dependency(monkeypatch):
    grants = grants_for(reports="today", invoices="today") - {"live_tour_view"}
    client, _ = date_client(monkeypatch, grants)
    response = client.get("/v2/live-tour/reports", params={"tab": "revenue", "preset": "today"})
    assert response.status_code == 200, response.text
    assert row_ids(response.json()["rows"]) == {f"r-{TODAY}"}


@pytest.mark.parametrize("section", SECTIONS)
def test_custom_grant_requires_explicit_custom_request_for_a_manual_range(monkeypatch, section):
    client, _ = date_client(monkeypatch, grants_for(**{section: "custom"}))
    query = {"date_from": "2026-08-01", "date_to": "2026-10-01"}
    path = f"/v2/live-tour/collections/{section}"
    assert client.get(path, params=query).status_code == 403
    accepted = client.get(path, params={**query, "preset": "custom"})
    assert accepted.status_code == 200, accepted.text
    assert len(section_rows(accepted.json()["data"], section)) == len(DAYS)


def test_hidden_historical_invoice_numbers_cannot_be_used_as_a_count_probe(monkeypatch):
    state = dated_state()
    state["invoice_changes"] = [{
        "id": "changed-invoice", "at": NOW.isoformat(),
        "before": deepcopy(next(row for row in state["invoices"] if row["id"] == f"i-{YESTERDAY}")),
        "after": deepcopy(next(row for row in state["invoices"] if row["id"] == f"i-{TODAY}")),
    }]
    client, _ = date_client(monkeypatch, grants_for(history="today", invoices="today"), state)
    response = client.get("/v2/live-tour/collections/history",
                          params={"preset": "today", "bill_no": f"I-{YESTERDAY}"})
    assert response.status_code == 200, response.text
    assert response.json()["total"] == 0
    assert response.json()["data"]["invoice_changes"] == []


@pytest.mark.parametrize("section", SECTIONS)
def test_explicit_month_can_be_narrowed_but_does_not_ignore_invalid_bound(monkeypatch, section):
    client, _ = date_client(monkeypatch, grants_for(**{section: "month"}))
    path = f"/v2/live-tour/collections/{section}"
    allowed = client.get(path, params={"preset": "month", "date_from": TODAY, "date_to": TODAY})
    assert allowed.status_code == 200, allowed.text
    assert row_ids(section_rows(allowed.json()["data"], section)) == {f"{prefix(section)}-{TODAY}"}
    assert client.get(path, params={"preset": "month", "date_from": "2026-08-31"}).status_code == 403


@pytest.mark.parametrize("section,parents,kind", [
    ("pending", {"live_tour_pending_view", "live_tour_invoice_view"}, "pending"),
    ("invoices", {"live_tour_paid_invoice_view"}, "paid"),
    ("reports", {"live_tour_reports_view"}, "reports"),
    ("reports", {"live_tour_reports_view"}, "revenue"),
    ("reports", {"live_tour_reports_view"}, "employee"),
    ("reports", {"live_tour_reports_view"}, "tip"),
    ("history", {"live_tour_history_view"}, "history"),
    ("history", {"live_tour_history_view"}, "breaks"),
])
def test_direct_exports_keep_their_existing_section_only_parent_access(monkeypatch, section, parents, kind):
    grants = parents | {"live_tour_export", date_feature(section, "today")}
    client, _ = date_client(monkeypatch, grants)
    response = client.get("/v2/live-tour/export.xlsx", params={"kind": kind, "preset": "today"})
    assert response.status_code == 200, response.text
    workbook = load_workbook(BytesIO(response.content))
    values = json.dumps([list(sheet.values) for sheet in workbook], default=str, ensure_ascii=False)
    if kind in {"paid", "revenue", "reports", "employee", "tip"}:
        assert f"I-{YESTERDAY}" not in values
    if kind in {"paid", "revenue"}:
        assert f"I-{TODAY}" in values  # Revenue source remains independent of receipt visibility.


@pytest.mark.parametrize("suffix", ["pdf", "png"])
def test_customer_count_exports_keep_standalone_report_access(monkeypatch, suffix):
    grants = {"live_tour_reports_view", "live_tour_export", date_feature("reports", "today")}
    client, _ = date_client(monkeypatch, grants)
    assert client.get(f"/v2/live-tour/customer-count.{suffix}", params={"preset": "today"}).status_code == 200


@pytest.mark.parametrize("export", [False, True])
def test_performance_report_and_export_require_report_date_scope_even_for_admin(monkeypatch, export):
    grants = set(BASE_GRANTS)
    client, _ = date_client(monkeypatch, grants, role="admin")
    path = "/v2/live-tour/export.xlsx" if export else "/v2/live-tour/reports"
    params = {"kind" if export else "tab": "performance", "preset": "today"}
    assert client.get(path, params=params).status_code == 403
    grants.add(date_feature("reports", "today"))
    accepted = client.get(path, params=params)
    assert accepted.status_code == 200, accepted.text


def test_legacy_financial_audit_snapshots_cannot_leak_top_level_invoice_fields(monkeypatch):
    state = dated_state()
    old_invoice = deepcopy(next(row for row in state["invoices"] if row["id"] == f"i-{YESTERDAY}"))
    state["audit"] = [{"id": "legacy-financial-event", "at": NOW.isoformat(), "action": "paid_invoice_update",
                       "actor": "tester", "bill_no": f"I-{YESTERDAY}", "total": 8675309,
                       "before": old_invoice, "after": old_invoice, "payload": {"invoice": old_invoice},
                       "detail": {"invoice": old_invoice}}]
    client, _ = date_client(monkeypatch, grants_for(history="today", invoices="today"), state)
    for path, params in (("/v2/live-tour", {}), ("/v2/live-tour/collections/history", {"preset": "today"})):
        response = client.get(path, params=params)
        assert response.status_code == 200, response.text
        data = response.json().get("data", response.json())
        assert row_ids(data["audit"]) == {"legacy-financial-event"}
        serialized = json.dumps(data["audit"])
        assert f"I-{YESTERDAY}" not in serialized
        assert "8675309" not in serialized


@pytest.mark.parametrize("kind", ["history", "breaks"])
def test_history_export_uses_vietnam_calendar_not_operational_rollover(monkeypatch, kind):
    state = dated_state()
    key = "audit" if kind == "history" else "break_events"
    state[key] = [{"id": "early-vietnam-morning", "at": "2026-09-04T18:00:00Z",
                   "created_at": "2026-09-04T18:00:00Z", "business_date": YESTERDAY,
                   "actor": "tester", "action": "set_vip", "event_type": "start_break",
                   "employee_name": "Early calendar row", "detail": {"marker": "calendar-today"}}]
    client, _ = date_client(monkeypatch, grants_for(history="today"), state)
    response = client.get("/v2/live-tour/export.xlsx", params={"kind": kind, "preset": "today"})
    assert response.status_code == 200, response.text
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.active.max_row == 2
    assert str(workbook.active.cell(2, 1).value).replace("/", "-") == "05-09-2026 01:00:00"


@pytest.mark.parametrize("action,key,payload", [
    ("start_break", "break_event", {"employee_id": "e1"}),
    ("backup", "backup", {"name": "Synthetic backup"}),
])
@pytest.mark.parametrize("revoke_parent", [False, True])
def test_historical_action_results_recheck_date_and_read_grants_on_replay(monkeypatch, action, key, payload, revoke_parent):
    state = state_with(employee("e1", "An"))
    grants = grants_for(history="today")
    client, shared = date_client(monkeypatch, grants, state)
    body = {"action": action, "payload": payload, "expected_revision": shared['revision'],
            "idempotency_key": "history-result-replay", "response_view": "receipt"}
    first = client.post("/v2/live-tour/action", json=body)
    assert first.status_code == 200, first.text
    original = first.json()['result'][key]
    assert 'id' in original and len(original) > 1
    if revoke_parent:
        # Revoking backup would block the action itself, so keep its mutation
        # authority and revoke only history read for the break-event scenario.
        if key == 'break_event':
            grants.remove('live_tour_history_view')
        else:
            grants.remove(date_feature('history', 'today'))
    else:
        grants.remove(date_feature('history', 'today'))
    replay = client.post("/v2/live-tour/action", json=body)
    assert replay.status_code == 200, replay.text
    assert replay.json()['duplicate'] is True
    assert replay.json()['result'][key] == {'id': original['id']}
    assert len(shared['state']['break_events' if key == 'break_event' else 'backups']) == 1


@pytest.mark.parametrize('can_export', [False, True])
def test_detail_capabilities_publish_current_export_permission(monkeypatch, can_export):
    grants = grants_for(invoices='today', reports='today', pending='today', history='today')
    if not can_export:
        grants.remove('live_tour_export')
    client, _ = date_client(monkeypatch, grants)
    monkeypatch.setattr(live, '_board_history_rows', lambda *_args, **_kwargs: [])
    for path, params in [('/v2/live-tour/customers/c1/history', {}),
                         ('/v2/live-tour/board-history', {'preset': 'today'})]:
        response = client.get(path, params=params)
        assert response.status_code == 200, response.text
        assert response.json()['capabilities']['export'] is can_export
