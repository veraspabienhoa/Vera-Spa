from types import SimpleNamespace
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import text
import pytest
import vera_web_v2_residence as residence
from test_live_tour_resource_postgres import database


def test_excel_preserves_ids_and_never_executes_formula_or_exposes_banks():
    row = {key: '' for key, _ in residence.COLUMNS}
    row.update(username='Synthetic', full_name='=1+1', cccd_number='001234567890', phone='0900000000',
               bank_name='PRIVATE', bank_account='PRIVATE')
    ws = load_workbook(residence.workbook([row])).active
    headers = [cell.value for cell in ws[1]]
    values = [cell.value for cell in ws[2]]
    assert 'PRIVATE' not in values
    assert all('ngân hàng' not in str(x).lower() for x in headers)
    assert ws.cell(2, headers.index('Họ và tên')+1).data_type == 's'
    assert values[headers.index('Số CCCD')] == '001234567890'
    assert ws.freeze_panes == 'C2'


def test_filter_uses_hr_department_and_accent_insensitive_name():
    data = [dict(username='Như', full_name='Ngọc Như', included=True, department_code='letan'),
            dict(username='A', full_name='Other', included=True, department_code='nhanvien'),
            dict(username='Như 2', full_name='Ngọc Như', included=False, department_code='letan')]
    assert residence.filtered(data, 'letan', 'ngoc nhu') == [data[0]]
    assert residence.display_date('2026-10-14') == '14-10-2026'
    assert residence.display_date('14/10/2026') == '14-10-2026'


def test_roster_membership_profile_refresh_export_and_authorization(database):
    with database.begin() as conn:
        for column in ('role','birth_date','phone','email','address','employment_start_date','work_shift','shift_start_date','rotation_cycle'):
            conn.execute(text(f'ALTER TABLE employees ADD COLUMN {column} text'))
        for column in ('monthly_generated','monthly_leave','annual_leave','stt'):
            conn.execute(text(f'ALTER TABLE employees ADD COLUMN {column} integer'))
        conn.execute(text("ALTER TABLE employees ADD COLUMN payload jsonb DEFAULT '{}'::jsonb, ADD COLUMN updated_at timestamptz"))
        conn.execute(text("INSERT INTO employees(username,full_name,role,phone,payload,bank_name,bank_account) VALUES ('Test','Synthetic','nhanvien','0900123456','{\"Số CCCD\":\"001234567890\",\"Ngày cấp CCCD\":\"14/10/2026\"}', 'PRIVATEBANK','PRIVATEACCOUNT')"))
        conn.execute(text("INSERT INTO vera_app_setting(category,setting_key,value_json) VALUES ('payroll','hr_registry','{\"assignments\":{\"Test\":\"letan\"}}')"))
    ident = SimpleNamespace(role='admin', employee_username='Admin')
    allowed = {'residence_view','residence_manage','residence_export'}
    def require(conn, actor, feature):
        if feature not in allowed: raise HTTPException(403, 'Forbidden')
    app = FastAPI()
    residence.install_residence_routes(app, engine_instance=lambda:database,
        current_identity=lambda:ident, require_feature=require, identity_type=SimpleNamespace)
    client = TestClient(app)
    assert client.get('/v2/residence').json()['records'] == []
    for _ in range(2): assert client.put('/v2/residence', json={'username':'Test'}).status_code == 200
    data = client.get('/v2/residence').json()
    assert len(data['records']) == 1
    assert data['records'][0]['department_code'] == 'letan'
    assert 'PRIVATE' not in str(data) and 'bank_' not in str(data)
    with database.begin() as conn:
        conn.execute(text("UPDATE employees SET full_name='Updated' WHERE username='Test'"))
    assert client.get('/v2/residence').json()['records'][0]['full_name'] == 'Updated'
    exported = client.get('/v2/residence/export.xlsx?department=nhanvien')
    from io import BytesIO
    assert load_workbook(BytesIO(exported.content)).active.max_row == 1
    allowed.remove('residence_manage')
    assert client.put('/v2/residence', json={'username':'Test','included':False}).status_code == 403
    assert client.get('/v2/residence/export?format=pdf').content.startswith(b'%PDF')
    assert client.get('/v2/residence/export?format=bad').status_code == 422
    allowed.remove('residence_export')
    assert client.get('/v2/residence/export.xlsx').status_code == 403
    assert client.get('/v2/residence/export?format=png').status_code == 403
    allowed.add('residence_manage')
    assert client.put('/v2/residence', json={'username':'Test','included':False}).status_code == 200
    assert client.get('/v2/residence').json()['records'] == []
    with database.connect() as conn:
        payload = conn.execute(text("SELECT payload FROM employees WHERE username='Test'")).scalar_one()
        assert payload['Số CCCD'] == '001234567890'
    allowed.clear()
    assert client.get('/v2/residence').status_code == 403


def test_admin_only_defaults_keep_explicit_grants_available():
    from vera_web_v2_permissions import DEFAULT_ROLE_FEATURES, FEATURES, permission_closure
    for role, features in DEFAULT_ROLE_FEATURES.items():
        assert ('residence_view' in features) == (role == 'admin')
    assert 'residence_export' in FEATURES
    assert 'residence_view' in permission_closure({'residence_export'})


def test_residence_pdf_png_all_fields_and_long_values():
    from vera_web_v2_residence_export import render_residence
    from io import BytesIO
    from PIL import Image
    from pypdf import PdfReader
    row = {key: f'VALUE_{index}' for index, (key, _) in enumerate(residence.COLUMNS)}
    row['bank_account'] = 'SECRET_BANK'
    row['address'] = 'Địa chỉ rất dài ' * 150
    pdf = render_residence([row], residence.COLUMNS, 'pdf')
    reader = PdfReader(BytesIO(pdf))
    text = '\n'.join(page.extract_text() for page in reader.pages)
    for index, (key, _) in enumerate(residence.COLUMNS):
        if key != 'address': assert f'VALUE_{index}' in text
    assert 'SECRET_BANK' not in text
    assert reader.metadata.title == 'VERA SPA - Danh sách tạm trú'
    png = render_residence([row], residence.COLUMNS, 'png')
    image = Image.open(BytesIO(png))
    assert image.width > 1000 and image.height > 1000
