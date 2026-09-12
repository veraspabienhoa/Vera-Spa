from datetime import timedelta
from io import BytesIO

import pytest
from openpyxl import load_workbook

import vera_web_v2_live_tour as live
from test_live_tour_backend import NOW, employee, state_with


@pytest.mark.parametrize('pending', [False, True])
def test_booking_user_survives_completion_and_payment_by_other_users(pending):
    state = state_with(employee('e1', 'An'))
    live._apply_action(state, 'booking', {'employee_id': 'e1', 'service': 'Body 90',
        'room': '1.1', 'booking_actor': 'spoofed'}, 'user1', NOW)
    state = live._normalize_state(state, NOW)
    live._apply_action(state, 'start', {'employee_id': 'e1'}, 'operator', NOW)
    live._apply_action(state, 'complete', {'employee_id': 'e1'}, 'operator', NOW + timedelta(minutes=90))
    payload = {'employee_id': 'e1', 'payment_method': 'TIỀN MẶT'}
    if pending:
        item = live._apply_action(state, 'move_pending', {'employee_id': 'e1'}, 'operator', NOW + timedelta(minutes=91))['pending']
        payload = {'pending_id': item['id'], 'payment_method': 'TIỀN MẶT'}
        assert not state['employees'][0]['booking_actor']
    paid_at = NOW + timedelta(days=1)
    payload['ticket_price'] = 100
    invoice = live._apply_action(state, 'checkout', payload, 'user2', paid_at)['invoice']
    assert invoice['entries'][0]['booking_actor'] == 'user1'
    assert invoice['entries'][0]['booked_at'] == live._iso(NOW)
    assert invoice['recorded_at'] == live._iso(paid_at)
    assert invoice['actor'] == 'user2'
    live._apply_action(state, 'paid_invoice_update', {'invoice_id': invoice['id'], 'reason': 'Sửa ghi chú', 'note': 'Đã đối soát'}, 'user3', paid_at)
    saved = state['invoices'][0]
    assert saved['actor'] == 'user2'
    assert saved['recorded_at'] == live._iso(paid_at)
    assert saved['entries'][0]['booking_actor'] == 'user1'


def test_revenue_export_preserves_invoice_totals_and_lists_all_employees():
    state = state_with()
    state['invoices'] = [dict(id='invoice', business_date=NOW.date().isoformat(), created_at=live._iso(NOW),
        bill_no='VERA-1', customer_name='Khách', customer_phone='0123456789', subtotal=300,
        discount=20, tip=50, total=330, actor='cashier', payment_method='TIỀN MẶT',
        entries=[{'employee_name': 'An', 'price': 100}, {'employee_name': 'Bình', 'price': 200}])]
    content, _ = live._excel_bytes(state, 'revenue', NOW)
    sheet = load_workbook(BytesIO(content))['Doanh_thu']
    assert sheet.max_row == 3
    values = dict(zip([cell.value for cell in sheet[1]], [cell.value for cell in sheet[2]]))
    assert values['Tên nhân viên'] == 'An, Bình'
    assert [values[label] for label in ['Tiền dịch vụ', 'Giảm giá', 'Tiền Tip', 'Tổng tiền']] == [300, 20, 50, 330]
    assert values['Điện thoại'] == '0123456789'
    assert values['Người tạo'] == 'cashier'
    assert sheet['A3'].value == 'Tổng cộng'
    assert [sheet[f'{col}3'].value for col in 'FGHI'] == [f'=SUBTOTAL(109,{col}2:{col}2)' for col in 'FGHI']
    assert sheet.auto_filter.ref == 'A1:K2'
    assert sheet.freeze_panes == 'A2'
