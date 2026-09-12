from io import BytesIO

from openpyxl import load_workbook

import vera_web_v2_live_tour as live
from test_live_tour_backend import NOW, state_with


def test_tip_summary_opens_first_and_agrees_with_filtered_detail():
    state = state_with()
    state['reports'] = [
        {'employee_id': 'a', 'employee_name': 'An An', 'business_date': '2026-09-05', 'tip': 100000},
        {'employee_id': 'b', 'employee_name': 'Bảo Trâm', 'business_date': '2026-09-05', 'tip': 50000},
        {'employee_id': 'a', 'employee_name': 'An An', 'business_date': '2026-09-05', 'tip': 200000},
        {'employee_id': 'a', 'employee_name': 'An An', 'business_date': '2026-09-04', 'tip': 900000},
    ]
    bounds = live._parse_export_bounds(date_from='2026-09-05', date_to='2026-09-05')
    content, _ = live._excel_bytes(state, 'tip', NOW, bounds=bounds)
    workbook = load_workbook(BytesIO(content), data_only=True)
    assert workbook.sheetnames == ['Tong_hop_Tip', 'Tip']
    sheet = workbook.active
    assert sheet.title == 'Tong_hop_Tip' and sheet.sheet_view.tabSelected
    assert sheet['A1'].value == 'Ngày 05/09/2026'
    assert list(sheet.iter_rows(min_row=3, values_only=True)) == [(1, 'An An', 300000), (2, 'Bảo Trâm', 50000)]
    assert sheet['C1'].value == sum(row[5] for row in workbook['Tip'].iter_rows(min_row=2, values_only=True)) == 350000
    assert sheet.freeze_panes == 'A3'
    assert sheet.auto_filter.ref == 'A2:C4'
    bounds['employee'] = 'Bảo Trâm'
    filtered, _ = live._excel_bytes(state, 'tip', NOW, bounds=bounds)
    assert load_workbook(BytesIO(filtered), data_only=True).active['C1'].value == 50000


def test_empty_tip_export_has_zero_total_and_keeps_detail_sheet():
    content, _ = live._excel_bytes(state_with(), 'tip', NOW)
    workbook = load_workbook(BytesIO(content), data_only=True)
    assert workbook.active['C1'].value == 0
    assert workbook.active.max_row == 2
    assert workbook['Tip'].max_row == 1
    assert workbook.active.freeze_panes == 'A3'
    assert workbook.active.auto_filter.ref == 'A2:C2'


def test_tip_summary_fits_long_names_and_formatted_amounts():
    name = 'Nguyễn Thị Nhân Viên Có Tên Rất Dài'
    amount = 1234567890123
    state = state_with()
    state['reports'] = [{'employee_name': name, 'tip': amount}]
    content, _ = live._excel_bytes(state, 'tip', NOW)
    sheet = load_workbook(BytesIO(content)).active
    assert sheet.column_dimensions['B'].width >= len(name) + 5
    assert sheet.column_dimensions['C'].width >= len(f'{amount:,.0f} đ') + 5
    assert all(sheet.column_dimensions[col].bestFit for col in 'ABC')
    assert sheet.freeze_panes == 'A3'
    assert sheet.auto_filter.ref == 'A2:C3'
