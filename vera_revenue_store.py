"""PostgreSQL-backed revenue/expense ledger and one-time workbook bootstrap."""
from __future__ import annotations

from datetime import date, datetime, time, timedelta, timezone
from io import BytesIO
from typing import Any, Callable

from openpyxl import load_workbook
import requests
from sqlalchemy import text


TABLE = "vera_revenue_entry"
SCHEMA_VERSION = 3
VN_TZ = timezone(timedelta(hours=7))
WORKBOOK_EXPORT_URL = "https://docs.google.com/spreadsheets/d/{spreadsheet_id}/export"
HEADERS = [
    "Dấu thời gian", "Loại giao dịch", "Số tiền", "Ngày giao dịch",
    "Ghi chú", "Địa chỉ email", "Tháng", "Ngày nhập", "Giờ nhập", "Người nhập",
]


def ensure_schema(conn) -> None:
    conn.execute(text(f"""
        CREATE TABLE IF NOT EXISTS {TABLE} (
            id bigserial PRIMARY KEY,
            transaction_type text NOT NULL CHECK (transaction_type IN ('Thu','Chi')),
            amount numeric(18,2) NOT NULL CHECK (amount >= 0),
            transaction_date date NULL,
            note text NOT NULL DEFAULT '',
            entered_at timestamptz NOT NULL DEFAULT NOW(),
            entered_by text NOT NULL DEFAULT '',
            entered_by_name text NOT NULL DEFAULT '',
            source_email text NOT NULL DEFAULT '',
            source_name text NOT NULL DEFAULT 'web_v2',
            source_row integer NULL,
            source_month date NULL,
            created_at timestamptz NOT NULL DEFAULT NOW(),
            CONSTRAINT vera_revenue_entry_source_row_uq UNIQUE(source_name, source_row)
        )
    """))
    conn.execute(text(f"CREATE INDEX IF NOT EXISTS idx_{TABLE}_transaction_date ON {TABLE}(transaction_date DESC, id DESC)"))
    conn.execute(text(f"CREATE INDEX IF NOT EXISTS idx_{TABLE}_entered_at ON {TABLE}(entered_at DESC, id DESC)"))
    conn.execute(text(f"ALTER TABLE {TABLE} ADD COLUMN IF NOT EXISTS is_deleted boolean NOT NULL DEFAULT false"))
    conn.execute(text(f"ALTER TABLE {TABLE} ADD COLUMN IF NOT EXISTS edit_revision integer NOT NULL DEFAULT 0"))
    conn.execute(text("""
        CREATE TABLE IF NOT EXISTS vera_revenue_entry_audit (
            id bigserial PRIMARY KEY,
            revenue_entry_id bigint NOT NULL,
            action text NOT NULL CHECK(action IN ('update','delete')),
            before_payload jsonb NOT NULL,
            after_payload jsonb,
            actor text NOT NULL DEFAULT '',
            audited_at timestamptz NOT NULL DEFAULT NOW()
        )
    """))
    conn.execute(text("CREATE INDEX IF NOT EXISTS idx_vera_revenue_entry_audit_time ON vera_revenue_entry_audit(audited_at DESC, id DESC)"))
    conn.execute(text("""
        INSERT INTO vera_schema_version(component,version,updated_at)
        VALUES('revenue_ledger',:version,NOW())
        ON CONFLICT(component) DO UPDATE SET
          version=GREATEST(vera_schema_version.version,EXCLUDED.version), updated_at=NOW()
    """), {"version": SCHEMA_VERSION})


def _as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = str(value or "").strip()
    if not raw:
        return None
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%m/%d/%Y"):
        try:
            return datetime.strptime(raw[:10], fmt).date()
        except ValueError:
            pass
    return None


def _as_entered_at(value: Any) -> datetime:
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime.combine(value, time.min)
    else:
        raw = str(value or "").strip()
        parsed = None
        for fmt in ("%d/%m/%Y %H:%M:%S", "%d-%m-%Y %H:%M:%S", "%d-%m-%Y %H:%M", "%m/%d/%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S"):
            try:
                parsed = datetime.strptime(raw.split(".", 1)[0], fmt)
                break
            except ValueError:
                pass
        if parsed is None:
            raise ValueError(f"Thời điểm nhập không hợp lệ: {raw[:30]}")
    return parsed.replace(tzinfo=VN_TZ) if parsed.tzinfo is None else parsed.astimezone(VN_TZ)


def _as_money(value: Any) -> float:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return round(float(value), 2)
    raw = "".join(ch for ch in str(value or "") if ch.isdigit() or ch in ",.-")
    if not raw:
        return 0.0
    if "," in raw and "." in raw:
        decimal = "," if raw.rfind(",") > raw.rfind(".") else "."
        thousands = "." if decimal == "," else ","
        raw = raw.replace(thousands, "").replace(decimal, ".")
    elif raw.count(",") == 1 and len(raw.rsplit(",", 1)[1]) <= 2:
        raw = raw.replace(",", ".")
    elif raw.count(".") == 1 and len(raw.rsplit(".", 1)[1]) <= 2:
        pass
    else:
        raw = raw.replace(",", "").replace(".", "")
    return round(float(raw), 2)


def parse_workbook(content: bytes, worksheet_name: str = "Input") -> list[dict[str, Any]]:
    workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    if worksheet_name not in workbook.sheetnames:
        raise RuntimeError(f"Workbook không có sheet '{worksheet_name}'.")
    worksheet = workbook[worksheet_name]
    iterator = worksheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(iterator, ())]
    required = ["Dấu thời gian", "Loại giao dịch", "Số tiền", "Ngày giao dịch", "Ghi chú", "Địa chỉ email"]
    missing = [name for name in required if name not in headers]
    if missing:
        raise RuntimeError("Sheet Input thiếu cột: " + ", ".join(missing))
    indexes = {name: headers.index(name) for name in headers if name}
    rows: list[dict[str, Any]] = []
    for source_row, raw_tuple in enumerate(iterator, 2):
        raw = list(raw_tuple)
        if not any(value is not None and str(value).strip() for value in raw):
            continue
        tx_type = str(raw[indexes["Loại giao dịch"]] or "").strip().title()
        if tx_type not in {"Thu", "Chi"}:
            raise RuntimeError(f"Dòng {source_row}: Loại giao dịch phải là Thu hoặc Chi.")
        entered_at = _as_entered_at(raw[indexes["Dấu thời gian"]])
        month_index = indexes.get("Tháng")
        rows.append({
            "transaction_type": tx_type,
            "amount": _as_money(raw[indexes["Số tiền"]]),
            "transaction_date": _as_date(raw[indexes["Ngày giao dịch"]]),
            "note": str(raw[indexes["Ghi chú"]] or "").strip(),
            "entered_at": entered_at,
            "source_email": str(raw[indexes["Địa chỉ email"]] or "").strip(),
            "source_month": _as_date(raw[month_index]) if month_index is not None and month_index < len(raw) else None,
            "source_row": source_row,
        })
    if not rows:
        raise RuntimeError("Sheet Input không có giao dịch hợp lệ.")
    return rows


def _actor_by_email(conn) -> dict[str, tuple[str, str]]:
    rows = conn.execute(text("""
        SELECT lower(btrim(COALESCE(email,''))) AS email_key,
               COALESCE(username,'') AS username, COALESCE(full_name,'') AS full_name
        FROM employees WHERE btrim(COALESCE(email,'')) <> ''
    """)).mappings().all()
    return {
        str(row["email_key"]): (str(row["username"]), str(row["full_name"] or row["username"]))
        for row in rows
    }


def import_rows(conn, rows: list[dict[str, Any]], *, source_name: str) -> dict[str, Any]:
    actors = _actor_by_email(conn)
    params = []
    for row in rows:
        email = str(row.get("source_email") or "").strip()
        username, full_name = actors.get(email.lower(), (email, email))
        params.append({
            **row,
            "source_name": source_name,
            "entered_by": username,
            "entered_by_name": full_name,
        })
    before = int(conn.execute(text(f"SELECT COUNT(*) FROM {TABLE}")).scalar_one())
    conn.execute(text(f"""
        INSERT INTO {TABLE}(
            transaction_type, amount, transaction_date, note, entered_at,
            entered_by, entered_by_name, source_email, source_name, source_row,
            source_month, created_at
        ) VALUES (
            :transaction_type, :amount, :transaction_date, :note, :entered_at,
            :entered_by, :entered_by_name, :source_email, :source_name, :source_row,
            :source_month, NOW()
        )
        ON CONFLICT(source_name, source_row) DO NOTHING
    """), params)
    after = int(conn.execute(text(f"SELECT COUNT(*) FROM {TABLE}")).scalar_one())
    totals = conn.execute(text(f"""
        SELECT COUNT(*) AS row_count,
               COALESCE(SUM(amount) FILTER (WHERE transaction_type='Thu'),0) AS total_income,
               COALESCE(SUM(amount) FILTER (WHERE transaction_type='Chi'),0) AS total_expense
        FROM {TABLE}
    """)).mappings().one()
    return {
        "source_rows": len(rows), "inserted_rows": after - before,
        "row_count": int(totals["row_count"]),
        "total_income": float(totals["total_income"]),
        "total_expense": float(totals["total_expense"]),
    }


def bootstrap_from_google_sheet(
    conn, *, spreadsheet_id: str, worksheet_name: str = "Input",
    http_get: Callable[..., Any] = requests.get,
) -> dict[str, Any]:
    ensure_schema(conn)
    existing = int(conn.execute(text(f"SELECT COUNT(*) FROM {TABLE}")).scalar_one())
    if existing:
        return {"skipped": True, "reason": "ledger_not_empty", "row_count": existing}
    response = http_get(
        WORKBOOK_EXPORT_URL.format(spreadsheet_id=spreadsheet_id),
        params={"format": "xlsx"}, timeout=60,
    )
    response.raise_for_status()
    rows = parse_workbook(bytes(response.content or b""), worksheet_name)
    result = import_rows(conn, rows, source_name=f"google-sheet:{spreadsheet_id}:{worksheet_name}")
    if result["row_count"] != result["source_rows"]:
        raise RuntimeError("Đối soát nhập Thu Chi thất bại: số dòng server khác số dòng nguồn.")
    return {"skipped": False, **result}


def values_from_db(conn) -> list[list[Any]]:
    rows = conn.execute(text(f"""
        SELECT transaction_type, amount, transaction_date, note, entered_at,
               entered_by, source_email, source_month
        FROM {TABLE}
        WHERE is_deleted=false
        ORDER BY COALESCE(transaction_date, (entered_at AT TIME ZONE 'Asia/Ho_Chi_Minh')::date), id
    """)).mappings().all()
    values: list[list[Any]] = [list(HEADERS)]
    for row in rows:
        entered = row["entered_at"].astimezone(VN_TZ)
        tx_date = row["transaction_date"]
        values.append([
            entered.strftime("%d/%m/%Y %H:%M:%S"), row["transaction_type"], float(row["amount"]),
            tx_date.strftime("%d/%m/%Y") if tx_date else "", row["note"], row["source_email"],
            row["source_month"].strftime("%d/%m/%Y") if row["source_month"] else "",
            entered.strftime("%d/%m/%Y"), entered.strftime("%H:%M:%S"), row["entered_by"],
        ])
    return values


def insert_web_entries(conn, *, transaction_date: date, entries: list[tuple[str, float, str]], ident) -> int:
    ensure_schema(conn)
    now = datetime.now(VN_TZ)
    actor = str(getattr(ident, "employee_username", "") or "").strip()
    full_name = str(getattr(ident, "full_name", "") or actor).strip() or actor
    email = str(getattr(ident, "email", "") or "").strip()
    params = [{
        "transaction_type": tx_type, "amount": round(float(amount), 2),
        "transaction_date": transaction_date, "note": str(note or "").strip(),
        "entered_at": now, "entered_by": actor, "entered_by_name": full_name,
        "source_email": email,
    } for tx_type, amount, note in entries]
    conn.execute(text(f"""
        INSERT INTO {TABLE}(
            transaction_type, amount, transaction_date, note, entered_at,
            entered_by, entered_by_name, source_email, source_name, created_at
        ) VALUES (
            :transaction_type, :amount, :transaction_date, :note, :entered_at,
            :entered_by, :entered_by_name, :source_email, 'web_v2', NOW()
        )
    """), params)
    return len(params)


def find_duplicate_web_entries(conn, *, entries: list[tuple[str, float, str]]) -> list[dict[str, Any]]:
    """Return prior rows with the same type, amount and normalized note."""
    ensure_schema(conn)
    duplicates: list[dict[str, Any]] = []
    for transaction_type, amount, note in entries:
        normalized_note = str(note or "").strip()
        row = conn.execute(text(f"""
            SELECT transaction_type, amount, transaction_date, note, entered_at
            FROM {TABLE}
            WHERE transaction_type = :transaction_type
              AND amount = :amount
              AND lower(btrim(note)) = lower(:note)
            ORDER BY entered_at DESC, id DESC
            LIMIT 1
        """), {
            "transaction_type": transaction_type,
            "amount": round(float(amount), 2),
            "note": normalized_note,
        }).mappings().first()
        if not row:
            continue
        entered_at = row["entered_at"]
        if entered_at and entered_at.tzinfo is None:
            entered_at = entered_at.replace(tzinfo=timezone.utc)
        duplicates.append({
            "transaction_type": str(row["transaction_type"]),
            "amount": float(row["amount"]),
            "note": str(row["note"] or ""),
            "transaction_date": row["transaction_date"].isoformat() if row["transaction_date"] else None,
            "entered_at": entered_at.astimezone(VN_TZ).isoformat() if entered_at else None,
        })
    return duplicates


def import_ledger_xlsx(conn, content: bytes, *, mode: str, actor: str) -> dict[str, int]:
    """Import the same seven-column workbook produced by the Revenue export.

    append: insert only rows whose business fields do not already exist.
    replace: soft-delete all active rows, then insert every imported row.
    """
    ensure_schema(conn)
    if mode not in {"append", "replace"}:
        raise ValueError("Chế độ import không hợp lệ.")
    if not content or len(content) > 15 * 1024 * 1024:
        raise ValueError("File Excel trống hoặc vượt quá 15 MB.")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("File Excel không hợp lệ hoặc không đọc được.") from exc
    sheet = workbook["Doanh thu-Chi phí"] if "Doanh thu-Chi phí" in workbook.sheetnames else workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError("File Excel không có dữ liệu.")
    normalized_headers = [str(value or "").strip().lower() for value in rows[0]]
    required = ["ngày", "loại giao dịch", "số tiền", "ghi chú", "ngày nhập", "giờ nhập", "người nhập"]
    if normalized_headers[:7] != required:
        raise ValueError("Excel phải có 7 cột: Ngày, Loại giao dịch, Số tiền, Ghi chú, Ngày nhập, Giờ nhập, Người nhập.")

    parsed: list[dict[str, Any]] = []
    for row_number, values in enumerate(rows[1:], 2):
        values = list(values) + [None] * (7 - len(values))
        if not any(value not in (None, "") for value in values[:7]):
            continue
        transaction_date = _as_date(values[0])
        transaction_type = str(values[1] or "").strip().title()
        if transaction_type not in {"Thu", "Chi"}:
            raise ValueError(f"Dòng {row_number}: Loại giao dịch phải là Thu hoặc Chi.")
        amount = _as_money(values[2])
        if amount < 0:
            raise ValueError(f"Dòng {row_number}: Số tiền không được âm.")
        entered_date = _as_date(values[4]) or transaction_date
        raw_time = values[5]
        if isinstance(raw_time, datetime):
            entered_time = raw_time.time()
        elif isinstance(raw_time, time):
            entered_time = raw_time
        else:
            text_time = str(raw_time or "00:00:00").strip()
            entered_time = None
            for fmt in ("%H:%M:%S", "%H:%M"):
                try:
                    entered_time = datetime.strptime(text_time, fmt).time()
                    break
                except ValueError:
                    pass
            if entered_time is None:
                raise ValueError(f"Dòng {row_number}: Giờ nhập không hợp lệ.")
        if not entered_date:
            raise ValueError(f"Dòng {row_number}: Ngày giao dịch/Ngày nhập không hợp lệ.")
        entered_at = datetime.combine(entered_date, entered_time).replace(tzinfo=VN_TZ)
        parsed.append({
            "transaction_type": transaction_type,
            "amount": round(float(amount), 2),
            "transaction_date": transaction_date,
            "note": str(values[3] or "").strip(),
            "entered_at": entered_at,
            "entered_by_name": str(values[6] or "").strip(),
        })
    if not parsed:
        raise ValueError("File Excel không có dòng dữ liệu hợp lệ.")

    if mode == "replace":
        conn.execute(text(f"""
            INSERT INTO vera_revenue_entry_audit(revenue_entry_id,action,before_payload,after_payload,actor)
            SELECT id, 'delete', to_jsonb(current_row), NULL, :actor
            FROM {TABLE} AS current_row WHERE is_deleted=false
        """), {"actor": actor})
        conn.execute(text(f"UPDATE {TABLE} SET is_deleted=true, edit_revision=edit_revision+1 WHERE is_deleted=false"))
        existing_keys: set[tuple[Any, ...]] = set()
    else:
        existing = conn.execute(text(f"""
            SELECT transaction_type, amount, transaction_date, note, entered_at, entered_by_name, entered_by
            FROM {TABLE} WHERE is_deleted=false
        """)).mappings().all()
        existing_keys = {
            (
                str(row["transaction_type"]), round(float(row["amount"]), 2), row["transaction_date"],
                str(row["note"] or "").strip().casefold(),
                row["entered_at"].astimezone(VN_TZ).replace(microsecond=0) if row["entered_at"] else None,
                str(row["entered_by_name"] or row["entered_by"] or "").strip().casefold(),
            )
            for row in existing
        }

    to_insert = []
    skipped = 0
    for row in parsed:
        key = (
            row["transaction_type"], row["amount"], row["transaction_date"], row["note"].casefold(),
            row["entered_at"].replace(microsecond=0), row["entered_by_name"].casefold(),
        )
        if mode == "append" and key in existing_keys:
            skipped += 1
            continue
        existing_keys.add(key)
        to_insert.append({**row, "actor": actor})
    if to_insert:
        conn.execute(text(f"""
            INSERT INTO {TABLE}(
                transaction_type, amount, transaction_date, note, entered_at,
                entered_by, entered_by_name, source_name, created_at
            ) VALUES (
                :transaction_type, :amount, :transaction_date, :note, :entered_at,
                :actor, :entered_by_name, 'excel_import', NOW()
            )
        """), to_insert)
    return {"read": len(parsed), "inserted": len(to_insert), "skipped": skipped}



def list_entries(conn, *, start_date: date | None = None, end_date: date | None = None) -> list[dict[str, Any]]:
    ensure_schema(conn)
    rows = conn.execute(text(f"""
        SELECT id, transaction_type, amount, transaction_date, note, entered_at,
               entered_by, entered_by_name, source_name, edit_revision
        FROM {TABLE}
        WHERE is_deleted=false
          AND (CAST(:start_date AS date) IS NULL OR COALESCE(transaction_date, (entered_at AT TIME ZONE 'Asia/Ho_Chi_Minh')::date) >= CAST(:start_date AS date))
          AND (CAST(:end_date AS date) IS NULL OR COALESCE(transaction_date, (entered_at AT TIME ZONE 'Asia/Ho_Chi_Minh')::date) <= CAST(:end_date AS date))
        ORDER BY COALESCE(transaction_date, (entered_at AT TIME ZONE 'Asia/Ho_Chi_Minh')::date) DESC, id DESC
    """), {"start_date": start_date, "end_date": end_date}).mappings().all()
    output = []
    for row in rows:
        entered = row["entered_at"]
        if entered and entered.tzinfo is None:
            entered = entered.replace(tzinfo=timezone.utc)
        tx_date = row["transaction_date"] or (entered.astimezone(VN_TZ).date() if entered else None)
        output.append({
            "id": int(row["id"]), "type": str(row["transaction_type"]), "amount": float(row["amount"]),
            "date": tx_date.isoformat() if tx_date else "", "date_label": tx_date.strftime("%d-%m-%Y") if tx_date else "",
            "note": str(row["note"] or ""), "entered_at": entered.astimezone(VN_TZ).isoformat() if entered else "",
            "entered_date": entered.astimezone(VN_TZ).date().isoformat() if entered else "",
            "entered_date_label": entered.astimezone(VN_TZ).strftime("%d-%m-%Y") if entered else "",
            "entered_time": entered.astimezone(VN_TZ).strftime("%H:%M:%S") if entered else "",
            "entered_by": str(row["entered_by_name"] or row["entered_by"] or ""),
            "source": str(row["source_name"] or ""), "revision": int(row["edit_revision"] or 0),
        })
    return output


def _audit_payload(row) -> dict[str, Any]:
    return {key: (value.isoformat() if hasattr(value, "isoformat") else value) for key, value in dict(row).items()}


def update_entry(conn, *, entry_id: int, transaction_type: str, amount: float, transaction_date: date | None, note: str, entered_by_name: str | None, entered_at: datetime | None, actor: str, required_entered_date: date | None = None) -> dict[str, Any]:
    import json
    ensure_schema(conn)
    before = conn.execute(text(f"SELECT * FROM {TABLE} WHERE id=:id AND is_deleted=false FOR UPDATE"), {"id": entry_id}).mappings().first()
    if not before:
        raise KeyError(entry_id)
    original_entered = before["entered_at"]
    if original_entered and original_entered.tzinfo is None:
        original_entered = original_entered.replace(tzinfo=timezone.utc)
    if required_entered_date and (not original_entered or original_entered.astimezone(VN_TZ).date() != required_entered_date):
        raise PermissionError(entry_id)
    conn.execute(text(f"""
        UPDATE {TABLE}
        SET transaction_type=:transaction_type, amount=:amount, transaction_date=:transaction_date,
            note=:note, entered_by_name=COALESCE(:entered_by_name, entered_by_name), entered_at=COALESCE(:entered_at, entered_at), edit_revision=edit_revision+1
        WHERE id=:id AND is_deleted=false
    """), {"id": entry_id, "transaction_type": transaction_type, "amount": round(float(amount), 2),
             "transaction_date": transaction_date, "note": str(note or "").strip(),
             "entered_by_name": None if entered_by_name is None else str(entered_by_name).strip(), "entered_at": entered_at})
    after = conn.execute(text(f"SELECT * FROM {TABLE} WHERE id=:id"), {"id": entry_id}).mappings().one()
    conn.execute(text("""
        INSERT INTO vera_revenue_entry_audit(revenue_entry_id,action,before_payload,after_payload,actor)
        VALUES(:id,'update',CAST(:before AS jsonb),CAST(:after AS jsonb),:actor)
    """), {"id": entry_id, "before": json.dumps(_audit_payload(before), ensure_ascii=False, default=str),
             "after": json.dumps(_audit_payload(after), ensure_ascii=False, default=str), "actor": actor})
    return {"id": entry_id, "revision": int(after["edit_revision"])}


def soft_delete_entry(conn, *, entry_id: int, actor: str, required_entered_date: date | None = None) -> None:
    import json
    ensure_schema(conn)
    before = conn.execute(text(f"SELECT * FROM {TABLE} WHERE id=:id AND is_deleted=false FOR UPDATE"), {"id": entry_id}).mappings().first()
    if not before:
        raise KeyError(entry_id)
    original_entered = before["entered_at"]
    if original_entered and original_entered.tzinfo is None:
        original_entered = original_entered.replace(tzinfo=timezone.utc)
    if required_entered_date and (not original_entered or original_entered.astimezone(VN_TZ).date() != required_entered_date):
        raise PermissionError(entry_id)
    conn.execute(text(f"UPDATE {TABLE} SET is_deleted=true, edit_revision=edit_revision+1 WHERE id=:id"), {"id": entry_id})
    conn.execute(text("""
        INSERT INTO vera_revenue_entry_audit(revenue_entry_id,action,before_payload,after_payload,actor)
        VALUES(:id,'delete',CAST(:before AS jsonb),NULL,:actor)
    """), {"id": entry_id, "before": json.dumps(_audit_payload(before), ensure_ascii=False, default=str), "actor": actor})


def list_audit_entries(conn, *, start_date: date | None = None, end_date: date | None = None) -> list[dict[str, Any]]:
    """Return the immutable edit/delete trail. This is exposed to Admin only."""
    ensure_schema(conn)
    rows = conn.execute(text("""
        SELECT id, revenue_entry_id, action, before_payload, after_payload, actor, audited_at
        FROM vera_revenue_entry_audit
        WHERE (CAST(:start_date AS date) IS NULL OR (audited_at AT TIME ZONE 'Asia/Ho_Chi_Minh')::date >= CAST(:start_date AS date))
          AND (CAST(:end_date AS date) IS NULL OR (audited_at AT TIME ZONE 'Asia/Ho_Chi_Minh')::date <= CAST(:end_date AS date))
        ORDER BY audited_at DESC, id DESC
    """), {"start_date": start_date, "end_date": end_date}).mappings().all()
    output: list[dict[str, Any]] = []
    for row in rows:
        audited_at = row["audited_at"]
        if audited_at and audited_at.tzinfo is None:
            audited_at = audited_at.replace(tzinfo=timezone.utc)
        before = row["before_payload"] if isinstance(row["before_payload"], dict) else {}
        after = row["after_payload"] if isinstance(row["after_payload"], dict) else None
        output.append({
            "id": int(row["id"]),
            "entry_id": int(row["revenue_entry_id"]),
            "action": str(row["action"]),
            "action_label": "Sửa" if row["action"] == "update" else "Xóa",
            "actor": str(row["actor"] or ""),
            "audited_at": audited_at.astimezone(VN_TZ).isoformat() if audited_at else "",
            "audited_at_label": audited_at.astimezone(VN_TZ).strftime("%d-%m-%Y %H:%M:%S") if audited_at else "",
            "before": before,
            "after": after,
        })
    return output


def duplicate_analysis(conn, *, start_date: date | None = None, end_date: date | None = None) -> dict[str, Any]:
    """Group active ledger rows that are exact business duplicates.

    A duplicate must share transaction date, type, amount and normalized note.
    This intentionally ignores the entry timestamp/user so repeated submissions
    from different devices remain visible to the reviewer.
    """
    ensure_schema(conn)
    rows = conn.execute(text(f"""
        SELECT transaction_date, transaction_type, amount, lower(btrim(note)) AS note_key,
               min(note) AS note, count(*) AS row_count,
               array_agg(id ORDER BY entered_at, id) AS entry_ids,
               min(entered_at) AS first_entered_at, max(entered_at) AS last_entered_at,
               array_agg(DISTINCT COALESCE(NULLIF(entered_by_name,''), entered_by)) AS entered_by
        FROM {TABLE}
        WHERE is_deleted=false
          AND (CAST(:start_date AS date) IS NULL OR COALESCE(transaction_date, (entered_at AT TIME ZONE 'Asia/Ho_Chi_Minh')::date) >= CAST(:start_date AS date))
          AND (CAST(:end_date AS date) IS NULL OR COALESCE(transaction_date, (entered_at AT TIME ZONE 'Asia/Ho_Chi_Minh')::date) <= CAST(:end_date AS date))
        GROUP BY transaction_date, transaction_type, amount, lower(btrim(note))
        HAVING count(*) > 1
        ORDER BY count(*) DESC, transaction_date DESC NULLS LAST, amount DESC
    """), {"start_date": start_date, "end_date": end_date}).mappings().all()
    groups: list[dict[str, Any]] = []
    duplicate_rows = 0
    duplicate_amount = 0.0
    for row in rows:
        tx_date = row["transaction_date"]
        count = int(row["row_count"] or 0)
        amount = float(row["amount"] or 0)
        duplicate_rows += max(0, count - 1)
        duplicate_amount += max(0, count - 1) * amount
        groups.append({
            "date": tx_date.isoformat() if tx_date else "",
            "date_label": tx_date.strftime("%d-%m-%Y") if tx_date else "—",
            "type": str(row["transaction_type"]), "amount": amount,
            "note": str(row["note"] or ""), "count": count,
            "extra_count": max(0, count - 1),
            "entry_ids": [int(value) for value in (row["entry_ids"] or [])],
            "entered_by": [str(value) for value in (row["entered_by"] or []) if str(value or "").strip()],
            "first_entered_at": row["first_entered_at"].astimezone(VN_TZ).isoformat() if row["first_entered_at"] else "",
            "last_entered_at": row["last_entered_at"].astimezone(VN_TZ).isoformat() if row["last_entered_at"] else "",
        })
    return {
        "group_count": len(groups), "duplicate_row_count": duplicate_rows,
        "duplicate_amount": round(duplicate_amount, 2), "groups": groups,
    }
