"""Read-only TimeSoft code catalogue. Never infers FaceGate profile IDs or changes attendance."""
from io import BytesIO
import re
from zipfile import ZipFile, BadZipFile
from openpyxl import load_workbook
from fastapi import Depends, HTTPException, Request
from sqlalchemy import text
from vera_web_v2_devices import norm

MAX_UPLOAD = 8 * 1024 * 1024
MAX_ROWS = 50000


def phone(value):
    digits = re.sub(r'\D', '', str(value or ''))
    return '0' + digits[2:] if digits.startswith('84') and len(digits) in (11, 12) else digits


def cell_text(cell):
    value = cell.value
    if value is None or cell.data_type == 'f':
        return ''
    if isinstance(value, (int, float)) and not isinstance(value, bool) and int(value) == value:
        result = str(int(value))
        return result.zfill(len(cell.number_format)) if re.fullmatch('0{1,64}', cell.number_format or '') else result
    return str(value).strip()


def workbook_codes(content):
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 40 * 1024 * 1024 or len(archive.infolist()) > 1000:
                raise ValueError('File Excel quá lớn sau khi giải nén.')
        wb = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except (BadZipFile, ValueError, OSError, KeyError) as exc:
        raise HTTPException(400, 'File Excel không hợp lệ hoặc vượt giới hạn đọc.') from exc
    rows, count = [], 0
    try:
        aliases = {
            'name': {'ten nhan vien','employee name','employeename'},
            'phone': {'so dien thoai','dien thoai','mobile'},
            'attendance_code': {'ma cham cong','enrollnumber'},
            'employee_code': {'ma nhan vien','employeecode'},
        }
        for ws in wb.worksheets:
            if (ws.max_column or 0) > 256:
                raise HTTPException(413, 'Báo cáo có quá nhiều cột.')
            headers = None
            for number, cells in enumerate(ws.iter_rows(), 1):
                count += 1
                if count > MAX_ROWS:
                    raise HTTPException(413, 'Báo cáo vượt 50.000 dòng; hãy xuất khoảng ngày nhỏ hơn.')
                if headers is None:
                    if number > 30:
                        break
                    labels = [norm(cell_text(cell)) for cell in cells]
                    found = {key: next((i for i, label in enumerate(labels) if label in options), None) for key, options in aliases.items()}
                    if found['name'] is not None and found['attendance_code'] is not None:
                        headers = found
                    continue
                item = {key: cell_text(cells[index]) if index is not None and index < len(cells) else '' for key, index in headers.items()}
                if item['name'] and re.fullmatch(r'[A-Za-z0-9_-]{1,64}', item['attendance_code']):
                    rows.append(item)
        if not rows:
            raise HTTPException(400, 'Không tìm thấy cột Tên nhân viên và Mã chấm công có dữ liệu trong file.')
        return rows
    finally:
        wb.close()


def match_codes(employees, records):
    by_name, by_phone = {}, {}
    for employee in employees:
        for name in {norm(employee['username']), norm(employee.get('full_name'))} - {''}:
            by_name.setdefault(name, set()).add(employee['username'])
        number = phone(employee.get('phone'))
        if number:
            by_phone.setdefault(number, set()).add(employee['username'])
    choices = {row['username']: set() for row in employees}
    unknown, owners, seen = [], {}, set()
    for row in records:
        code = str(row.get('attendance_code') or '').strip()
        if not re.fullmatch(r'[A-Za-z0-9_-]{1,64}', code):
            continue
        key = (norm(row.get('name')), phone(row.get('phone')), code, str(row.get('employee_code') or ''))
        if key in seen:
            continue
        seen.add(key)
        names, phones = by_name.get(key[0], set()), by_phone.get(key[1], set())
        candidates = names & phones if names and phones else names or phones
        if len(candidates) == 1:
            username = next(iter(candidates))
            choices[username].add((code, key[3]))
            owners.setdefault(code, set()).add(username)
        else:
            owners.setdefault(code, set()).add(('unmatched', key[0], key[1]))
            unknown.append({'name': str(row.get('name') or ''), 'attendance_code': code,
                            'employee_code': key[3], 'status': 'ambiguous' if names or phones else 'unmatched'})
    result = []
    for employee in employees:
        values = sorted(choices[employee['username']])
        codes = {code for code, _ in values}
        status = 'matched' if len(codes) == 1 and len(owners[next(iter(codes))]) == 1 else 'conflict' if values else 'missing'
        result.append({'username': employee['username'], 'full_name': employee.get('full_name') or '',
            'status': status, 'codes': [{'attendance_code': code, 'employee_code': emp} for code, emp in values]})
    return {'employees': result, 'unmatched': unknown, 'matched_count': sum(row['status'] == 'matched' for row in result),
            'total_count': len(result), 'record_count': len(seen), 'attendance_calculation_enabled': False}


def read_employees(conn):
    return [dict(row) for row in conn.execute(text("""SELECT username, full_name, phone FROM employees
        WHERE COALESCE(payload->>'__deleted','false') <> 'true' AND lower(COALESCE(role,'')) <> 'admin'
        ORDER BY lower(username)""")).mappings().all()]


def saved_codes(conn):
    # Project only four fields in SQL; do not load attendance history or run its projection.
    return [dict(row) for row in conn.execute(text("""SELECT DISTINCT
        COALESCE(NULLIF(item->>'employeeInfo.Name',''),item->>'EmployeeName','') AS name,
        COALESCE(NULLIF(item->>'Mobile',''),item->>'employeeInfo.Mobile','') AS phone,
        COALESCE(item->>'EnrollNumber','') AS attendance_code,
        COALESCE(NULLIF(item->>'EmployeeCode',''),item->>'employeeInfo.EmployeeCode','') AS employee_code
        FROM vera_dataset_cache CROSS JOIN LATERAL jsonb_array_elements(
            CASE WHEN jsonb_typeof(payload::jsonb)='array' THEN payload::jsonb ELSE '[]'::jsonb END) AS item
        WHERE (dataset_key='timesoft_employee_checkin_today' OR dataset_key LIKE 'timesoft_employee_checkin_20%')
          AND COALESCE(item->>'EnrollNumber','') <> '' LIMIT 10001""")).mappings().all()]


def install_attendance_code_routes(app, *, engine_instance, current_identity, identity_type, require_feature):
    def check(conn, ident):
        require_feature(conn, ident, 'device_facegate_mapping_manage')

    @app.get('/v2/devices/attendance-codes')
    def catalogue(ident: identity_type = Depends(current_identity)):
        with engine_instance().connect() as conn:
            check(conn, ident)
            employees, rows = read_employees(conn), saved_codes(conn)
        if len(rows) > 10000:
            raise HTTPException(409, 'Nguồn mã vượt giới hạn; hãy nhập báo cáo TimeSoft mới nhất.')
        return dict(match_codes(employees, rows), source='TimeSoft đã đồng bộ vào VERA')

    @app.post('/v2/devices/attendance-codes/preview')
    async def preview(request: Request, ident: identity_type = Depends(current_identity)):
        with engine_instance().connect() as conn:
            check(conn, ident)
            employees = read_employees(conn)
        content = bytearray()
        async for chunk in request.stream():
            content.extend(chunk)
            if len(content) > MAX_UPLOAD:
                raise HTTPException(413, 'Chỉ nhận file Excel tối đa 8 MB.')
        from starlette.concurrency import run_in_threadpool
        rows = await run_in_threadpool(workbook_codes, bytes(content))
        return dict(match_codes(employees, rows), source='Báo cáo TimeSoft vừa chọn · chưa ghi ánh xạ')
