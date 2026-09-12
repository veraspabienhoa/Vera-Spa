"""Independent, transactional Live Tour board for Web V2.

No workbook, filesystem or Google service is used as an operational data source.
Live Tour keeps its own canonical state in ``vera_app_setting`` so a whole
booking/payment operation is committed as one PostgreSQL transaction.
"""
from __future__ import annotations

import hashlib
import json
import math
import re
import unicodedata
from collections.abc import Callable
from copy import deepcopy
from datetime import date, datetime, time, timedelta
from io import BytesIO
from typing import Any
from urllib.parse import quote
from uuid import NAMESPACE_URL, uuid4, uuid5
from zoneinfo import ZoneInfo

from fastapi import Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageDraw, ImageFont
from pydantic import BaseModel, Field
from sqlalchemy import text
from vera_web_v2_live_tour_payment import default_settings as _default_payment_settings, settings_update as _payment_settings_update, payment_values as _payment_values
from vera_web_v2_live_tour_payment import profile_bank as _profile_bank, selected_bank as _selected_bank
from vera_web_v2_live_tour_daily import sync_daily as _sync_daily
from vera_web_v2_live_tour_checkin import with_checkin as _directory_with_checkin
from vera_web_v2_combo_import import import_terms as _combo_import_terms
from vera_web_v2_live_tour_roster import shift_label as _directory_shift
from vera_web_v2_live_tour_roster import eligible as _roster_eligible, reconcile as _reconcile_roster
from vera_web_v2_service_catalog import catalog_details, component_debits, purchase_terms, require_available
from vera_web_v2_live_tour_permissions import CAPABILITY_FEATURES, EXPORT_FEATURES
from vera_web_v2_live_tour_changes import change_customer
from vera_web_v2_live_tour_invoice import change_paid_invoice

VN_TZ = ZoneInfo("Asia/Ho_Chi_Minh")
STATE_CATEGORY = "live_tour"
STATE_KEY = "state"
STATE_LOCK = "vera:v2:live_tour:state"
STATE_VERSION = 1
BUSINESS_DAY_CUTOFF = time(11, 10)
MAX_AUDIT = 3000
MAX_BACKUPS = 20
MAX_IDEMPOTENCY = 3000
MAX_MONEY = 10_000_000_000
MAX_SERVICE_DURATION_MINUTES = 1_440
MAX_TICKET_UNITS = 100_000
MAX_PURCHASE_QUANTITY = 1_000
CUSTOMER_PII_KEYS = frozenset({"customer_id", "customer_name", "customer_phone", "phone", "combo_purchase_id", "combo_reserved_units", "combo_reserved_components"})
PROTECTED_IDEMPOTENCY_ACTIONS = frozenset({
    "customer_delete", "customer_combo_update", "customer_combo_delete", "report_invoice_update", "report_invoice_delete",
    "paid_invoice_update", "paid_invoice_delete",
    "pending_update", "pending_delete",
    "checkout", "quick_checkout", "combo_purchase", "sync_leaves", "finish_to_pending", "start_room", "finish_room",
})
BACKDATE_ACTIONS = frozenset({"checkout", "quick_checkout", "combo_purchase"})
CLIENT_FINANCIAL_TIME_FIELDS = frozenset({
    "business_date", "created_at", "effective_at", "recorded_at",
})
IDEMPOTENCY_REQUIRED_ACTIONS = {
    "customer_delete", "customer_combo_update", "customer_combo_delete", "report_invoice_update", "report_invoice_delete",
    "paid_invoice_update", "paid_invoice_delete",
    "pending_update", "pending_delete",
    "booking", "multi_booking", "start", "add_minutes", "complete", "move_pending",
    "checkout", "quick_checkout", "set_work_status", "set_shift", "start_break", "end_break",
    "reorder", "admin_reorder", "sync_daily_status",
    "set_vip", "replace_service", "add_service", "room_upsert", "room_delete", "service_upsert",
    "service_delete", "combo_upsert", "combo_delete", "combo_purchase", "combo_import", "backup",
    "restore", "clear_expired", "customer_upsert", "service_area_upsert", "service_area_delete",
    "update_booking", "cancel_booking", "change_employee", "update_appointment", "finish_to_pending", "payment_settings_update", "start_room", "finish_room",
}
BOARD_COLUMNS = [
    "STT", "Tên nhân viên", "Lịch hẹn", "Trạng thái", "Phòng", "TG CÒN LẠI", "Yêu cầu",
    "Dịch vụ", "Đi làm", "Vào ca", "Breaktime", "TG nghỉ còn lại",
    "Giờ ra", "Giờ vào", "Ghi chú", "Thời lượng", "TG bắt đầu thực hiện",
    "TG bắt đầu thực hiện YC", "TT thanh toán", "Kết quả hoàn thành", "SL tua", "SL yêu cầu",
    "Tổng SL", "VIP", "Giờ Booking", "TG khách chờ", "TG Xông Hơi",
]

DEFAULT_ROOM_BEDS = [
    "1.1", "1.2", "1.3", "1.4", "1.5", "2.1", "2.2", "4", "6.1", "6.2", "6.3",
    "8.1", "8.2", "8.3", "8.4", "9.1", "9.2", "10.1", "10.2", "10.3", "10.4",
    "11.1", "11.2", "11.3", "12.1", "12.2", "12.3", "12B", "14.1", "14.2",
    "15.1", "15.2", "16.1", "16.2", "17.1", "17.2", "18.1", "18.2",
    "19.1", "19.2", "19.3", "20.1", "20.2", "21.1", "21.2", "21.3",
]

DEFAULT_SERVICES = [
    {"name": "P.Riêng", "duration": 90, "price": 0, "ticket_units": 1},
    {"name": "Body 90", "duration": 90, "price": 0, "ticket_units": 1},
    {"name": "Body 70", "duration": 70, "price": 0, "ticket_units": 1},
    {"name": "VIP 90 PR", "duration": 90, "price": 0, "ticket_units": 1},
    {"name": "VIP 90", "duration": 90, "price": 0, "ticket_units": 1},
    {"name": "VIP 70 PR", "duration": 70, "price": 0, "ticket_units": 1},
    {"name": "VIP 70", "duration": 70, "price": 0, "ticket_units": 1},
    {"name": "Mua thêm 30", "duration": 30, "price": 0, "ticket_units": 0},
    {"name": "Xông hơi", "duration": None, "price": 0, "ticket_units": 0},
]

DEFAULT_COMBOS = [
    {"name": "Combo 13", "tickets": 13, "price": 2_500_000},
    {"name": "Combo 14", "tickets": 14, "price": 2_500_000},
    {"name": "Combo VIP 90", "tickets": 13, "price": 3_000_000},
    {"name": "Combo VIP PR", "tickets": 13, "price": 3_500_000},
]


class LiveTourAction(BaseModel):
    action: str = Field(min_length=1, max_length=80)
    expected_revision: int | None = Field(default=None, ge=0)
    idempotency_key: str | None = Field(default=None, min_length=8, max_length=160)
    payload: dict[str, Any] = Field(default_factory=dict)


def _norm(value: Any) -> str:
    raw = unicodedata.normalize("NFD", str(value or "").strip().lower())
    raw = "".join(ch for ch in raw if unicodedata.category(ch) != "Mn").replace("đ", "d")
    return " ".join(raw.replace("_", " ").split())


def _canonical_shift(value: Any, *, allow_blank: bool = False) -> str:
    token = _norm(value).replace(" ", "")
    if allow_blank and not token:
        return ""
    if token in {"ca1", "1", "10", "10h", "10h00"}:
        return "Ca 1"
    if token in {"ca2", "2", "12", "12h", "12h00", "14", "14h", "14h00"}:
        return "Ca 2"
    raise HTTPException(400, "Ca làm việc chỉ nhận Ca 1 hoặc Ca 2.")


def _canonical_work_status(value: Any) -> str:
    token = _norm(value)
    mapping = {
        "di lam": "Đi làm", "dang lam viec": "Đi làm",
        "nghi phep": "Nghỉ phép", "nghi co phep": "Nghỉ phép",
        "nghi": "Nghỉ", "khong di lam": "Nghỉ",
    }
    if token not in mapping:
        raise HTTPException(400, "Trạng thái chỉ nhận Đi làm, Nghỉ phép hoặc Nghỉ.")
    return mapping[token]


def _canonical_payment_method(value: Any, *, quick: bool = False) -> str:
    token = _norm(value)
    if not token and quick:
        token = "tien mat"
    mapping = {
        "tien mat": "TIỀN MẶT", "cash": "TIỀN MẶT",
        "chuyen khoan": "CHUYỂN KHOẢN", "transfer": "CHUYỂN KHOẢN",
        "bank transfer": "CHUYỂN KHOẢN", "the": "THẺ", "card": "THẺ",
        "combo": "COMBO",
    }
    if token not in mapping:
        raise HTTPException(400, "Phương thức thanh toán chỉ nhận TIỀN MẶT, CHUYỂN KHOẢN, THẺ hoặc COMBO.")
    return mapping[token]


def _phone_key(value: Any) -> str:
    return re.sub(r"\D", "", str(value or ""))


def _redact_customer_pii(value: Any) -> Any:
    """Return a deep redacted copy, including dictionaries nested in lists."""
    if isinstance(value, dict):
        return {
            key: _redact_customer_pii(item)
            for key, item in value.items()
            if str(key).casefold() not in CUSTOMER_PII_KEYS
        }
    if isinstance(value, list):
        return [_redact_customer_pii(item) for item in value]
    if isinstance(value, tuple):
        return tuple(_redact_customer_pii(item) for item in value)
    return deepcopy(value)


def _redact_customer_name_alias(value: Any) -> Any:
    if isinstance(value, dict):
        return {
            key: _redact_customer_name_alias(item)
            for key, item in value.items() if str(key).casefold() != "name"
        }
    if isinstance(value, list):
        return [_redact_customer_name_alias(item) for item in value]
    return deepcopy(value)


def _contains_customer_pii(value: Any) -> bool:
    if isinstance(value, dict):
        for key, item in value.items():
            if str(key).casefold() in CUSTOMER_PII_KEYS and item not in (None, ""):
                return True
            if _contains_customer_pii(item):
                return True
    elif isinstance(value, (list, tuple)):
        return any(_contains_customer_pii(item) for item in value)
    return False


def _canonical_payload_hash(action: str, payload: dict[str, Any]) -> str:
    canonical_payload = {
        str(key): value for key, value in payload.items()
        if str(key) != "idempotency_key" and not str(key).startswith("_")
    }
    encoded = json.dumps(
        {"action": str(action or "").strip().lower(), "payload": canonical_payload},
        ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str,
    )
    return hashlib.sha256(encoded.encode("utf-8")).hexdigest()


def _common_customer_identity(items: list[dict[str, Any]]) -> dict[str, str]:
    identity_tokens: list[set[tuple[str, str]]] = []
    has_anonymous = False
    for item in items:
        tokens = {
            ("customer_id", str(item.get("customer_id") or "").strip()),
            ("customer_phone", _phone_key(item.get("customer_phone"))),
            ("customer_name", _norm(item.get("customer_name"))),
        }
        tokens = {(key, value) for key, value in tokens if value}
        if tokens:
            identity_tokens.append(tokens)
        else:
            has_anonymous = True
    if identity_tokens and has_anonymous:
        raise HTTPException(409, "Không được thanh toán chung dịch vụ có khách và dịch vụ khách vãng lai.")

    ids = {str(item.get("customer_id") or "").strip() for item in items if item.get("customer_id")}
    phones = {_phone_key(item.get("customer_phone")) for item in items if _phone_key(item.get("customer_phone"))}
    names = {_norm(item.get("customer_name")) for item in items if _norm(item.get("customer_name"))}
    if len(ids) > 1 or len(phones) > 1 or len(names) > 1:
        raise HTTPException(409, "Các dịch vụ đã chọn không thuộc cùng một khách hàng.")
    # Sparse rows are accepted only when their identity fields form one
    # connected identity (for example an ID-only row plus an ID/name/phone row).
    if len(identity_tokens) > 1:
        connected = set(identity_tokens[0])
        remaining = identity_tokens[1:]
        while remaining:
            newly_connected = [tokens for tokens in remaining if connected & tokens]
            if not newly_connected:
                raise HTTPException(409, "Thông tin khách hàng giữa các dịch vụ không nhất quán.")
            for tokens in newly_connected:
                connected.update(tokens)
                remaining.remove(tokens)
    first_name = next((str(item.get("customer_name") or "").strip() for item in items if _norm(item.get("customer_name"))), "")
    first_phone = next((str(item.get("customer_phone") or "").strip() for item in items if _phone_key(item.get("customer_phone"))), "")
    return {
        "customer_id": next(iter(ids), ""),
        "customer_phone": first_phone if phones else "",
        "customer_name": first_name if names else "",
    }


def _protect_customer_identity(payload: dict[str, Any], canonical: dict[str, str]) -> dict[str, Any]:
    protected = dict(payload)
    requested_id = str(payload.get("customer_id") or "").strip()
    requested_phone = _phone_key(payload.get("customer_phone") or payload.get("phone"))
    requested_name = _norm(payload.get("customer_name"))
    if canonical.get("customer_id") and requested_id and requested_id != canonical["customer_id"]:
        raise HTTPException(409, "Không được đổi khách hàng của dịch vụ khi thanh toán.")
    if _phone_key(canonical.get("customer_phone")) and requested_phone and requested_phone != _phone_key(canonical["customer_phone"]):
        raise HTTPException(409, "Số điện thoại thanh toán không khớp với khách đã đặt dịch vụ.")
    if canonical.get("customer_name") and requested_name and requested_name != _norm(canonical["customer_name"]):
        raise HTTPException(409, "Tên khách thanh toán không khớp với khách đã đặt dịch vụ.")
    for key in ("customer_id", "customer_name", "customer_phone"):
        if canonical.get(key):
            protected[key] = canonical[key]
    protected.pop("phone", None)
    return protected


def _auto_yc_ca1(now: datetime, employee: dict[str, Any], request: Any, enabled: bool) -> tuple[str, bool]:
    requested = str(request or "").strip()
    if requested and _norm(requested) != "yc":
        raise HTTPException(400, "Yêu cầu chỉ nhận YC hoặc để trống.")
    requested = "YC" if requested else ""
    local_clock = now.astimezone(VN_TZ).time().replace(tzinfo=None)
    in_window = local_clock >= time(23, 0) or local_clock <= time(3, 0)
    auto = bool(enabled) and not requested and _shift_bucket(employee) == "ca1" and in_window
    return ("YC" if auto else requested), auto


def _number(value: Any, default: float = 0) -> float:
    if value in (None, "") or isinstance(value, bool):
        return default
    if isinstance(value, (int, float)):
        return float(value)
    raw = str(value).strip().replace(" ", "")
    if "," in raw and "." not in raw:
        raw = raw.replace(",", ".")
    else:
        raw = raw.replace(",", "")
    try:
        return float(raw)
    except ValueError:
        return default


def _money(value: Any) -> int:
    return int(round(_number(value)))


def _bounded_number(
    value: Any, *, label: str, minimum: float = 0, maximum: float,
    integer: bool = False, allow_blank: bool = False,
) -> float | int | None:
    if value in (None, ""):
        if allow_blank:
            return None
        raise HTTPException(400, f"{label} không được để trống.")
    number = _number(value, math.nan)
    if not math.isfinite(number):
        raise HTTPException(400, f"{label} phải là số hữu hạn hợp lệ.")
    if number < minimum or number > maximum:
        raise HTTPException(400, f"{label} phải từ {minimum:g} đến {maximum:g}.")
    if integer and not float(number).is_integer():
        raise HTTPException(400, f"{label} phải là số nguyên.")
    return int(number) if integer else number


def _bounded_money(value: Any, *, label: str, allow_blank_as_zero: bool = False) -> int:
    normalized = 0 if allow_blank_as_zero and value in (None, "") else value
    number = _bounded_number(
        normalized, label=label, minimum=0, maximum=MAX_MONEY,
    )
    return int(round(float(number)))


def _duration_value(value: Any) -> float | None:
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        number = abs(float(value))
        if not math.isfinite(number):
            return None
        return number * 1440 if 0 < number < 1 else number
    raw = str(value).strip()
    match = re.fullmatch(r"(\d{1,3}):(\d{1,2})(?::(\d{1,2}))?", raw)
    if match:
        return int(match.group(1)) * 60 + int(match.group(2)) + int(match.group(3) or 0) / 60
    match = re.search(r"[-+]?\d+(?:[.,]\d+)?", raw)
    if not match:
        return None
    number = abs(float(match.group(0).replace(",", ".")))
    return number if math.isfinite(number) else None


def _source_start(value: Any, now: datetime, duration: float | None) -> datetime | None:
    if value in (None, ""):
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        number = float(value)
        if 0 <= number < 1:
            seconds = int(round(number * 86400)) % 86400
            parsed = now.replace(hour=seconds // 3600, minute=(seconds % 3600) // 60, second=seconds % 60, microsecond=0)
            if parsed > now and (now - (parsed - timedelta(days=1))).total_seconds() / 60 <= max(float(duration or 0) + 240, 480):
                parsed -= timedelta(days=1)
            return parsed
        if number >= 1:
            return datetime(1899, 12, 30, tzinfo=now.tzinfo) + timedelta(days=number)
    raw = str(value).strip()
    clock = re.fullmatch(r"(\d{1,2}):(\d{2})(?::(\d{2}))?", raw)
    if clock:
        try:
            parsed = now.astimezone(VN_TZ).replace(
                hour=int(clock.group(1)), minute=int(clock.group(2)),
                second=int(clock.group(3) or 0), microsecond=0,
            )
        except ValueError:
            return None
        if parsed > now.astimezone(VN_TZ) + timedelta(hours=3):
            parsed -= timedelta(days=1)
        return parsed
    return _parse_datetime(value)


def _iso(now: datetime) -> str:
    return now.astimezone(VN_TZ).replace(microsecond=0).isoformat()


def _business_date(now: datetime) -> date:
    local = now.astimezone(VN_TZ)
    return local.date() if local.time().replace(tzinfo=None) >= BUSINESS_DAY_CUTOFF else local.date() - timedelta(days=1)


def _counter_business_date(now: datetime) -> date:
    local = now.astimezone(VN_TZ)
    return local.date() - timedelta(days=1) if local.time() < time(10) else local.date()


def _ensure_counter_day(state: dict[str, Any], now: datetime) -> None:
    today = _counter_business_date(now).isoformat()
    prior = state.get("counter_business_date", state.get("business_date", today))
    if str(prior) < today:
        for employee in state["employees"]:
            employee.update({"tour_count": 0, "request_count": 0, "break_count": 0})
    state["counter_business_date"] = max(str(prior), today)


def _backdate_requested(payload: dict[str, Any]) -> bool:
    value = payload.get("backdate_one_day", False)
    if value is True:
        return True
    if value is False or value in (None, ""):
        return False
    raise HTTPException(400, "backdate_one_day chỉ nhận giá trị true hoặc false.")


def _financial_timing(payload: dict[str, Any], now: datetime) -> dict[str, Any]:
    """Return server-owned recorded/effective timestamps for a payment action."""
    overridden = sorted(
        key for key in CLIENT_FINANCIAL_TIME_FIELDS
        if key in payload and payload.get(key) not in (None, "")
    )
    if overridden:
        raise HTTPException(
            400,
            "Thời điểm giao dịch do máy chủ xác định; không được gửi " + ", ".join(overridden) + ".",
        )
    backdated = _backdate_requested(payload)
    reason = str(payload.get("correction_reason") or "").strip()
    if backdated and len(reason) < 3:
        raise HTTPException(400, "Lùi 1 ngày cần lý do điều chỉnh tối thiểu 3 ký tự.")
    if not backdated and reason:
        raise HTTPException(400, "Lý do điều chỉnh chỉ dùng khi chọn Lùi 1 ngày.")

    recorded = now.astimezone(VN_TZ).replace(microsecond=0)
    effective = recorded
    if backdated:
        # VBA displays the current business date (11:10 cutoff), then its
        # "Lùi 1 ngày" button subtracts one more date and fixes the time at
        # 23:59.  During the overnight window this is two calendar dates
        # behind ``recorded.date()``, by design.
        effective = datetime.combine(
            _business_date(recorded) - timedelta(days=1), time(23, 59), tzinfo=VN_TZ,
        )
    return {
        "backdate_one_day": backdated,
        "correction_reason": reason,
        "recorded_at": _iso(recorded),
        "effective_at": _iso(effective),
        "business_date": _business_date(effective).isoformat(),
        "effective_datetime": effective,
    }


def _booking_timing(entries, now, pending=None):
    pending = pending or {}
    dates = [_parse_datetime(row.get("booked_at")) for row in entries]
    effective = (_parse_datetime(pending.get("effective_at")) or _parse_datetime(pending.get("booked_at"))
                 or min((value for value in dates if value), default=None)
                 or _parse_datetime(pending.get("created_at")) or now.astimezone(VN_TZ))
    return {"recorded_at": _iso(now), "effective_at": _iso(effective),
            "business_date": effective.astimezone(VN_TZ).date().isoformat(),
            "effective_datetime": effective, "backdate_one_day": False, "correction_reason": ""}


def _invoice_date(value):
    parsed = _parse_datetime(value)
    if not isinstance(value, str) or not parsed or not 2000 <= parsed.year <= 2100:
        raise HTTPException(400, "Ngày giờ hóa đơn không hợp lệ (2000–2100).")
    return {"effective_at": _iso(parsed), "business_date": parsed.astimezone(VN_TZ).date().isoformat()}


def _parse_datetime(value: Any) -> datetime | None:
    if not value:
        return None
    raw = str(value).strip()
    try:
        parsed = datetime.fromisoformat(raw)
        return parsed.replace(tzinfo=VN_TZ) if parsed.tzinfo is None else parsed.astimezone(VN_TZ)
    except ValueError:
        pass
    for fmt in ("%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%H:%M:%S", "%H:%M"):
        try:
            parsed = datetime.strptime(raw, fmt)
            if fmt.startswith("%H"):
                today = datetime.now(VN_TZ)
                parsed = parsed.replace(year=today.year, month=today.month, day=today.day)
            return parsed.replace(tzinfo=VN_TZ)
        except ValueError:
            continue
    return None


def _display_datetime(value: Any) -> str:
    parsed = _parse_datetime(value)
    return parsed.strftime("%d/%m/%Y %H:%M:%S") if parsed else str(value or "")


def _room_group(value: Any) -> str:
    room = _norm(value)
    room = re.sub(r"^(phong|vip)\s*", "", room).strip()
    return room.split(".", 1)[0]


def _catalog_room_group(state: dict[str, Any], name: Any) -> str:
    item = next((row for row in state["rooms"] if _norm(row.get("name")) == _norm(name)), {})
    return str(item.get("area_name") or _room_group(name))


def _room_action_members(state: dict[str, Any], room: str, action: str) -> list[dict[str, Any]]:
    statuses = {"dang cho"} if action == "start_room" else {"dang thuc hien", "dang su dung"}
    return [employee for employee in state["employees"]
            if employee.get("service") and employee.get("room")
            and _norm(_catalog_room_group(state, employee["room"])) == _norm(room)
            and _norm(employee.get("status")) in statuses]


def _room_action_counts(state: dict[str, Any]) -> dict[str, dict[str, int]]:
    counts: dict[str, dict[str, int]] = {}
    for employee in state["employees"]:
        if not employee.get("service") or not employee.get("room"):
            continue
        room = _catalog_room_group(state, employee["room"])
        bucket = counts.setdefault(room, {"waiting": 0, "doing": 0})
        status = _norm(employee.get("status"))
        if status == "dang cho":
            bucket["waiting"] += 1
        elif status in {"dang thuc hien", "dang su dung"}:
            bucket["doing"] += 1
    return counts


def _service_areas(state: dict[str, Any]) -> list[dict[str, Any]]:
    """Project existing booking places into areas without changing their IDs."""
    areas: dict[str, dict[str, Any]] = {}
    for room in state["rooms"]:
        group = _catalog_room_group(state, room.get("name"))
        area_id = str(room.get("area_id") or f"legacy:{group}")
        area = areas.setdefault(area_id, {
            "id": area_id, "name": group, "kind": room.get("area_kind", "room"), "beds": [],
        })
        area["beds"].append({
            "id": room["id"], "name": room.get("bed_name") or room.get("name", ""),
            "booking_name": room.get("name", ""), "active": room.get("active", True),
        })
    for name in state.get("physical_rooms", []):
        if not any(_norm(area["name"]) == _norm(name) for area in areas.values()):
            areas[f"legacy:{name}"] = {"id": f"legacy:{name}", "name": name, "kind": "room", "beds": []}
    return list(areas.values())


def _service_area_change(state: dict[str, Any], payload: dict[str, Any], *, delete: bool = False) -> dict[str, Any]:
    areas = _service_areas(state)
    area_id = str(payload.get("id") or "").strip()
    current = _find_by_id(areas, area_id, "khu vực dịch vụ") if area_id else None
    old_ids = {bed["id"] for bed in (current or {}).get("beds", [])}
    old_rooms = [row for row in state["rooms"] if row["id"] in old_ids]
    if any(_catalog_referenced(state, "rooms", row) for row in old_rooms):
        raise HTTPException(409, "Khu vực còn dịch vụ chưa thanh toán; hãy hoàn tất trước khi thay đổi hoặc xóa.")
    if delete:
        if current is None:
            raise HTTPException(400, "Thiếu mã khu vực dịch vụ.")
        state["rooms"] = [row for row in state["rooms"] if row["id"] not in old_ids]
        if current and "physical_rooms" in state:
            state["physical_rooms"] = [name for name in state["physical_rooms"] if _norm(name) != _norm(current["name"])]
        return {"deleted_id": area_id}
    name = str(payload.get("name") or "").strip()
    kind = payload.get("kind")
    if not name or len(name) > 100:
        raise HTTPException(400, "Tên khu vực cần từ 1 đến 100 ký tự.")
    if kind not in {"room", "bed", "table"}:
        raise HTTPException(400, "Loại khu vực chỉ nhận Phòng, Giường hoặc Bàn.")
    if any(area["id"] != area_id and _norm(area["name"]) == _norm(name) for area in areas):
        raise HTTPException(409, "Tên khu vực dịch vụ đã tồn tại.")
    beds = payload.get("beds") if kind == "room" else [{"name": name}]
    if not isinstance(beds, list) or not 1 <= len(beds) <= 100:
        raise HTTPException(400, "Mỗi phòng cần từ 1 đến 100 giường.")
    area_id = area_id or str(uuid4())
    remaining = [row for row in state["rooms"] if row["id"] not in old_ids]
    names = {_norm(row["name"]) for row in remaining}
    bed_names: set[str] = set()
    used_ids: set[str] = set()
    replacements = []
    for bed in beds:
        if not isinstance(bed, dict):
            raise HTTPException(400, "Thông tin giường không hợp lệ.")
        bed_name = str(bed.get("name") or "").strip()
        if not bed_name or len(bed_name) > 100 or _norm(bed_name) in bed_names:
            raise HTTPException(400, "Tên giường cần từ 1 đến 100 ký tự và không được trùng trong phòng.")
        bed_names.add(_norm(bed_name))
        bed_id = str(bed.get("id") or "")
        if kind != "room" and current and current["kind"] == kind:
            bed_id = current["beds"][0]["id"]
        if bed_id and (bed_id not in old_ids or bed_id in used_ids):
            raise HTTPException(400, "Mã giường không thuộc khu vực hoặc bị trùng.")
        bed_id = bed_id or str(uuid4())
        used_ids.add(bed_id)
        old = next((row for row in old_rooms if row["id"] == bed_id), {})
        old_bed = next((row for row in (current or {}).get("beds", []) if row["id"] == bed_id), {})
        booking_name = f"{name} · {bed_name}" if kind == "room" else name
        # A no-op edit of a legacy room must keep existing booking names.
        if current and current["name"] == name and current["kind"] == kind and old_bed.get("name") == bed_name:
            booking_name = old["name"]
        if _norm(booking_name) in names:
            raise HTTPException(409, "Tên vị trí phục vụ đã tồn tại.")
        names.add(_norm(booking_name))
        replacements.append({
            **old, "id": bed_id, "name": booking_name, "group": name,
            "area_id": area_id, "area_name": name, "area_kind": kind, "bed_name": bed_name,
            "type": "vip" if kind == "room" and _room_group(name) in {str(n) for n in range(16, 22)} else "standard",
            "active": old.get("active", True),
        })
    state["rooms"] = remaining + replacements
    if current and "physical_rooms" in state:
        state["physical_rooms"] = [value for value in state["physical_rooms"] if _norm(value) != _norm(current["name"])]
    return {"service_area": next(area for area in _service_areas(state) if area["id"] == area_id)}


def _is_private_service(value: Any) -> bool:
    token = _norm(value)
    return bool(
        re.search(r"(^|[^a-z0-9])pr($|[^a-z0-9])", token)
        or re.search(r"(^|[^a-z0-9])p\s*\.?\s*rieng($|[^a-z0-9])", token)
    )


def _stable_id(prefix: str, value: Any) -> str:
    return str(uuid5(NAMESPACE_URL, f"vera-live-tour:{prefix}:{_norm(value)}"))


def _catalog_defaults() -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    rooms = [
        {
            "id": _stable_id("room", name), "name": name, "group": _room_group(name),
            "type": "vip" if _room_group(name).isdigit() and 16 <= int(_room_group(name)) <= 21 else "standard",
            "active": True,
        }
        for name in DEFAULT_ROOM_BEDS
    ]
    services = [
        {
            "id": _stable_id("service", item["name"]), **item,
            "private": _is_private_service(item["name"]), "active": True,
        }
        for item in DEFAULT_SERVICES
    ]
    combos = [
        {"id": _stable_id("combo", item["name"]), **item, "active": True}
        for item in DEFAULT_COMBOS
    ]
    return rooms, services, combos


def _empty_state(now: datetime) -> dict[str, Any]:
    rooms, services, combos = _catalog_defaults()
    return {
        "version": STATE_VERSION,
        "business_date": _business_date(now).isoformat(),
        "counter_business_date": _counter_business_date(now).isoformat(),
        "created_at": _iso(now), "updated_at": _iso(now),
        "employees": [], "rooms": rooms, "services": services, "combos": combos,
        "customers": [], "pending": [], "invoices": [], "reports": [], "combo_usage": [],
        "break_events": [], "payment_settings": _default_payment_settings(),
        "audit": [], "backups": [], "pending_changes": [], "invoice_changes": [], "customer_changes": [], "bill_counters": {}, "idempotency": {},
    }


def _normalize_state(raw: Any, now: datetime) -> dict[str, Any]:
    state = deepcopy(raw) if isinstance(raw, dict) else _empty_state(now)
    defaults = _empty_state(now)
    if not isinstance(state.get("customer_changes"), list):
        state["customer_changes"] = []
    if not isinstance(state.get("invoice_changes"), list):
        state["invoice_changes"] = []
    for key in ("employees", "rooms", "services", "combos", "customers", "pending", "invoices", "reports", "combo_usage", "break_events", "audit", "backups", "pending_changes"):
        if not isinstance(state.get(key), list):
            state[key] = deepcopy(defaults[key])
    # An explicitly empty catalog is a saved choice, not an uninitialized state.
    state["version"] = STATE_VERSION
    state.setdefault("payment_settings", _default_payment_settings())
    state.setdefault("business_date", _business_date(now).isoformat())
    state.setdefault("created_at", _iso(now))
    state.setdefault("updated_at", _iso(now))
    if not isinstance(state.get("bill_counters"), dict):
        state["bill_counters"] = {}
    if not isinstance(state.get("idempotency"), dict):
        state["idempotency"] = {}
    for index, employee in enumerate(state["employees"]):
        employee.setdefault("id", str(uuid4()))
        employee.setdefault("sort_index", index)
        employee["hidden"] = False
        employee.setdefault("vip", False)
        employee.setdefault("tour_count", 0)
        employee.setdefault("request_count", 0)
        employee.setdefault("work_status", "Đi làm")
        employee["work_status"] = _source_work_status(employee.get("work_status"))
        employee["shift"] = _source_shift(employee.get("shift"))
    return state


def _source_value(record: dict[str, Any], *names: str) -> Any:
    lookup = {_norm(key): value for key, value in record.items()}
    for name in names:
        if _norm(name) in lookup:
            return lookup[_norm(name)]
    return ""


def _source_work_status(value: Any) -> str:
    if value in (None, ""):
        return "Đi làm"
    try:
        return _canonical_work_status(value)
    except HTTPException:
        # Unknown source states must never silently make somebody available.
        return "Nghỉ"


def _source_shift(value: Any) -> str:
    try:
        return _canonical_shift(value, allow_blank=True)
    except HTTPException:
        return ""


def _source_employee(record: dict[str, Any], index: int, now: datetime) -> dict[str, Any] | None:
    source_name = str(_source_value(record, "Tên nhân viên", "Nhân viên", "Họ tên") or "").strip()
    vip = bool(re.search(r"\*+\s*$", source_name))
    name = re.sub(r"\s*\*+\s*$", "", source_name).strip()
    if not name:
        return None
    stt = str(_source_value(record, "STT", "Số thứ tự") or index + 1).strip()
    duration_raw = _source_value(record, "Thời lượng", "Thời lượng (phút)")
    duration = _duration_value(duration_raw)
    raw_request = str(_source_value(record, "Yêu cầu") or "").strip()
    request = "YC" if _norm(raw_request) == "yc" else ""
    request_source = "tour_import" if not raw_request or request else "tour_import_invalid"
    start_raw = _source_value(
        record,
        "TG bắt đầu thực hiện YC" if _norm(request) == "yc" else "TG bắt đầu thực hiện",
        "Bắt đầu thực hiện YC" if _norm(request) == "yc" else "Bắt đầu thực hiện",
    )
    started = _source_start(start_raw, now, duration) if start_raw else None
    booked_raw = _source_value(record, "Giờ Booking", "TG Booking")
    booked = _source_start(booked_raw, now, None) if booked_raw else None
    break_flag = _source_value(record, "Break", "Breaktime")
    clock_in_raw = _source_value(record, "Giờ vào")
    break_active = (
        bool(str(break_flag or "").strip())
        and _norm(break_flag) not in {"0", "false", "khong", "no"}
        and not str(clock_in_raw or "").strip()
    )
    clock_out_raw = _source_value(record, "Giờ ra")
    break_start_raw = clock_out_raw or (break_flag if not isinstance(break_flag, str) else "")
    break_started = _source_start(break_start_raw, now, None) if break_active and break_start_raw else None
    service = str(_source_value(record, "Dịch vụ") or "").strip()
    column_x_elapsed = _duration_value(_source_value(record, "TG Xông Hơi", "TG Xong Hoi"))
    is_steam_service = _norm(service) in {"xong hoi", "steam"}
    steam_elapsed = column_x_elapsed if is_steam_service else None
    wait_elapsed = None if is_steam_service else column_x_elapsed
    break_remaining = _number(_source_value(record, "Thời gian", "TG Break"), math.nan)
    break_remaining = break_remaining if math.isfinite(break_remaining) else math.nan
    tour_count = _number(_source_value(record, "SL tua"), 0)
    request_count = _number(_source_value(record, "SL yêu cầu"), 0)
    tour_count = tour_count if math.isfinite(tour_count) else 0
    request_count = request_count if math.isfinite(request_count) else 0
    status = str(_source_value(record, "Trạng thái") or "").strip()
    payment_raw = str(_source_value(record, "Chờ thanh toán", "TT thanh toán") or "").strip()
    payment_pending = _norm(payment_raw) in {"cho thanh toan", "chua thanh toan"}
    completion_note = "" if payment_pending else payment_raw
    if payment_pending:
        status = "CHO THANH TOÁN"
        payment_raw = "CHO THANH TOÁN"
    else:
        payment_raw = ""
    return {
        "id": _stable_id("employee", f"{stt}:{name}"), "stt": stt, "name": name,
        "appointment": str(_source_value(record, "Lịch hẹn") or "").strip(),
        "service": service,
        "request": request, "request_source": request_source,
        "room": str(_source_value(record, "Phòng") or "").strip(),
        "status": status,
        "duration": int(round(duration)) if duration is not None else None,
        "service_price": None, "service_price_source": "tour_import",
        "booked_at": _iso(booked) if booked else "", "started_at": _iso(started) if started else "", "completed_at": "",
        "payment_status": payment_raw,
        "completion_note": completion_note,
        "tour_count": int(round(tour_count)),
        "request_count": int(round(request_count)),
        "work_status": _source_work_status(_source_value(record, "Đi làm")),
        "shift": _source_shift(_source_value(record, "Vào ca", "Ca")),
        "break_started_at": _iso(break_started) if break_started else "",
        "clock_out": str(clock_out_raw or "").strip(),
        "clock_in": str(clock_in_raw or "").strip(),
        # Input!T is the VBA break countdown, while Input!X is elapsed steam
        # time.  They intentionally remain distinct from service wait_minutes.
        "break_remaining_minutes": None if math.isnan(break_remaining) else break_remaining,
        "steam_elapsed_minutes": steam_elapsed, "wait_minutes": wait_elapsed, "completion_delta_minutes": None,
        "note": str(_source_value(record, "Ghi chú") or "").strip(),
        "hidden": False, "vip": vip, "sort_index": index,
    }


def _employee_directory(conn, now=None) -> list[dict[str, Any]]:
    rows = [dict(row) for row in conn.execute(text("""
        SELECT username, COALESCE(full_name, '') AS full_name,
               lower(btrim(COALESCE(role,''))) AS role, COALESCE(payload,'{}'::jsonb) AS payload, COALESCE(work_shift,'') AS work_shift,
               shift_start_date, rotation_cycle,
               (SELECT value_json FROM vera_app_setting WHERE category='shift' AND setting_key='shift_definitions' LIMIT 1) AS shift_definitions
        FROM employees
        ORDER BY lower(username)
    """)).mappings().all()]
    return _directory_with_checkin(conn, rows, (now or datetime.now(VN_TZ)).astimezone(VN_TZ))


def _new_directory_employee(row, index):
    return {
        "id": _stable_id("employee", row["username"]),
        "username": row["username"], "name": row["username"], "role": row["role"],
        "stt": str(index + 1), "roster_eligible": True,
        "appointment": "", "service": "", "request": "", "room": "", "status": "",
        "duration": None, "service_price": 0, "booked_at": "", "started_at": "",
        "completed_at": "", "payment_status": "", "wait_minutes": None,
        "completion_delta_minutes": None, "steam_elapsed_minutes": None,
        "tour_count": 0, "request_count": 0, "work_status": "Nghỉ", "shift": _directory_shift(row.get("work_shift"), row.get("shift_definitions")),
        "break_started_at": "", "clock_out": "", "clock_in": "", "note": "",
        "hidden": False, "vip": False, "sort_index": index,
    }


def _database_employees(conn) -> list[dict[str, Any]]:
    return [_new_directory_employee(row, i) for i, row in enumerate(_employee_directory(conn)) if _roster_eligible(row)]


def _bootstrap_state(conn, now: datetime) -> dict[str, Any]:
    state = _empty_state(now)
    _reconcile_roster(state, _employee_directory(conn, now), _new_directory_employee)
    state["bootstrap_source"] = "employees"
    state["storage_mode"] = "server"
    return state


def _find_by_id(items: list[dict[str, Any]], item_id: Any, label: str) -> dict[str, Any]:
    wanted = str(item_id or "").strip()
    item = next((row for row in items if str(row.get("id") or "") == wanted), None)
    if not item:
        raise HTTPException(404, f"Không tìm thấy {label}.")
    return item


def _employee(state: dict[str, Any], employee_id: Any) -> dict[str, Any]:
    return _find_by_id(state["employees"], employee_id, "nhân viên")


def _catalog_item(state: dict[str, Any], kind: str, payload: dict[str, Any]) -> dict[str, Any] | None:
    items = state[kind]
    singular = {"services": "service", "rooms": "room", "combos": "combo"}[kind]
    item_id = payload.get(f"{singular}_id")
    name = payload.get(singular) or payload.get("name")
    if item_id:
        return _find_by_id(items, item_id, singular)
    if name:
        return next((item for item in items if _norm(item.get("name")) == _norm(name)), None)
    return None


def _active_booking(employee: dict[str, Any]) -> bool:
    # Completion releases the physical room while the row remains available
    # for payment/move-pending bookkeeping.
    return _norm(employee.get("status")) in {"dang cho", "dang thuc hien", "dang su dung"} and bool(employee.get("room"))


def _has_unsettled_work(employee: dict[str, Any]) -> bool:
    return bool(
        employee.get("service") or employee.get("break_started_at")
        or _norm(employee.get("status")) in {"dang cho", "dang thuc hien", "dang su dung", "cho thanh toan"}
        or _norm(employee.get("payment_status")) == "cho thanh toan"
    )


def _catalog_private_service(state: dict[str, Any], service: Any) -> bool:
    name = _norm(service)
    return _is_private_service(service) or any(
        item.get("private") and re.search(r"(?:^|\s*&\s*)" + re.escape(_norm(item.get("name"))) + r"(?:\s*&\s*|$)", name)
        for item in state["services"]
    )


def _catalog_referenced(state: dict[str, Any], kind: str, item: dict[str, Any]) -> bool:
    entries = state["employees"] + [entry for pending in state["pending"] for entry in pending.get("entries", [])]
    field = "room" if kind == "rooms" else "service"
    name = _norm(item.get("name"))
    return any(name == _norm(entry.get(field)) or name in {_norm(value) for value in str(entry.get(field) or "").split("&")} or (kind == "services" and any(part.get("service_id") == item["id"] for part in entry.get("service_items", []))) for entry in entries)


def _check_room_collision(state: dict[str, Any], candidate: dict[str, Any], room: str, service: str) -> None:
    wanted_room = _norm(room)
    wanted_group = _norm(_catalog_room_group(state, room))
    wanted_private = _catalog_private_service(state, service)
    for employee in state["employees"]:
        if employee is candidate or not _active_booking(employee):
            continue
        current_room = str(employee.get("room") or "")
        same_bed = _norm(current_room) == wanted_room
        same_group = _norm(_catalog_room_group(state, current_room)) == wanted_group
        current_private = _catalog_private_service(state, employee.get("service"))
        if same_bed or (same_group and (wanted_private or current_private)):
            if wanted_private or current_private:
                raise HTTPException(409, f"Phòng {wanted_group} đang bị khóa toàn phòng bởi dịch vụ PR.")
            raise HTTPException(409, f"Giường/phòng {room} đang được sử dụng.")


RETAINED_ASSIGNMENT_COLUMNS = (
    "Thời lượng", "TG bắt đầu thực hiện", "TG bắt đầu thực hiện YC",
    "TT thanh toán", "Kết quả hoàn thành", "Giờ Booking", "TG khách chờ", "TG Xông Hơi",
)


def _clear_assignment(employee: dict[str, Any], now: datetime) -> None:
    # Retain display history separately from the live booking. Old timestamps,
    # prices and combo reservations must never become a second payable service.
    record = _employee_record(employee, now)
    employee["last_assignment_display"] = {column: record[column] for column in RETAINED_ASSIGNMENT_COLUMNS}
    employee.pop("service_items", None)
    employee.pop("pre_start_tour_position", None)
    employee.pop("employee_change_minutes", None)
    for key in ("combo_purchase_id", "combo_reserved_units", "combo_reserved_components"):
        employee.pop(key, None)
    for key in (
        "service", "request", "request_source", "room", "status", "booked_at",
        "started_at", "completed_at", "payment_status", "customer_id", "customer_name",
        "customer_phone", "booking_id",
        "service_price_source", "completion_note",
    ):
        employee[key] = ""
    for key in ("duration", "service_price", "wait_minutes", "completion_delta_minutes", "steam_elapsed_minutes"):
        employee[key] = None


def _audit(
    state: dict[str, Any], action: str, payload: dict[str, Any], actor: str,
    now: datetime, *, timing: dict[str, Any] | None = None,
) -> None:
    safe_payload = _redact_customer_pii({
        str(key): value for key, value in payload.items()
        if not str(key).startswith("_") and key not in {"snapshot"}
    })
    if action in {
        "booking", "multi_booking", "checkout", "quick_checkout",
        "move_pending", "combo_purchase", "combo_import", "customer_upsert",
    }:
        safe_payload = _redact_customer_name_alias(safe_payload)
    try:
        encoded = json.dumps(safe_payload, ensure_ascii=False, default=str)
        if len(encoded) > 8000:
            safe_payload = {"summary": f"Payload {len(encoded)} ký tự", "count": len(payload)}
    except (TypeError, ValueError):
        safe_payload = {"summary": "Payload không tuần tự hóa được"}
    effective = (timing or {}).get("effective_datetime") or now
    state["audit"].append({
        "id": str(uuid4()), "at": _iso(now), "created_at": _iso(now),
        "recorded_at": _iso(now), "effective_at": _iso(effective),
        "business_date": _business_date(effective).isoformat(),
        "actor": actor, "action": action, "detail": safe_payload,
    })
    state["audit"] = state["audit"][-MAX_AUDIT:]


def _service_values(state: dict[str, Any], payload: dict[str, Any], now: datetime | None = None) -> tuple[str, int | None, int]:
    item = _catalog_item(state, "services", payload)
    name = str((item or {}).get("name") or payload.get("service") or "").strip()
    if not name:
        raise HTTPException(400, "Vui lòng chọn dịch vụ.")
    if item is None:
        raise HTTPException(400, f"Dịch vụ '{name}' chưa có trong danh mục.")
    if now is not None:
        require_available(item, now.astimezone(VN_TZ).date(), "Dịch vụ")
    duration_value = (item or {}).get("duration")
    is_request = _norm(payload.get("request")) == "yc"
    eligibility = "request_eligible" if is_request else "non_request_eligible"
    if item.get(eligibility) is False:
        raise HTTPException(409, "Dịch vụ không áp dụng cho loại yêu cầu đã chọn.")
    if is_request and item.get("request_duration") not in (None, ""):
        duration_value = item["request_duration"]
    duration_number = _bounded_number(
        duration_value, label="Thời lượng dịch vụ", minimum=0,
        maximum=MAX_SERVICE_DURATION_MINUTES, allow_blank=True,
    )
    duration = int(round(duration_number)) if duration_number is not None else None
    # Catalog pricing is canonical; operator payloads never define money.
    price = item.get("price", 0)
    return name, duration, _bounded_money(price, label="Giá dịch vụ", allow_blank_as_zero=True)


def _service_selection(state, payload, now):
    requested = payload.get("service_items")
    if requested is None:
        name, duration, price = _service_values(state, payload, now)
        service = next(row for row in state["services"] if row["name"] == name)
        return name, duration, price, [{"service_id": service["id"], "name": name, "quantity": 1, "unit_price": price, "duration": duration, "ticket_units": service.get("ticket_units", 1)}]
    if not isinstance(requested, list) or not 1 <= len(requested) <= 30:
        raise HTTPException(400, "Chọn từ 1 đến 30 dịch vụ.")
    items, seen = [], set()
    for item in requested:
        if not isinstance(item, dict) or not isinstance(item.get("service_id"), str) or not item["service_id"] or item["service_id"] in seen or isinstance(item.get("quantity"), bool):
            raise HTTPException(400, "Dịch vụ chưa chọn hoặc bị trùng; hãy tăng số lượng ở dòng đã có.")
        seen.add(item["service_id"])
        quantity = int(_bounded_number(item.get("quantity", 1), label="Số lượng dịch vụ", minimum=1, maximum=30, integer=True))
        name, duration, price = _service_values(state, {"service_id": item["service_id"], "request": payload.get("request", "")}, now)
        items.append({"service_id": item["service_id"], "name": name, "quantity": quantity, "unit_price": price, "duration": duration, "ticket_units": _service_ticket_units(state, name)})
    total_quantity = sum(row["quantity"] for row in items)
    duration = sum((row["duration"] or 0) * row["quantity"] for row in items)
    price = sum(row["unit_price"] * row["quantity"] for row in items)
    if total_quantity > 30 or duration > MAX_SERVICE_DURATION_MINUTES or price > MAX_MONEY:
        raise HTTPException(400, "Tổng dịch vụ vượt giới hạn 30 lượt, 1440 phút hoặc giá trị cho phép.")
    name = " & ".join(row["name"] for row in items for _ in range(row["quantity"]))
    return name, duration if any(row["duration"] is not None for row in items) else None, price, items


def _entry_ticket_units(state, entry):
    if entry.get("price_source") == "combo_ticket":
        return 1  # Server-created redemption of one legacy generic ticket.
    if entry.get("service_items"):
        return sum(row.get("ticket_units", 1) * row["quantity"] for row in entry["service_items"])
    return _service_ticket_units(state, entry.get("service"))


def _available_combo(state, customer_id, purchase, *, employee_ids=(), pending_id=""):
    """Subtract open reservations without changing the purchased ticket ledger.

    Called under the same state lock as booking/checkout. Include hidden and
    retained assignments; exclude only the transaction currently being edited.
    """
    available = deepcopy(purchase)
    reservations = [row for row in state["employees"]
                    if row.get("service") and row.get("id") not in employee_ids
                    and row.get("customer_id") == customer_id]
    for pending in state["pending"]:
        if pending.get("id") != pending_id and pending.get("customer_id") == customer_id:
            reservations.extend(pending.get("entries") or [])
    reserved = 0
    by_service = {}
    for row in reservations:
        if row.get("combo_purchase_id") != purchase.get("id"):
            continue
        reserved += int(row.get("combo_reserved_units") or 0)
        for part in row.get("combo_reserved_components") or []:
            key = part["service_id"]
            by_service[key] = by_service.get(key, 0) + part["units"]
    if purchase.get("deleted_at"):
        available["remaining"] = 0
        for part in available.get("component_balances", []):
            part["remaining"] = 0
        available["booking_reserved"] = reserved
        return available
    available["booking_reserved"] = reserved
    available["remaining"] = max(0, int(purchase.get("remaining") or 0) - reserved)
    for part in available.get("component_balances") or []:
        part["remaining"] = max(0, part["remaining"] - by_service.get(part["service_id"], 0))
    return available


def _booking_combo(state, customer, payload, entry, now, employee=None):
    employee = employee or {}
    purchase_id = payload.get("combo_purchase_id", employee.get("combo_purchase_id", ""))
    if purchase_id is not None and not isinstance(purchase_id, str):
        raise HTTPException(400, "Mã combo đã mua không hợp lệ.")
    purchase_id = (purchase_id or "").strip()
    if not purchase_id:
        return {"combo_purchase_id": "", "combo_reserved_units": 0, "combo_reserved_components": []}
    if not customer:
        raise HTTPException(400, "Phải chọn khách hàng khi đặt lịch bằng combo.")
    purchase = next((row for row in customer.get("combo_purchases", []) if row.get("id") == purchase_id), None)
    if purchase is None:
        raise HTTPException(404, "Không tìm thấy combo đã mua của khách hàng.")
    available = _available_combo(state, customer["id"], purchase, employee_ids=[employee.get("id")])
    if available["remaining"] <= 0:
        raise HTTPException(409, "Combo còn 0 vé có thể đặt lịch (đã hết hoặc đã được giữ chỗ).")
    plan = component_debits(available, [entry], state["services"], now.astimezone(VN_TZ).date())
    units = sum(row["units"] for row in plan) if "component_balances" in purchase else _entry_ticket_units(state, entry)
    if units <= 0 or units > available["remaining"] or units > MAX_TICKET_UNITS:
        raise HTTPException(409, f"Combo chỉ còn {available['remaining']} vé có thể đặt lịch, không đủ dùng {units} vé.")
    return {"combo_purchase_id": purchase_id, "combo_reserved_units": units, "combo_reserved_components": plan}


def _service_catalog_price(state: dict[str, Any], service_name: Any) -> int | None:
    name = str(service_name or "").strip()
    if not name:
        return None
    exact = next((item for item in state["services"] if _norm(item.get("name")) == _norm(name)), None)
    if exact is not None:
        return _bounded_money(exact.get("price"), label="Giá dịch vụ", allow_blank_as_zero=True)
    parts = [part.strip() for part in re.split(r"\s*&\s*", name) if part.strip()]
    if len(parts) <= 1:
        return None
    prices = [_service_catalog_price(state, part) for part in parts]
    return sum(price for price in prices if price is not None) if all(price is not None for price in prices) else None


def _service_catalog_duration(state: dict[str, Any], service_name: Any) -> int | None:
    name = str(service_name or "").strip()
    if not name:
        return None
    exact = next((item for item in state["services"] if _norm(item.get("name")) == _norm(name)), None)
    if exact is not None:
        raw = exact.get("duration")
        return int(round(_number(raw))) if raw not in (None, "") else None
    parts = [part.strip() for part in re.split(r"\s*&\s*", name) if part.strip()]
    if len(parts) <= 1:
        return None
    durations = [_service_catalog_duration(state, part) for part in parts]
    return sum(value for value in durations if value is not None) if all(value is not None for value in durations) else None


def _resolved_service_price(state: dict[str, Any], entry: dict[str, Any]) -> int:
    stored = entry.get("price", entry.get("service_price"))
    source = str(entry.get("price_source") or entry.get("service_price_source") or "").strip()
    catalog_price = _service_catalog_price(state, entry.get("service"))
    # Tour imports never carry a trustworthy price. Old imported rows did not
    # have a source marker, so zero/missing is also repaired from the catalog.
    stored_value = None
    if stored not in (None, ""):
        stored_value = _bounded_money(stored, label="Giá dịch vụ")
    needs_catalog = source == "tour_import" or stored_value is None or (not source and stored_value == 0)
    if needs_catalog:
        if catalog_price is None:
            raise HTTPException(409, f"Dịch vụ '{entry.get('service') or ''}' chưa có giá trong danh mục.")
        return catalog_price
    return int(stored_value)


def _booking(state: dict[str, Any], payload: dict[str, Any], now: datetime) -> dict[str, Any]:
    employee = _employee(state, payload.get("employee_id"))
    if employee.get("roster_eligible") is False:
        raise HTTPException(409, "Chỉ xếp tua cho Leader/Nhân viên đang làm việc trong danh sách nhân viên.")
    if _norm(employee.get("work_status")) != "di lam":
        raise HTTPException(409, "Chỉ có thể xếp tua cho nhân viên đang Đi làm.")
    if _shift_bucket(employee) not in {"ca1", "ca2"}:
        raise HTTPException(409, "Nhân viên phải được xếp Ca 1 hoặc Ca 2.")
    if employee.get("break_started_at"):
        raise HTTPException(409, "Nhân viên đang nghỉ giữa ca nên chưa thể xếp tua.")
    if employee.get("service") or _norm(employee.get("status")) in {"dang cho", "dang thuc hien", "cho thanh toan"}:
        raise HTTPException(409, "Nhân viên đang có dịch vụ chưa kết thúc thanh toán.")
    room_item = _catalog_item(state, "rooms", payload)
    room = str((room_item or {}).get("name") or payload.get("room") or "").strip()
    if not room:
        raise HTTPException(400, "Vui lòng chọn phòng/giường.")
    if room_item is None:
        raise HTTPException(400, f"Phòng/giường '{room}' chưa có trong danh mục.")
    if room_item.get("active") is False:
        raise HTTPException(400, "Vị trí phục vụ đã ngừng sử dụng.")
    request, auto_request = _auto_yc_ca1(now, employee, payload.get("request"), bool(payload.get("auto_yc_ca1")))
    service, duration, price, service_items = _service_selection(state, {**payload, "request": request}, now)
    _check_room_collision(state, employee, room, service)
    customer = None
    if any(payload.get(key) for key in ("customer_id", "customer_name", "customer_phone", "phone")):
        customer = _customer(state, payload)
    combo = _booking_combo(state, customer, payload, {"service": service, "service_items": service_items}, now)
    employee.pop("last_assignment_display", None)
    employee.update({
        **combo,
        "appointment": str(payload.get("appointment", employee.get("appointment")) or "").strip(),
        "service": service, "service_price": price, "service_items": service_items,
        "service_price_source": "catalog",
        "request": request, "request_source": "auto_yc_ca1" if auto_request else "manual",
        "room": room,
        "status": "Đang chờ", "duration": duration, "booked_at": _iso(now),
        "started_at": "", "completed_at": "", "wait_minutes": None,
        "completion_note": "", "completion_delta_minutes": None, "steam_elapsed_minutes": None,
        "payment_status": "", "customer_id": str((customer or {}).get("id") or ""),
        "customer_name": str((customer or {}).get("name") or ""),
        "customer_phone": str((customer or {}).get("phone") or ""),
        "note": str(payload.get("note", employee.get("note")) or ""),
        "vip": bool(payload.get("vip", employee.get("vip", False))),
    })
    if payload.get("start_now"):
        _start_employee(state, employee, now)
    return employee


def _capture_tour_position(employee: dict[str, Any], now: datetime, state: dict | None = None) -> dict[str, Any]:
    record = _employee_record(employee, now)
    return {
        "board_index": _ordered_employees(state["employees"], now).index(employee) if state else None,
        "sort_index": employee.get("sort_index", 0),
        "manual_order": any(row.get("manual_order") for row in state["employees"]) if state else False,
        "display": {column: record.get(column, "") for column in ("TG bắt đầu thực hiện", "TG bắt đầu thực hiện YC")},
        "counter_key": "request_count" if _norm(employee.get("request")) == "yc" else "tour_count",
        "counter_day": _counter_business_date(now).isoformat(),
    }


def _employee_change_until(employee: dict[str, Any]) -> datetime | None:
    started = _parse_datetime(employee.get("started_at"))
    if (_norm(employee.get("status")) != "dang thuc hien" or not started
            or employee.get("completed_at") or not employee.get("pre_start_tour_position")):
        return None
    return started + timedelta(minutes=employee.get("employee_change_minutes", 10))


def _start_employee(state: dict[str, Any], employee: dict[str, Any], now: datetime) -> None:
    if _norm(employee.get("work_status")) != "di lam" or _shift_bucket(employee) not in {"ca1", "ca2"}:
        raise HTTPException(409, "Nhân viên phải đang Đi làm và được xếp Ca 1/Ca 2.")
    if employee.get("break_started_at"):
        raise HTTPException(409, "Nhân viên đang nghỉ giữa ca nên chưa thể bắt đầu dịch vụ.")
    if not employee.get("service") or not employee.get("room"):
        raise HTTPException(400, "Nhân viên chưa được xếp dịch vụ và phòng.")
    if _norm(employee.get("status")) != "dang cho":
        raise HTTPException(409, "Chỉ dịch vụ ở trạng thái Đang chờ mới được bắt đầu.")
    _check_room_collision(state, employee, str(employee["room"]), str(employee["service"]))
    if employee.get("combo_purchase_id"):
        customer = _customer(state, {"customer_id": employee.get("customer_id")}, create=False)
        employee.update(_booking_combo(state, customer, {}, employee, now, employee))
    employee["pre_start_tour_position"] = _capture_tour_position(employee, now, state)
    if _norm(employee.get("request")) == "yc":
        employee["request_count"] = int(employee.get("request_count") or 0) + 1
    else:
        employee["tour_count"] = int(employee.get("tour_count") or 0) + 1
        # A manual move lasts until the next standard tour starts. Clear the
        # whole board override so API, browser and exports use start-time order.
        for row in state["employees"]:
            row.pop("manual_order", None)
    employee["status"] = "Đang thực hiện"
    employee["started_at"] = _iso(now)
    employee["employee_change_minutes"] = (state.get("payment_settings") or {}).get("employee_change_minutes", 10)
    booked_at = _parse_datetime(employee.get("booked_at"))
    employee["wait_minutes"] = max(0, int(round((now.astimezone(VN_TZ) - booked_at).total_seconds() / 60))) if booked_at else 0
    employee["completed_at"] = ""
    employee["payment_status"] = ""


def _customer(
    state: dict[str, Any], payload: dict[str, Any], create: bool = True,
    *, allow_identity_update: bool = False,
) -> dict[str, Any] | None:
    customer_id = str(payload.get("customer_id") or "").strip()
    phone = str(payload.get("customer_phone") or payload.get("phone") or "").strip()
    name = str(payload.get("customer_name") or payload.get("name") or "").strip()
    phone_key = _phone_key(phone)
    by_id = next((item for item in state["customers"] if customer_id and str(item.get("id")) == customer_id), None)
    by_phone = next((item for item in state["customers"] if phone_key and not item.get("deleted_at") and _phone_key(item.get("phone")) == phone_key), None)
    if by_id and by_id.get("deleted_at"):
        raise HTTPException(409, "Khách hàng đã xóa, không thể tạo giao dịch mới.")
    if customer_id and by_id is None:
        raise HTTPException(404, "Không tìm thấy khách hàng đã chọn.")
    if by_id is not None and by_phone is not None and by_id is not by_phone:
        raise HTTPException(409, "Số điện thoại đã thuộc về một khách hàng khác.")
    found = by_id or by_phone
    if found is None and not create:
        return None
    if found is None:
        if not name and not phone:
            return None
        found = {"id": str(uuid4()), "name": name, "phone": phone, "combo_purchases": [], "created_at": _iso(datetime.now(VN_TZ))}
        state["customers"].append(found)
    else:
        current_name = str(found.get("name") or "").strip()
        current_phone = str(found.get("phone") or "").strip()
        if name and current_name and _norm(name) != _norm(current_name) and not allow_identity_update:
            raise HTTPException(409, "Không được đổi tên khách hàng trong thao tác giao dịch.")
        if phone and current_phone and _phone_key(phone) != _phone_key(current_phone) and not allow_identity_update:
            raise HTTPException(409, "Không được đổi số điện thoại khách hàng trong thao tác giao dịch.")
        if name and (allow_identity_update or not current_name):
            found["name"] = name
        if phone and (allow_identity_update or not current_phone):
            found["phone"] = phone
    found.setdefault("combo_purchases", [])
    return found


def _next_bill_no(state: dict[str, Any], payload: dict[str, Any], now: datetime, *, calendar_day: bool = False) -> str:
    manual = str(payload.get("bill_no") or "").strip()
    # Voiding never frees a document number, including manually supplied bills.
    issued = state["invoices"] + [row["before"] for row in state.get("invoice_changes", []) if isinstance(row.get("before"), dict)]
    existing = {_norm(item.get("bill_no")) for item in issued if item.get("bill_no")}
    if manual:
        if _norm(manual) in existing:
            raise HTTPException(409, f"Số hóa đơn '{manual}' đã tồn tại.")
        return manual
    day = (now.astimezone(VN_TZ).date() if calendar_day else _business_date(now)).isoformat()
    counters = state.setdefault("bill_counters", {})
    current = int(counters.get(day) or 0)
    prefix = f"VERA-{day.replace('-', '')}-"
    for invoice in issued:
        bill_no = str(invoice.get("bill_no") or "")
        for known_prefix in (prefix, prefix.replace("VERA-", "LIVE-", 1)):
            if bill_no.startswith(known_prefix) and bill_no[len(known_prefix):].isdigit():
                current = max(current, int(bill_no[len(known_prefix):]))
    while True:
        current += 1
        candidate = f"{prefix}{current:04d}"
        if _norm(candidate) not in existing:
            counters[day] = current
            return candidate


def _service_ticket_units(state: dict[str, Any], service_name: Any) -> int:
    name = str(service_name or "").strip()
    if not name:
        return 0
    exact = next((item for item in state["services"] if _norm(item.get("name")) == _norm(name)), None)
    if exact is not None:
        unit_value = 1 if exact.get("ticket_units") in (None, "") else exact.get("ticket_units")
        return int(_bounded_number(
            unit_value, label="Số vé trừ theo dịch vụ", minimum=0,
            maximum=MAX_TICKET_UNITS, integer=True,
        ))
    parts = [part.strip() for part in re.split(r"\s*&\s*", name) if part.strip()]
    if len(parts) > 1:
        return sum(_service_ticket_units(state, part) for part in parts)
    # VBA/custom services consume one ticket unless the catalog explicitly
    # maps them to zero (for example Xông hơi or Mua thêm 30).
    return 1


def _quick_booking_at(booking, now):
    if not isinstance(booking, dict) or set(booking) - {"employee_id", "room", "service_items", "booked_at", "correction_reason"}:
        raise HTTPException(400, "Thông tin nhập thanh toán nhanh không hợp lệ.")
    raw = booking.get("booked_at")
    booked = _parse_datetime(raw)
    if not isinstance(raw, str) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}(?::\d{2})?(?:Z|[+-]\d{2}:\d{2})", raw) or not booked or not 2000 <= booked.year <= 2100:
        raise HTTPException(400, "Hãy chọn ngày và giờ booking hợp lệ.")
    if booked > now.astimezone(VN_TZ):
        raise HTTPException(400, "Không thể thanh toán booking trong tương lai.")
    if booked.date() < now.astimezone(VN_TZ).date() and len(str(booking.get("correction_reason") or "").strip()) < 3:
        raise HTTPException(400, "Nhập booking trước hôm nay cần lý do điều chỉnh.")
    return booked


def _quick_combo_selection(state, payload, booked):
    """Resolve owned, unreserved entitlements; never infer services from a name."""
    if _canonical_payment_method(payload.get("payment_method"), quick=True) != "COMBO":
        raise HTTPException(400, "Dùng dịch vụ từ combo cần phương thức COMBO.")
    customer = _customer(state, payload, create=False)
    if not customer:
        raise HTTPException(400, "Phải chọn khách hàng khi dùng combo.")
    purchase = next((row for row in customer.get("combo_purchases", [])
                     if row.get("id") == payload.get("combo_purchase_id")), None)
    if purchase is None:
        raise HTTPException(404, "Không tìm thấy combo đã mua của khách hàng.")
    available = _available_combo(state, customer["id"], purchase)
    require_available(available, booked.date(), "Combo đã mua")
    if available["remaining"] <= 0:
        raise HTTPException(409, "Combo đã hết lượt hoặc các lượt còn lại đã được giữ chỗ.")
    if "component_balances" not in available:
        # Imported generic tickets do not identify a catalog service. Record one
        # prepaid redemption explicitly instead of inventing a service/price.
        return f"Sử dụng combo: {purchase.get('combo_name') or 'Combo vé'}", None, 0, []
    requested = []
    for part in available["component_balances"]:
        if part.get("remaining", 0) <= 0:
            continue
        service = next((row for row in state["services"] if row["id"] == part["service_id"]), None)
        if service is None:
            continue
        try:
            require_available(service, booked.date(), "Dịch vụ trong combo")
        except HTTPException:
            continue
        requested.append({"service_id": service["id"], "quantity": 1})
    if not requested:
        raise HTTPException(409, "Combo không còn dịch vụ khả dụng trong ngày booking.")
    # Same default as booking: one use per available component, not its entire
    # balance. The browser previews every component and the total debit.
    selection = _service_selection(state, {"service_items": requested}, booked)
    plan = component_debits(available, [{"service_items": selection[3]}], state["services"], booked.date())
    if sum(part["units"] for part in plan) > available["remaining"]:
        raise HTTPException(409, "Combo không đủ lượt chưa giữ chỗ để thanh toán.")
    return selection


def _quick_booking_entry(state, booking, now, checkout_payload=None):
    booked = _quick_booking_at(booking, now)
    checkout_payload = checkout_payload or {}
    from_combo = booking.get("service_items") in (None, []) and checkout_payload.get("combo_purchase_id")
    service, duration, price, items = (_quick_combo_selection(state, checkout_payload, booked)
                                      if from_combo else _service_selection(state, booking, booked))
    if items and all(re.match(r"^xong hoi(?:\b|$)", _norm(item['name'])) for item in items):
        return {"employee_id": "", "employee_name": "", "room": "", "service": service,
                "service_items": items, "duration": duration, "price": price,
                "price_source": "catalog", "booked_at": _iso(booked), "request": ""}
    if not booking.get("booked_at"):
        raise HTTPException(400, "Hãy chọn ngày và giờ booking hợp lệ.")
    employee = _employee(state, booking.get("employee_id"))
    if employee.get("hidden") or employee.get("roster_eligible") is False:
        raise HTTPException(409, "Chọn nhân viên đang có trong danh sách phục vụ.")
    room = _catalog_item(state, "rooms", booking)
    if not room or room.get("active") is False:
        raise HTTPException(400, "Hãy chọn phòng/giường đang sử dụng trong danh mục.")
    return {"employee_id": employee["id"], "employee_name": employee["name"],
            "room": room["name"], "service": service, "service_items": items,
            "duration": duration, "price": price, "price_source": "combo_ticket" if from_combo and not items else "catalog",
            "booked_at": _iso(booked), "request": ""}


def _checkout_mutating(
    state: dict[str, Any], payload: dict[str, Any], actor: str, now: datetime,
    quick: bool, timing: dict[str, Any],
) -> dict[str, Any]:
    pending_id = str(payload.get("pending_id") or "")
    direct_ids = payload.get("employee_ids") or ([payload.get("employee_id")] if payload.get("employee_id") else [])
    if pending_id and any(direct_ids):
        raise HTTPException(400, "Không được trộn nguồn chờ thanh toán với dòng nhân viên trực tiếp.")
    if len({str(item) for item in direct_ids}) != len(direct_ids):
        raise HTTPException(400, "Danh sách nhân viên thanh toán bị trùng.")
    pending = next((item for item in state["pending"] if str(item.get("id")) == pending_id), None) if pending_id else None
    if pending_id and pending is None:
        raise HTTPException(404, "Không tìm thấy khoản chờ thanh toán.")
    employees = [_employee(state, item_id) for item_id in direct_ids]
    manual_booking = payload.get("quick_booking") if "quick_booking" in payload else None
    if "quick_booking" in payload:
        if not quick or pending_id or direct_ids:
            raise HTTPException(400, "Không được trộn nhập thanh toán nhanh với hóa đơn hoặc dòng đang chờ thanh toán.")
        entries = [_quick_booking_entry(state, manual_booking, now, payload)]
    elif pending:
        entries = deepcopy(list(pending.get("entries") or []))
    else:
        if not employees:
            raise HTTPException(400, "Chưa chọn dòng nhân viên hoặc khoản chờ thanh toán.")
        if any(_norm(item.get("status")) != "cho thanh toan" for item in employees):
            raise HTTPException(409, "Chỉ thanh toán dịch vụ đã Hoàn thành/CHO THANH TOÁN.")
        entries = [{
            "employee_id": item.get("id"), "employee_name": item.get("name"), "booked_at": item.get("booked_at", ""),
            "started_at": item.get("started_at", ""), "completed_at": item.get("completed_at", ""),
            "combo_purchase_id": item.get("combo_purchase_id", ""),
            "combo_reserved_units": item.get("combo_reserved_units", 0),
            "combo_reserved_components": deepcopy(item.get("combo_reserved_components", [])),
            "service": item.get("service"), "room": item.get("room"), "service_items": deepcopy(item.get("service_items", [])),
            "request": item.get("request"), "duration": item.get("duration"),
            "completion_delta_minutes": item.get("completion_delta_minutes"),
            "completion_note": item.get("completion_note", ""),
            "price": _resolved_service_price(state, item),
            "price_source": str(item.get("service_price_source") or "catalog_reconciled"),
        } for item in employees]
    if not entries:
        raise HTTPException(400, "Không có dịch vụ để thanh toán.")
    timing = _booking_timing(entries, now, pending)
    if manual_booking:
        timing["correction_reason"] = str(manual_booking.get("correction_reason") or "").strip()
    for entry in entries:
        entry["price"] = _resolved_service_price(state, entry)
        entry.setdefault("price_source", "catalog_reconciled")

    combo_purchase_id = str(payload.get("combo_purchase_id") or "").strip()
    payment_method = _canonical_payment_method(payload.get("payment_method"), quick=quick)
    pays_with_combo = payment_method == "COMBO"
    if bool(combo_purchase_id) != pays_with_combo:
        raise HTTPException(
            400,
            "Phương thức COMBO phải đi cùng combo đã mua; tiền mặt/thẻ không được trừ vé combo.",
        )
    customer_payload = dict(payload)
    if pending:
        canonical_identity = {
            "customer_id": str(pending.get("customer_id") or ""),
            "customer_name": str(pending.get("customer_name") or ""),
            "customer_phone": str(pending.get("customer_phone") or ""),
        }
    else:
        canonical_identity = _common_customer_identity(employees)
    customer_payload = _protect_customer_identity(customer_payload, canonical_identity)
    customer = _customer(state, customer_payload)
    subtotal = sum(int(item.get("price") or 0) for item in entries)
    if subtotal > MAX_MONEY:
        raise HTTPException(400, f"Tạm tính không được vượt quá {MAX_MONEY}.")
    manual_subtotal_raw = payload.get("ticket_price", payload.get("subtotal_override"))
    pricing_source = "service_catalog"
    zero_price_indexes = [index for index, item in enumerate(entries) if int(item.get("price") or 0) == 0]
    if pays_with_combo:
        if manual_subtotal_raw not in (None, ""):
            raise HTTPException(400, "Thanh toán COMBO không nhận giá vé nhập tay.")
        pricing_source = "combo_debit"
    elif zero_price_indexes and subtotal > 0:
        raise HTTPException(
            400,
            "Không thể gộp dịch vụ đã có giá với dịch vụ giá 0; hãy cấu hình giá hoặc tách thanh toán.",
        )
    elif subtotal == 0:
        if manual_subtotal_raw in (None, ""):
            raise HTTPException(400, "Dịch vụ chưa có giá danh mục; vui lòng nhập giá vé.")
        subtotal = _bounded_money(manual_subtotal_raw, label="Giá vé")
        if subtotal <= 0 or subtotal > 1_000_000_000:
            raise HTTPException(400, "Giá vé phải lớn hơn 0 và không vượt quá 1000000000.")
        pricing_source = "manual"
        base, remainder = divmod(subtotal, len(entries))
        for index, entry in enumerate(entries):
            entry["price"] = base + (1 if index < remainder else 0)
            entry["price_source"] = "manual"
    elif manual_subtotal_raw not in (None, ""):
        raise HTTPException(400, "Dịch vụ đã có giá danh mục nên không được ghi đè giá vé.")
    discount, tip, payment_details = _payment_values(state, payload, subtotal, _bounded_money, MAX_MONEY)
    default_total = max(0, subtotal - discount) + tip
    if default_total > MAX_MONEY:
        raise HTTPException(400, f"Tổng thanh toán không được vượt quá {MAX_MONEY}.")
    total = default_total

    combo_units = sum(_entry_ticket_units(state, item) for item in entries) if combo_purchase_id else 0
    combo_purchase = None
    debit_plan = []
    if combo_purchase_id:
        if not customer:
            raise HTTPException(400, "Phải chọn khách hàng khi thanh toán bằng combo.")
        combo_purchase = next((item for item in customer["combo_purchases"] if str(item.get("id")) == combo_purchase_id), None)
        if not combo_purchase:
            raise HTTPException(404, "Không tìm thấy combo đã mua của khách hàng.")
        available_combo = _available_combo(state, customer["id"], combo_purchase,
                                          employee_ids=direct_ids, pending_id=pending_id)
        debit_plan = component_debits(available_combo, entries, state["services"], timing["effective_datetime"].astimezone(VN_TZ).date())
        if "component_balances" in combo_purchase:
            combo_units = sum(item["units"] for item in debit_plan)
            if discount:
                raise HTTPException(400, "Combo dịch vụ đã thanh toán khi mua; không áp dụng giảm giá lần nữa khi dùng lượt.")
            total = tip  # Service revenue was recorded at purchase, not at redemption.
        if combo_units > MAX_TICKET_UNITS:
            raise HTTPException(400, f"Số lượt combo trừ không được vượt quá {MAX_TICKET_UNITS}.")
        remaining = int(combo_purchase.get("remaining") or 0)
        if combo_units <= 0 or available_combo["remaining"] < combo_units:
            raise HTTPException(409, f"Combo chỉ còn {available_combo['remaining']} vé chưa giữ chỗ cho booking khác, không đủ trừ {combo_units} vé.")
        for debit in debit_plan:
            balance = next(row for row in combo_purchase["component_balances"] if row["service_id"] == debit["service_id"])
            balance["used"] += debit["units"]
            balance["remaining"] -= debit["units"]
        combo_purchase["used"] = int(combo_purchase.get("used") or 0) + combo_units
        combo_purchase["remaining"] = remaining - combo_units
        combo_purchase["updated_at"] = _iso(now)

    invoice_id = str(uuid4())
    invoice = {
        "id": invoice_id, "business_date": timing["business_date"],
        "created_at": timing["recorded_at"], "recorded_at": timing["recorded_at"],
        "effective_at": timing["effective_at"],
        "backdate_one_day": timing["backdate_one_day"],
        "correction_reason": timing["correction_reason"],
        "actor": actor, "customer_id": (customer or {}).get("id", ""),
        "customer_name": (customer or {}).get("name", customer_payload.get("customer_name", "")),
        "customer_phone": (customer or {}).get("phone", customer_payload.get("customer_phone", "")),
        "bill_no": _next_bill_no(state, payload, timing["effective_datetime"], calendar_day=True),
        "ticket_no": str(payload.get("ticket_no") or ""),
        "payment_method": payment_method,
        "subtotal": subtotal, "discount": discount, "tip": tip, "total": total, **payment_details,
        "pricing_source": pricing_source,
        "combo_purchase_id": combo_purchase_id, "combo_units": combo_units,
        "combo_units_source": "server_purchase_components" if combo_purchase is not None and "component_balances" in combo_purchase else "server_service_catalog",
        "combo_covered_amount": subtotal if combo_purchase is not None and "component_balances" in combo_purchase else 0,
        "combo_component_debits": deepcopy(debit_plan),
        "entries": entries, "note": str(payload.get("note") or ""), "quick": quick,
        **({"source": "quick_booking"} if manual_booking else {}),
    }
    state["invoices"].append(invoice)
    if combo_purchase is not None:
        state["combo_usage"].append({
            "id": str(uuid4()), "invoice_id": invoice_id,
            "business_date": invoice["business_date"], "created_at": invoice["created_at"],
            "recorded_at": invoice["recorded_at"], "effective_at": invoice["effective_at"],
            "backdate_one_day": invoice["backdate_one_day"],
            "correction_reason": invoice["correction_reason"],
            "actor": actor, "customer_id": invoice["customer_id"],
            "combo_purchase_id": combo_purchase_id, "units": combo_units,
            "remaining_before": remaining, "remaining_after": int(combo_purchase.get("remaining") or 0),
            "component_debits": deepcopy(debit_plan),
            "entries": [{"service": item["service_name"], "service_id": item["service_id"], "units": item["units"]} for item in debit_plan] if "component_balances" in combo_purchase else [
                {"service": item.get("service", ""), "units": _entry_ticket_units(state, item)}
                for item in entries
            ],
        })
    for index, entry in enumerate(entries):
        allocated_tip = tip // len(entries) + (1 if index < tip % len(entries) else 0)
        allocated_total = total // len(entries) + (1 if index < total % len(entries) else 0)
        state["reports"].append({
            "id": str(uuid4()), "invoice_id": invoice_id, "business_date": invoice["business_date"],
            "created_at": invoice["created_at"], "recorded_at": invoice["recorded_at"],
            "effective_at": invoice["effective_at"],
            "backdate_one_day": invoice["backdate_one_day"],
            "correction_reason": invoice["correction_reason"], "actor": actor,
            "employee_id": entry.get("employee_id", ""), "employee_name": entry.get("employee_name", ""),
            "service": entry.get("service", ""), "room": entry.get("room", ""),
            "request": entry.get("request", ""), "customer_id": invoice["customer_id"],
            "completion_delta_minutes": entry.get("completion_delta_minutes"),
            "completion_note": entry.get("completion_note", ""),
            "customer_name": invoice["customer_name"], "customer_phone": invoice["customer_phone"],
            "bill_no": invoice["bill_no"], "ticket_no": invoice["ticket_no"],
            "payment_method": invoice["payment_method"], "tip": allocated_tip,
            "total": allocated_total, "note": invoice["note"],
            "combo_units": (
                sum(item["units"] for item in component_debits(combo_purchase, [entry], state["services"], timing["effective_datetime"].astimezone(VN_TZ).date(), check_balance=False))
                if combo_purchase is not None and "component_balances" in combo_purchase
                else _entry_ticket_units(state, entry) if combo_purchase_id else 0
            ),
        })
    if pending:
        state["pending"] = [item for item in state["pending"] if str(item.get("id")) != pending_id]
        for employee in state["employees"]:
            retained = employee.get("last_assignment_display") or {}
            if retained.get("pending_id") == pending_id and not employee.get("service") and not employee.get("status"):
                retained["TT thanh toán"] = "ĐÃ THANH TOÁN"
    else:
        for employee in employees:
            employee["payment_status"] = "ĐÃ THANH TOÁN"
            _clear_assignment(employee, now)
    return invoice


def _checkout(
    state: dict[str, Any], payload: dict[str, Any], actor: str, now: datetime,
    quick: bool, timing: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Apply invoice/report/combo/customer changes as one in-memory unit."""
    if _backdate_requested(payload):
        raise HTTPException(400, "Hóa đơn lấy ngày giờ booking. Hãy dùng Sửa hóa đơn để điều chỉnh ngày giờ.")
    financial_timing = timing or _financial_timing(payload, now)
    snapshot = deepcopy(state)
    try:
        return _checkout_mutating(state, payload, actor, now, quick, financial_timing)
    except Exception:
        state.clear()
        state.update(snapshot)
        raise


def _snapshot_for_backup(state: dict[str, Any]) -> dict[str, Any]:
    snapshot = deepcopy(state)
    snapshot["backups"] = []
    snapshot["audit"] = []
    return snapshot


def _change_pending(state, action, payload, actor, now):
    """Edit/void only an unpaid draft; preserve a complete, append-only snapshot.

    Customer/staff identity and completion counters are historical facts here.
    Service selections, line prices and notes may be corrected. Nothing debits
    tickets or rewrites a paid invoice until the normal checkout transaction.
    """
    allowed = {"pending_id", "reason", "note", "entries", "invoice_at"} if action == "pending_update" else {"pending_id", "reason"}
    if set(payload) - allowed:
        raise HTTPException(400, "Chỉ được sửa dịch vụ, giá và ghi chú của hóa đơn chờ thanh toán.")
    reason = payload.get("reason")
    if not isinstance(reason, str) or not reason.strip() or len(reason) > 1000:
        raise HTTPException(400, "Nhập lý do sửa/xóa hóa đơn (tối đa 1000 ký tự).")
    working = deepcopy(state)
    pending = _find_by_id(working["pending"], payload.get("pending_id"), "hóa đơn chờ thanh toán")
    before = deepcopy(pending)
    if action == "pending_delete":
        working["pending"].remove(pending)
        after = None
    else:
        if "invoice_at" in payload:
            pending.update(_invoice_date(payload["invoice_at"]))
        if "note" in payload:
            if not isinstance(payload["note"], str) or len(payload["note"]) > 2000:
                raise HTTPException(400, "Ghi chú không hợp lệ hoặc vượt quá 2000 ký tự.")
            pending["note"] = payload["note"].strip()
        edits = payload.get("entries", [])
        if not isinstance(edits, list) or len(edits) > len(pending["entries"]):
            raise HTTPException(400, "Danh sách dòng hóa đơn không hợp lệ.")
        entries = deepcopy(pending["entries"])
        seen, changed_services = set(), set()
        for edit in edits:
            if not isinstance(edit, dict) or set(edit) - {"index", "service_items", "price"}:
                raise HTTPException(400, "Nội dung dòng hóa đơn không hợp lệ.")
            index = edit.get("index")
            if type(index) is not int or not 0 <= index < len(entries) or index in seen:
                raise HTTPException(400, "Dòng hóa đơn không tồn tại hoặc bị trùng.")
            seen.add(index)
            entry = entries[index]
            if "service_items" in edit:
                requested = edit["service_items"]
                if not isinstance(requested, list) or not 1 <= len(requested) <= 30 or any(not isinstance(row, dict) or isinstance(row.get("quantity"), bool) for row in requested):
                    raise HTTPException(400, "Chọn từ 1 đến 30 dịch vụ với số lượng hợp lệ.")
                current = [{"service_id": row["service_id"], "quantity": row.get("quantity", 1)} for row in entry.get("service_items", [])]
                if edit["service_items"] != current:
                    name, duration, price, items = _service_selection(working, {**edit, "request": entry.get("request", "")}, now)
                    entry.update(service=name, duration=duration, price=price, service_items=items, price_source="catalog")
                    changed_services.add(index)
            if "price" in edit:
                entry["price"] = _bounded_money(edit["price"], label="Giá dòng hóa đơn")
                entry["price_source"] = "pending_correction"
        _bounded_money(sum(int(row.get("price") or 0) for row in entries), label="Tổng hóa đơn")
        # Retain all unchanged reservations first. Revalidate changed services
        # against that ledger and other drafts, never against their old selves.
        pending["entries"] = [row for index, row in enumerate(entries) if index not in changed_services]
        customer = next((row for row in working["customers"] if row["id"] == pending.get("customer_id")), None)
        for index in sorted(changed_services):
            entry = entries[index]
            entry.update(_booking_combo(working, customer, {"combo_purchase_id": entry.get("combo_purchase_id", "")}, entry, now))
            pending["entries"].append(entry)
        pending["entries"] = entries
        pending.update(updated_at=_iso(now), updated_by=actor)
        after = deepcopy(pending)
    change = {"id": str(uuid4()), "pending_id": before["id"], "action": action,
              "actor": actor, "at": _iso(now), "reason": reason.strip(), "before": before, "after": after}
    working.setdefault("pending_changes", []).append(change)
    state.clear()
    state.update(working)
    return {"pending": after, "pending_id": before["id"], "change_id": change["id"], "deleted": after is None}


def _required_action_feature(action: str) -> str:
    if action in {"reorder", "admin_reorder"}:
        return "live_tour_reorder"
    if action == "update_appointment":
        return "live_tour_view"
    if action in {"report_invoice_update", "report_invoice_delete"}:
        return "live_tour_reports_edit" if action.endswith("update") else "live_tour_reports_delete"
    if action in {"customer_delete", "customer_combo_update", "customer_combo_delete"}:
        return {"customer_delete": "live_tour_customers_delete", "customer_combo_update": "live_tour_customer_combo_edit", "customer_combo_delete": "live_tour_customer_combo_delete"}[action]
    if action in {"paid_invoice_update", "paid_invoice_delete"}:
        return "live_tour_paid_invoice_edit" if action == "paid_invoice_update" else "live_tour_paid_invoice_delete"
    if action in {"booking", "multi_booking"}:
        return "live_tour_booking"
    if action in {"pending_update", "pending_delete"}:
        return "live_tour_invoice_edit" if action == "pending_update" else "live_tour_invoice_delete"
    if action in {"backup", "restore"}:
        return "live_tour_backup"
    if action in {"checkout", "quick_checkout", "move_pending", "combo_purchase", "customer_upsert"}:
        return "live_tour_payment"
    if action in {
        "room_upsert", "room_delete", "service_upsert",
        "service_delete", "combo_upsert", "combo_delete", "backup", "restore",
        "clear_expired", "clear_expired_preview", "combo_import", "set_vip", "service_area_upsert", "service_area_delete", "payment_settings_update",
    }:
        return "live_tour_admin"
    return "live_tour_operate"


def _reject_external_action(action: str) -> None:
    if action in {"sync_leaves", "merge_current_tour", "merge_current_tour_preview"}:
        raise HTTPException(410, detail={
            "code": "LIVE_TOUR_SERVER_ONLY",
            "message": "Live Tour chỉ sử dụng dữ liệu trên máy chủ. Kết nối file và đồng bộ nguồn ngoài đã được gỡ bỏ.",
        })


def _apply_action(state: dict[str, Any], action: str, payload: dict[str, Any], actor: str, now: datetime) -> dict[str, Any]:
    action = str(action or "").strip().lower()
    _reject_external_action(action)
    _ensure_counter_day(state, now)
    financial_timing = _financial_timing(payload, now) if action in BACKDATE_ACTIONS else None
    batch_actions = {
        "start", "add_minutes", "complete", "set_work_status", "set_shift", "start_break", "reorder", "admin_reorder",
        "end_break", "set_vip",
        "replace_service", "add_service", "finish_to_pending", "cancel_booking",
    }
    employee_ids = [str(item) for item in (payload.get("employee_ids") or []) if str(item or "").strip()]
    if len(set(employee_ids)) != len(employee_ids):
        raise HTTPException(400, "Danh sách nhân viên bị trùng.")
    if action in batch_actions and employee_ids:
        working = deepcopy(state)
        results = []
        for employee_id in employee_ids:
            single_payload = {key: value for key, value in payload.items() if key != "employee_ids"}
            single_payload["employee_id"] = employee_id
            results.append(_apply_action(working, action, single_payload, actor, now))
        state.clear()
        state.update(working)
        return {"items": results, "count": len(results)}
    result: dict[str, Any] = {}
    if action in {"start_room", "finish_room"}:
        room = str(payload.get("room") or "").strip()
        known_rooms = {str(area["name"]) for area in _service_areas(state)}
        known_rooms.update(_catalog_room_group(state, row["room"]) for row in state["employees"] if row.get("room"))
        canonical_room = next((name for name in known_rooms if _norm(name) == _norm(room)), None)
        if not room or canonical_room is None:
            raise HTTPException(400, "Hãy chọn phòng/khu vực có trong Live Tour.")
        members = _room_action_members(state, canonical_room, action)
        if not members:
            raise HTTPException(409, "Phòng không còn dịch vụ đang chờ." if action == "start_room" else "Phòng không còn dịch vụ đang thực hiện.")
        # Resolve every bed from the stored state, regardless of table filters or
        # hidden rows. A single failed transition rolls back the entire room.
        result = _apply_action(state, "start" if action == "start_room" else "finish_to_pending",
                               {"employee_ids": [row["id"] for row in members]}, actor, now)
        result["room"] = canonical_room
    elif action == "sync_daily_status":
        result = _sync_daily(state, payload["directory"], payload["leaves"])
        payload = {"updated": result["updated"]}
    elif action == "booking":
        result["employee"] = _booking(state, payload, now)
    elif action == "update_appointment":
        if len(employee_ids) > 1 or (employee_ids and employee_ids[0] != str(payload.get("employee_id") or "")):
            raise HTTPException(400, "Hãy chọn đúng một nhân viên để lưu lịch hẹn.")
        appointment = payload.get("appointment")
        if not isinstance(appointment, str) or len(appointment) > 200 or any(ord(char) < 32 for char in appointment):
            raise HTTPException(400, "Lịch hẹn phải là một dòng, tối đa 200 ký tự.")
        employee = _employee(state, payload.get("employee_id"))
        employee["appointment"] = appointment.strip()
        result["employee"] = employee
    elif action == "change_employee":
        if len(employee_ids) > 1:
            raise HTTPException(400, "Chỉ đổi một nhân viên mỗi lần.")
        source = _employee(state, payload.get("employee_id"))
        target = _employee(state, payload.get("target_employee_id"))
        if source["id"] == target["id"]:
            raise HTTPException(400, "Hãy chọn nhân viên thay thế khác.")
        deadline = _employee_change_until(source)
        started = _parse_datetime(source.get("started_at"))
        if not deadline or not started or not started <= now.astimezone(VN_TZ) <= deadline:
            raise HTTPException(409, f"Chỉ đổi nhân viên trong {source.get('employee_change_minutes', 10)} phút đầu; phiên cần có vị trí tua trước khi bắt đầu.")
        if (target.get("roster_eligible") is False or _norm(target.get("work_status")) != "di lam"
                or _shift_bucket(target) not in {"ca1", "ca2"} or target.get("break_started_at")
                or _has_unsettled_work(target)):
            raise HTTPException(409, "Nhân viên thay thế phải đang đi làm, có ca và đang rảnh.")
        # Keep the original service clock: repeated replacement cannot reopen the
        # ten-minute window. Each employee retains their own pre-service position.
        original_position = deepcopy(source["pre_start_tour_position"])
        target_position = _capture_tour_position(target, now, state)
        counter_key = original_position["counter_key"]
        target_position["counter_key"] = counter_key
        target_position["counter_day"] = original_position["counter_day"]
        fields = ("service", "service_items", "service_price", "service_price_source", "duration",
                  "request", "request_source", "room", "status", "booked_at", "booking_id",
                  "started_at", "employee_change_minutes", "completed_at", "wait_minutes", "payment_status", "completion_note",
                  "completion_delta_minutes", "steam_elapsed_minutes", "customer_id", "customer_name",
                  "customer_phone", "combo_purchase_id", "combo_reserved_units", "combo_reserved_components", "note")
        assignment = {key: deepcopy(source[key]) for key in fields if key in source}
        _clear_assignment(target, now)
        target.pop("last_assignment_display", None)
        target.update(assignment)
        target["pre_start_tour_position"] = target_position
        _clear_assignment(source, now)
        source["sort_index"] = original_position["sort_index"]
        source["last_assignment_display"] = original_position["display"]
        if original_position.get("board_index") is not None and (original_position.get("manual_order") or any(row.get("manual_order") for row in state["employees"])):
            restored_order = [row for row in _ordered_employees(state["employees"], now) if row["id"] != source["id"]]
            restored_order.insert(min(original_position["board_index"], len(restored_order)), source)
            for index, row in enumerate(restored_order):
                row["sort_index"] = index
                row["manual_order"] = True
        if original_position["counter_day"] == _counter_business_date(now).isoformat():
            source[counter_key] = max(0, int(source.get(counter_key) or 0) - 1)
            target[counter_key] = int(target.get(counter_key) or 0) + 1
        source["note"] = ""
        payload = {**payload, "booking_id": assignment.get("booking_id", "")}
        result = {"employee": target, "previous_employee_id": source["id"]}
    elif action == "cancel_booking":
        employee = _employee(state, payload.get("employee_id"))
        if _norm(employee.get("status")) != "dang cho" or employee.get("started_at") or employee.get("completed_at"):
            raise HTTPException(409, "Chỉ hủy Booking đang chờ, chưa thực hiện.")
        booking_id = employee.get("booking_id", "")
        payload = {**payload, "booking_id": booking_id, "room": employee.get("room", ""), "service": employee.get("service", "")}
        _clear_assignment(employee, now)
        employee.pop("last_assignment_display", None)
        result = {"employee": employee, "cancelled_booking_id": booking_id}
    elif action == "update_booking":
        employee = _employee(state, payload.get("employee_id"))
        if _norm(employee.get("status")) not in {"dang cho", "dang thuc hien", "dang su dung"}:
            raise HTTPException(409, "Chỉ sửa booking đang chờ hoặc đang thực hiện.")
        request = str(payload.get("request", employee.get("request", "")))
        if request not in {"", "YC"}:
            raise HTTPException(400, "Yêu cầu chỉ nhận trống hoặc YC.")
        if request != employee.get("request", "") and employee.get("started_at"):
            raise HTTPException(409, "Dịch vụ đã bắt đầu; giữ nguyên loại yêu cầu để không đổi số tua đã ghi nhận.")
        name, duration, price, items = _service_selection(state, {**payload, "request": request}, now)
        room_item = _catalog_item(state, "rooms", {"room": payload.get("room") or employee.get("room")})
        if room_item is None or room_item.get("active") is False:
            raise HTTPException(400, "Hãy chọn phòng/giường đang sử dụng.")
        _check_room_collision(state, employee, room_item["name"], name)
        identity = _common_customer_identity([employee])
        customer_payload = _protect_customer_identity(payload, identity)
        customer = _customer(state, customer_payload) if _contains_customer_pii(customer_payload) else None
        combo = _booking_combo(state, customer, payload, {"service": name, "service_items": items}, now, employee)
        employee.update(service=name, duration=duration, service_price=price, service_price_source="catalog", service_items=items,
                        room=room_item["name"], request=request, note=str(payload.get("note") or ""), **combo)
        if customer:
            employee.update(customer_id=customer["id"], customer_name=customer.get("name", ""), customer_phone=customer.get("phone", ""))
        if payload.get("start_now") and _norm(employee["status"]) == "dang cho":
            _start_employee(state, employee, now)
        result["employee"] = employee
    elif action == "finish_to_pending":
        working = deepcopy(state)
        if "service_items" in payload:
            _apply_action(working, "update_booking", payload, actor, now)
        _apply_action(working, "complete", payload, actor, now)
        pending_result = _apply_action(working, "move_pending", payload, actor, now)
        state.clear()
        state.update(working)
        result = {"pending": pending_result["pending"], "employee": _employee(state, payload.get("employee_id"))}
    elif action == "payment_settings_update":
        state["payment_settings"] = _payment_settings_update(payload, _bounded_money)
        result["payment_settings"] = deepcopy(state["payment_settings"])
    elif action == "multi_booking":
        rows = payload.get("bookings")
        if not rows:
            employee_ids = payload.get("employee_ids") or []
            requested_room = str(payload.get("room") or "")
            room_candidates = [requested_room] + [
                str(item.get("name")) for item in state["rooms"]
                if _room_group(item.get("name")) == _room_group(requested_room)
                and _norm(item.get("name")) != _norm(requested_room)
                and item.get("active", True)
            ]
            rows = [
                {**payload, "employee_id": employee_id, "room": room_candidates[index] if index < len(room_candidates) else requested_room}
                for index, employee_id in enumerate(employee_ids)
            ]
        if not rows:
            raise HTTPException(400, "Chưa chọn nhân viên để xếp tua.")
        booked = []
        # Mutate a working copy first so any collision aborts the complete operation.
        working = deepcopy(state)
        for row in rows:
            row_payload = dict(row)
            row_payload.setdefault("auto_yc_ca1", bool(payload.get("auto_yc_ca1")))
            booked.append(_booking(working, row_payload, now))
        state.clear()
        state.update(working)
        result["employees"] = booked
    elif action == "start":
        employee = _employee(state, payload.get("employee_id"))
        _start_employee(state, employee, now)
        result["employee"] = employee
    elif action == "add_minutes":
        employee = _employee(state, payload.get("employee_id"))
        if _norm(employee.get("status")) != "dang thuc hien" or not employee.get("started_at"):
            raise HTTPException(409, "Chỉ cộng thời gian cho dịch vụ đang thực hiện.")
        minutes = int(_bounded_number(
            payload.get("minutes"), label="Số phút điều chỉnh", minimum=-1_440,
            maximum=1_440, integer=True,
        ))
        if minutes == 0:
            raise HTTPException(400, "Số phút cộng thêm phải khác 0.")
        next_duration = int(employee.get("duration") or 0) + minutes
        if not 0 <= next_duration <= MAX_SERVICE_DURATION_MINUTES:
            raise HTTPException(400, f"Thời lượng sau điều chỉnh phải từ 0 đến {MAX_SERVICE_DURATION_MINUTES} phút.")
        employee["duration"] = next_duration
        result["employee"] = employee
    elif action == "complete":
        employee = _employee(state, payload.get("employee_id"))
        if not employee.get("service") or _norm(employee.get("status")) not in {"dang thuc hien", "dang su dung"}:
            raise HTTPException(409, "Chỉ hoàn thành dịch vụ đang thực hiện.")
        employee["status"] = "CHO THANH TOÁN"
        employee["payment_status"] = "CHO THANH TOÁN"
        employee["completed_at"] = _iso(now)
        started_at = _parse_datetime(employee.get("started_at"))
        elapsed = (now.astimezone(VN_TZ) - started_at).total_seconds() / 60 if started_at else 0
        employee["completion_delta_minutes"] = int(round(elapsed - float(employee.get("duration") or 0)))
        delta = employee["completion_delta_minutes"]
        employee["completion_note"] = (
            f"Sớm {abs(delta)} phút" if delta < 0
            else f"Trễ {delta} phút" if delta > 0
            else "Đúng giờ"
        )
        result["employee"] = employee
    elif action == "move_pending":
        ids = payload.get("employee_ids") or [payload.get("employee_id")]
        if len({str(item) for item in ids if item}) != len([item for item in ids if item]):
            raise HTTPException(400, "Danh sách nhân viên bị trùng.")
        employees = [_employee(state, item_id) for item_id in ids if item_id]
        if not employees:
            raise HTTPException(400, "Chưa chọn dịch vụ chuyển sang chờ thanh toán.")
        if any(_norm(item.get("status")) != "cho thanh toan" for item in employees):
            raise HTTPException(409, "Chỉ chuyển các dịch vụ đã Hoàn thành sang chờ thanh toán.")
        canonical_identity = _common_customer_identity(employees)
        protected_customer = _protect_customer_identity(payload, canonical_identity)
        customer = _customer(state, protected_customer)
        combo_ids = {item.get("combo_purchase_id", "") for item in employees}
        pending = {
            "id": str(uuid4()), "created_at": _iso(now), "business_date": _business_date(now).isoformat(),
            "combo_purchase_id": next(iter(combo_ids)) if len(combo_ids) == 1 else "",
            "customer_id": str((customer or {}).get("id") or canonical_identity.get("customer_id") or ""),
            "customer_name": str((customer or {}).get("name") or canonical_identity.get("customer_name") or ""),
            "customer_phone": str((customer or {}).get("phone") or canonical_identity.get("customer_phone") or ""),
            "entries": [{
                "employee_id": item.get("id"), "employee_name": item.get("name"), "booked_at": item.get("booked_at", ""),
                "started_at": item.get("started_at", ""), "completed_at": item.get("completed_at", ""),
                "combo_purchase_id": item.get("combo_purchase_id", ""),
                "combo_reserved_units": item.get("combo_reserved_units", 0),
                "combo_reserved_components": deepcopy(item.get("combo_reserved_components", [])),
                "service": item.get("service"), "room": item.get("room"), "service_items": deepcopy(item.get("service_items", [])),
                "request": item.get("request"), "duration": item.get("duration"),
                "completion_delta_minutes": item.get("completion_delta_minutes"),
                "completion_note": item.get("completion_note", ""),
                "price": _resolved_service_price(state, item),
                "price_source": str(item.get("service_price_source") or "catalog_reconciled"),
            } for item in employees],
            "note": str(payload.get("note") or ""),
        }
        booking_time = _booking_timing(pending["entries"], now)
        pending.update(effective_at=booking_time["effective_at"], booked_at=booking_time["effective_at"], business_date=booking_time["business_date"])
        state["pending"].append(pending)
        for employee in employees:
            _clear_assignment(employee, now)
            employee["last_assignment_display"]["pending_id"] = pending["id"]
        result["pending"] = pending
    elif action in {"pending_update", "pending_delete"}:
        result = _change_pending(state, action, payload, actor, now)
    elif action in {"paid_invoice_update", "paid_invoice_delete"}:
        result = change_paid_invoice(state, action, payload, actor, now, money=_bounded_money,
                                     payment_values=_payment_values, canonical_method=_canonical_payment_method,
                                     available_combo=_available_combo, iso=_iso, max_money=MAX_MONEY, invoice_date=_invoice_date)
    elif action in {"checkout", "quick_checkout"}:
        result["invoice"] = _checkout(
            state, payload, actor, now, action == "quick_checkout", financial_timing,
        )
    elif action == "set_work_status":
        employee = _employee(state, payload.get("employee_id"))
        next_status = _canonical_work_status(payload.get("work_status") or payload.get("status"))
        if _norm(next_status) != "di lam" and (
            employee.get("service") or _active_booking(employee) or employee.get("break_started_at")
        ):
            raise HTTPException(409, "Hãy hoàn tất dịch vụ/giờ nghỉ trước khi đổi trạng thái đi làm.")
        employee["work_status"] = next_status
        result["employee"] = employee
    elif action == "set_shift":
        employee = _employee(state, payload.get("employee_id"))
        if employee.get("shift_checkin_date"):
            raise HTTPException(409, "Ca được tự động cập nhật theo lịch làm và check-in hôm nay.")
        next_shift = _canonical_shift(payload.get("shift"), allow_blank=True)
        if _norm(employee.get("work_status")) != "di lam":
            raise HTTPException(409, "Chỉ xếp ca cho nhân viên đang Đi làm.")
        if employee.get("break_started_at"):
            raise HTTPException(409, "Không được đổi ca khi nhân viên đang nghỉ giữa ca.")
        if (employee.get("service") or _active_booking(employee)) and next_shift != _canonical_shift(employee.get("shift"), allow_blank=True):
            raise HTTPException(409, "Không được đổi ca khi nhân viên đang có dịch vụ/chưa thanh toán.")
        employee["shift"] = next_shift
        result["employee"] = employee
    elif action == "start_break":
        employee = _employee(state, payload.get("employee_id"))
        if employee.get("break_started_at"):
            raise HTTPException(409, "Nhân viên đang trong giờ nghỉ giữa ca.")
        if _norm(employee.get("work_status")) != "di lam" or _shift_bucket(employee) not in {"ca1", "ca2"}:
            raise HTTPException(409, "Nhân viên phải đang Đi làm và được xếp Ca 1/Ca 2.")
        if _active_booking(employee) or employee.get("service"):
            raise HTTPException(409, "Nhân viên đang có dịch vụ/chưa thanh toán nên chưa thể bắt đầu nghỉ.")
        started_at = _iso(now)
        employee["break_started_at"] = started_at
        employee["clock_out"] = started_at
        employee["clock_in"] = ""
        employee["break_remaining_minutes"] = 90
        start_event = {
            "id": str(uuid4()), "event_type": "start", "at": started_at,
            "created_at": started_at, "business_date": _business_date(now).isoformat(),
            "actor": actor, "employee_id": str(employee.get("id") or ""),
            "employee_name": str(employee.get("name") or ""),
            "started_at": started_at, "allowance_minutes": 90,
            "minutes": 0, "remaining_minutes": 90, "late_minutes": 0,
            "outcome": "Đang nghỉ",
        }
        state.setdefault("break_events", []).append(start_event)
        result["employee"] = employee
        result["break_event"] = start_event
    elif action == "end_break":
        employee = _employee(state, payload.get("employee_id"))
        if not employee.get("break_started_at"):
            raise HTTPException(409, "Nhân viên chưa bắt đầu nghỉ giữa ca.")
        break_started = _parse_datetime(employee.get("break_started_at"))
        if not break_started:
            raise HTTPException(409, "Giờ bắt đầu nghỉ giữa ca không hợp lệ.")
        break_minutes = max(0, int(round((now.astimezone(VN_TZ) - break_started).total_seconds() / 60)))
        outcome = "Đúng giờ" if break_minutes <= 90 else "Quá 90 phút"
        employee["break_started_at"] = ""
        employee["clock_in"] = _iso(now)
        employee["break_count"] = int(employee.get("break_count") or 0) + 1
        employee.setdefault("break_history", []).append({
            "started_at": _iso(break_started), "ended_at": _iso(now),
            "minutes": break_minutes, "limit_minutes": 90, "outcome": outcome,
        })
        employee["break_history"] = employee["break_history"][-100:]
        employee["last_break_minutes"] = break_minutes
        employee["last_break_outcome"] = outcome
        employee["last_break_remaining_minutes"] = max(0, 90 - break_minutes)
        employee["break_remaining_minutes"] = ""
        start_event = next((
            item for item in reversed(state.setdefault("break_events", []))
            if item.get("event_type") == "start"
            and str(item.get("employee_id") or "") == str(employee.get("id") or "")
            and str(item.get("started_at") or "") == _iso(break_started)
        ), None)
        end_event = {
            "id": str(uuid4()), "event_type": "end", "at": _iso(now),
            "created_at": _iso(now), "business_date": _business_date(now).isoformat(),
            "actor": actor, "employee_id": str(employee.get("id") or ""),
            "employee_name": str(employee.get("name") or ""),
            "start_event_id": str((start_event or {}).get("id") or ""),
            "started_at": _iso(break_started), "ended_at": _iso(now),
            "allowance_minutes": 90, "minutes": break_minutes,
            "remaining_minutes": max(0, 90 - break_minutes),
            "late_minutes": max(0, break_minutes - 90), "outcome": outcome,
        }
        state["break_events"].append(end_event)
        result["employee"] = employee
        result["break_event"] = end_event
    elif action in {"reorder", "admin_reorder"}:
        employee = _employee(state, payload.get("employee_id"))
        # Positions refer to visible roster members; retained assignments are not rows.
        retained = [row for row in state["employees"] if row.get("roster_eligible") is False]
        ordered = _ordered_employees([row for row in state["employees"] if row.get("roster_eligible") is not False], now)
        if employee not in ordered:
            raise HTTPException(409, "Nhân viên không còn trong danh sách bảng tua.")
        peers = list(ordered) if action == "admin_reorder" else [item for item in ordered if _employee_time_key(item, now) == _employee_time_key(employee, now)]
        current = peers.index(employee)
        direction = str(payload.get("direction") or "").lower()
        steps = int(_bounded_number(
            payload.get("steps", 1), label="Bước di chuyển", minimum=1,
            maximum=5, integer=True,
        ))
        if steps not in {1, 3, 5}:
            raise HTTPException(400, "Bước di chuyển chỉ nhận 1, 3 hoặc 5.")
        target = 0 if direction == "top" else len(peers) - 1 if direction == "bottom" else current - steps if direction == "up" else current + steps if direction == "down" else current
        if direction == "position" and action == "admin_reorder":
            target = int(_bounded_number(payload.get("position"), label="STT", minimum=1, maximum=len(peers), integer=True)) - 1
        elif direction not in {"top", "bottom", "up", "down"}:
            raise HTTPException(400, "Hướng sắp xếp không hợp lệ.")
        peers.pop(current)
        peers.insert(max(0, min(target, len(peers))), employee)
        peer_ids = {item["id"] for item in peers}
        peer_iter = iter(peers)
        ordered = [next(peer_iter) if item["id"] in peer_ids else item for item in ordered]
        ordered = [item for item in ordered if _norm(item.get("work_status")) != "nghi phep"] + [item for item in ordered if _norm(item.get("work_status")) == "nghi phep"]
        for index, item in enumerate(ordered):
            item["sort_index"] = index
            if action == "admin_reorder":
                item["manual_order"] = True
        state["employees"] = ordered + retained
        result["employee"] = employee
    elif action == "set_vip":
        employee = _employee(state, payload.get("employee_id"))
        employee["vip"] = bool(payload.get("vip", payload.get("enabled", True)))
        result["employee"] = employee
    elif action in {"replace_service", "add_service"}:
        original_employee = _employee(state, payload.get("employee_id"))
        employee = deepcopy(original_employee)
        if _norm(employee.get("status")) not in {"dang cho", "dang thuc hien", "dang su dung"}:
            raise HTTPException(409, "Chỉ sửa dịch vụ đang chờ hoặc đang thực hiện.")
        name, duration, price = _service_values(state, {**payload, "request": employee.get("request")}, now)
        if action == "replace_service" or not employee.get("service"):
            employee["service"] = name
            employee["duration"] = duration
            employee["service_price"] = price
        else:
            existing_price = _resolved_service_price(state, employee)
            employee["service"] = f"{employee['service']} & {name}"
            employee["duration"] = int(employee.get("duration") or 0) + int(duration or 0)
            employee["service_price"] = existing_price + price
        employee.pop("service_items", None)
        employee["service_price_source"] = (
            "catalog" if _service_catalog_price(state, employee.get("service")) is not None else "custom"
        )
        if employee.get("room"):
            _check_room_collision(state, original_employee, str(employee["room"]), str(employee["service"]))
        if employee.get("combo_purchase_id"):
            customer = _customer(state, {"customer_id": employee.get("customer_id")}, create=False)
            employee.update(_booking_combo(state, customer, {}, employee, now, employee))
        original_employee.update(employee)
        original_employee.pop("service_items", None)
        employee = original_employee
        result["employee"] = employee
    elif action in {"customer_delete", "customer_combo_update", "customer_combo_delete"}:
        result = change_customer(state, action, payload, actor, now, iso=_iso,
                                 bounded_number=_bounded_number)
    elif action in {"report_invoice_update", "report_invoice_delete"}:
        result = _apply_action(state, action.replace("report_invoice_", "paid_invoice_"), payload, actor, now)
    elif action == "customer_upsert":
        name = str(payload.get("customer_name") or "").strip()
        phone = str(payload.get("customer_phone") or "").strip()
        if not name or len(name) > 150 or len(phone) > 30:
            raise HTTPException(400, "Tên khách hàng cần từ 1 đến 150 ký tự; điện thoại tối đa 30 ký tự.")
        if phone and not 6 <= len(_phone_key(phone)) <= 15:
            raise HTTPException(400, "Số điện thoại cần từ 6 đến 15 chữ số.")
        customer_id = str(payload.get("customer_id") or "").strip()
        if not customer_id and phone and any(not row.get("deleted_at") and _phone_key(row.get("phone")) == _phone_key(phone) for row in state["customers"]):
            raise HTTPException(409, "Số điện thoại đã tồn tại. Hãy mở khách hàng đó để sửa.")
        before = deepcopy(next((c for c in state["customers"] if c.get("id") == customer_id), None))
        if customer_id:
            reason = payload.get("reason")
            if not isinstance(reason, str) or not reason.strip() or len(reason) > 1000:
                raise HTTPException(400, "Nhập lý do sửa khách hàng (tối đa 1000 ký tự).")
        customer = _customer(state, {"customer_id": customer_id, "customer_name": name, "customer_phone": phone}, allow_identity_update=True)
        customer.update(name=name, phone=phone, updated_at=_iso(now))
        # Keep open transactions usable after contact edits; paid ledgers remain historical.
        open_entries = state["employees"] + state["pending"] + [entry for row in state["pending"] for entry in row.get("entries", [])]
        for entry in open_entries:
            if entry.get("customer_id") == customer["id"]:
                entry.update(customer_name=name, customer_phone=phone)
        if before:
            state.setdefault("customer_changes", []).append({"id": str(uuid4()), "customer_id": customer_id, "action": action, "at": _iso(now), "actor": actor, "reason": payload["reason"].strip(), "before": before, "after": deepcopy(customer)})
        result = {"customer": deepcopy(customer)}
    elif action in {"service_area_upsert", "service_area_delete"}:
        result = _service_area_change(state, payload, delete=action == "service_area_delete")
    elif action in {"room_upsert", "service_upsert", "combo_upsert"}:
        kind = {"room_upsert": "rooms", "service_upsert": "services", "combo_upsert": "combos"}[action]
        singular = kind[:-1]
        incoming = dict(payload.get("item") or payload)
        name = str(incoming.get("name") or incoming.get(singular) or "").strip()
        if not name:
            raise HTTPException(400, "Thiếu tên danh mục.")
        item_id = str(incoming.get("id") or incoming.get(f"{singular}_id") or "")
        current = next((item for item in state[kind] if item_id and str(item.get("id")) == item_id), None)
        if item_id and current is None:
            raise HTTPException(404, "Danh mục đã bị xóa; hãy làm mới danh sách.")
        current = current or next((item for item in state[kind] if _norm(item.get("name")) == _norm(name)), None)
        if incoming.pop("create_only", False) and current is not None:
            raise HTTPException(409, "Tên danh mục đã tồn tại. Hãy mở danh mục đó để sửa.")
        if kind == "rooms" and (current or {}).get("area_id"):
            raise HTTPException(409, "Vị trí này thuộc khu vực dịch vụ. Hãy sửa trong Cài đặt → Cài đặt khu vực dịch vụ.")
        if any(item is not current and _norm(item.get("name")) == _norm(name) for item in state[kind]):
            raise HTTPException(409, "Tên danh mục đã tồn tại.")
        if "active" in incoming and not isinstance(incoming["active"], bool):
            raise HTTPException(400, "Trạng thái sử dụng phải là giá trị đúng/sai.")
        if kind == "rooms" and any(key in incoming for key in ("area_id", "area_name", "area_kind", "bed_name")):
            raise HTTPException(400, "Hãy dùng Cài đặt khu vực dịch vụ để quản lý phòng, giường và bàn.")
        validated_fields: dict[str, Any] = {}
        if kind in {"services", "combos"}:
            validated_fields.update(catalog_details(kind, incoming, current, state["services"]))
        if kind == "services":
            duration_value = _bounded_number(
                incoming.get("duration"), label="Thời lượng dịch vụ", minimum=0,
                maximum=MAX_SERVICE_DURATION_MINUTES, allow_blank=True,
            )
            validated_fields["duration"] = int(round(duration_value)) if duration_value is not None else None
            validated_fields["price"] = _bounded_money(
                incoming.get("price"), label="Giá dịch vụ", allow_blank_as_zero=True,
            )
            unit_source = incoming.get("ticket_units")
            if unit_source in (None, ""):
                unit_source = (current or {}).get("ticket_units", 1)
            validated_fields["ticket_units"] = int(_bounded_number(
                unit_source, label="Số vé trừ theo dịch vụ", minimum=0,
                maximum=MAX_TICKET_UNITS, integer=True,
            ))
            if "request_duration" in incoming:
                validated_fields["request_duration"] = _bounded_number(
                    incoming["request_duration"], label="Thời lượng YC", minimum=0,
                    maximum=MAX_SERVICE_DURATION_MINUTES, allow_blank=True,
                )
            for flag in ("private", "request_eligible", "non_request_eligible"):
                if flag in incoming and not isinstance(incoming[flag], bool):
                    raise HTTPException(400, f"{flag} phải là giá trị đúng/sai.")
        if kind == "combos":
            ticket_source = validated_fields.get("tickets", incoming.get("tickets", incoming.get("quantity", (current or {}).get("tickets"))))
            if ticket_source in (None, ""):
                ticket_source = 0
            validated_fields["tickets"] = int(_bounded_number(
                ticket_source, label="Số vé combo", minimum=1,
                maximum=MAX_TICKET_UNITS, integer=True,
            ))
            validated_fields["price"] = _bounded_money(
                incoming.get("price"), label="Giá combo", allow_blank_as_zero=True,
            )
        if current and kind in {"rooms", "services"} and _catalog_referenced(state, kind, current):
            protected = {"name", "active", "duration", "ticket_units", "private", "request_eligible", "non_request_eligible", "request_duration"}
            defaults = {"active": True, "private": _is_private_service(current.get("name")),
                        "request_eligible": True, "non_request_eligible": True, "ticket_units": 1}
            normalized = {**incoming, **validated_fields, "name": name}
            if any(key in normalized and normalized[key] != current.get(key, defaults.get(key)) for key in protected):
                raise HTTPException(409, "Danh mục đang được dùng bởi dịch vụ chưa thanh toán; chưa thể đổi tên hoặc quy tắc.")
        if current is None:
            current = {"id": str(uuid4())}
            state[kind].append(current)
        current.update({key: value for key, value in incoming.items() if key not in {"id", f"{singular}_id"}})
        current.update(validated_fields)
        current["name"] = name
        current.setdefault("active", True)
        if kind == "rooms":
            current["group"] = _room_group(name)
            inferred_type = (
                "vip" if current["group"].isdigit() and 16 <= int(current["group"]) <= 21
                else "standard"
            )
            current["type"] = inferred_type
            current["is_vip"] = current["type"] == "vip"
        if kind == "services":
            current["private"] = bool(incoming.get("private", _is_private_service(name)))
        result[singular] = current
    elif action in {"room_delete", "service_delete", "combo_delete"}:
        kind = {"room_delete": "rooms", "service_delete": "services", "combo_delete": "combos"}[action]
        singular = kind[:-1]
        item = _find_by_id(state[kind], payload.get(f"{singular}_id") or payload.get("id"), singular)
        if kind == "rooms" and item.get("area_id"):
            raise HTTPException(409, "Hãy quản lý vị trí này trong Cài đặt → Cài đặt khu vực dịch vụ.")
        if kind in {"rooms", "services"} and _catalog_referenced(state, kind, item):
            raise HTTPException(409, "Không thể xóa danh mục còn gắn với dịch vụ chưa thanh toán.")
        if kind == "services":
            component_refs = [part for combo in state["combos"] for part in combo.get("components", [])]
            component_refs += [part for customer in state["customers"] for purchase in customer.get("combo_purchases", []) for part in purchase.get("component_balances", []) if part.get("remaining", 0) > 0]
            if any(part.get("service_id") == item["id"] for part in component_refs):
                raise HTTPException(409, "Dịch vụ đang thuộc combo hoặc còn lượt khách đã mua; chưa thể xóa.")
        state[kind] = [row for row in state[kind] if row is not item]
        result["deleted_id"] = item.get("id")
    elif action == "combo_purchase":
        assert financial_timing is not None
        combo = _catalog_item(state, "combos", payload)
        if not combo:
            raise HTTPException(404, "Không tìm thấy combo trong danh mục.")
        require_available(combo, financial_timing["effective_datetime"].astimezone(VN_TZ).date(), "Combo")
        quantity_source = payload.get("quantity")
        if quantity_source in (None, ""):
            quantity_source = 1
        quantity = int(_bounded_number(
            quantity_source, label="Số lượng combo", minimum=1,
            maximum=MAX_PURCHASE_QUANTITY, integer=True,
        ))
        catalog_tickets = int(_bounded_number(
            combo.get("tickets"), label="Số vé combo", minimum=1,
            maximum=MAX_TICKET_UNITS, integer=True,
        ))
        catalog_price = _bounded_money(
            combo.get("price"), label="Giá combo", allow_blank_as_zero=True,
        )
        # Quantity is the only client input. Ticket count and money always
        # come from the catalog, regardless of legacy `tickets`/`amount` fields.
        tickets = catalog_tickets * quantity
        if tickets > MAX_TICKET_UNITS:
            raise HTTPException(400, f"Tổng số vé mua không được vượt quá {MAX_TICKET_UNITS}.")
        purchase_price = catalog_price * quantity
        if purchase_price > MAX_MONEY:
            raise HTTPException(400, f"Tổng giá combo không được vượt quá {MAX_MONEY}.")
        terms = purchase_terms(combo, quantity, state["services"], financial_timing["effective_datetime"].astimezone(VN_TZ).date())
        payment_method = _canonical_payment_method(payload.get("payment_method"), quick=True)
        if payment_method == "COMBO":
            raise HTTPException(400, "Không thể dùng chính combo để thanh toán giao dịch mua combo.")
        customer = _customer(state, payload)
        if not customer:
            raise HTTPException(400, "Thiếu thông tin khách hàng.")
        purchase = {
            **terms,
            "id": str(uuid4()), "combo_id": combo.get("id"), "combo_name": combo.get("name"),
            "total": tickets, "used": 0, "remaining": tickets,
            "price": purchase_price,
            "purchased_at": financial_timing["recorded_at"],
            "created_at": financial_timing["recorded_at"],
            "recorded_at": financial_timing["recorded_at"],
            "effective_at": financial_timing["effective_at"],
            "business_date": financial_timing["business_date"],
            "backdate_one_day": financial_timing["backdate_one_day"],
            "correction_reason": financial_timing["correction_reason"],
            "note": str(payload.get("note") or ""),
        }
        customer["combo_purchases"].append(purchase)
        invoice = {
            "id": str(uuid4()), "business_date": financial_timing["business_date"],
            "created_at": financial_timing["recorded_at"],
            "recorded_at": financial_timing["recorded_at"],
            "effective_at": financial_timing["effective_at"],
            "backdate_one_day": financial_timing["backdate_one_day"],
            "correction_reason": financial_timing["correction_reason"],
            "actor": actor, "customer_id": customer.get("id"), "customer_name": customer.get("name", ""),
            "customer_phone": customer.get("phone", ""),
            "bill_no": _next_bill_no(state, payload, financial_timing["effective_datetime"]),
            "ticket_no": "", "payment_method": payment_method,
            "subtotal": purchase["price"], "discount": 0, "tip": 0, "total": purchase["price"],
            "combo_purchase_id": "", "combo_units": 0,
            "purchased_combo_id": purchase["id"], "combo_purchased_units": tickets,
            "entries": [{"type": "combo_purchase", "combo_id": combo.get("id"), "service": combo.get("name"), "price": purchase["price"]}],
            "note": purchase["note"], "quick": False,
        }
        state["invoices"].append(invoice)
        state["reports"].append({
            "id": str(uuid4()), "invoice_id": invoice["id"], "business_date": invoice["business_date"],
            "created_at": invoice["created_at"], "recorded_at": invoice["recorded_at"],
            "effective_at": invoice["effective_at"],
            "backdate_one_day": invoice["backdate_one_day"],
            "correction_reason": invoice["correction_reason"],
            "actor": actor, "type": "combo_purchase",
            "employee_id": "", "employee_name": "", "service": combo.get("name"), "room": "",
            "customer_id": customer.get("id"), "customer_name": customer.get("name", ""),
            "customer_phone": customer.get("phone", ""), "bill_no": invoice["bill_no"],
            "payment_method": invoice["payment_method"], "tip": 0, "total": invoice["total"],
            "note": invoice["note"], "combo_units": tickets,
        })
        result.update({"customer": customer, "purchase": purchase, "invoice": invoice})
    elif action == "combo_import":
        imported = []
        source_rows = payload.get("purchases")
        if not source_rows and (payload.get("customer_name") or payload.get("phone")):
            remaining = int(_bounded_number(
                payload.get("remaining"), label="Số vé combo còn lại", minimum=0,
                maximum=MAX_TICKET_UNITS, integer=True,
            ))
            total = int(_bounded_number(
                payload.get("total", remaining), label="Tổng số vé combo", minimum=1,
                maximum=MAX_TICKET_UNITS, integer=True,
            ))
            source_rows = [{
                **payload,
                "total": total, "used": total - remaining,
            }]
        if not source_rows:
            raise HTTPException(400, "Không có dữ liệu combo để import.")
        for row in source_rows:
            row_payload = dict(row)
            combo = _catalog_item(state, "combos", row_payload)
            if not combo:
                raise HTTPException(400, "Dòng import combo thiếu khách hàng hoặc combo hợp lệ.")
            total = int(_bounded_number(
                row_payload.get("total", row_payload.get("tickets", combo.get("tickets"))),
                label="Tổng số vé combo", minimum=1, maximum=MAX_TICKET_UNITS, integer=True,
            ))
            used = int(_bounded_number(
                row_payload.get("used", 0), label="Số vé combo đã dùng", minimum=0,
                maximum=MAX_TICKET_UNITS, integer=True,
            ))
            if used > total:
                raise HTTPException(400, "Số vé combo đã dùng không thể lớn hơn tổng vé.")
            terms = _combo_import_terms(combo, row_payload, state["services"], now.astimezone(VN_TZ).date(), total, used, _bounded_number)
            customer = _customer(state, row_payload)
            if not customer:
                raise HTTPException(400, "Dòng import combo thiếu khách hàng hợp lệ.")
            purchase = {
                **terms,
                "id": str(row_payload.get("id") or uuid4()), "combo_id": combo.get("id"),
                "combo_name": combo.get("name"), "total": total, "used": used,
                "remaining": total - used, "price": _bounded_money(
                    row_payload.get("price", combo.get("price")), label="Giá combo",
                    allow_blank_as_zero=True,
                ),
                "purchased_at": str(row_payload.get("purchased_at") or _iso(now)),
                "note": str(row_payload.get("note") or ""),
            }
            customer["combo_purchases"].append(purchase)
            imported.append(purchase)
        result["imported"] = len(imported)
    elif action == "backup":
        backup = {
            "id": str(uuid4()), "name": str(payload.get("name") or payload.get("label") or f"Backup {now.strftime('%d/%m/%Y %H:%M')}")[:160],
            "created_at": _iso(now), "actor": actor, "snapshot": _snapshot_for_backup(state),
        }
        state["backups"].append(backup)
        state["backups"] = state["backups"][-MAX_BACKUPS:]
        result["backup"] = {key: value for key, value in backup.items() if key != "snapshot"}
    elif action == "restore":
        backup = _find_by_id(state["backups"], payload.get("backup_id"), "bản sao lưu")
        restored = _normalize_state(backup.get("snapshot"), now)
        if any(board.get("pending") or any(_has_unsettled_work(item) for item in board["employees"])
               for board in (state, restored)):
            raise HTTPException(409, "Chỉ khôi phục khi cả bảng hiện tại và bản sao không có dịch vụ, khoản chờ thanh toán hoặc nghỉ giữa ca; tránh phục hồi giao dịch đã thu tiền.")
        # Backups restore board/catalog operation only. Financial/customer
        # ledgers and audit history are append-only and can never be rolled
        # back to create duplicate tickets, invoices or combo balances.
        for append_only_key in (
            "customers", "pending", "invoices", "reports", "combo_usage", "break_events", "audit", "backups", "pending_changes",
            "invoice_changes",
            "bill_counters", "idempotency", "created_at", "sync_status", "customer_changes",
        ):
            restored[append_only_key] = deepcopy(state.get(append_only_key))
        state.clear()
        state.update(restored)
        result["restored_backup_id"] = backup.get("id")
    elif action == "clear_expired":
        preview = _expired_preview(state, payload, now)
        if payload.get("confirm_token") != preview["preview_token"]:
            raise HTTPException(409, "Danh sách phiên quá hạn đã thay đổi. Hãy xem trước và xác nhận lại.")
        ids = {item["employee_id"] for item in preview["employees"]}
        for employee in state["employees"]:
            if str(employee.get("id")) in ids:
                employee["status"] = "CHO THANH TOÁN"
                employee["payment_status"] = "CHO THANH TOÁN"
                employee["completed_at"] = _iso(now)
        result.update({"marked_for_payment": len(ids), "grace_minutes": preview["grace_minutes"], "employee_ids": sorted(ids)})
    else:
        raise HTTPException(400, f"Thao tác Live Tour không hợp lệ: {action}")

    state["updated_at"] = _iso(now)
    state["business_date"] = _business_date(now).isoformat()
    _audit(state, action, payload, actor, now, timing=financial_timing)
    return result


def _expired_preview(state: dict[str, Any], payload: dict[str, Any], now: datetime) -> dict[str, Any]:
    grace = _bounded_number(payload.get("grace_minutes", 15), label="Ngưỡng quá hạn", minimum=0, maximum=1440, integer=True)
    employees = []
    for employee in state["employees"]:
        started = _parse_datetime(employee.get("started_at"))
        duration = employee.get("duration")
        if _norm(employee.get("status")) not in {"dang thuc hien", "dang su dung"} or not started or duration in (None, ""):
            continue
        ends_at = started + timedelta(minutes=float(duration))
        if now.astimezone(VN_TZ) >= ends_at + timedelta(minutes=grace):
            employees.append({"employee_id": str(employee.get("id")), "employee_name": employee.get("name", ""),
                              "room": employee.get("room", ""), "service": employee.get("service", ""), "ends_at": _iso(ends_at)})
    token = _canonical_payload_hash("clear_expired", {"grace_minutes": grace, "employees": employees})
    return {"grace_minutes": grace, "employees": employees, "count": len(employees), "preview_token": token}


def _remaining(employee: dict[str, Any], now: datetime) -> tuple[int | None, str]:
    started = _parse_datetime(employee.get("started_at"))
    duration = employee.get("duration")
    if not started or duration in (None, "") or _norm(employee.get("status")) not in {"dang thuc hien", "dang su dung"}:
        return None, ""
    deadline = started + timedelta(minutes=float(duration))
    return int(math.ceil((deadline - now.astimezone(VN_TZ)).total_seconds() / 60)), _iso(deadline)


def _employee_time_key(employee: dict[str, Any], now: datetime) -> tuple[int, int, int]:
    # Use precisely the standard-start column, including retained history after
    # completion/payment. YC start times and countdown expiry do not reorder it.
    raw = str(_employee_record(employee, now).get("TG bắt đầu thực hiện") or "").strip()
    started = _parse_datetime(raw)
    rank = 0 if not raw else 1 if started else 2
    return (int(_norm(employee.get("work_status")) == "nghi phep"),
            rank, int(started.timestamp()) if started else 0)


def _ordered_employees(employees: list[dict[str, Any]], now: datetime) -> list[dict[str, Any]]:
    if any(item.get("manual_order") for item in employees):
        return sorted(employees, key=lambda item: (_norm(item.get("work_status")) == "nghi phep", int(item.get("sort_index") or 0)))
    return sorted(employees, key=lambda item: (
        *_employee_time_key(item, now), int(item.get("sort_index") or 0), _norm(item.get("name")),
    ))


def _row_style(employee: dict[str, Any], remaining: int | None) -> tuple[str, list[str]]:
    status = _norm(employee.get("status"))
    work = _norm(employee.get("work_status"))
    groups: list[str] = []
    if status in {"dang thuc hien", "dang su dung"}:
        groups.append("doing")
    if status == "dang cho":
        groups.append("waiting")
    if work == "di lam":
        groups.append("working")
    if work == "nghi phep":
        groups.append("leave")
    if employee.get("break_started_at"):
        groups.append("break")
        return "break", groups
    if remaining is not None and -15 < remaining <= 30:
        groups.extend(["finishing", "available"])
    idle = status not in {"dang cho", "dang thuc hien", "dang su dung"} and work == "di lam" and bool(employee.get("shift"))
    if idle:
        groups.append("available")
    if work == "nghi phep":
        return "leave", groups
    if remaining is not None:
        if remaining >= 15:
            return "green", groups
        if remaining >= 0:
            return "yellow", groups
        if remaining > -15:
            return "red", groups
    if idle:
        return "idle", groups
    return "work" if work == "di lam" else "default", groups


def _employee_record(employee: dict[str, Any], now: datetime) -> dict[str, Any]:
    remaining, deadline = _remaining(employee, now)
    style, groups = _row_style(employee, remaining)
    request_start = employee.get("started_at") if _norm(employee.get("request")) == "yc" else ""
    standard_start = employee.get("started_at") if _norm(employee.get("request")) != "yc" else ""
    break_remaining: int | str = ""
    break_started = _parse_datetime(employee.get("break_started_at"))
    if break_started:
        break_elapsed = max(0, int((now.astimezone(VN_TZ) - break_started).total_seconds() // 60))
        break_remaining = 90 - break_elapsed
    elif employee.get("last_break_remaining_minutes") not in (None, ""):
        break_remaining = int(employee["last_break_remaining_minutes"])
    elif employee.get("break_remaining_minutes") not in (None, ""):
        break_remaining = int(employee["break_remaining_minutes"])
    record = {
        "STT": employee.get("stt", ""), "Tên nhân viên": employee.get("name", ""),
        "Trạng thái": employee.get("status", ""), "Phòng": employee.get("room", ""),
        "TG CÒN LẠI": "" if remaining is None or remaining <= -15 else remaining,
        "Yêu cầu": employee.get("request", ""), "Lịch hẹn": employee.get("appointment", ""),
        "Dịch vụ": employee.get("service", ""), "Thời lượng": employee.get("duration", "") if employee.get("duration") is not None else "",
        "TG bắt đầu thực hiện": _display_datetime(standard_start),
        "TG bắt đầu thực hiện YC": _display_datetime(request_start),
        "TT thanh toán": employee.get("payment_status", ""),
        "Kết quả hoàn thành": employee.get("completion_note", ""),
        "SL tua": int(employee.get("tour_count") or 0), "SL yêu cầu": int(employee.get("request_count") or 0),
        "Tổng SL": int(employee.get("tour_count") or 0) + int(employee.get("request_count") or 0),
        "Đi làm": employee.get("work_status", ""), "Vào ca": "" if _norm(employee.get("work_status")) == "nghi phep" else employee.get("shift", ""),
        "Breaktime": _display_datetime(employee.get("break_started_at")),
        "TG nghỉ còn lại": break_remaining,
        "Giờ ra": employee.get("clock_out", ""), "Giờ vào": employee.get("clock_in", ""),
        "Ghi chú": employee.get("note", ""), "VIP": "VIP" if employee.get("vip") else "",
        "Giờ Booking": _display_datetime(employee.get("booked_at")),
        "TG khách chờ": employee.get("wait_minutes", "") if employee.get("wait_minutes") is not None else "",
        "TG Xông Hơi": employee.get("steam_elapsed_minutes", "") if employee.get("steam_elapsed_minutes") is not None else "",
        "_id": employee.get("id"), "id": employee.get("id"),
        "employee_id": employee.get("id"), "_employee_id": employee.get("id"),
        "_row_style": style, "_tour_groups": groups,
        "_employee_change_until": _iso(_employee_change_until(employee)) if _employee_change_until(employee) else "",
        "_employee_change_started_at": employee.get("started_at", ""),
        "_countdown_deadline": deadline, "_attendance_break_active": bool(employee.get("break_started_at")),
        "_manual_order": bool(employee.get("manual_order")), "_sort_index": employee.get("sort_index", 0), "_hidden": bool(employee.get("hidden")), "_active_booking": _active_booking(employee),
        "_payment_pending": _norm(employee.get("status")) == "cho thanh toan",
        "_private_service": _is_private_service(employee.get("service")),
    }
    if not employee.get("service") and not employee.get("status"):
        retained = employee.get("last_assignment_display") or {}
        record.update({column: retained[column] for column in RETAINED_ASSIGNMENT_COLUMNS if column in retained})
    if record["_payment_pending"]:
        for column in ("Trạng thái", "Phòng", "TG CÒN LẠI", "Yêu cầu", "Dịch vụ"):
            record[column] = ""
    return record


def _shift_bucket(employee: dict[str, Any]) -> str:
    token = _norm(employee.get("shift")).replace(" ", "")
    if token in {"ca1", "10", "10h", "10h00"}:
        return "ca1"
    if token in {"ca2", "12", "12h", "12h00", "14", "14h", "14h00"}:
        return "ca2"
    match = re.search(r"(\d{1,2})\s*[:h]", str(employee.get("shift") or "").lower())
    return "ca1" if match and int(match.group(1)) < 12 else "ca2" if match else ""


def _metric_bucket(employees: list[dict[str, Any]], now: datetime) -> dict[str, int]:
    records = [_employee_record(employee, now) for employee in employees]
    total_quantity = sum(int(record.get("Tổng SL") or 0) for record in records)
    waiting = sum("waiting" in record["_tour_groups"] for record in records)
    breaks = sum(bool(record.get("_attendance_break_active")) for record in records)
    return {
        "total_quantity": total_quantity, "waiting_count": waiting,
        "customer_count": total_quantity + waiting,
        "break_count": breaks, "break_total_count": sum(int(item.get("break_count") or 0) for item in employees),
        "break_active_count": breaks,
    }


def _public_backup(backup: dict[str, Any]) -> dict[str, Any]:
    return {key: value for key, value in backup.items() if key != "snapshot"}


def _history_sort_key(item: dict[str, Any]) -> str:
    return str(
        item.get("effective_at")
        or item.get("created_at")
        or item.get("purchased_at")
        or item.get("at")
        or ""
    )


def _customer_history(
    state: dict[str, Any], customer_id: Any, *, bounds: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Return ledgers belonging to exactly one stable customer id.

    Names and phone numbers are deliberately not used as fallbacks: they are
    mutable identifiers and could join two different customer histories.
    """
    wanted = str(customer_id or "").strip()
    if not wanted:
        raise HTTPException(400, "Thiếu mã khách hàng.")
    customer = next(
        (item for item in state.get("customers", []) if str(item.get("id") or "") == wanted),
        None,
    )
    if customer is None:
        raise HTTPException(404, "Không tìm thấy khách hàng.")
    export_bounds = bounds or _parse_export_bounds()

    def matching(items: list[dict[str, Any]]) -> list[dict[str, Any]]:
        rows = [
            deepcopy(item) for item in items
            if str(item.get("customer_id") or "") == wanted
            and _event_in_export_bounds(item, export_bounds)
        ]
        return sorted(rows, key=_history_sort_key, reverse=True)

    invoices = matching(state.get("invoices", []))
    reports = matching(state.get("reports", []))
    combo_usage = matching(state.get("combo_usage", []))
    pending = matching(state.get("pending", []))
    combo_purchases = [
        deepcopy(item) for item in (customer.get("combo_purchases") or [])
        if _event_in_export_bounds(item, export_bounds)
    ]
    combo_purchases.sort(key=_history_sort_key, reverse=True)

    services: list[dict[str, Any]] = []
    for invoice in invoices:
        invoice_entries = list(invoice.get("entries") or [])
        for index, entry in enumerate(invoice_entries):
            services.append({
                "id": f"{invoice.get('id') or ''}:{index}",
                "invoice_id": invoice.get("id", ""),
                "bill_no": invoice.get("bill_no", ""),
                "business_date": invoice.get("business_date", ""),
                "created_at": invoice.get("created_at", ""),
                "effective_at": invoice.get("effective_at", invoice.get("created_at", "")),
                "recorded_at": invoice.get("recorded_at", invoice.get("created_at", "")),
                "employee_id": entry.get("employee_id", ""),
                "employee_name": entry.get("employee_name", ""),
                "service": entry.get("service", ""), "room": entry.get("room", ""),
                "request": entry.get("request", ""), "duration": entry.get("duration", ""),
                "price": entry.get("price", 0),
                "completion_delta_minutes": entry.get("completion_delta_minutes"),
                "completion_note": entry.get("completion_note", ""),
                "payment_method": invoice.get("payment_method", ""),
                "customer_id": wanted,
            })
    services.sort(key=_history_sort_key, reverse=True)

    all_purchases = list(customer.get("combo_purchases") or [])
    summary = {
        "invoice_count": len(invoices), "service_count": len(services),
        "combo_purchase_count": len(combo_purchases), "combo_usage_count": len(combo_usage),
        "pending_count": len(pending),
        "total_revenue": sum(int(item.get("total") or 0) for item in invoices),
        "total_tip": sum(int(item.get("tip") or 0) for item in invoices),
        "combo_purchased_units": sum(int(item.get("total") or 0) for item in combo_purchases),
        "combo_used_units": sum(int(item.get("units") or 0) for item in combo_usage),
        "combo_remaining_units": sum(int(item.get("remaining") or 0) for item in all_purchases if not item.get("deleted_at")),
    }
    return {
        "customer": deepcopy(customer), "summary": summary,
        "invoices": invoices, "services": services, "reports": reports,
        "combo_purchases": combo_purchases, "purchases": deepcopy(combo_purchases),
        "combo_usage": combo_usage, "pending": pending,
    }


def _unbilled_entries(state: dict[str, Any], bounds: dict[str, Any] | None = None) -> list[dict[str, Any]]:
    entries = [dict(item, employee_name=item.get("name"), created_at=item.get("booked_at"))
               for item in state["employees"] if item.get("service")]
    entries.extend(dict(entry, created_at=pending.get("created_at"), status="Chờ thanh toán")
                   for pending in state["pending"] for entry in pending.get("entries", []))
    result = []
    for entry in entries:
        if not _event_in_export_bounds(entry, bounds or {}, fallback_business_date=state.get("business_date")):
            continue
        try:
            value = _resolved_service_price(state, entry)
        except HTTPException:
            value = None
        result.append({**entry, "expected_amount": value, "unpriced": value in (None, 0)})
    return result


def _readable_audit(events, *, invoice_view, paid_invoice_view, customers_view):
    result = _redact_customer_pii(events)
    for event in result:
        if (event.get("action") in {"pending_update", "pending_delete"} and not invoice_view
                or event.get("action") in {"checkout", "quick_checkout", "paid_invoice_update", "paid_invoice_delete"} and not paid_invoice_view
                or event.get("action") in {"combo_purchase", "combo_import", "customer_upsert"} and not customers_view):
            event["detail"] = {"summary": "Nội dung cần quyền xem tương ứng."}
    return result


def _report_rows_with_combo_kind(state):
    sales = {row["id"] for row in state["invoices"] if row.get("purchased_combo_id")}
    return [dict(row, combo_sale=row.get("invoice_id") in sales) for row in state["reports"]]


def _state_response(
    state: dict[str, Any], revision: int, now: datetime, *, include_hidden: bool = False,
    can_admin: bool = False, can_operate: bool = False, can_appointment_edit: bool = False,
    can_payment: bool = False, can_export: bool = False,
    can_booking: bool | None = None, can_invoice_view: bool | None = None,
    can_paid_invoice_view: bool | None = None,
    can_paid_invoice_edit: bool = False, can_paid_invoice_delete: bool = False,
    can_invoice_edit: bool = False, can_invoice_delete: bool = False,
    can_pending_view: bool | None = None, can_customers_view: bool | None = None,
    can_reports_view: bool | None = None, can_history_view: bool | None = None,
    can_backup: bool | None = None,
    can_reorder: bool = False,
    viewer_bank: dict | None = None,
    can_invoice_date_edit: bool = False,
    can_customers_edit: bool = False,
    can_customers_delete: bool = False,
    can_customer_combo_edit: bool = False,
    can_customer_combo_delete: bool = False,
    can_reports_edit: bool = False,
    can_reports_delete: bool = False,

) -> dict[str, Any]:
    # Defaults support internal legacy callers; HTTP routes always pass explicit
    # effective grants, including denials, from the canonical permission store.
    can_booking = can_operate if can_booking is None else can_booking
    can_invoice_view = can_payment if can_invoice_view is None else can_invoice_view
    can_paid_invoice_view = can_payment if can_paid_invoice_view is None else can_paid_invoice_view
    can_pending_view = can_payment if can_pending_view is None else can_pending_view
    can_customers_view = can_payment if can_customers_view is None else can_customers_view
    can_reports_view = can_payment if can_reports_view is None else can_reports_view
    can_history_view = can_admin if can_history_view is None else can_history_view
    can_backup = can_admin if can_backup is None else can_backup
    pending_access = can_pending_view and can_invoice_view
    customer_pii = can_customers_view or can_invoice_view or can_paid_invoice_view
    state = deepcopy(state)
    _ensure_counter_day(state, now)
    ordered = _ordered_employees([row for row in state["employees"] if row.get("roster_eligible") is not False], now)
    can_recover_hidden = False
    for item in ordered:
        item["hidden"] = False
    visible = ordered
    records = [_employee_record(employee, now) for employee in visible]
    for index, record in enumerate(records, 1):
        record["STT"] = index
        record["_private_service"] = _catalog_private_service(state, record.get("Dịch vụ"))
    occupied = sorted({str(item.get("room")) for item in state["employees"] if _active_booking(item) and item.get("room")})
    all_rooms = [str(item.get("name")) for item in state["rooms"] if item.get("active", True)]
    available = [room for room in all_rooms if _room_available(state, room)]
    physical_rooms = sorted(set(state.get("physical_rooms", [])) | {_catalog_room_group(state, room) for room in all_rooms})
    occupied_groups = sorted({_catalog_room_group(state, room) for room in occupied})
    available_groups = sorted({_catalog_room_group(state, room) for room in available})
    active_records = [record for record in records if not record.get("_hidden")]
    groups = lambda key: sum(key in (item.get("_tour_groups") or []) for item in active_records)
    metrics = {
        "all": _metric_bucket(ordered, now),
        "ca1": _metric_bucket([item for item in ordered if _shift_bucket(item) == "ca1"], now),
        "ca2": _metric_bucket([item for item in ordered if _shift_bucket(item) == "ca2"], now),
    }
    customers = []
    if can_customers_view:
        for item in state["customers"]:
            if item.get("deleted_at"):
                continue
            customer = deepcopy(item)
            customer["combo_purchases"] = [p for p in customer.get("combo_purchases", []) if not p.get("deleted_at")]
            purchases = list(customer.get("combo_purchases") or [])
            for purchase in purchases:
                available_purchase = _available_combo(state, customer["id"], purchase)
                purchase["booking_remaining"] = available_purchase["remaining"]
                purchase["booking_reserved"] = available_purchase["booking_reserved"]
                for part, available_part in zip(purchase.get("component_balances", []), available_purchase.get("component_balances", [])):
                    part["booking_remaining"] = available_part["remaining"]
            customer["combos"] = purchases
            customer["combo_balance"] = sum(int(purchase.get("remaining") or 0) for purchase in purchases)
            customers.append(customer)
    report_summary = {
        "invoice_count": len(state["invoices"]),
        "pending_count": len(state["pending"]),
        "total_revenue": sum(int(item.get("total") or 0) for item in state["invoices"]),
        "total_tip": sum(int(item.get("tip") or 0) for item in state["invoices"]),
    } if can_reports_view else {}
    if can_reports_view:
        unbilled = _unbilled_entries(state)
        report_summary.update({
            "expected_unbilled_revenue": sum(item["expected_amount"] or 0 for item in unbilled),
            "unbilled_count": len(unbilled),
            "unbilled_unpriced_count": sum(item["unpriced"] for item in unbilled),
        })
    public_backups = [_public_backup(item) for item in state["backups"]] if can_backup else []
    # Re-redact historical audit entries too: older persisted state may predate
    # write-time redaction and must never leak through an admin response.
    public_audit = _readable_audit(state["audit"][-1000:], invoice_view=pending_access,
                                  paid_invoice_view=can_paid_invoice_view, customers_view=can_customers_view) if can_history_view else []
    public_break_events = _redact_customer_pii(state["break_events"]) if can_history_view else []
    public_employees = deepcopy(visible) if customer_pii else _redact_customer_pii(visible)
    public_reports = _report_rows_with_combo_kind(state) if can_reports_view else []
    if not customer_pii:
        public_reports = _redact_customer_pii(public_reports)
    pending_changes = deepcopy(state.get("pending_changes", [])[-500:]) if can_history_view and pending_access else []
    invoice_changes = deepcopy(state.get("invoice_changes", [])[-500:]) if can_history_view and can_paid_invoice_view else []
    if not can_customers_view:
        for change in invoice_changes:
            change.pop("purchase_before", None)
            change.pop("purchase_after", None)
    return {
        "storage_mode": "server", "columns": BOARD_COLUMNS, "records": records, "count": len(records), "employee_count": len(records),
        "available": groups("available"), "working_count": groups("working"), "leave_count": groups("leave"),
        "doing_count": groups("doing"), "waiting_count": groups("waiting"), "finishing_count": groups("finishing"),
        "break_count": groups("break"), "stats": [{"label": "Có thể lên tua", "value": groups("available"), "detail": "Sắp xong + Đang rảnh"}],
        "rooms": {"all": physical_rooms, "available": available_groups, "occupied": occupied_groups, "total_count": len(physical_rooms), "available_count": len(available_groups), "occupied_count": len(occupied_groups), "source_sheet": "Live Tour"},
        "available_rooms": available_groups, "available_beds": available, "metric_snapshots": metrics,
        # Occupancy must include hidden/retained staff without exposing customer data.
        "room_assignments": [{"id": row["id"], "room": row.get("room", ""),
                              "service": row.get("service", ""), "status": row.get("status", ""),
                              "private": _catalog_private_service(state, row.get("service"))}
                             for row in state["employees"] if _active_booking(row)],
        "room_groups": {room["name"]: _catalog_room_group(state, room["name"]) for room in state["rooms"]},
        "room_action_counts": _room_action_counts(state) if can_operate else {},
        "service_areas": _service_areas(state),
        "metrics_business_date": state.get("counter_business_date"), "metrics_retained_until_10": now.astimezone(VN_TZ).hour < 10,
        "metrics_rollover_hour": 10, "metrics_rollover_minute": 0,
        "break_metrics_business_date": state.get("counter_business_date"), "break_metrics_format": "cumulative-active",
        "break_metrics_source": "live_tour", "countdown_at": _iso(now), "countdown_error": "",
        "source_updated_at": state.get("updated_at", ""), "revision": revision,
        # Root aliases keep the API convenient for both the copied Tour UI and
        # the richer Live Tour operator drawers.
        "services": state["services"], "combo_catalog": state["combos"],
        "payment_settings": {**deepcopy(state.get("payment_settings") or _default_payment_settings()), **({"user_bank": deepcopy(viewer_bank)} if viewer_bank else {})} if can_payment or can_admin or can_paid_invoice_view else {},
        "employee_directory": deepcopy(state.get("employee_directory", [])) if can_admin else [],
        "retained_assignments": [deepcopy(row) if customer_pii else _redact_customer_pii(row) for row in state["employees"] if row.get("roster_eligible") is False and (row.get("service") or row.get("break_started_at"))] if can_operate or can_payment or can_admin else [],
        "customers": customers, "pending_payments": state["pending"] if pending_access else [],
        "pending": state["pending"] if pending_access else [], "reports": report_summary,
        "pending_count": len(state["pending"]) if can_pending_view else 0,
        "pending_changes": pending_changes,
        "invoice_changes": invoice_changes,
        "customer_changes": deepcopy(state.get("customer_changes", [])) if can_customers_view and can_history_view else [],
        "report_rows": public_reports,
        "audit": public_audit, "history": public_audit, "backups": public_backups,
        "break_events": public_break_events,
        "catalogs": {"rooms": state["rooms"], "services": state["services"], "combos": state["combos"]},
        "state": {
            "version": state.get("version"), "business_date": state.get("business_date"),
            "updated_at": state.get("updated_at"), "storage_mode": "server",
            "employees": public_employees, "hidden_count": sum(bool(item.get("hidden")) for item in ordered),
            "rooms": state["rooms"], "services": state["services"], "combos": state["combos"],
            "customers": customers,
            "pending": state["pending"] if pending_access else [],
            "invoices": state["invoices"] if can_paid_invoice_view else [],
            "reports": public_reports,
            "combo_usage": state["combo_usage"][-1000:] if can_customers_view else [],
            "audit": public_audit, "break_events": public_break_events,
            "backups": public_backups,
        },
        "capabilities": {
            "appointment_edit": can_appointment_edit, "reorder": can_reorder,
            "admin": can_admin, "catalog_admin": can_admin, "manage_catalog": can_admin,
            "operate": can_operate, "payment": can_payment, "export": can_export,
            "booking": can_booking, "invoice_view": can_invoice_view,
            "paid_invoice_view": can_paid_invoice_view,
            "paid_invoice_edit": can_paid_invoice_edit and can_paid_invoice_view,
            "paid_invoice_delete": can_paid_invoice_delete and can_paid_invoice_view,
            "invoice_edit": can_invoice_edit and pending_access, "invoice_delete": can_invoice_delete and pending_access,
            "pending_view": can_pending_view, "customers_view": can_customers_view,
            "reports_view": can_reports_view, "history_view": can_history_view, "backup": can_backup,
            "hide_recovery": can_recover_hidden,
            "invoice_date_edit": can_invoice_date_edit,
            "customers_edit": can_customers_edit,
            "customers_delete": can_customers_delete,
            "customer_combo_edit": can_customer_combo_edit,
            "customer_combo_delete": can_customer_combo_delete,
            "reports_edit": can_reports_edit,
            "reports_delete": can_reports_delete,

        },
    }


def _room_available(state: dict[str, Any], room: str) -> bool:
    wanted_group = _norm(_catalog_room_group(state, room))
    wanted = _norm(room)
    for employee in state["employees"]:
        if not _active_booking(employee):
            continue
        current_room = str(employee.get("room") or "")
        if _norm(current_room) == wanted:
            return False
        if _norm(_catalog_room_group(state, current_room)) == wanted_group and _catalog_private_service(state, employee.get("service")):
            return False
    return True


def _read_state(conn, now: datetime, *, for_update: bool = False) -> tuple[dict[str, Any], int]:
    suffix = " FOR UPDATE" if for_update else ""
    row = conn.execute(text(f"""
        SELECT value_json, revision FROM vera_app_setting
        WHERE category=:category AND setting_key=:key{suffix}
    """), {"category": STATE_CATEGORY, "key": STATE_KEY}).mappings().first()
    if row:
        state = _normalize_state(row.get("value_json"), now)
        revision = int(row.get("revision") or 0)
        if _reconcile_roster(state, _employee_directory(conn, now), _new_directory_employee):
            revision = _write_state(conn, state, revision, "live_tour_employee_directory")
        return state, revision
    state = _bootstrap_state(conn, now)
    conn.execute(text("""
        INSERT INTO vera_app_setting(
          category,setting_key,value_json,source,updated_by,revision,created_at,updated_at
        ) VALUES (
          :category,:key,CAST(:value AS jsonb),'web_v2',:actor,1,NOW(),NOW()
        ) ON CONFLICT(category,setting_key) DO NOTHING
    """), {
        "category": STATE_CATEGORY, "key": STATE_KEY,
        "value": json.dumps(state, ensure_ascii=False), "actor": "live_tour_bootstrap",
    })
    row = conn.execute(text("""
        SELECT value_json, revision FROM vera_app_setting
        WHERE category=:category AND setting_key=:key
    """), {"category": STATE_CATEGORY, "key": STATE_KEY}).mappings().first()
    if not row:
        raise HTTPException(500, "Không khởi tạo được trạng thái Live Tour.")
    return _normalize_state(row.get("value_json"), now), int(row.get("revision") or 1)


def _write_state(conn, state: dict[str, Any], revision: int, actor: str) -> int:
    result = conn.execute(text("""
        UPDATE vera_app_setting SET value_json=CAST(:value AS jsonb), source='web_v2',
          updated_by=:actor, revision=revision+1, updated_at=NOW()
        WHERE category=:category AND setting_key=:key AND revision=:revision
    """), {
        "value": json.dumps(state, ensure_ascii=False), "actor": actor,
        "category": STATE_CATEGORY, "key": STATE_KEY, "revision": revision,
    })
    if getattr(result, "rowcount", 1) != 1:
        raise HTTPException(409, "Live Tour đã thay đổi ở thiết bị khác. Hãy làm mới rồi thao tác lại.")
    return revision + 1


def _parse_export_bounds(
    *, date_from: str = "", date_to: str = "", time_from: str = "", time_to: str = "",
) -> dict[str, Any]:
    try:
        start_date = date.fromisoformat(date_from) if date_from else None
        end_date = date.fromisoformat(date_to) if date_to else None
    except ValueError as exc:
        raise HTTPException(400, "Khoảng ngày export phải theo định dạng YYYY-MM-DD.") from exc
    time_pattern = re.compile(r"^(?:[01]\d|2[0-3]):[0-5]\d$")
    if time_from and not time_pattern.fullmatch(time_from):
        raise HTTPException(400, "Giờ bắt đầu export phải theo định dạng HH:MM.")
    if time_to and not time_pattern.fullmatch(time_to):
        raise HTTPException(400, "Giờ kết thúc export phải theo định dạng HH:MM.")
    start_time = time.fromisoformat(time_from) if time_from else None
    end_time = time.fromisoformat(time_to) if time_to else None
    if start_date and end_date and start_date > end_date:
        raise HTTPException(400, "Ngày bắt đầu export không được sau ngày kết thúc.")
    return {"date_from": start_date, "date_to": end_date, "time_from": start_time, "time_to": end_time}


def _event_in_export_bounds(
    item: dict[str, Any], bounds: dict[str, Any], *, fallback_business_date: Any = None,
) -> bool:
    if not any(bounds.values()):
        return True
    if bounds.get("customer") and _norm(bounds["customer"]) not in _norm(f"{item.get('customer_name', '')} {item.get('customer_phone', '')}"):
        return False
    entries = item.get("entries") or [item]
    if not any((not bounds.get("employee") or _norm(bounds["employee"]) in _norm(row.get("employee_name")))
               and (not bounds.get("service") or _norm(bounds["service"]) in _norm(row.get("service"))) for row in entries):
        return False
    if bounds.get("report_kind") == "combos" and not (item.get("combo_sale") or item.get("purchased_combo_id") or item.get("combo_units") or any("combo" in _norm(r.get("service")) for r in entries)):
        return False
    event = None
    for key in ("effective_at", "created_at", "completed_at", "started_at", "booked_at", "purchased_at", "at", "updated_at"):
        event = _parse_datetime(item.get(key))
        if event:
            break
    raw_day = item.get("business_date") or fallback_business_date
    try:
        event_day = event.astimezone(VN_TZ).date() if event and bounds.get("calendar_date") else date.fromisoformat(str(raw_day)) if raw_day else (event.astimezone(VN_TZ).date() if event else None)
    except ValueError:
        event_day = event.astimezone(VN_TZ).date() if event else None
    if bounds.get("date_from") and (event_day is None or event_day < bounds["date_from"]):
        return False
    if bounds.get("date_to") and (event_day is None or event_day > bounds["date_to"]):
        return False
    if bounds.get("time_from") or bounds.get("time_to"):
        if event is None:
            return False
        event_time = event.astimezone(VN_TZ).time().replace(tzinfo=None, second=0, microsecond=0)
        start_time, end_time = bounds.get("time_from"), bounds.get("time_to")
        if start_time and end_time and start_time > end_time:
            if not (event_time >= start_time or event_time <= end_time):
                return False
        elif (start_time and event_time < start_time) or (end_time and event_time > end_time):
            return False
    return True


def _export_rows(
    state: dict[str, Any], kind: str, now: datetime, *, include_hidden: bool = False,
    bounds: dict[str, Any] | None = None,
) -> tuple[str, list[str], list[list[Any]]]:
    bounds = bounds or _parse_export_bounds()
    if kind == "board":
        employees = [
            item for item in _ordered_employees(state["employees"], now)
            if item.get("roster_eligible") is not False
            and _event_in_export_bounds(item, bounds, fallback_business_date=state.get("business_date"))
        ]
        records = [{**_employee_record(item, now), "STT": index} for index, item in enumerate(employees, 1)]
        return "Bang_tua", BOARD_COLUMNS, [[record.get(column, "") for column in BOARD_COLUMNS] for record in records]
    if kind == "custom":
        title, headers, rows = _export_rows(state, "board", now, include_hidden=include_hidden, bounds=bounds)
        return "Tuy_chinh", headers, rows
    if kind == "revenue":
        headers = ["Ngày", "Số bill", "Khách hàng", "Điện thoại", "Tổng tiền", "Giảm giá", "Tip", "Thanh toán", "Người tạo"]
        rows = [[item.get("business_date"), item.get("bill_no"), item.get("customer_name"), item.get("customer_phone"), item.get("total"), item.get("discount"), item.get("tip"), item.get("payment_method"), item.get("actor")] for item in state["invoices"] if _event_in_export_bounds(item, bounds)]
        return "Doanh_thu", headers, rows
    if kind == "tip":
        headers = ["Ngày", "Nhân viên", "Dịch vụ", "Phòng", "Số bill", "Tip", "Người tạo"]
        rows = [[item.get("business_date"), item.get("employee_name"), item.get("service"), item.get("room"), item.get("bill_no"), item.get("tip"), item.get("actor")] for item in _report_rows_with_combo_kind(state) if _event_in_export_bounds(item, bounds)]
        return "Tip", headers, rows
    if kind == "reports":
        fields = ["business_date", "effective_at", "employee_name", "service", "room", "bill_no", "customer_name", "customer_phone", "total", "tip", "payment_method", "actor"]
        headers = ["Ngày", "Ngày giờ hóa đơn", "Nhân viên", "Dịch vụ", "Phòng", "Số bill", "Khách hàng", "Điện thoại", "Tổng tiền", "Tip", "Thanh toán", "Người tạo"]
        rows = [[item.get(key) for key in fields] for item in _report_rows_with_combo_kind(state) if _event_in_export_bounds(item, bounds)]
        return "Bao_cao", headers, rows
    if kind == "customers":
        headers = ["Khách hàng", "Điện thoại", "Combo", "Tổng vé", "Đã dùng", "Còn lại", "Ngày mua"]
        rows = [[customer.get("name"), customer.get("phone"), purchase.get("combo_name"), purchase.get("total"), purchase.get("used"), purchase.get("remaining"), purchase.get("lk") or purchase.get("purchased_at")] for customer in state["customers"] if not customer.get("deleted_at") for purchase in ([p for p in customer.get("combo_purchases", []) if not p.get("deleted_at")] or [{}]) if _event_in_export_bounds(purchase, bounds)]
        return "Khach_hang", headers, rows
    if kind == "pending":
        headers = ["Ngày giờ", "Khách hàng", "Điện thoại", "Nhân viên", "Dịch vụ", "Phòng", "Ghi chú"]
        rows = [[item.get("effective_at") or item.get("booked_at") or item.get("created_at"), item.get("customer_name"), item.get("customer_phone"), entry.get("employee_name"), entry.get("service"), entry.get("room"), item.get("note")] for item in state["pending"] if _event_in_export_bounds(item, bounds) for entry in item.get("entries", [])]
        return "Cho_thanh_toan", headers, rows
    if kind == "history":
        headers = ["Ngày giờ", "Ngày kinh doanh", "Người thao tác", "Hành động", "Chi tiết"]
        rows = [[item.get("at"), item.get("business_date"), item.get("actor"), item.get("action"), json.dumps(_redact_customer_pii(item.get("detail")), ensure_ascii=False)] for item in state["audit"] if _event_in_export_bounds(item, bounds)]
        return "Lich_su", headers, rows
    if kind == "breaks":
        headers = [
            "Ngày giờ", "Ngày kinh doanh", "Loại sự kiện", "Nhân viên", "Mã nhân viên",
            "Bắt đầu", "Kết thúc", "Định mức", "Số phút", "Còn lại", "Trễ", "Kết quả",
            "Người thao tác",
        ]
        rows = [[
            item.get("at"), item.get("business_date"), item.get("event_type"),
            item.get("employee_name"), item.get("employee_id"), item.get("started_at"),
            item.get("ended_at"), item.get("allowance_minutes"), item.get("minutes"),
            item.get("remaining_minutes"), item.get("late_minutes"), item.get("outcome"),
            item.get("actor"),
        ] for item in state["break_events"] if _event_in_export_bounds(item, bounds)]
        return "Nghi_giua_ca", headers, rows
    raise HTTPException(400, "Loại báo cáo Live Tour không hợp lệ.")


def _excel_literal(value: Any) -> Any:
    """Force formula-like user text to remain text in every XLSX export."""
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _fill_excel_sheet(sheet, headers: list[Any], rows: list[list[Any]]) -> None:
    sheet.append([_excel_literal(value) for value in headers])
    for row in rows:
        sheet.append([_excel_literal(value) for value in row])
    header_fill = PatternFill("solid", fgColor="174E3B")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = min(42, max(10, max(len(str(cell.value or "")) for cell in column) + 2))
        sheet.column_dimensions[column[0].column_letter].width = width


def _customer_detail_excel_bytes(
    state: dict[str, Any], customer_id: Any, now: datetime, *,
    bounds: dict[str, Any] | None = None,
) -> tuple[bytes, str]:
    history = _customer_history(state, customer_id, bounds=bounds)
    customer = history["customer"]
    workbook = Workbook()
    workbook.remove(workbook.active)
    overview_rows = [
        ["customer_id", customer.get("id", "")],
        ["customer_name", customer.get("name", "")],
        ["customer_phone", customer.get("phone", "")],
        *[[key, value] for key, value in history["summary"].items()],
    ]
    sheets: list[tuple[str, list[str], list[list[Any]]]] = [
        ("Tong_quan", ["Trường", "Giá trị"], overview_rows),
        ("Hoa_don", [
            "Ngày hiệu lực", "Ngày ghi nhận", "Ngày kinh doanh", "Số bill", "Thanh toán",
            "Tạm tính", "Giảm giá", "Tip", "Tổng tiền", "Người tạo", "Ghi chú",
        ], [[
            item.get("effective_at", item.get("created_at")),
            item.get("recorded_at", item.get("created_at")), item.get("business_date"),
            item.get("bill_no"), item.get("payment_method"), item.get("subtotal"),
            item.get("discount"), item.get("tip"), item.get("total"), item.get("actor"),
            item.get("note"),
        ] for item in history["invoices"]]),
        ("Dich_vu", [
            "Ngày hiệu lực", "Ngày kinh doanh", "Số bill", "Nhân viên", "Dịch vụ",
            "Phòng", "Yêu cầu", "Thời lượng", "Giá", "Kết quả hoàn thành", "Chênh lệch phút",
            "Thanh toán",
        ], [[
            item.get("effective_at", item.get("created_at")), item.get("business_date"),
            item.get("bill_no"), item.get("employee_name"), item.get("service"),
            item.get("room"), item.get("request"), item.get("duration"), item.get("price"),
            item.get("completion_note"), item.get("completion_delta_minutes"),
            item.get("payment_method"),
        ] for item in history["services"]]),
        ("Combo_da_mua", [
            "Ngày hiệu lực", "Combo", "Tổng vé", "Đã dùng", "Còn lại", "Giá", "Ghi chú",
        ], [[
            item.get("effective_at", item.get("purchased_at")), item.get("combo_name"),
            item.get("total"), item.get("used"), item.get("remaining"), item.get("price"),
            item.get("note"),
        ] for item in history["combo_purchases"]]),
        ("Combo_su_dung", [
            "Ngày hiệu lực", "Ngày kinh doanh", "Mã hóa đơn", "Combo", "Số vé dùng",
            "Trước", "Sau", "Người tạo",
        ], [[
            item.get("effective_at", item.get("created_at")), item.get("business_date"),
            item.get("invoice_id"), item.get("combo_purchase_id"), item.get("units"),
            item.get("remaining_before"), item.get("remaining_after"), item.get("actor"),
        ] for item in history["combo_usage"]]),
        ("Cho_thanh_toan", [
            "Ngày giờ", "Mã phiếu", "Nhân viên", "Dịch vụ", "Phòng", "Ghi chú",
        ], [[
            item.get("created_at"), item.get("id"), entry.get("employee_name"),
            entry.get("service"), entry.get("room"), item.get("note"),
        ] for item in history["pending"] for entry in (item.get("entries") or [{}])]),
    ]
    for title, headers, rows in sheets:
        sheet = workbook.create_sheet(title[:31])
        _fill_excel_sheet(sheet, headers, rows)
    output = BytesIO()
    workbook.save(output)
    safe_id = re.sub(r"[^A-Za-z0-9_-]+", "_", str(customer.get("id") or "customer"))[:80]
    return output.getvalue(), f"Live_Tour_Khach_hang_{safe_id}_{_business_date(now).strftime('%Y%m%d')}.xlsx"


def _excel_bytes(
    state: dict[str, Any], kind: str, now: datetime, *, include_hidden: bool = False,
    bounds: dict[str, Any] | None = None, customer_id: Any = "",
    selected_columns: list[str] | None = None, employee_ids: list[str] | None = None,
) -> tuple[bytes, str]:
    if kind == "custom":
        if selected_columns is not None and (not selected_columns or len(selected_columns) != len(set(selected_columns)) or any(column not in BOARD_COLUMNS for column in selected_columns)):
            raise HTTPException(400, "Danh sách cột xuất phải thuộc Bảng tua, không trùng và không rỗng.")
        if employee_ids is not None:
            if not employee_ids or len(employee_ids) > 5000:
                raise HTTPException(400, "Hãy chọn từ 1 đến 5000 nhân viên để xuất.")
            allowed = {str(item.get("id")) for item in state["employees"] if item.get("roster_eligible") is not False}
            if not set(employee_ids).issubset(allowed):
                raise HTTPException(409, "Danh sách nhân viên đã thay đổi hoặc có dòng đang ẩn. Hãy tải lại và chọn lại.")
            state = {**state, "employees": [item for item in state["employees"] if str(item.get("id")) in set(employee_ids)]}
    elif selected_columns is not None or employee_ids is not None:
        raise HTTPException(400, "Chọn cột/nhân viên chỉ áp dụng cho xuất tùy chỉnh.")
    if kind == "customer_detail":
        return _customer_detail_excel_bytes(state, customer_id, now, bounds=bounds)
    title, headers, rows = _export_rows(state, kind, now, include_hidden=include_hidden, bounds=bounds)
    if kind == "custom" and selected_columns is not None:
        indices = [headers.index(column) for column in selected_columns]
        rows = [[row[index] for index in indices] for row in rows]
        headers = selected_columns
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31]
    _fill_excel_sheet(sheet, headers, rows)
    if kind == "revenue":
        # A separate worksheet prevents estimated service values from being
        # mistaken for collected revenue or booked invoice transactions.
        expected = workbook.create_sheet("Du_kien_chua_xuat_bill")
        _fill_excel_sheet(expected, ["Nhân viên", "Dịch vụ", "Phòng", "Trạng thái", "Dự kiến", "Chưa xác định giá"], [
            [item.get("employee_name"), item.get("service"), item.get("room"), item.get("status"),
             item["expected_amount"], "Có" if item["unpriced"] else ""]
            for item in _unbilled_entries(state, bounds)
        ])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue(), f"Live_Tour_{title}_{_business_date(now).strftime('%Y%m%d')}.xlsx"


COPY_BOARD_COLUMNS = [
    "STT", "Tên nhân viên", "Thao tác", "Lịch hẹn", "Trạng thái", "Phòng",
    "TG CÒN LẠI", "Yêu cầu", "Dịch vụ", "Đi làm", "Vào ca", "Breaktime", "TG nghỉ còn lại",
]


def _png_bytes(state: dict[str, Any], now: datetime, *, include_hidden: bool = False) -> bytes:
    import textwrap
    employees = [item for item in _ordered_employees(state["employees"], now) if item.get("roster_eligible") is not False]
    records = [{**_employee_record(item, now), "STT": i} for i, item in enumerate(employees, 1)]
    font_path = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    try:
        font = ImageFont.truetype(font_path, 14)
        bold = ImageFont.truetype(font_path.replace(".ttf", "-Bold.ttf"), 14)
    except OSError:
        font = ImageFont.load_default(size=14)
        bold = font
    widths = [65 if c == "STT" else 210 if c in {"Tên nhân viên", "Dịch vụ", "Ghi chú"} else 155 for c in COPY_BOARD_COLUMNS]
    def cells(values):
        return ["\n".join(textwrap.wrap(str(value if value is not None else ""), max(6, (width-16)//9)) or [""])
                for value, width in zip(values, widths)]
    rendered = [cells(COPY_BOARD_COLUMNS)] + [cells([("Thực hiện · Hoàn thành" if c == "Thao tác" else r.get(c, "")) for c in COPY_BOARD_COLUMNS]) for r in records]
    heights = [max(len(c.split("\n")) for c in row)*20+16 for row in rendered]
    image = Image.new("RGB", (sum(widths), sum(heights)), "white")
    draw = ImageDraw.Draw(image)
    palette = {"green": "#caedb2", "yellow": "#ffe477", "red": "#ffaaa2", "break": "#f6b27d", "waiting": "#dcc3ee", "idle": "#dcebd8", "leave": "#eeeeee"}
    y = 0
    for i, (row, height) in enumerate(zip(rendered, heights)):
        fill = "#174e3b" if i == 0 else palette.get(records[i-1]["_row_style"], "#ffffff")
        x = 0
        for value, width in zip(row, widths):
            draw.rectangle((x, y, x+width-1, y+height-1), fill=fill, outline="#d5dfda")
            draw.multiline_text((x+8, y+8), value, font=bold if i == 0 else font, fill="white" if i == 0 else "#15251f", spacing=3)
            x += width
        y += height
    output = BytesIO()
    image.save(output, format="PNG", optimize=True)
    return output.getvalue()


def _idempotency_entry(
    state: dict[str, Any], key: str, *, action: str, actor: str, payload_hash: str,
) -> dict[str, Any] | None:
    previous = state.get("idempotency", {}).get(key) if key else None
    if not previous:
        return None
    if (
        previous.get("actor") != actor
        or previous.get("action") != action
        or not previous.get("payload_hash")
        or previous.get("payload_hash") != payload_hash
    ):
        raise HTTPException(409, "idempotency_key đã được dùng cho một nội dung thao tác khác.")
    return previous


def _idempotency_status(entry: dict[str, Any] | None) -> str:
    if not entry:
        return ""
    if entry.get("status"):
        return str(entry["status"])
    if isinstance(entry.get("result"), dict) and entry["result"].get("recovery_required"):
        return "recovery_required"
    # Entries written before phase-aware sync were completed operations.
    return "completed"


def _idempotency_replay(
    state: dict[str, Any], key: str, *, action: str, actor: str, payload_hash: str,
) -> dict[str, Any] | None:
    previous = _idempotency_entry(
        state, key, action=action, actor=actor, payload_hash=payload_hash,
    )
    status = _idempotency_status(previous)
    if previous is not None and status != "completed":
        raise HTTPException(
            503,
            detail={
                "code": "LIVE_TOUR_SYNC_RECOVERY_REQUIRED",
                "status": status,
                "message": "Yêu cầu đồng bộ chưa hoàn tất; hãy thử lại đúng idempotency_key để đối soát.",
                "retry_same_key": True,
            },
        )
    return previous


def _remember_idempotency(
    state: dict[str, Any], key: str, *, action: str, actor: str,
    payload_hash: str, result: dict[str, Any], now: datetime,
) -> None:
    state.setdefault("idempotency", {})[key] = {
        "action": action, "actor": actor, "payload_hash": payload_hash,
        "status": "completed", "at": _iso(now), "result": deepcopy(result),
    }
    if len(state["idempotency"]) > MAX_IDEMPOTENCY:
        overflow = len(state["idempotency"]) - MAX_IDEMPOTENCY
        prunable = sorted(
            (
                item_key for item_key, entry in state["idempotency"].items()
                if not (
                    (entry.get("action") == "sync_leaves" and _idempotency_status(entry) not in {"failed", "abandoned", "completed"})
                    or (
                        entry.get("action") in PROTECTED_IDEMPOTENCY_ACTIONS
                        and _idempotency_status(entry) == "completed"
                    )
                )
            ),
            key=lambda item_key: (item_key == key, str(state["idempotency"][item_key].get("at") or "")),
        )
        removed = prunable[:overflow]
        for old_key in removed:
            state["idempotency"].pop(old_key, None)
        # Financial/sync receipts must not expire just because unrelated
        # traffic passed a count threshold: an old payment key could otherwise
        # be replayed against a new service on the same employee. MAX_IDEMPOTENCY
        # is a soft cap; durable financial receipts share the ledger lifetime.


def install_live_tour_routes(
    app, *, engine_instance: Callable[[], Any], current_identity, require_feature,
    feature_allowed: Callable[..., bool], identity_type, vn_tz=VN_TZ,
) -> None:
    if getattr(app.state, "live_tour_installed", False):
        return
    timezone = vn_tz or VN_TZ

    def permissions(conn, ident) -> dict[str, bool]:
        viewer_bank = None
        if str(getattr(ident, "role", "")).strip().lower() == "letan" and feature_allowed(conn, ident, "live_tour_payment"):
            bank_row = conn.execute(text("SELECT full_name, bank_name, bank_account FROM employees WHERE username=:username"), {"username": ident.employee_username}).mappings().first()
            viewer_bank = _profile_bank(dict(bank_row)) if bank_row else None
        can_admin = bool(feature_allowed(conn, ident, "live_tour_admin"))
        can_operate = bool(feature_allowed(conn, ident, "live_tour_operate"))
        return {
            "can_appointment_edit": str(getattr(ident, "role", "") or "").strip().lower() in {"admin", "quanly", "letan"} and bool(feature_allowed(conn, ident, "live_tour_view")),
            "viewer_bank": viewer_bank,
            "can_admin": can_admin, "can_operate": can_operate,
            "can_payment": bool(feature_allowed(conn, ident, "live_tour_payment")),
            "can_export": bool(feature_allowed(conn, ident, "live_tour_export")),
            **{f"can_{name}": bool(feature_allowed(conn, ident, feature)) for name, feature in CAPABILITY_FEATURES.items()},
        }

    def action_response(
        *, state: dict[str, Any], revision: int, now: datetime, action: str,
        result: dict[str, Any], grants: dict[str, bool], duplicate: bool = False,
    ) -> dict[str, Any]:
        def readable_result(value):
            if isinstance(value, list):
                return [readable_result(row) for row in value]
            if not isinstance(value, dict):
                return deepcopy(value)
            public = {}
            for key, nested in value.items():
                if key == "invoice":
                    if not grants.get("can_paid_invoice_view"):
                        continue
                    # Never print a superseded/voided receipt on a retry.
                    invoice_id = nested.get("id") if isinstance(nested, dict) else None
                    public[key] = deepcopy(next((row for row in state["invoices"] if row.get("id") == invoice_id), None))
                    if invoice_id and public[key] is None:
                        public["voided"] = True
                elif key == "pending" and isinstance(nested, dict):
                    pending = next((row for row in state["pending"] if row.get("id") == nested.get("id")), None)
                    public[key] = deepcopy(pending) if pending and grants.get("can_invoice_view") and grants.get("can_pending_view") else {"id": nested.get("id")}
                elif key in {"customer", "purchase", "customers", "combo_purchase"} and not grants.get("can_customers_view"):
                    continue
                else:
                    public[key] = readable_result(nested)
            return public
        public_result = readable_result(result)
        if not (grants.get("can_customers_view") or grants.get("can_invoice_view") or grants.get("can_paid_invoice_view")):
            public_result = _redact_customer_pii(public_result)
        return {
            "ok": True, "duplicate": duplicate, "action": action,
            "revision": revision, "result": public_result,
            **({"capabilities": {name: grants[f"can_{name}"] for name in CAPABILITY_FEATURES}} if action.startswith("report_invoice_") else _state_response(state, revision, now, **grants)),
        }

    @app.get("/v2/live-tour")
    def live_tour(
        include_hidden: bool = Query(default=False),
        ident: identity_type = Depends(current_identity),
    ):
        now = datetime.now(timezone)
        with engine_instance().begin() as conn:
            require_feature(conn, ident, "live_tour_view")
            # Creation and reads share the same lock only for the first bootstrap.
            conn.execute(text("SELECT pg_advisory_xact_lock(hashtext(:key))"), {"key": STATE_LOCK})
            state, revision = _read_state(conn, now)
            grants = permissions(conn, ident)
        return _state_response(
            state, revision, now, include_hidden=include_hidden,
            **grants,
        )

    @app.get("/v2/live-tour/reports")
    def live_tour_reports(ident: identity_type = Depends(current_identity)):
        now = datetime.now(timezone)
        with engine_instance().begin() as conn:
            require_feature(conn, ident, "live_tour_reports_view")
            conn.execute(text("SELECT pg_advisory_xact_lock(hashtext(:key))"), {"key": STATE_LOCK})
            state, revision = _read_state(conn, now)
            grants = permissions(conn, ident)
        public = _state_response(state, revision, now, **grants)
        return {"revision": revision, "invoices": public["state"]["invoices"],
                "reports": public["report_rows"], "capabilities": public["capabilities"], "payment_settings": public["payment_settings"]}

    @app.get("/v2/live-tour/customers")
    def spa_customers(ident: identity_type = Depends(current_identity)):
        now = datetime.now(timezone)
        with engine_instance().begin() as conn:
            require_feature(conn, ident, "live_tour_customers_view")
            conn.execute(text("SELECT pg_advisory_xact_lock(hashtext(:key))"), {"key": STATE_LOCK})
            state, revision = _read_state(conn, now)
            can_export = bool(feature_allowed(conn, ident, "live_tour_export"))
        return {"revision": revision, "customers": [dict(deepcopy(c), combo_purchases=[deepcopy(p) for p in c.get("combo_purchases", []) if not p.get("deleted_at")]) for c in state["customers"] if not c.get("deleted_at")], "can_export": can_export}

    @app.get("/v2/live-tour/settings")
    def spa_settings(ident: identity_type = Depends(current_identity)):
        now = datetime.now(timezone)
        with engine_instance().begin() as conn:
            require_feature(conn, ident, "live_tour_admin")
            conn.execute(text("SELECT pg_advisory_xact_lock(hashtext(:key))"), {"key": STATE_LOCK})
            state, revision = _read_state(conn, now)
        return {"revision": revision, "services": deepcopy(state["services"]), "combos": deepcopy(state["combos"]), "service_areas": _service_areas(state)}

    @app.get("/v2/live-tour/customers/{customer_id}/history")
    def live_tour_customer_history(
        customer_id: str,
        ident: identity_type = Depends(current_identity),
    ):
        now = datetime.now(timezone)
        with engine_instance().begin() as conn:
            require_feature(conn, ident, "live_tour_customers_view")
            conn.execute(text("SELECT pg_advisory_xact_lock(hashtext(:key))"), {"key": STATE_LOCK})
            state, _ = _read_state(conn, now)
            grants = permissions(conn, ident)
        readable = deepcopy(state)
        if not grants["can_paid_invoice_view"]:
            readable["invoices"] = []
        if not (grants["can_pending_view"] and grants["can_invoice_view"]):
            readable["pending"] = []
        if not grants["can_reports_view"]:
            readable["reports"] = []
        history = _customer_history(readable, customer_id)
        if not grants["can_paid_invoice_view"]:
            for key in ("invoice_count", "service_count", "total_revenue", "total_tip"):
                history["summary"].pop(key, None)
        if not (grants["can_pending_view"] and grants["can_invoice_view"]):
            history["summary"].pop("pending_count", None)
        return {**history, "capabilities": {
            name: grants[f"can_{name}"] for name in CAPABILITY_FEATURES
        }}

    @app.post("/v2/live-tour/action")
    def live_tour_action(body: LiveTourAction, ident: identity_type = Depends(current_identity)):
        now = datetime.now(timezone)
        action = body.action.strip().lower()
        _reject_external_action(action)
        if action == "update_appointment" and str(getattr(ident, "role", "") or "").strip().lower() not in {"admin", "quanly", "letan"}:
            raise HTTPException(403, "Chỉ Lễ tân, Quản lý và Admin được sửa lịch hẹn.")
        if action == "combo_import" and str(getattr(ident, "role", "") or "").strip().lower() != "admin":
            raise HTTPException(403, "Chỉ Admin được nhập combo.")
        payload = deepcopy(body.payload)
        actor = str(ident.employee_username or ident.full_name or "web_v2")
        idempotency_key = str(body.idempotency_key or payload.get("idempotency_key") or "").strip()
        if idempotency_key and not 8 <= len(idempotency_key) <= 160:
            raise HTTPException(400, "idempotency_key phải có từ 8 đến 160 ký tự.")
        if action in IDEMPOTENCY_REQUIRED_ACTIONS and not idempotency_key:
            raise HTTPException(400, "Mọi thao tác thay đổi Live Tour cần idempotency_key để chống ghi trùng.")
        payload_hash = _canonical_payload_hash(action, payload)
        with engine_instance().begin() as conn:
            require_feature(conn, ident, "live_tour_customers_edit" if action == "customer_upsert" and payload.get("customer_id") else _required_action_feature(action))
            if action in {"booking", "multi_booking"}:
                require_feature(conn, ident, "live_tour_view")
            if "quick_booking" in payload:
                if action != "quick_checkout":
                    raise HTTPException(400, "Nhập booking trực tiếp chỉ dùng trong Thanh toán nhanh.")
                require_feature(conn, ident, "live_tour_booking")
                require_feature(conn, ident, "live_tour_view")
                booked = _quick_booking_at(payload["quick_booking"], now)
                if booked.date() < now.astimezone(VN_TZ).date():
                    require_feature(conn, ident, "live_tour_admin")
            if action in {"booking", "multi_booking", "update_booking"} and (
                payload.get("start_now") or any(row.get("start_now") for row in (payload.get("bookings") or []) if isinstance(row, dict))
            ):
                require_feature(conn, ident, "live_tour_operate")
            if action in {"pending_update", "pending_delete"} or (action in {"checkout", "quick_checkout"} and payload.get("pending_id")):
                require_feature(conn, ident, "live_tour_pending_view")
                require_feature(conn, ident, "live_tour_invoice_view")
            if action in {"report_invoice_update", "report_invoice_delete"}:
                require_feature(conn, ident, "live_tour_reports_view")
                require_feature(conn, ident, "live_tour_paid_invoice_view")
            if "invoice_at" in payload:
                require_feature(conn, ident, "live_tour_invoice_date_edit")
            if action == "customer_upsert" and payload.get("customer_id"):
                require_feature(conn, ident, "live_tour_customers_edit")
            if action in {"customer_delete", "customer_combo_update", "customer_combo_delete"}:
                require_feature(conn, ident, "live_tour_customers_view")
            if action in {"paid_invoice_update", "paid_invoice_delete"}:
                require_feature(conn, ident, "live_tour_paid_invoice_view")
            if action in BACKDATE_ACTIONS and _backdate_requested(payload):
                # Lùi ngày changes the financial ledger date and therefore
                # requires both the normal payment grant and the admin grant.
                require_feature(conn, ident, "live_tour_admin")
            if action in {"booking", "multi_booking", "update_booking", "finish_to_pending"} and (
                _contains_customer_pii(payload) or "combo_purchase_id" in payload
                or any("combo_purchase_id" in row for row in (payload.get("bookings") or []) if isinstance(row, dict))
            ):
                require_feature(conn, ident, "live_tour_customers_view")
            if action == "combo_import":
                require_feature(conn, ident, "live_tour_payment")
            if action in {"combo_import", "combo_purchase", "customer_upsert"}:
                require_feature(conn, ident, "live_tour_customers_view")
            if action in {"checkout", "quick_checkout"} and _contains_customer_pii(payload):
                require_feature(conn, ident, "live_tour_customers_view")
            conn.execute(text("SELECT pg_advisory_xact_lock(hashtext(:key))"), {"key": STATE_LOCK})
            state, revision = _read_state(conn, now, for_update=True)
            previous = _idempotency_replay(
                state, idempotency_key, action=action, actor=actor, payload_hash=payload_hash,
            )
            grants = permissions(conn, ident)
            if previous:
                return action_response(
                    state=state, revision=revision, now=now, action=action,
                    result=deepcopy(previous.get("result") or {}), grants=grants, duplicate=True,
                )
            if body.expected_revision is None:
                raise HTTPException(428, "Thiếu phiên bản Live Tour. Hãy tải lại bảng trước khi thao tác.")
            if body.expected_revision != revision:
                raise HTTPException(409, "Live Tour đã thay đổi ở thiết bị khác. Hãy làm mới rồi thao tác lại.")
            if action == "clear_expired_preview":
                return {"ok": True, "base_revision": revision, **_expired_preview(state, payload, now)}
            # A defensive copy guarantees multi-step actions never leak a partial
            # mutation into the value written after an exception.
            working = deepcopy(state)
            if action == "sync_daily_status":
                payload = {**payload, "directory": _employee_directory(conn), "leaves": [dict(row) for row in conn.execute(text("SELECT employee_name, leave_reason, leave_type FROM leave_records WHERE leave_date=:day ORDER BY id"), {"day": now.astimezone(VN_TZ).date()}).mappings().all()]}
            result = _apply_action(working, action, payload, actor, now)
            if action in {"checkout", "quick_checkout", "combo_purchase"} and result.get("invoice"):
                bank = _selected_bank(working.get("payment_settings") or {}, grants.get("viewer_bank"), payload.get("bank_selection", "auto"))
                result["invoice"]["payment_bank"] = bank
                for invoice in working["invoices"]:
                    if invoice["id"] == result["invoice"]["id"]:
                        invoice["payment_bank"] = deepcopy(bank)
            if idempotency_key:
                _remember_idempotency(
                    working, idempotency_key, action=action, actor=actor,
                    payload_hash=payload_hash, result=result, now=now,
                )
            next_revision = _write_state(conn, working, revision, actor)
        return action_response(
            state=working, revision=next_revision, now=now, action=action,
            result=result, grants=grants,
        )

    @app.get("/v2/live-tour/export.xlsx")
    def live_tour_export_excel(
        kind: str = Query(default="board"), include_hidden: bool = Query(default=False),
        date_from: str = Query(default=""), date_to: str = Query(default=""),
        time_from: str = Query(default=""), time_to: str = Query(default=""),
        customer_id: str = Query(default=""),
        employee: str = "", customer: str = "", service: str = "",
        report_kind: str = "",
        columns: list[str] | None = Query(default=None),
        employee_ids: list[str] | None = Query(default=None),
        ident: identity_type = Depends(current_identity),
    ):
        now = datetime.now(timezone)
        export_kind = kind.strip().lower()
        if export_kind not in {
            "board", "custom", "revenue", "tip", "reports", "customers", "pending", "history",
            "breaks", "customer_detail",
        }:
            raise HTTPException(400, "Loại báo cáo Live Tour không hợp lệ.")
        customer_id_value = str(customer_id or "").strip()
        if export_kind == "customer_detail" and not customer_id_value:
            raise HTTPException(400, "Xuất lịch sử khách hàng cần customer_id.")
        bounds = _parse_export_bounds(
            date_from=date_from.strip(), date_to=date_to.strip(),
            time_from=time_from.strip(), time_to=time_to.strip(),
        )
        bounds.update(employee=employee.strip(), customer=customer.strip(), service=service.strip(), report_kind=report_kind.strip(), calendar_date=export_kind in {"revenue", "tip", "reports", "pending"})
        with engine_instance().begin() as conn:
            require_feature(conn, ident, "live_tour_export")
            for feature in EXPORT_FEATURES.get(export_kind, ()):
                require_feature(conn, ident, feature)
            conn.execute(text("SELECT pg_advisory_xact_lock(hashtext(:key))"), {"key": STATE_LOCK})
            state, _ = _read_state(conn, now)
            can_admin = feature_allowed(conn, ident, "live_tour_admin")
            can_recover_hidden = can_admin or feature_allowed(conn, ident, "live_tour_operate")
            grants = permissions(conn, ident)
        if export_kind == "history":
            state = deepcopy(state)
            state["audit"] = _readable_audit(state["audit"], invoice_view=grants["can_invoice_view"] and grants["can_pending_view"],
                                             paid_invoice_view=grants["can_paid_invoice_view"], customers_view=grants["can_customers_view"])
        elif export_kind == "revenue" and not (grants["can_customers_view"] or grants["can_paid_invoice_view"]):
            state = deepcopy(state)
            state["invoices"] = _redact_customer_pii(state["invoices"])
        elif export_kind == "reports" and not (grants["can_customers_view"] or grants["can_paid_invoice_view"]):
            state = deepcopy(state)
            state["reports"] = _redact_customer_pii(state["reports"])
        content, filename = _excel_bytes(
            state, export_kind, now, include_hidden=bool(include_hidden and can_recover_hidden),
            bounds=bounds, customer_id=customer_id_value, selected_columns=columns, employee_ids=employee_ids,
        )
        return StreamingResponse(
            BytesIO(content),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
        )

    @app.get("/v2/live-tour/export.png")
    def live_tour_export_png(
        kind: str = Query(default="board"), include_hidden: bool = Query(default=False),
        ident: identity_type = Depends(current_identity),
    ):
        if kind.strip().lower() != "board":
            raise HTTPException(400, "Export PNG hiện chỉ hỗ trợ Bảng tua.")
        now = datetime.now(timezone)
        with engine_instance().begin() as conn:
            require_feature(conn, ident, "live_tour_export")
            conn.execute(text("SELECT pg_advisory_xact_lock(hashtext(:key))"), {"key": STATE_LOCK})
            state, _ = _read_state(conn, now)
            can_admin = feature_allowed(conn, ident, "live_tour_admin")
            can_recover_hidden = can_admin or feature_allowed(conn, ident, "live_tour_operate")
        filename = f"Live_Tour_{_business_date(now).strftime('%Y%m%d')}.png"
        return StreamingResponse(
            BytesIO(_png_bytes(state, now, include_hidden=bool(include_hidden and can_recover_hidden))), media_type="image/png",
            headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
        )

    app.state.live_tour_installed = True
