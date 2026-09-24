"""Admin device inventory and bounded check-in reporting.

Inventory addresses are metadata only. Only the deployed FaceGate adapter can
perform I/O; neither inventory writes nor exports probe arbitrary destinations.
"""
from __future__ import annotations

from datetime import date, datetime, timedelta, timezone
from io import BytesIO
import json
import os
import re
from typing import Literal
import unicodedata

from fastapi import Depends, HTTPException, Query, Request
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from pydantic import BaseModel, ConfigDict, Field, model_validator
from sqlalchemy import text

VN = timezone(timedelta(hours=7))
KINDS = {'faceid': 'FaceID / Chấm công', 'printer': 'Máy in', 'scanner': 'Máy quét', 'screen': 'Màn hình', 'other': 'Thiết bị khác'}
CONNECTIONS = {'network': 'Mạng LAN / TCP/IP', 'usb': 'USB', 'bluetooth': 'Bluetooth', 'serial': 'Cổng nối tiếp', 'agent': 'Qua máy trạm', 'other': 'Khác'}
SOURCES = {'facegate_saved': 'FaceGate · Đã lưu trong VERA', 'facegate': 'FaceGate · Control Log', 'capture': 'FaceGate · Capture Log', 'timesoft': 'TimeSoft · Đã đồng bộ VERA'}


class Device(BaseModel):
    model_config = ConfigDict(extra='forbid', str_strip_whitespace=True)
    id: str = Field(pattern=r'^[A-Za-z0-9_-]{1,64}$')
    name: str = Field(min_length=1, max_length=160)
    kind: Literal['faceid', 'printer', 'scanner', 'screen', 'other']
    connection: Literal['network', 'usb', 'bluetooth', 'serial', 'agent', 'other'] = 'network'
    manufacturer: str = Field(default='', max_length=100)
    model: str = Field(default='', max_length=100)
    serial: str = Field(default='', max_length=100)
    location: str = Field(default='', max_length=160)
    address: str = Field(default='', max_length=253, pattern=r'^[A-Za-z0-9_.:\[\]-]*$')
    port: int | None = Field(default=None, ge=1, le=65535)
    notes: str = Field(default='', max_length=1000)
    enabled: bool = True
    adapter: Literal['pending', 'facegate_server'] = 'pending'

    @model_validator(mode='after')
    def validate_adapter(self):
        if self.adapter == 'facegate_server' and (self.id != 'facegate-current' or self.kind != 'faceid'):
            raise ValueError('Bộ kết nối FaceGate hiện tại chỉ dành cho hồ sơ FaceGate đang sử dụng.')
        if self.id == 'facegate-current' and self.adapter != 'facegate_server':
            raise ValueError('Giữ bộ kết nối của hồ sơ FaceGate đang sử dụng.')
        return self


class RegistryInput(BaseModel):
    model_config = ConfigDict(extra='forbid')
    expected_revision: int = Field(ge=0)
    devices: list[Device] = Field(max_length=100)

    @model_validator(mode='after')
    def unique_devices(self):
        if len({d.id for d in self.devices}) != len(self.devices):
            raise ValueError('ID thiết bị bị trùng.')
        if not any(d.id == 'facegate-current' for d in self.devices):
            raise ValueError('Giữ hồ sơ FaceGate hiện tại; có thể ngừng sử dụng thay vì xóa.')
        serials = [d.serial.casefold() for d in self.devices if d.serial]
        if len(set(serials)) != len(serials):
            raise ValueError('Số serial thiết bị bị trùng.')
        return self


def default_devices():
    # This is the already configured device, not a claimed live health result.
    return [Device(id='facegate-current', name='Máy nhận diện chấm công', kind='faceid',
                   serial='2023044', adapter='facegate_server').model_dump()]


def read_registry(conn, *, lock=False):
    row = conn.execute(text("SELECT value_json, revision FROM vera_app_setting WHERE category='devices' AND setting_key='registry'" + (' FOR UPDATE' if lock else ''))).mappings().first()
    if not row:
        return {'devices': default_devices(), 'revision': 0}
    value = row['value_json']
    if isinstance(value, str):
        value = json.loads(value)
    return {'devices': value['devices'], 'revision': int(row['revision'])}


def registry_result(data):
    configured = all(os.getenv(key, '').strip() for key in ('VERA_FACEGATE_BASE_URL', 'VERA_FACEGATE_USERNAME', 'VERA_FACEGATE_PASSWORD'))
    return {**data, 'kinds': KINDS, 'connections': CONNECTIONS, 'adapters': {
        'facegate_server': {'label': 'FaceGate đang sử dụng', 'configured': configured,
                            'capabilities': ['control_log', 'capture_log', 'capture_image', 'profile_mapping']},
        'pending': {'label': 'Chờ tích hợp bộ kết nối', 'configured': False, 'capabilities': []},
    }}


def norm(value):
    return ' '.join(''.join(c for c in unicodedata.normalize('NFD', str(value or '').casefold()) if unicodedata.category(c) != 'Mn').replace('đ', 'd').split())


def record_day(item):
    value = str(item.get('occurred_at') or item.get('date') or '')
    try:
        if 'T' in value:
            dt = datetime.fromisoformat(value.replace('Z', '+00:00'))
            return dt.replace(tzinfo=dt.tzinfo or VN).astimezone(VN).date().isoformat()
        if re.fullmatch(r'\d{2}/\d{2}/\d{4}', value):
            return datetime.strptime(value, '%d/%m/%Y').date().isoformat()
        return date.fromisoformat(value).isoformat()
    except ValueError:
        return ''


def record_status(item, source):
    value = item.get({'facegate_saved': 'status_code', 'facegate': 'status_code', 'capture': 'event_status', 'timesoft': 'arrival_status'}[source])
    return str(value if value is not None else '')


def record_type(item, source):
    value = item.get({'facegate_saved': 'type_code', 'facegate': 'type_code', 'capture': 'event_text', 'timesoft': 'departure_status'}[source])
    return str(value if value is not None else '')


def filtered_records(records, source, employee='', event_id='', event_date=None, status='', event_type=''):
    return [r for r in records
            if (not employee or norm(employee) in norm(' '.join(str(r.get(k) or '') for k in ('employee_name', 'employee_code', 'device_name'))))
            and (not event_id or event_id.casefold() in str(r.get('event_id', '')).casefold())
            and (not event_date or record_day(r) == event_date.isoformat())
            and (not status or record_status(r, source) == status)
            and (not event_type or record_type(r, source) == event_type)]


def display_time(value):
    try:
        dt = datetime.fromisoformat(str(value).replace('Z', '+00:00'))
        return dt.replace(tzinfo=dt.tzinfo or VN).astimezone(VN).strftime('%d-%m-%Y %H:%M:%S')
    except ValueError:
        return ''


def history_workbook(payload):
    source = payload['source']
    if source in ('facegate', 'facegate_saved'):
        columns = [('event_id', 'Mã sự kiện'), ('occurred_at', 'Thời điểm'), ('device_name', 'Tên trên máy'), ('status_code', 'Mã trạng thái'), ('type_code', 'Mã loại trên máy')]
        if source == 'facegate_saved':
            columns += [('employee_name', 'Nhân viên đã đối chiếu'), ('employee_code', 'Mã TimeSoft'), ('mapping_status', 'Trạng thái ánh xạ')]
    elif source == 'capture':
        columns = [('event_id', 'Mã sự kiện'), ('occurred_at', 'Thời điểm'), ('event_text', 'Loại sự kiện'), ('event_status', 'Trạng thái trên máy'), ('image_available', 'Có ảnh')]
    else:
        columns = [('date', 'Ngày'), ('employee_code', 'Mã nhân viên'), ('employee_name', 'Nhân viên'), ('check_in', 'Giờ vào'), ('punch_times', 'Các lần chấm'), ('check_out', 'Giờ ra'), ('arrival_status', 'Trạng thái vào'), ('departure_status', 'Trạng thái ra')]
    wb = Workbook(); ws = wb.active; ws.title = 'Lịch sử checkin'
    ws.append([label for _, label in columns])
    border = Border(*( [Side(style='thin', color='BDD0C4')] * 4 ))
    for record in payload['records']:
        values = []
        for key, _ in columns:
            value = record.get(key, '')
            if key == 'date':
                iso = record_day(record)
                value = date.fromisoformat(iso).strftime('%d-%m-%Y') if iso else ''
            elif key == 'occurred_at':
                value = display_time(value)
            elif key == 'image_available':
                value = 'Có' if value else 'Không'
            elif isinstance(value, list):
                value = ' · '.join(str(v) for v in value)
            values.append(str(value if value is not None else ''))
        ws.append(values)
        # Device names/codes are untrusted; always emit literal strings, never formulas.
        for cell in ws[ws.max_row]:
            cell.data_type = 's'
    for row in ws:
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='23543F')
    from openpyxl.utils import get_column_letter
    for index, _ in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(index)].width = 27
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    info = wb.create_sheet('Thông tin')
    for row in [('Nguồn', SOURCES[source]), ('Từ ngày', date.fromisoformat(payload['start']).strftime('%d-%m-%Y')),
                ('Đến ngày', date.fromisoformat(payload['end']).strftime('%d-%m-%Y')),
                ('Số bản ghi', len(payload['records'])), ('Xuất lúc', datetime.now(VN).strftime('%d-%m-%Y %H:%M:%S')),
                ('Phạm vi', 'Dữ liệu tra cứu; không xác nhận công/lương. Không xuất ảnh hoặc thông tin xác thực.')]:
        info.append(row)
    for key, value in payload['filters'].items():
        labels = {'employee': 'Nhân viên / tên trên máy', 'event_id': 'Mã sự kiện', 'event_date': 'Ngày cụ thể', 'status': 'Trạng thái', 'event_type': 'Loại sự kiện / trạng thái ra'}
        if key == 'event_date' and value:
            value = date.fromisoformat(value).strftime('%d-%m-%Y')
        info.append([labels[key], str(value or 'Tất cả')]); info.cell(info.max_row, 2).data_type = 's'
    info.column_dimensions['A'].width = 32; info.column_dimensions['B'].width = 90
    stream = BytesIO(); wb.save(stream)
    return stream.getvalue()


def saved_facegate_history(conn, start, end):
    """Read archived evidence without probing the device or changing attendance."""
    from vera_facegate_control_log import mapping_device_id
    from vera_facegate_readiness import reference
    device_id = mapping_device_id()
    rows = conn.execute(text('''SELECT event_id, occurred_at, payload_json FROM vera_facegate_event
        WHERE device_id=:device_id AND work_date BETWEEN :start AND :end
        ORDER BY occurred_at, event_id LIMIT 10001'''),
        {'device_id': device_id, 'start': start.isoformat(), 'end': end.isoformat()}).mappings().all()
    raw = conn.execute(text("SELECT value_json FROM vera_app_setting WHERE category='facegate' AND setting_key=:key"),
                       {'key': 'mapping_' + device_id}).scalar()
    mappings = json.loads(raw) if isinstance(raw, str) else (raw or [])
    by_ref = {}
    for mapping in mappings if isinstance(mappings, list) else []:
        if not isinstance(mapping, dict) or not mapping.get('confirmed_by'):
            continue
        key = reference(mapping.get('registration_ref'))
        if key:
            by_ref.setdefault(key, []).append(mapping)
    records = []
    for row in rows[:10000]:
        payload = row['payload_json']
        payload = json.loads(payload) if isinstance(payload, str) else payload
        if not isinstance(payload, dict):
            continue
        key = reference(payload.get('registration_ref'))
        matched = by_ref.get(key, []) if key else []
        records.append({'event_id': row['event_id'], 'occurred_at': row['occurred_at'],
                        'device_name': payload.get('device_name', ''),
                        'status_code': payload.get('status_code', ''),
                        'type_code': payload.get('type_code', ''),
                        'registration_ref': payload.get('registration_ref'),
                        'mapping_status': 'reference_match' if len(matched) == 1 else 'unmapped',
                        'employee_name': matched[0].get('username', '') if len(matched) == 1 else '',
                        'employee_code': matched[0].get('employee_code', '') if len(matched) == 1 else ''})
    return {'records': records, 'truncated': len(rows) > 10000, 'total_count': len(rows)}


def install_device_routes(app, *, engine_instance, current_identity, require_feature, identity_type, read_timesoft):
    def admin(ident):
        if str(getattr(ident, 'role', '')).strip().lower() != 'admin':
            raise HTTPException(403, 'Chỉ Admin được quản lý thiết bị và lịch sử checkin.')

    @app.get('/v2/devices/registry')
    def registry(ident: identity_type = Depends(current_identity)):
        admin(ident)
        with engine_instance().connect() as conn:
            data = read_registry(conn)
        return registry_result(data)

    @app.put('/v2/devices/registry')
    def save_registry(body: RegistryInput, ident: identity_type = Depends(current_identity)):
        admin(ident)
        actor = str(getattr(ident, 'employee_username', '') or getattr(ident, 'username', ''))
        with engine_instance().begin() as conn:
            conn.execute(text("""INSERT INTO vera_app_setting(category,setting_key,value_json,source,updated_by,revision,created_at,updated_at)
                VALUES ('devices','registry',CAST(:initial AS jsonb),'web_v2',:actor,0,NOW(),NOW())
                ON CONFLICT(category,setting_key) DO NOTHING"""), {'initial': json.dumps({'devices': default_devices()}), 'actor': actor})
            current = read_registry(conn, lock=True)
            if current['revision'] != body.expected_revision:
                raise HTTPException(409, 'Danh sách thiết bị vừa được sửa ở phiên khác. Hủy chỉnh sửa và tải lại danh sách trước khi lưu.')
            devices = [d.model_dump() for d in body.devices]
            conn.execute(text("""UPDATE vera_app_setting SET value_json=CAST(:value AS jsonb), updated_by=:actor,
                updated_at=NOW(),revision=revision+1 WHERE category='devices' AND setting_key='registry'"""),
                {'value': json.dumps({'devices': devices}, ensure_ascii=False), 'actor': actor})
        return registry_result({'devices': devices, 'revision': current['revision'] + 1})

    def query_history(start, end, source, employee, event_id, event_date, status, event_type, ident, *, exporting=False):
        admin(ident)
        if end < start or end - start > timedelta(days=62):
            raise HTTPException(400, 'Chọn khoảng ngày hợp lệ, tối đa 63 ngày mỗi lần.')
        if event_date and not start <= event_date <= end:
            raise HTTPException(400, 'Ngày cụ thể phải nằm trong khoảng tra cứu.')
        if (source == 'capture' and employee) or (source == 'timesoft' and event_id):
            raise HTTPException(400, 'Bộ lọc không phù hợp nguồn dữ liệu đã chọn.')
        if source == 'facegate_saved':
            with engine_instance().connect() as conn:
                data = saved_facegate_history(conn, start, end)
        elif source == 'timesoft':
            with engine_instance().begin() as conn:
                require_feature(conn, ident, 'snapshot_export' if exporting else 'snapshot_today')
                records = read_timesoft(conn, start, end)
            data = {'records': records, 'truncated': False, 'total_count': len(records)}
        else:
            from vera_facegate_control_log import fetch_control_log, fetch_capture_log
            try:
                data = (fetch_control_log if source == 'facegate' else fetch_capture_log)(start.isoformat(), end.isoformat())
            except RuntimeError as exc:
                raise HTTPException(503, str(exc)) from exc
            except ValueError as exc:
                raise HTTPException(400, str(exc)) from exc
            except ConnectionError as exc:
                raise HTTPException(502, str(exc)) from exc
        records = data['records']
        return {**data, 'source': source, 'start': start.isoformat(), 'end': end.isoformat(),
                'records': filtered_records(records, source, employee, event_id, event_date, status, event_type),
                'filters': {'employee': employee, 'event_id': event_id, 'event_date': event_date.isoformat() if event_date else '', 'status': status, 'event_type': event_type},
                'options': {'statuses': sorted({record_status(r, source) for r in records} - {''}),
                            'types': sorted({record_type(r, source) for r in records} - {''})}}

    @app.get('/v2/devices/checkin-history')
    @app.get('/v2/devices/checkin-history/export.xlsx')
    def history(request: Request,
                start: date = Query(...), end: date = Query(...),
                source: Literal['facegate_saved', 'facegate', 'capture', 'timesoft'] = 'facegate',
                employee: str = Query('', max_length=200), event_id: str = Query('', max_length=64),
                event_date: date | None = None, status: str = Query('', max_length=160), event_type: str = Query('', max_length=160),
                ident: identity_type = Depends(current_identity)):
        exporting = request.url.path.endswith('/export.xlsx')
        payload = query_history(start, end, source, employee.strip(), event_id.strip(), event_date, status, event_type, ident, exporting=exporting)
        if not exporting:
            return payload
        if payload.get('truncated'):
            raise HTTPException(409, 'Dữ liệu vượt giới hạn của thiết bị. Thu hẹp khoảng ngày trước khi xuất Excel.')
        return Response(history_workbook(payload), media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                        headers={'Content-Disposition': f'attachment; filename="VERA_Checkin_{source}_{start}_{end}.xlsx"',
                                 'Cache-Control': 'private, no-store'})
