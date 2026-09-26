"""Revenue summary and leave-list read enhancements for VERA SPA Web V2."""
from __future__ import annotations
from vera_notification_delivery import route_event as route_notification, enqueue as enqueue_notification

import csv
from datetime import date, datetime, timedelta, timezone
import json
import os
import re
from typing import Any, Callable, Literal

from fastapi import BackgroundTasks, Depends, HTTPException, Query, Request
from fastapi.responses import StreamingResponse
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, Field
import requests
from sqlalchemy import text

from vera_progressive_penalty import applies as progressive_penalty_applies
from vera_progressive_penalty import load_weekend_unpaid_enabled

import vera_web_v2_permissions as permissions
import vera_revenue_store as revenue_store
import vera_revenue_auto as revenue_auto


RELEASE = "revenue-shared-auto-daily-2026-09-26-v1"
REVENUE_FEATURE = "revenue_view"
REVENUE_TIP_FEATURE = "revenue_tip_edit"
REVENUE_ENTRY_FEATURE = "revenue_entry_create"
REVENUE_ENTRY_EDIT_FEATURE = "revenue_entry_edit"
REVENUE_ENTRY_DELETE_FEATURE = "revenue_entry_delete"
REVENUE_TIP_SETTING = "current_period_tip"
REVENUE_SPREADSHEET_ID = os.getenv(
    "VERA_REVENUE_SHEET_ID",
    "1KLYz2iQSfNU0xOfrl8V9iz9-dQKMkwyUicuuiGsqC4U",
)
REVENUE_WORKSHEET = os.getenv("VERA_REVENUE_SHEET_NAME", "Input")
REVENUE_REPORT_WORKSHEET = os.getenv("VERA_REVENUE_REPORT_SHEET_NAME", "Report").strip() or "Report"
REVENUE_INPUT_GID = os.getenv("VERA_REVENUE_INPUT_GID", "2058724516").strip() or "2058724516"
REVENUE_REPORT_URL = os.getenv(
    "VERA_REVENUE_REPORT_URL",
    "https://docs.google.com/spreadsheets/d/1KLYz2iQSfNU0xOfrl8V9iz9-dQKMkwyUicuuiGsqC4U/edit?usp=drivesdk",
)
REVENUE_ENTRY_FORM_URL = os.getenv(
    "VERA_REVENUE_ENTRY_FORM_URL",
    "https://docs.google.com/forms/d/e/1FAIpQLSeJp1bLrl8zSyESu_K0eo6NxdKsm85p4fxGXPXigPlmgkAs7w/viewform",
)
VN_TZ = timezone(timedelta(hours=7))
REVENUE_PERIOD_START = date(2026, 9, 5)
REVENUE_CURRENT_DATE_COLUMN_INDEX = 4  # Sheet Input, column E.
DATE_IN_TEXT_RE = re.compile(
    r"(?<!\d)(\d{1,2})\s*([./-])\s*(\d{1,2})\s*\2\s*(\d{4})(?!\d)"
)


class RevenueSourceUpdate(BaseModel):
    source: Literal["manual", "auto", "manual_tip_auto"]
    revision: int = Field(ge=0)


class RevenueTipPeriod(BaseModel):
    start_date: date
    end_date: date


class RevenueTipUpdate(BaseModel):
    amount: float = Field(ge=0, le=10_000_000_000_000)
    start_date: date | None = None
    end_date: date | None = None


class RevenueEntryCreate(BaseModel):
    transaction_date: date
    income_amount: float = Field(default=0, ge=0, le=10_000_000_000_000)
    income_note: str = Field(default="", max_length=1000)
    expense_amount: float = Field(default=0, ge=0, le=10_000_000_000_000)
    expense_note: str = Field(default="", max_length=1000)
    confirm_duplicate: bool = False


class RevenueEntryUpdate(BaseModel):
    transaction_type: Literal["Thu", "Chi"]
    amount: float = Field(ge=0, le=10_000_000_000_000)
    transaction_date: date | None = None
    note: str = Field(default="", max_length=1000)
    entered_by_name: str | None = Field(default=None, max_length=200)
    entered_date: date | None = None
    entered_time: str | None = Field(default=None, max_length=8)


def _range_bounds(time_range: str, start: date | None = None, end: date | None = None) -> tuple[date | None, date | None]:
    today = datetime.now(VN_TZ).date()
    token = str(time_range or "all").strip().lower()
    if token == "all":
        return None, None
    if token == "yesterday":
        day = today - timedelta(days=1); return day, day
    if token == "today":
        return today, today
    if token == "this_week":
        first = today - timedelta(days=today.weekday()); return first, first + timedelta(days=6)
    if token == "last_week":
        last = today - timedelta(days=today.weekday() + 1); return last - timedelta(days=6), last
    if token == "this_month":
        first = today.replace(day=1)
        next_month = (first.replace(day=28) + timedelta(days=4)).replace(day=1)
        return first, next_month - timedelta(days=1)
    if token == "last_month":
        end_day = today.replace(day=1) - timedelta(days=1); return end_day.replace(day=1), end_day
    if token == "custom":
        if not start or not end or start > end:
            raise HTTPException(400, "Khoảng ngày tùy chỉnh không hợp lệ.")
        return start, end
    raise HTTPException(400, "Bộ lọc thời gian không hợp lệ.")


def _revenue_audit_workbook(rows: list[dict[str, Any]]) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Lịch sử sửa-xóa"
    headers = [
        "Thời điểm", "Hành động", "Mã dòng", "Người thao tác", "Phiên bản",
        "Ngày", "Loại giao dịch", "Số tiền", "Ghi chú", "Ngày nhập", "Giờ nhập", "Người nhập",
    ]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F513F")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for audit in rows:
        versions = [("Trước", audit.get("before") or {})]
        if audit.get("after"):
            versions.append(("Sau", audit["after"]))
        for version_label, payload in versions:
            transaction_date = _parse_date(payload.get("transaction_date"))
            entered_at = payload.get("entered_at")
            try:
                entered = datetime.fromisoformat(str(entered_at).replace("Z", "+00:00")) if entered_at else None
                if entered and entered.tzinfo is None:
                    entered = entered.replace(tzinfo=VN_TZ)
                entered = entered.astimezone(VN_TZ) if entered else None
            except ValueError:
                entered = None
            sheet.append([
                audit.get("audited_at_label") or "", audit.get("action_label") or "",
                audit.get("entry_id"), audit.get("actor") or "", version_label,
                transaction_date.strftime("%d-%m-%Y") if transaction_date else "",
                payload.get("transaction_type") or "", float(payload.get("amount") or 0),
                payload.get("note") or "", entered.strftime("%d-%m-%Y") if entered else "",
                entered.strftime("%H:%M:%S") if entered else "",
                payload.get("entered_by_name") or payload.get("entered_by") or "",
            ])
    for cell in sheet["H"][1:]:
        cell.number_format = '#,##0"đ"'
    widths = [21, 13, 11, 20, 11, 14, 18, 18, 48, 14, 12, 26]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream


def _revenue_duplicate_workbook(result: dict[str, Any]) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Dữ liệu trùng"
    headers = ["Ngày", "Loại giao dịch", "Số tiền", "Ghi chú", "Số dòng", "Số dòng dư", "Người nhập", "Mã dòng"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F513F")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for group in result.get("groups") or []:
        sheet.append([
            group.get("date_label") or "", group.get("type") or "", float(group.get("amount") or 0),
            group.get("note") or "", int(group.get("count") or 0), int(group.get("extra_count") or 0),
            ", ".join(group.get("entered_by") or []), ", ".join(str(value) for value in (group.get("entry_ids") or [])),
        ])
    for cell in sheet["C"][1:]:
        cell.number_format = '#,##0"đ"'
    widths = [14, 18, 18, 48, 12, 14, 30, 24]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    stream = BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream


def _dispatch_revenue_admin_push(*, engine_instance, api_module, event: str, detail: dict[str, Any]) -> None:
    """Deliver after commit; notification failure never rolls back revenue data."""
    if api_module is None:
        return
    try:
        import vera_web_v2_notification_settings as notification_settings
        with engine_instance().connect() as conn:
            if (not notification_settings.is_enabled(conn, "revenue_manual_changes") or not notification_settings.is_channel_enabled(conn, "revenue_manual_changes", "push")):
                return
            private_key = api_module._vault_secret(conn, "vera_v2_vapid_private_key")
            subject = api_module._vault_secret(conn, "vera_v2_vapid_subject") or "https://app.veraspa.vn/"
            subscriptions = conn.execute(text("""
                SELECT s.subscription_id::text AS subscription_id,s.endpoint,s.p256dh,s.auth_secret
                FROM vera_v2_push_subscription s
                JOIN vera_v2_user_profile p ON p.auth_user_id=s.auth_user_id
                WHERE s.is_active=true AND p.is_active=true AND lower(COALESCE(p.role,''))='admin'
                ORDER BY s.updated_at DESC
            """)).mappings().all()
        action_label = {"create": "Nhập mới", "update": "Sửa", "delete": "Xóa", "import": "Import"}.get(event, event)
        amount = round(float(detail.get("amount") or 0))
        body = f"{detail.get('actor') or 'Hệ thống'} · {detail.get('type') or ''} {amount:,}đ · {detail.get('note') or ''}".replace(",", ".")
        results = []
        payload = {
                "title": f"VERA SPA · {action_label} Thu Chi thủ công",
                "body": body[:900], "url": "https://app.veraspa.vn/", "kind": "revenue-manual-change",
                "tag": f"vera-revenue-{event}-{detail.get('entry_id') or detail.get('nonce') or datetime.now().timestamp()}",
                "dismissible": True,
            }
        from uuid import uuid4
        payload['event_id'] = uuid4().hex
        if route_notification(engine_instance(), 'revenue_manual_changes', payload): return
        if not private_key: return
        for subscription in subscriptions:
            delivery = {**dict(subscription), 'payload': payload}
            ok, status, error_text = api_module._send_web_push(delivery, private_key, subject)
            results.append({"subscription_id": subscription["subscription_id"], "ok": bool(ok),
                            "inactive": (not ok) and status in {404, 410}, "last_error": str(error_text or "")[:1000]})
        if results:
            with engine_instance().begin() as conn:
                for result in results:
                    conn.execute(text("""
                        UPDATE vera_v2_push_subscription
                        SET is_active=CASE WHEN :inactive THEN false ELSE is_active END,
                            last_success_at=CASE WHEN :ok THEN NOW() ELSE last_success_at END,
                            failure_count=CASE WHEN :ok THEN 0 ELSE failure_count+1 END,
                            last_error=CASE WHEN :ok THEN NULL ELSE :last_error END,updated_at=NOW()
                        WHERE subscription_id=CAST(:subscription_id AS uuid)
                    """), result)
    except Exception:
        return


def _auto_revenue(conn, start_date: date | None, end_date: date | None) -> dict[str, Any]:
    days = revenue_auto.daily(conn, start_date, end_date)
    return {**revenue_auto.totals(days), "transaction_count": sum(row["payments"] for row in days),
            "entries": revenue_auto.ledger_rows(days)}


def _find_route(app, path: str, method: str):
    wanted = method.upper()
    for route in list(app.router.routes):
        methods = set(getattr(route, "methods", set()) or set())
        if getattr(route, "path", "") == path and wanted in methods:
            app.router.routes.remove(route)
            return route.endpoint
    raise RuntimeError(f"Cannot find {wanted} {path} to enhance")


def _money(value: Any) -> float:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    raw = str(value or "").strip()
    if not raw:
        return 0.0
    cleaned = re.sub(r"[^0-9,.\-]", "", raw)
    if not cleaned:
        return 0.0
    if "," in cleaned and "." in cleaned:
        decimal = "," if cleaned.rfind(",") > cleaned.rfind(".") else "."
        thousands = "." if decimal == "," else ","
        cleaned = cleaned.replace(thousands, "").replace(decimal, ".")
    elif cleaned.count(",") == 1 and len(cleaned.rsplit(",", 1)[1]) <= 2:
        cleaned = cleaned.replace(",", ".")
    elif cleaned.count(".") == 1 and len(cleaned.rsplit(".", 1)[1]) <= 2:
        pass
    else:
        cleaned = cleaned.replace(",", "").replace(".", "")
    try:
        return float(cleaned)
    except ValueError:
        return 0.0


def _dates_in_text(value: Any) -> list[date]:
    raw = str(value or "").strip()
    if not raw:
        return []
    parsed_dates: list[date] = []
    for match in DATE_IN_TEXT_RE.finditer(raw):
        try:
            parsed_dates.append(date(int(match.group(4)), int(match.group(3)), int(match.group(1))))
        except (TypeError, ValueError):
            pass
    return parsed_dates


def _parse_date(value: Any) -> date | None:
    raw = str(value or "").strip()
    if not raw:
        return None
    try:
        return datetime.strptime(raw.split(" ", 1)[0], "%Y-%m-%d").date()
    except ValueError:
        parsed_dates = _dates_in_text(raw)
        return parsed_dates[0] if parsed_dates else None


def _transaction_date_index(values: list[list[Any]], norm: Callable[[Any], str]) -> int:
    if not values:
        return -1
    headers = [norm(value) for value in values[0]]
    try:
        return headers.index(norm("Ngày giao dịch"))
    except ValueError:
        return -1


def _report_totals(values: list[list[Any]]) -> dict[str, float]:
    numeric_values: list[Any] = []
    for row in values:
        for value in row:
            if re.search(r"\d", str(value or "")):
                numeric_values.append(value)
    if len(numeric_values) < 2:
        raise HTTPException(503, "Sheet Report phải có Tổng thu tại B2 và Tổng chi tại B3.")
    return {
        "total_income": round(_money(numeric_values[0]), 2),
        "total_expense": round(_money(numeric_values[1]), 2),
    }


def _revenue_summary(
    values: list[list[Any]],
    norm: Callable[[Any], str],
    *,
    period_start: date | None = None,
) -> dict[str, Any]:
    if not values:
        raise HTTPException(503, "Sheet Input chưa có dữ liệu.")
    headers = [norm(value) for value in values[0]]
    try:
        type_index = headers.index(norm("Loại giao dịch"))
        amount_index = headers.index(norm("Số tiền"))
    except ValueError as exc:
        raise HTTPException(503, "Sheet Input phải có cột 'Loại giao dịch' và 'Số tiền'.") from exc

    date_index = _transaction_date_index(values, norm)

    income = expense = 0.0
    transaction_count = 0
    transaction_dates: list[date] = []
    input_column_e_dates: list[date] = []
    for row in values[1:]:
        tx_type = norm(row[type_index] if type_index < len(row) else "")
        if tx_type not in {"thu", "chi"}:
            continue
        parsed = _parse_date(row[date_index] if 0 <= date_index < len(row) else "")
        if period_start is not None and parsed is not None and parsed < period_start:
            continue
        if REVENUE_CURRENT_DATE_COLUMN_INDEX < len(row):
            input_column_e_dates.extend(_dates_in_text(row[REVENUE_CURRENT_DATE_COLUMN_INDEX]))
        amount = _money(row[amount_index] if amount_index < len(row) else 0)
        if tx_type == "thu":
            income += amount
        else:
            expense += amount
        transaction_count += 1
        if parsed:
            transaction_dates.append(parsed)

    start_date = period_start or (min(transaction_dates) if transaction_dates else None)
    if input_column_e_dates:
        current_date = max(input_column_e_dates)
        current_date_source = "input_column_e"
    elif transaction_dates:
        current_date = max(transaction_dates)
        current_date_source = "transaction_date_fallback"
    else:
        current_date = datetime.now(VN_TZ).date()
        current_date_source = "server_date_fallback"
    return {
        "total_income": round(income, 2),
        "total_expense": round(expense, 2),
        "transaction_count": transaction_count,
        "start_date": start_date.isoformat() if start_date else "",
        "current_date": current_date.isoformat(),
        "start_date_label": start_date.strftime("%d-%m-%Y") if start_date else "—",
        "current_date_label": current_date.strftime("%d-%m-%Y"),
        "current_date_source": current_date_source,
    }


def _period_tip(conn, default_start_date_text: str = "", default_end_date_text: str = "", *, auto: bool = False) -> dict[str, Any]:
    payload = conn.execute(text("""
        SELECT value_json
        FROM vera_app_setting
        WHERE category='revenue' AND setting_key=:key
        LIMIT 1
    """), {"key": REVENUE_TIP_SETTING + ("_auto" if auto else "")}).scalar_one_or_none()
    if not isinstance(payload, dict):
        payload = {}
    return {
        "amount": max(0.0, _money(payload.get("amount", 0))),
        "period_start": str(payload.get("period_start") or default_start_date_text or ""),
        "period_end": str(payload.get("period_end") or default_end_date_text or ""),
    }


def _save_period_tip(conn, start_date_text: str, end_date_text: str, amount: float, actor: str, *, auto: bool = False) -> None:
    payload = {
        "period_start": str(start_date_text or ""),
        "period_end": str(end_date_text or ""),
        "amount": round(max(0.0, float(amount)), 2),
    }
    conn.execute(text("""
        INSERT INTO vera_app_setting(
            category, setting_key, value_json, source, updated_by,
            revision, created_at, updated_at
        ) VALUES (
            'revenue', :key, CAST(:payload AS jsonb), 'web_v2', :actor,
            1, NOW(), NOW()
        )
        ON CONFLICT(category, setting_key) DO UPDATE SET
            value_json=EXCLUDED.value_json,
            source='web_v2',
            updated_by=EXCLUDED.updated_by,
            revision=vera_app_setting.revision+1,
            updated_at=NOW()
    """), {
        "key": REVENUE_TIP_SETTING + ("_auto" if auto else ""),
        "payload": json.dumps(payload, ensure_ascii=False),
        "actor": str(actor or ""),
    })


def _read_public_revenue_values() -> list[list[str]]:
    """Download every Input row, including rows hidden by a Sheet filter."""
    response = requests.get(
        f"https://docs.google.com/spreadsheets/d/{REVENUE_SPREADSHEET_ID}/export",
        params={"format": "csv", "gid": REVENUE_INPUT_GID},
        timeout=30,
    )
    response.raise_for_status()
    response.encoding = "utf-8"
    values = list(csv.reader(response.text.splitlines()))
    if not values:
        raise RuntimeError("Full Google Sheets CSV response is empty")
    return values


def _read_public_report_values() -> list[list[str]]:
    """Read the two official summary cells without depending on Sheet filters."""
    response = requests.get(
        f"https://docs.google.com/spreadsheets/d/{REVENUE_SPREADSHEET_ID}/gviz/tq",
        params={
            "tqx": "out:csv",
            "sheet": REVENUE_REPORT_WORKSHEET,
            "range": "B2:B3",
            "headers": "0",
        },
        timeout=30,
    )
    response.raise_for_status()
    response.encoding = "utf-8"
    values = list(csv.reader(response.text.splitlines()))
    _report_totals(values)
    return values


def _revenue_period_start(
    norm: Callable[[Any], str],
    full_values: list[list[Any]],
) -> date | None:
    del norm, full_values
    return REVENUE_PERIOD_START


def _read_revenue_values(google_client, engine_instance=None) -> list[list[Any]]:
    if engine_instance is not None:
        with engine_instance().connect() as conn:
            revenue_store.ensure_schema(conn)
            return revenue_store.values_from_db(conn)
    credential_error: Exception | None = None
    try:
        worksheet = google_client().open_by_key(REVENUE_SPREADSHEET_ID).worksheet(REVENUE_WORKSHEET)
        return worksheet.get_all_values()
    except Exception as exc:
        credential_error = exc

    # Quản lý Thu Chi is intentionally shared read-only. The Visualization
    # endpoint omits rows hidden by the active Sheet filter, so the production
    # fallback must use the full-workbook CSV export instead.
    try:
        return _read_public_revenue_values()
    except Exception as public_exc:
        raise HTTPException(
            503,
            f"Không đọc được Quản lý Thu Chi · sheet {REVENUE_WORKSHEET}: "
            f"{type(credential_error).__name__} / {type(public_exc).__name__}.",
        ) from public_exc


def _read_revenue_report_values(google_client) -> list[list[Any]]:
    credential_error: Exception | None = None
    try:
        worksheet = google_client().open_by_key(REVENUE_SPREADSHEET_ID).worksheet(REVENUE_REPORT_WORKSHEET)
        values = worksheet.get("B2:B3")
        _report_totals(values)
        return values
    except Exception as exc:
        credential_error = exc

    try:
        return _read_public_report_values()
    except Exception as public_exc:
        raise HTTPException(
            503,
            f"Không đọc được Tổng thu/Tổng chi tại {REVENUE_REPORT_WORKSHEET}!B2:B3: "
            f"{type(credential_error).__name__} / {type(public_exc).__name__}.",
        ) from public_exc


def _progressive_detail_map(conn, start_date: date, end_date: date, progressive_key) -> dict[str, str]:
    rows = conn.execute(text("""
        SELECT record_uid, leave_date, leave_reason, COALESCE(detail,'') AS detail,
               source_row, id
        FROM leave_records
        WHERE leave_date BETWEEN :start_date AND :end_date
        ORDER BY leave_date, COALESCE(source_row, 2147483647), id
    """), {"start_date": start_date, "end_date": end_date}).mappings().all()
    counters: dict[tuple[date, str], int] = {}
    output: dict[str, str] = {}
    weekend_unpaid_enabled = load_weekend_unpaid_enabled(conn)
    for row in rows:
        key = str(progressive_key(row.get("leave_reason")) or "")
        if not key or not progressive_penalty_applies(
            row.get("leave_date"),
            row.get("leave_reason"),
            weekend_unpaid_enabled=weekend_unpaid_enabled,
        ):
            continue
        bucket = (row["leave_date"], key)
        counters[bucket] = counters.get(bucket, 0) + 1
        ordinal = counters[bucket]
        detail = str(row.get("detail") or "").strip()
        if re.match(r"^\s*Người\s+Thứ\s+\d+", detail, flags=re.IGNORECASE):
            output[str(row["record_uid"])] = detail
        else:
            prefix = f"Người Thứ {ordinal}"
            output[str(row["record_uid"])] = f"{prefix} | {detail}" if detail else prefix
    return output


def install_revenue_leave_list_routes(
    app, *, engine_instance, current_identity, require_feature, feature_allowed,
    norm, progressive_key, google_client, api_module=None,
) -> None:
    if getattr(app.state, "revenue_leave_list_installed", False):
        return

    revenue_group = permissions.FEATURE_GROUPS.setdefault("Doanh thu", {})
    revenue_group[REVENUE_FEATURE] = "Xem Doanh thu"
    revenue_group[REVENUE_TIP_FEATURE] = "Nhập Tiền TIP trong kỳ"
    revenue_group[REVENUE_ENTRY_FEATURE] = "Nhập Thu Chi"
    revenue_group[REVENUE_ENTRY_EDIT_FEATURE] = "Sửa bản ghi nhập trong ngày hiện tại"
    revenue_group[REVENUE_ENTRY_DELETE_FEATURE] = "Xóa bản ghi nhập trong ngày hiện tại"
    permissions.FEATURES[REVENUE_FEATURE] = "Xem Doanh thu"
    permissions.FEATURES[REVENUE_TIP_FEATURE] = "Nhập Tiền TIP trong kỳ"
    permissions.FEATURES[REVENUE_ENTRY_FEATURE] = "Nhập Thu Chi"
    permissions.FEATURES[REVENUE_ENTRY_EDIT_FEATURE] = "Sửa bản ghi nhập trong ngày hiện tại"
    permissions.FEATURES[REVENUE_ENTRY_DELETE_FEATURE] = "Xóa bản ghi nhập trong ngày hiện tại"
    permissions.DEFAULT_ROLE_FEATURES.setdefault("admin", set()).update({REVENUE_FEATURE, REVENUE_TIP_FEATURE, REVENUE_ENTRY_FEATURE, REVENUE_ENTRY_EDIT_FEATURE, REVENUE_ENTRY_DELETE_FEATURE})

    original_records = _find_route(app, "/v2/leave/records", "GET")
    original_daily_stats = _find_route(app, "/v2/leave/daily-stats", "GET")

    @app.get("/v2/revenue/health")
    def revenue_health():
        return {
            "ok": True,
            "release": RELEASE,
            "storage": "postgresql",
            "transaction_table": revenue_store.TABLE,
            "period_metadata": True,
            "period_start_source": "fixed 2026-09-05",
            "summary_scope": "PostgreSQL revenue ledger",
            "current_date_source": "ledger transaction date / note",
            "period_tip": True,
            "net_formula": "total_income-total_expense",
            "balance_formula": "(total_income-total_expense)-period_tip",
            "entry_form": True,
            "web_entry": True,
            "report_link": False,
        }

    @app.get("/v2/leave/list-enhancements/health")
    def leave_list_enhancements_health():
        return {
            "ok": True,
            "release": RELEASE,
            "progressive_detail": True,
            "stats_scope": "all_registered_employees",
            "penalty_visibility": "permission_gated",
        }

    @app.get("/v2/revenue/source")
    def revenue_source(ident=Depends(current_identity)):
        with engine_instance().connect() as conn:
            require_feature(conn, ident, REVENUE_FEATURE)
            return {"ok": True, **revenue_auto.mode(conn)}

    @app.put("/v2/revenue/source")
    def update_revenue_source(body: RevenueSourceUpdate, ident=Depends(current_identity)):
        _require_revenue_admin(ident)
        with engine_instance().begin() as conn:
            require_feature(conn, ident, REVENUE_FEATURE)
            return {"ok": True, **revenue_auto.set_mode(conn, body.source, body.revision,
                str(getattr(ident, "employee_username", "") or ""))}

    @app.get("/v2/revenue/tip-summary")
    def revenue_tip_summary(start: date, end: date, ident=Depends(current_identity)):
        if start > end:
            raise HTTPException(400, "Ngày bắt đầu Tiền TIP không được sau Đến ngày.")
        with engine_instance().connect() as conn:
            require_feature(conn, ident, REVENUE_FEATURE)
            shared = revenue_auto.mode(conn)
            amount = revenue_auto.tip_total(conn, start, end, auto=shared["source"] == "auto")
            return {"ok": True, "period_tip": amount, "source": shared["source"],
                    "period_tip_start": start.isoformat(), "period_tip_end": end.isoformat()}

    @app.put("/v2/revenue/tip-period")
    def save_auto_tip_period(body: RevenueTipPeriod, ident=Depends(current_identity)):
        if (body.start_date > body.end_date or body.start_date < REVENUE_PERIOD_START
                or body.end_date > datetime.now(VN_TZ).date()):
            raise HTTPException(400, "Kỳ TIP Auto phải nằm trong khoảng 05-09-2026 đến ngày hiện tại.")
        with engine_instance().begin() as conn:
            require_feature(conn, ident, REVENUE_TIP_FEATURE)
            revenue_auto.lock_mode(conn)
            if revenue_auto.mode(conn)["source"] != "auto":
                raise HTTPException(409, "Chế độ Doanh thu đã thay đổi. Hãy tải lại.")
            tip = revenue_auto.tip_total(conn, body.start_date, body.end_date, auto=True)
            _save_period_tip(conn, body.start_date.isoformat(), body.end_date.isoformat(), tip,
                             getattr(ident, "employee_username", ""), auto=True)
            totals = revenue_auto.totals(revenue_auto.daily(conn))
            return {"ok": True, "period_tip": tip, "period_tip_start": body.start_date.isoformat(),
                    "period_tip_end": body.end_date.isoformat(), "balance": round(totals["net_income"] - tip, 2),
                    "message": "Đã lưu kỳ TIP; số tiền được tính từ dữ liệu hệ thống."}

    @app.get("/v2/revenue/summary")
    def revenue_summary(
        source: Literal["manual", "auto", "manual_tip_auto"] = Query("manual"),
        time_range: str = Query("all"),
        start: date | None = Query(None),
        end: date | None = Query(None),
        ident=Depends(current_identity),
    ):
        start_date, end_date = _range_bounds(time_range, start, end)
        with engine_instance().connect() as conn:
            require_feature(conn, ident, REVENUE_FEATURE)
            shared = revenue_auto.mode(conn)
            source = shared["source"]  # Never let a client override the global mode.
            can_edit_tip = bool(feature_allowed(conn, ident, REVENUE_TIP_FEATURE))
            can_create_entry = bool(feature_allowed(conn, ident, REVENUE_ENTRY_FEATURE))
            can_edit_entry = bool(feature_allowed(conn, ident, REVENUE_ENTRY_EDIT_FEATURE))
            can_delete_entry = bool(feature_allowed(conn, ident, REVENUE_ENTRY_DELETE_FEATURE))
            if source == "auto":
                auto = _auto_revenue(conn, start_date, end_date)
                today = datetime.now(VN_TZ).date()
                first, last = revenue_auto.bounds(start_date, end_date)
                default_tip_start = max(today.replace(day=1 if today.day <= 15 else 16), REVENUE_PERIOD_START)
                tip_setting = _period_tip(conn, default_tip_start.isoformat(), today.isoformat(), auto=True)
                tip_start = max(_parse_date(tip_setting["period_start"]) or default_tip_start, REVENUE_PERIOD_START)
                tip_end = min(_parse_date(tip_setting["period_end"]) or today, today)
                period_tip = revenue_auto.tip_total(conn, tip_start, tip_end, auto=True)
                return {
                    **auto, "ok": True, "release": RELEASE, "source": "auto", "source_revision": shared["revision"],
                    "source_label": "Tự động từ hệ thống", "storage": "postgresql", "time_range": time_range,
                    "start_date": first.isoformat(), "start_date_label": first.strftime("%d-%m-%Y"),
                    "end_date": last.isoformat(), "business_date": today.isoformat(),
                    "current_date": today.isoformat(), "current_date_label": today.strftime("%d-%m-%Y"),
                    "can_edit_tip": can_edit_tip, "can_create_entry": False, "can_edit_entry": False,
                    "can_delete_entry": False, "can_admin_crud": False,
                    "period_tip": period_tip, "period_tip_start": tip_start.isoformat(), "period_tip_end": tip_end.isoformat(),
                    "balance": round(auto["net_income"] - period_tip, 2),
                }
            entries = revenue_store.list_entries(conn, start_date=start_date, end_date=end_date)
            report_dates = [_parse_date(row.get("date")) for row in entries]
            report_date = max((value for value in report_dates if value is not None), default=None)
            total_income = round(sum(row["amount"] for row in entries if row["type"] == "Thu"), 2)
            total_expense = round(sum(row["amount"] for row in entries if row["type"] == "Chi"), 2)
            tip_setting = _period_tip(conn, start_date.isoformat() if start_date else "", end_date.isoformat() if end_date else "")
            auto_tip = _auto_revenue(conn, start_date, end_date)["tip_revenue"] if source == "manual_tip_auto" else None
        tip = float(tip_setting["amount"])
        return {
            "ok": True, "release": RELEASE, "source": source, "source_revision": shared["revision"],
            "source_label": "Dịch vụ Manual · Tip Auto" if source == "manual_tip_auto" else "Thủ công",
            "storage": "postgresql", "transaction_table": revenue_store.TABLE, "time_range": time_range,
            "start_date": (start_date or REVENUE_PERIOD_START).isoformat(),
            "start_date_label": (start_date or REVENUE_PERIOD_START).strftime("%d-%m-%Y"),
            "end_date": end_date.isoformat() if end_date else "",
            "current_date": (report_date or datetime.now(VN_TZ).date()).isoformat(),
            "current_date_label": (report_date or datetime.now(VN_TZ).date()).strftime("%d-%m-%Y"),
            "current_date_source": "last_ledger_transaction" if report_date else "server_date_fallback",
            "business_date": datetime.now(VN_TZ).date().isoformat(),
            "can_edit_tip": can_edit_tip, "can_create_entry": can_create_entry,
            "can_edit_entry": can_edit_entry, "can_delete_entry": can_delete_entry,
            "can_admin_crud": can_edit_entry or can_delete_entry, "entries": entries, "transaction_count": len(entries),
            "total_income": total_income, "total_expense": total_expense,
            "period_tip": round(auto_tip if auto_tip is not None else tip, 2),
            "tip_revenue": round(auto_tip if auto_tip is not None else tip, 2),
            "service_revenue": total_income,
            "total_revenue": round(total_income + (auto_tip if auto_tip is not None else tip), 2),
            "period_tip_start": tip_setting["period_start"],
            "period_tip_end": tip_setting["period_end"],
            "net_income": round(total_income - total_expense, 2),
            "balance": round(total_income - total_expense - tip, 2),
        }

    @app.post("/v2/revenue/entry")
    def create_revenue_entry(body: RevenueEntryCreate, background_tasks: BackgroundTasks, ident=Depends(current_identity)):
        income_amount = round(float(body.income_amount or 0), 2)
        expense_amount = round(float(body.expense_amount or 0), 2)
        if income_amount <= 0 and expense_amount <= 0:
            raise HTTPException(400, "Hãy nhập ít nhất một số tiền Thu hoặc Chi lớn hơn 0.")

        entries: list[tuple[str, float, str]] = []
        if income_amount > 0:
            entries.append(("Thu", income_amount, body.income_note))
        if expense_amount > 0:
            entries.append(("Chi", expense_amount, body.expense_note))
        with engine_instance().begin() as conn:
            require_feature(conn, ident, REVENUE_ENTRY_FEATURE)
            revenue_auto.require_manual(conn)
            duplicate_rows = revenue_store.find_duplicate_web_entries(conn, entries=entries)
            if duplicate_rows and not body.confirm_duplicate:
                labels = ", ".join(
                    f"{row['transaction_type']} {round(float(row['amount'])):,}đ".replace(",", ".")
                    for row in duplicate_rows
                )
                raise HTTPException(
                    status_code=409,
                    detail={
                        "code": "duplicate_revenue_entry",
                        "message": (
                            f"Phát hiện giao dịch có cùng số tiền và cùng nội dung: {labels}. "
                            "Bạn có muốn vẫn lưu giao dịch này không?"
                        ),
                        "duplicates": duplicate_rows,
                    },
                )
            saved_rows = revenue_store.insert_web_entries(
                conn, transaction_date=body.transaction_date, entries=entries, ident=ident,
            )

        actor = str(getattr(ident, "employee_username", "") or getattr(ident, "full_name", "") or "")
        for transaction_type, amount, note in entries:
            background_tasks.add_task(
                _dispatch_revenue_admin_push, engine_instance=engine_instance, api_module=api_module,
                event="create", detail={"type": transaction_type, "amount": amount, "note": note, "actor": actor},
            )

        parts = []
        if income_amount > 0:
            parts.append(f"Thu {round(income_amount):,}đ".replace(",", "."))
        if expense_amount > 0:
            parts.append(f"Chi {round(expense_amount):,}đ".replace(",", "."))
        return {
            "ok": True,
            "saved_rows": saved_rows,
            "message": "Đã lưu cùng thời điểm " + " và ".join(parts) + " trên server.",
        }

    def _require_revenue_admin(ident):
        if str(getattr(ident, "role", "") or "").strip().lower() != "admin":
            raise HTTPException(403, "Chỉ Admin được sửa hoặc xóa báo cáo doanh thu.")

    @app.post("/v2/revenue/import.xlsx")
    async def import_revenue_excel(request: Request, mode: Literal["append", "replace"] = Query("append"), ident=Depends(current_identity)):
        _require_revenue_admin(ident)
        content = await request.body()
        with engine_instance().begin() as conn:
            require_feature(conn, ident, REVENUE_FEATURE)
            revenue_auto.require_manual(conn)
            try:
                result = revenue_store.import_ledger_xlsx(
                    conn, content, mode=mode,
                    actor=str(getattr(ident, "employee_username", "") or ""),
                )
            except ValueError as exc:
                raise HTTPException(400, str(exc))
        action = "thay thế toàn bộ dữ liệu" if mode == "replace" else "cập nhật dữ liệu mới"
        return {"ok": True, **result, "message": f"Đã {action}: thêm {result['inserted']} dòng, bỏ qua {result['skipped']} dòng trùng."}

    @app.patch("/v2/revenue/entries/{entry_id}")
    def update_revenue_entry(entry_id: int, body: RevenueEntryUpdate, background_tasks: BackgroundTasks, ident=Depends(current_identity)):
        admin_unrestricted = str(getattr(ident, "role", "") or "").strip().lower() == "admin"
        entered_at = None
        if body.entered_date is not None or body.entered_time is not None:
            if body.entered_date is None or not body.entered_time:
                raise HTTPException(400, "Phải nhập đủ Ngày nhập và Giờ nhập.")
            try:
                entered_clock = datetime.strptime(body.entered_time, "%H:%M:%S").time()
            except ValueError:
                try:
                    entered_clock = datetime.strptime(body.entered_time, "%H:%M").time()
                except ValueError:
                    raise HTTPException(400, "Giờ nhập phải đúng HH:MM hoặc HH:MM:SS.")
            entered_at = datetime.combine(body.entered_date, entered_clock).replace(tzinfo=VN_TZ)
        with engine_instance().begin() as conn:
            require_feature(conn, ident, REVENUE_ENTRY_EDIT_FEATURE)
            revenue_auto.require_manual(conn)
            try:
                result = revenue_store.update_entry(
                    conn, entry_id=entry_id, transaction_type=body.transaction_type,
                    amount=body.amount, transaction_date=body.transaction_date, note=body.note,
                    entered_by_name=body.entered_by_name, entered_at=entered_at,
                    actor=str(getattr(ident, "employee_username", "") or ""),
                    required_entered_date=None if admin_unrestricted else datetime.now(VN_TZ).date(),
                )
            except KeyError:
                raise HTTPException(404, "Không tìm thấy bản ghi doanh thu.")
            except PermissionError:
                raise HTTPException(403, "Chỉ được sửa bản ghi đã nhập trong ngày hiện tại.")
        background_tasks.add_task(
            _dispatch_revenue_admin_push, engine_instance=engine_instance, api_module=api_module,
            event="update", detail={"entry_id": entry_id, "type": body.transaction_type, "amount": body.amount,
                                    "note": body.note, "actor": str(getattr(ident, "employee_username", "") or "")},
        )
        return {"ok": True, **result, "message": "Đã sửa bản ghi doanh thu."}

    @app.delete("/v2/revenue/entries/{entry_id}")
    def delete_revenue_entry(entry_id: int, background_tasks: BackgroundTasks, ident=Depends(current_identity)):
        admin_unrestricted = str(getattr(ident, "role", "") or "").strip().lower() == "admin"
        notification_detail: dict[str, Any] = {"entry_id": entry_id, "actor": str(getattr(ident, "employee_username", "") or "")}
        with engine_instance().begin() as conn:
            require_feature(conn, ident, REVENUE_ENTRY_DELETE_FEATURE)
            revenue_auto.require_manual(conn)
            try:
                current = next((row for row in revenue_store.list_entries(conn) if row.get("id") == entry_id), None)
                if current:
                    notification_detail.update({"type": current.get("type"), "amount": current.get("amount"), "note": current.get("note")})
                revenue_store.soft_delete_entry(
                    conn, entry_id=entry_id,
                    actor=str(getattr(ident, "employee_username", "") or ""),
                    required_entered_date=None if admin_unrestricted else datetime.now(VN_TZ).date(),
                )
            except KeyError:
                raise HTTPException(404, "Không tìm thấy bản ghi doanh thu.")
            except PermissionError:
                raise HTTPException(403, "Chỉ được xóa bản ghi đã nhập trong ngày hiện tại.")
        background_tasks.add_task(
            _dispatch_revenue_admin_push, engine_instance=engine_instance, api_module=api_module,
            event="delete", detail=notification_detail,
        )
        return {"ok": True, "message": "Đã xóa bản ghi khỏi báo cáo; timestamp lịch sử gốc không thay đổi."}

    @app.get("/v2/revenue/audit")
    def revenue_audit(
        time_range: str = Query("this_month"), start: date | None = Query(None),
        end: date | None = Query(None), ident=Depends(current_identity),
    ):
        _require_revenue_admin(ident)
        start_date, end_date = _range_bounds(time_range, start, end)
        with engine_instance().connect() as conn:
            require_feature(conn, ident, REVENUE_FEATURE)
            rows = revenue_store.list_audit_entries(conn, start_date=start_date, end_date=end_date)
        return {"ok": True, "rows": rows, "count": len(rows),
                "start_date": start_date.isoformat() if start_date else "",
                "end_date": end_date.isoformat() if end_date else ""}

    @app.get("/v2/revenue/audit/export.xlsx")
    def revenue_audit_export(
        time_range: str = Query("this_month"), start: date | None = Query(None),
        end: date | None = Query(None), ident=Depends(current_identity),
    ):
        _require_revenue_admin(ident)
        start_date, end_date = _range_bounds(time_range, start, end)
        with engine_instance().connect() as conn:
            require_feature(conn, ident, REVENUE_FEATURE)
            rows = revenue_store.list_audit_entries(conn, start_date=start_date, end_date=end_date)
        filename = f"VERA_LichSu_SuaXoa_ThuChi_{datetime.now(VN_TZ):%d-%m-%Y}.xlsx"
        return StreamingResponse(
            _revenue_audit_workbook(rows),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )

    @app.get("/v2/revenue/duplicates")
    def revenue_duplicates(
        time_range: str = Query("all"), start: date | None = Query(None),
        end: date | None = Query(None), ident=Depends(current_identity),
    ):
        _require_revenue_admin(ident)
        start_date, end_date = _range_bounds(time_range, start, end)
        with engine_instance().connect() as conn:
            require_feature(conn, ident, REVENUE_FEATURE)
            result = revenue_store.duplicate_analysis(conn, start_date=start_date, end_date=end_date)
        return {"ok": True, **result}

    @app.get("/v2/revenue/duplicates/export.xlsx")
    def revenue_duplicates_export(
        time_range: str = Query("all"), start: date | None = Query(None),
        end: date | None = Query(None), ident=Depends(current_identity),
    ):
        _require_revenue_admin(ident)
        start_date, end_date = _range_bounds(time_range, start, end)
        with engine_instance().connect() as conn:
            require_feature(conn, ident, REVENUE_FEATURE)
            result = revenue_store.duplicate_analysis(conn, start_date=start_date, end_date=end_date)
        filename = f"VERA_KiemTra_DuLieuTrung_ThuChi_{datetime.now(VN_TZ):%d-%m-%Y}.xlsx"
        return StreamingResponse(
            _revenue_duplicate_workbook(result),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )

    @app.put("/v2/revenue/tip")
    def save_revenue_tip(body: RevenueTipUpdate, ident=Depends(current_identity)):
        with engine_instance().begin() as conn:
            require_feature(conn, ident, REVENUE_TIP_FEATURE)
            revenue_auto.require_manual(conn)
            entries = revenue_store.list_entries(conn)
            report_date = max((_parse_date(row["date"]) for row in entries if _parse_date(row["date"])),
                              default=datetime.now(VN_TZ).date())
            tip_start = body.start_date or report_date.replace(day=1 if report_date.day <= 15 else 16)
            tip_end = body.end_date or report_date
            if tip_start > tip_end:
                raise HTTPException(400, "Ngày bắt đầu Tiền TIP không được sau Đến ngày.")
            _save_period_tip(conn, tip_start.isoformat(), tip_end.isoformat(), body.amount,
                             getattr(ident, "employee_username", ""))
            total_income = sum(row["amount"] for row in entries if row["type"] == "Thu")
            total_expense = sum(row["amount"] for row in entries if row["type"] == "Chi")
        tip = round(float(body.amount), 2)
        return {"ok": True, "release": RELEASE, "period_tip": tip,
                "balance": round(total_income - total_expense - tip, 2),
                "period_tip_start": tip_start.isoformat(), "period_tip_end": tip_end.isoformat(),
                "period_start": tip_start.isoformat(), "period_end": tip_end.isoformat(),
                "message": f"Đã lưu Tiền TIP trong kỳ {tip_start:%d-%m-%Y} đến {tip_end:%d-%m-%Y}."}

    @app.get("/v2/leave/records")
    def leave_records_enhanced(
        date_value: date | None = Query(default=None, alias="date"),
        start_date: date | None = Query(default=None, alias="start"),
        end_date: date | None = Query(default=None, alias="end"),
        ident=Depends(current_identity),
    ):
        payload = original_records(date_value=date_value, start_date=start_date, end_date=end_date, ident=ident)
        records = list(payload.get("records") or [])
        if not records:
            return payload
        effective_start = date_value or start_date
        effective_end = date_value or end_date
        if not effective_start or not effective_end:
            return payload
        with engine_instance().connect() as conn:
            detail_map = _progressive_detail_map(conn, effective_start, effective_end, progressive_key)
        for item in records:
            enriched = detail_map.get(str(item.get("record_uid") or ""))
            if enriched:
                item["detail"] = enriched
        return {**payload, "records": records}

    @app.get("/v2/leave/daily-stats")
    def leave_daily_stats_scoped(
        start_date: date = Query(alias="start"),
        end_date: date = Query(alias="end"),
        employee: str = Query(default="", max_length=200),
        ident=Depends(current_identity),
    ):
        requested_employee = str(employee or "").strip()
        return original_daily_stats(start_date=start_date, end_date=end_date, employee=requested_employee, ident=ident)

    app.state.revenue_leave_list_installed = True
    app.state.revenue_leave_list_release = RELEASE
