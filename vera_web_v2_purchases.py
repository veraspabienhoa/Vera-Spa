"""Permission-checked purchase entry, import, audit and export API."""
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from uuid import UUID
import hashlib
import json
from fastapi import Depends, HTTPException, Query, Request
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import text
import vera_purchase_store as store
import vera_web_v2_permissions as permissions
from vera_web_v2_purchase_reconcile import _resolve_range


class PurchaseRow(BaseModel):
    purchase_date: date
    item: str = Field(min_length=1, max_length=2000)
    quantity: Decimal = Field(gt=0, le=100000000, max_digits=16, decimal_places=4)
    unit_price: Decimal = Field(ge=0, le=100000000000, max_digits=16, decimal_places=2)
    note: str = Field(default='', max_length=4000)


class PurchaseBatch(BaseModel):
    request_id: UUID
    rows: list[PurchaseRow] = Field(min_length=1, max_length=100)


class PurchaseEdit(PurchaseRow):
    revision: int = Field(ge=0)
    quantity: str = Field(min_length=1, max_length=50)
    amount: Decimal = Field(ge=0, lt=10000000000000000, max_digits=18, decimal_places=2)


def row_values(row):
    values = row.model_dump()
    values['item'] = values['item'].strip()
    if not values['item']:
        raise HTTPException(400, 'Phải nhập chi tiết hàng hóa.')
    if isinstance(row, PurchaseEdit):
        try:
            store.decimal_value(row.quantity.rstrip('%'))
        except ValueError as exc:
            raise HTTPException(400,str(exc)) from None
        values['amount'] = row.amount
    else:
        values['amount'] = (row.quantity * row.unit_price).quantize(Decimal('.01'), rounding=ROUND_HALF_UP)
    if values['amount'] >= Decimal('1e16'):
        raise HTTPException(400, 'Thành tiền vượt giới hạn.')
    values['quantity'] = str(row.quantity)
    return values


def install_purchase_routes(app, *, engine_instance, current_identity, require_feature, feature_allowed):
    if getattr(app.state, 'purchase_routes_installed', False):
        return
    group = {'purchase_view':'Xem Nhập mua', 'purchase_create':'Nhập mua hàng',
             'purchase_edit':'Sửa bản ghi nhập trong ngày hiện tại',
             'purchase_delete':'Xóa bản ghi nhập trong ngày hiện tại'}
    permissions.FEATURE_GROUPS['Nhập mua'] = group
    permissions.FEATURES.update(group)
    for key in ('purchase_create','purchase_edit','purchase_delete'):
        permissions.FEATURE_DEPENDENCIES[key] = {'purchase_view'}
    permissions.DEFAULT_ROLE_FEATURES.setdefault('admin', set()).update(group)

    def actor(ident):
        return str(ident.employee_username or '')

    def admin(ident):
        if ident.role != 'admin':
            raise HTTPException(403, 'Chỉ Admin được import và xem lịch sử.')

    def prepare(conn, ident, feature):
        require_feature(conn, ident, feature)
        store.ensure_schema(conn)

    def replay(conn, key, fingerprint):
        conn.execute(text('SELECT pg_advisory_xact_lock(hashtextextended(:key,0))'),dict(key=key))
        previous=conn.execute(text('SELECT fingerprint,result FROM vera_purchase_request WHERE key=:key'),dict(key=key)).mappings().first()
        if previous and previous['fingerprint'] != fingerprint:
            raise HTTPException(409,'Nội dung lần gửi trước khác lần này. Hãy tải lại để kiểm tra dữ liệu đã lưu.')
        return previous['result'] if previous else None

    def remember(conn,key,fingerprint,result):
        conn.execute(text('INSERT INTO vera_purchase_request(key,fingerprint,result) VALUES(:key,:fingerprint,CAST(:result AS jsonb))'),
                     dict(key=key,fingerprint=fingerprint,result=json.dumps(result)))

    def entries(conn, preset, start, end):
        first, last = _resolve_range(preset, start, end)
        rows = store.list_entries(conn, first, last)
        return rows, first, last

    @app.get('/v2/purchases')
    def listing(preset: str='this_month', start: date | None=None, end: date | None=None, ident=Depends(current_identity)):
        with engine_instance().begin() as conn:
            prepare(conn, ident, 'purchase_view')
            rows, first, last = entries(conn, preset, start, end)
            allowed = {key:feature_allowed(conn, ident, key) for key in group}
        return dict(rows=rows, total=sum(r['amount'] for r in rows), start=first, end=last, permissions=allowed)

    @app.post('/v2/purchases')
    def create(body: PurchaseBatch, ident=Depends(current_identity)):
        prepared = [row_values(r) for r in body.rows]
        with engine_instance().begin() as conn:
            prepare(conn, ident, 'purchase_create')
            key=f'create:{actor(ident)}:{body.request_id}'
            fingerprint=hashlib.sha256(body.model_dump_json().encode()).hexdigest()
            previous=replay(conn,key,fingerprint)
            if previous is not None:
                return previous
            for index, values in enumerate(prepared):
                values.update(actor=actor(ident), key=f'{actor(ident)}:{body.request_id}:{index}')
                conn.execute(text('''INSERT INTO vera_purchase_entry
                  (purchase_date,item,quantity,unit_price,amount,note,entered_at,entered_by,request_key)
                  VALUES(:purchase_date,:item,:quantity,:unit_price,:amount,:note,now(),:actor,:key)
                  ON CONFLICT(request_key) DO NOTHING'''), values)
            remember(conn,key,fingerprint,{'ok':True})
        return {'ok':True}

    def locked(conn, ident, entry_id, revision):
        row = conn.execute(text('SELECT * FROM vera_purchase_entry WHERE id=:id FOR UPDATE'),dict(id=entry_id)).mappings().first()
        try:
            store.require_editable(row, ident.role, revision)
        except KeyError as exc:
            raise HTTPException(404, str(exc)) from None
        except PermissionError as exc:
            raise HTTPException(403, str(exc)) from None
        except ValueError as exc:
            raise HTTPException(409, str(exc)) from None
        return dict(row)

    @app.patch('/v2/purchases/{entry_id}')
    def edit(entry_id:int, body:PurchaseEdit, ident=Depends(current_identity)):
        values = row_values(body)
        with engine_instance().begin() as conn:
            prepare(conn, ident, 'purchase_edit')
            old = locked(conn,ident,entry_id,body.revision)
            values['id'] = entry_id
            conn.execute(text('''UPDATE vera_purchase_entry SET purchase_date=:purchase_date,item=:item,
              quantity=:quantity,unit_price=:unit_price,amount=:amount,note=:note,revision=revision+1 WHERE id=:id'''),values)
            store.audit(conn,entry_id,'update',old,values,actor(ident))
        return {'ok':True}

    @app.delete('/v2/purchases/{entry_id}')
    def delete(entry_id:int, revision:int=Query(ge=0), ident=Depends(current_identity)):
        with engine_instance().begin() as conn:
            prepare(conn,ident,'purchase_delete')
            old = locked(conn,ident,entry_id,revision)
            conn.execute(text('UPDATE vera_purchase_entry SET deleted=true,revision=revision+1 WHERE id=:id'),dict(id=entry_id))
            store.audit(conn,entry_id,'delete',old,None,actor(ident))
        return {'ok':True}

    @app.post('/v2/purchases/import')
    async def import_file(request:Request, request_id:UUID, mode:str=Query('append', pattern='^(append|replace)$'), ident=Depends(current_identity)):
        admin(ident)
        content = bytearray()
        async for chunk in request.stream():
            content.extend(chunk)
            if len(content) > store.MAX_FILE:
                raise HTTPException(413,'File tối đa 20 MB.')
        try:
            # Validate before acquiring a pooled database connection.
            parsed = store.parse_workbook(bytes(content))
        except Exception:
            raise HTTPException(400,'File Input không hợp lệ. Kiểm tra ngày, số lượng và thành tiền.') from None
        with engine_instance().begin() as conn:
            prepare(conn,ident,'purchase_view')
            key=f'import:{actor(ident)}:{request_id}'
            fingerprint=hashlib.sha256(bytes(content)+mode.encode()).hexdigest()
            previous=replay(conn,key,fingerprint)
            if previous is not None:
                return previous
            result = dict(ok=True,**store.import_workbook(conn,bytes(content),actor(ident),mode,rows=parsed))
            remember(conn,key,fingerprint,result)
        return result

    @app.get('/v2/purchases/export.xlsx')
    def export(preset:str='this_month', start:date | None=None, end:date | None=None, ident=Depends(current_identity)):
        from openpyxl import Workbook
        with engine_instance().begin() as conn:
            prepare(conn,ident,'purchase_view')
            rows, _, _ = entries(conn,preset,start,end)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Input'
        sheet.append(['Ngày mua hàng','Chi tiết hàng hóa','Số lượng','Đơn giá','Thành Tiền','Ghi chú/ Người đặt mua hàng','Ngày nhập','Giờ nhập','User'])
        for row in rows:
            entered = row['entered_at'].astimezone(store.VN_TZ) if row['entered_at'] else None
            values = [row['purchase_date'],row['item'],row['quantity'],row['unit_price'],row['amount'],row['note'],
                      entered.date() if entered else None,entered.strftime('%H:%M:%S') if entered else '',row['entered_by']]
            sheet.append(values)
            for col in (2,3,6,9):
                sheet.cell(sheet.max_row,col).data_type='s'
            for col in (1,7):
                sheet.cell(sheet.max_row,col).number_format='dd-mm-yyyy'
        sheet.freeze_panes='A2'
        sheet.auto_filter.ref=sheet.dimensions
        for col,width in {'A':16,'B':45,'C':14,'D':18,'E':20,'F':40,'G':16,'H':14,'I':22}.items():
            sheet.column_dimensions[col].width=width
        output=BytesIO(); workbook.save(output); output.seek(0)
        return StreamingResponse(output,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                                 headers={'Content-Disposition':'attachment; filename="NhapMua.xlsx"'})

    @app.get('/v2/purchases/audit')
    def history(ident=Depends(current_identity)):
        admin(ident)
        with engine_instance().begin() as conn:
            prepare(conn,ident,'purchase_view')
            rows=[dict(r) for r in conn.execute(text('SELECT * FROM vera_purchase_audit ORDER BY id DESC LIMIT 500')).mappings()]
        return {'rows':rows}

    app.state.purchase_routes_installed=True
